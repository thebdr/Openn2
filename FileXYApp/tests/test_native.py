"""Parity gate for the native (Rust) engine - runnable with a bare `python tests/test_native.py`.
Skips cleanly when the .pyd is not built (rust/install_native.py). With it built, this proves the
two engines are behaviour-identical where it matters:

1. the shared golden vectors return the SAME results from the Python implementations
   (core.PY_IMPLS) and the native module;
2. the escape hatches work (Python-only regex syntax routes to Python, non-str cells fall back);
3. the calamine xlsx loader reproduces the openpyxl loader CELL-BY-CELL on the repo's real
   workbooks - plus a load-time comparison, since that is the point of the port."""
import os
import sys
import time

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

if os.environ.get("FILEXY_RUST", "1") == "0":
    print("  SKIP  native engine disabled (FILEXY_RUST=0)")
    sys.exit(0)
try:
    import filexy_core as native
except ImportError:
    print("  SKIP  native engine not built (run: py rust/install_native.py)")
    sys.exit(0)

from filexy import core, files  # noqa: E402

failures = []


def check(cond, label):
    if not cond:
        failures.append(label)
    print(("  PASS  " if cond else "  FAIL  ") + label)


check(core.ENGINE == "rust", "core swapped to the native engine")

# --- 1. both engines, same golden vectors ------------------------------------------------------- #
rows = [["S1", "I2.0", "door"], ["S1", "I10.1", "estop"], ["S2", "I2.2", "door open"], ["", "", ""]]
py = core.PY_IMPLS

for label, filters in [
    ("column substring", [{"col": 0, "text": "s1"}]),
    ("cascade AND", [{"col": 0, "text": "S1"}, {"col": 2, "values": {"door"}}]),
    ("quick search", [{"col": None, "text": "door"}]),
    ("regex", [{"col": 1, "text": r"^I\d+\.0$", "regex": True}]),
    ("broken regex matches nothing", [{"col": 0, "text": "([", "regex": True}]),
    ("empty spec passes all", [{"col": 0, "text": ""}]),
]:
    check(py["apply_filters"](rows, filters) == core.apply_filters(rows, filters),
          f"apply_filters parity: {label}")

for label, sort in [("asc", (1, "asc")), ("desc", (1, "desc")), ("released", None)]:
    check(py["sorted_view"](rows, [0, 1, 2, 3], sort) == core.sorted_view(rows, [0, 1, 2, 3], sort),
          f"sorted_view parity: {label} (blanks last)")

check(py["distinct_values"](rows, [0, 1, 2], 2) == core.distinct_values(rows, [0, 1, 2], 2),
      "distinct_values parity")
check(py["to_tsv"](["A", "B"], [["1", "x\ty"], ["2", "z"]], [1, 0])
      == core.to_tsv(["A", "B"], [["1", "x\ty"], ["2", "z"]], [1, 0]), "to_tsv parity")
stats_rows = [["1.5"], ["2.5"], [""], ["x"], [" "]]
check(py["column_stats"](stats_rows, [0, 1, 2, 3, 4], 0)
      == core.column_stats(stats_rows, [0, 1, 2, 3, 4], 0),
      "column_stats parity (blank = '' exactly, whitespace is non-blank)")
multiline = [["one\ntwo", "plain"]]
check(py["sanitize_rows"](multiline) == core.sanitize_rows(multiline), "sanitize_rows parity")

# --- 2. the escape hatches ----------------------------------------------------------------------- #
# regex routing is EXACT (re.compile pre-check + native regex_ok), not a token list: Python-only
# syntax must run on Python (else the native engine silently matches nothing), and Python-broken
# patterns must match nothing even when Rust considers them valid.
for label, data, pattern in [
    ("lookahead", [["I2.0"]], r"^I(?=\d)"),
    (r"\Z anchor", [["door"]], r"door\Z"),
    (r"backreference \4", [["abcdd"]], r"(a)(b)(c)(d)\4"),
    ("named backreference", [["aa"]], r"(?P<x>a)(?P=x)"),
    ("conditional group", [["ab"]], r"(a)?(?(1)b|c)"),
    ("atomic group", [["abc"]], r"(?>ab)c"),
    ("Python-broken mid-pattern (?i)", [["xay"]], r"x(?i)a"),
    (r"Python-broken \x{41}", [["A"]], r"\x{41}"),
]:
    spec = [{"col": 0, "text": pattern, "regex": True}]
    check(core.apply_filters(data, spec) == py["apply_filters"](data, spec),
          f"regex routing parity: {label}")
mixed = [[None, 42, "text\nmore"]]
check(core.sanitize_rows(mixed) == py["sanitize_rows"](mixed),
      "non-str cells fall back to Python (TypeError hatch)")

# --- 2b. full Unicode case folding (Python casefold, not lowercase) ------------------------------- #
de = [["straße"], ["PRESS"], ["strassf"]]
for label, spec in [("'ss' finds ß", [{"col": 0, "text": "ss"}]),
                    ("'ß' finds SS", [{"col": 0, "text": "ß"}])]:
    check(core.apply_filters(de, spec) == py["apply_filters"](de, spec),
          f"casefold parity: {label}")
check(core.sorted_view(de, [0, 1, 2], (0, "asc")) == py["sorted_view"](de, [0, 1, 2], (0, "asc")),
      "casefold parity: natural sort of ß vs ss")

# --- 3. xlsx loader parity on the repo's real workbooks ------------------------------------------ #
repo = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
workbooks = [p for p in (
    os.path.join(repo, "Shared", "HardwareConfigBuilderData", "DeviceTypesDatabase.xlsx"),
    os.path.join(repo, "Pipeline3App", "assets", "ButtonsLayout.xlsx"),
) if os.path.exists(p)]
if not workbooks:
    print("  SKIP  no real workbooks found for the loader parity check")

def trim_trailing(cols, rows):
    """Drop trailing ALL-empty rows - openpyxl yields them for formatting-only cells inside the
    sheet's declared dimension, calamine's used range (values only) does not. The one accepted
    loader divergence; every remaining cell must be identical."""
    while rows and not any(rows[-1]):
        rows = rows[:-1]
    return cols, rows


for path in workbooks:
    name = os.path.basename(path)
    sheets_py, sheets_rs = files._xlsx_sheets_py(path), native.xlsx_sheets(path)
    check(sheets_py == sheets_rs, f"{name}: sheet names")
    for sheet in sheets_py:
        cols_py, rows_py = trim_trailing(*files._read_xlsx_py(path, sheet))
        cols_rs, rows_rs = trim_trailing(*native.read_xlsx(path, sheet))
        ok = (cols_py, rows_py) == (cols_rs, rows_rs)
        check(ok, f"{name} [{sheet}]: cell-identical ({len(rows_py)} rows)")
        if not ok:  # pinpoint the first divergence for debugging
            if cols_py != cols_rs:
                print(f"         header differs: {cols_py!r} vs {cols_rs!r}")
            if len(rows_py) != len(rows_rs):
                print(f"         row count: openpyxl {len(rows_py)} vs calamine {len(rows_rs)}")
            for r, (a, b) in enumerate(zip(rows_py, rows_rs)):
                if a != b:
                    for c, (x, y) in enumerate(zip(a, b)):
                        if x != y:
                            print(f"         first diff row {r} col {c}: {x!r} vs {y!r}")
                            break
                    else:
                        print(f"         row {r} width: {len(a)} vs {len(b)}")
                    break

# --- 4. loader parity on the value classes the real workbooks don't exercise --------------------- #
try:
    import datetime as dt
    import tempfile
    from openpyxl import Workbook

    book = Workbook()
    ws = book.active
    for row in (["kind", "value"], ["int", 4], ["float_integral", 4.0], ["float", 4.5],
                ["float_tiny", 1.5e-7], ["float_big", 2.5e16], ["bool_t", True],
                ["bool_f", False], ["multiline", "a\nb"],
                ["datetime", dt.datetime(2026, 7, 5, 14, 30, 15)], ["date", dt.date(2026, 7, 5)],
                ["time", dt.time(14, 30, 15)], ["micro", dt.datetime(2026, 7, 5, 14, 30, 15, 123456)],
                ["unicode", "ü é 漢"], ["none", None]):
        ws.append(row)
    # duration formats -> openpyxl yields timedelta; the yyyymmdd-in-a-date-column accident
    # overflows Python's datetime -> openpyxl substitutes '#VALUE!'
    for kind, value, fmt in (["duration_days", 1.5, "[h]:mm:ss"],
                             ["duration_hms", 0.10416666666666667, "[hh]:mm:ss"],
                             ["date_overflow", 20260705, "yyyy-mm-dd"]):
        ws.append([kind, value])
        ws.cell(row=ws.max_row, column=2).number_format = fmt
    with tempfile.TemporaryDirectory() as tmp:
        probe = os.path.join(tmp, "probe.xlsx")
        book.save(probe)
        got_py, got_rs = files._read_xlsx_py(probe), native.read_xlsx(probe)
        check(got_py == got_rs,
              "synthetic workbook: dates/times/durations/bools/floats/unicode cell-identical")
        if got_py != got_rs:
            for a, b in zip(got_py[1], got_rs[1]):
                if a != b:
                    print(f"         {a!r} vs {b!r}")
        as_dict = dict(got_py[1])
        check(as_dict.get("duration_days") == "1 day, 12:00:00"
              and as_dict.get("date_overflow") == "#VALUE!",
              "synthetic workbook: the duration/overflow spot values are the expected ones")
except ImportError:
    print("  SKIP  openpyxl not available for the synthetic-workbook check")

# --- 5. a calamine LIMITATION must fall back to openpyxl, not make the file unopenable ------------ #
# calamine 0.26 rejects modern cached error literals (#SPILL! & friends) at sheet level; craft a
# minimal workbook with one and check files.read_xlsx (the PUBLIC api) still returns the openpyxl
# result. (openpyxl shows the cached error text like any error cell.)
try:
    import tempfile
    import zipfile

    def _write_spill_xlsx(path):
        sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                 '<sheetData>'
                 '<row r="1"><c r="A1" t="inlineStr"><is><t>col</t></is></c></row>'
                 '<row r="2"><c r="A2" t="e"><v>#SPILL!</v></c></row>'
                 '</sheetData></worksheet>')
        workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                    '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
        wb_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument'
                   '/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument'
                '/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                 '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                 '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package'
                 '.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
                 '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats'
                 '-officedocument.spreadsheetml.sheet.main+xml"/>'
                 '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.'
                 'openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("[Content_Types].xml", types)
            z.writestr("_rels/.rels", rels)
            z.writestr("xl/workbook.xml", workbook)
            z.writestr("xl/_rels/workbook.xml.rels", wb_rels)
            z.writestr("xl/worksheets/sheet1.xml", sheet)

    with tempfile.TemporaryDirectory() as tmp:
        spill = os.path.join(tmp, "spill.xlsx")
        _write_spill_xlsx(spill)
        native_errored = False
        try:
            native.read_xlsx(spill)
        except OSError:
            native_errored = True
        expected = files._read_xlsx_py(spill)
        check(files.read_xlsx(spill) == expected,
              "a #SPILL! cell: files.read_xlsx falls back to the openpyxl result"
              + ("" if native_errored else " (calamine handled it natively)"))
except ImportError:
    print("  SKIP  openpyxl not available for the #SPILL! fallback check")

if workbooks:
    big = workbooks[0]
    t0 = time.perf_counter()
    files._read_xlsx_py(big)
    t_py = time.perf_counter() - t0
    t0 = time.perf_counter()
    native.read_xlsx(big)
    t_rs = time.perf_counter() - t0
    print(f"  INFO  load {os.path.basename(big)}: openpyxl {t_py * 1000:.0f} ms, "
          f"calamine {t_rs * 1000:.0f} ms ({t_py / max(t_rs, 1e-9):.1f}x)")

print("-> all passed" if not failures else f"-> FAILURES: {failures}")
sys.exit(1 if failures else 0)
