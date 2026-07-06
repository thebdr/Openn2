"""File loading for the standalone viewer: CSV (delimiter-sniffed) and xlsx (openpyxl read-only,
lazily imported so the package works without it for CSV-only use). Returns (columns, rows) of
STRINGS - the grid's input shape. Caps guard against runaway sheets, not real data (projects of
10-20k rows are the normal case).

When the native engine is built (rust/install_native.py), xlsx loading goes through calamine -
the actual reason for the Rust port: openpyxl re-parses the workbook XML in Python. The native
loader's output is cell-identical to the openpyxl path (gated by tests/test_native.py) with ONE
accepted difference: trailing rows that hold only formatting (no values) are dropped instead of
showing as blank rows. Opt out with FILEXY_RUST=0."""
from __future__ import annotations

import csv
import os
import sys

MAX_ROWS, MAX_COLS = 200_000, 256


def _native():
    """The compiled `filexy_core` module, or None (not built / disabled via FILEXY_RUST=0)."""
    if os.environ.get("FILEXY_RUST", "1") == "0":
        return None
    try:
        import filexy_core
        return filexy_core
    except ImportError:
        return None


def _native_failed(path: str, exc: Exception) -> None:
    """A native-loader error is a calamine LIMITATION (e.g. the modern cached-error literals like
    #SPILL! that calamine 0.26 rejects), not a reason the file can't open: warn once and let the
    openpyxl path take it."""
    if not getattr(_native_failed, "warned", False):
        _native_failed.warned = True
        print(f"filexy: native xlsx loader failed on {os.path.basename(path)} ({exc}); "
              "falling back to openpyxl", file=sys.stderr)


def read_csv(path: str) -> tuple:
    """(columns, rows) from a CSV - the delimiter sniffed from the first line (',' ';' '\\t' '|'),
    the first row taken as the header."""
    with open(path, encoding="utf-8-sig", newline="") as handle:
        sample = handle.readline()
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(handle, dialect)
        table = [row for _i, row in zip(range(MAX_ROWS + 1), reader)]
    if not table:
        return [], []
    return [str(c) for c in table[0][:MAX_COLS]], [r[:MAX_COLS] for r in table[1:]]


def xlsx_sheets(path: str) -> list:
    """The workbook's sheet names (calamine when the native engine is built, else openpyxl)."""
    native = _native()
    if native is not None:
        try:
            return list(native.xlsx_sheets(path))
        except Exception as exc:
            _native_failed(path, exc)
    return _xlsx_sheets_py(path)


def _xlsx_sheets_py(path: str) -> list:
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_xlsx(path: str, sheet: str | None = None) -> tuple:
    """(columns, rows) from one xlsx sheet (cached values; `sheet=None` = the first). The first
    row is the header; empty cells become ''."""
    native = _native()
    if native is not None:
        try:
            return native.read_xlsx(path, sheet)
        except Exception as exc:
            _native_failed(path, exc)
    return _read_xlsx_py(path, sheet)


def _read_xlsx_py(path: str, sheet: str | None = None) -> tuple:
    """The openpyxl loader - the native path's executable spec (test_native.py compares them
    cell-by-cell on real workbooks)."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet] if sheet else book[book.sheetnames[0]]
        table = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS:
                break
            table.append(["" if c is None else str(c) for c in row[:MAX_COLS]])
    finally:
        book.close()
    if not table:
        return [], []
    return table[0], table[1:]


def read_table(path: str, sheet: str | None = None) -> tuple:
    """(columns, rows) by extension: .csv or .xlsx/.xlsm."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path, sheet)
    return read_csv(path)
