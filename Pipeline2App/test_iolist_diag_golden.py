#!/usr/bin/env python3
"""Acceptance test for the iolist_diag populator against the golden workbook (spec §13).

Part A - the suggested-type port reproduces the golden's AC column EXACTLY (all rows), and the
         canonical script_type (AB) matches where the formula resolves a real type (bare-KI series
         and human-only fire-alarm are documented exceptions - the formula can't emit those).
Part B - on a fixture with AD:AF blanked + DiagnosisBlocks deleted (AB kept - the realistic flow:
         the engineer authors script_type, the tool fills index/diag), `populate` reproduces the
         object index-sharing, the per-block diag layout and the DiagnosisBlocks structure (compared
         THROUGH the block FullName, not the golden's trimmed absolute IDs/node bits, per §13), and
         the DR/DQ/DL members land in _UnresolvedIndex rather than being auto-filled.

Skips cleanly (exit 0) if the golden or the configured catalogue (object_families.csv + Z#/F1/2
in_diag) isn't present, so the suite stays green on a machine without them.

Run:  python test_iolist_diag_golden.py
"""
import os
import sys
import shutil
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import openpyxl
from openpyxl.utils import column_index_from_string as _ci

from pipeline2.core import config, staging
from pipeline2 import iolist_diag
from pipeline2.iolist_diag import reader, script_type as st
from pipeline2.iolist_diag.columns import ColumnResolver
from pipeline2.iolist_diag.models import INPUT_REQUIRED

GOLDEN = r"C:\Source\Repos\_Openn2\Shared\DocumentsValidationData\Passing\8XXX_IOyyyy_R0.0_EmergencyIoListValidationTemplate.xlsx"

failures = 0


def check(name, ok, detail=""):
    global failures
    print(("PASS  " if ok else "FAIL  ") + name + (f"  [{detail}]" if detail else ""))
    if not ok:
        failures += 1


def _skip(msg):
    print(f"SKIP  {msg}")
    print("ALL CHECKS PASS (skipped)")
    raise SystemExit(0)


def _preconditions():
    if not os.path.isfile(GOLDEN):
        _skip(f"golden workbook not found: {GOLDEN}")
    types = config.load_signal_types()
    if not (types.get("Z#", {}).get("in_diagnosis") and types.get("F1/2", {}).get("in_diagnosis")):
        _skip("signal_types.csv: Z# / F1/2 not yet flagged in_diag (catalogue not configured)")
    from pipeline2.iolist_diag import families
    if len(families.load_families()) < 7:
        _skip("object_families.csv missing/incomplete")


def main():
    _preconditions()
    tmp = tempfile.mkdtemp(prefix="idiaggold_")
    pristine = os.path.join(tmp, "golden.xlsx")
    shutil.copy2(GOLDEN, pristine)                 # the golden is open in Excel; a read-copy is fine
    cr = ColumnResolver("IoList")

    # ---- Part A: AC port reproduces the golden exactly --------------------- #
    wb = openpyxl.load_workbook(pristine, data_only=True)
    rows, _ = reader.read_iolist(wb, {"io_list": {"sheet": "NET SAFETY", "header_row": 1}}, cr)
    wb.close()
    ac_bad = ab_bad = 0
    for io in rows:
        sug = st.suggested_type(io)
        gold = "" if io.ex_suggested in (None, "") else io.ex_suggested
        mine = "" if sug in (None, "") else sug
        if str(mine) != str(gold):
            ac_bad += 1
        canon = st.to_canonical(sug, io)
        if io.ex_script_type and canon and canon != INPUT_REQUIRED and canon != io.ex_script_type:
            # documented exceptions: bare 'KI' (series channel is a human refinement)
            if not (canon == "KI" and io.ex_script_type.upper().startswith("KI")):
                ab_bad += 1
    check("AC reproduced exactly for every row", ac_bad == 0, f"{ac_bad} mismatch(es)")
    check("AB matches where the formula resolves a type", ab_bad == 0, f"{ab_bad} mismatch(es)")

    # ---- Part B: index/diag/blocks reproduction (blank AD:AF, keep AB) ------ #
    fixture = os.path.join(tmp, "fixture.xlsx")
    shutil.copy2(pristine, fixture)
    wb = openpyxl.load_workbook(fixture)
    ws = wb["NET SAFETY 50"]
    for r in range(2, ws.max_row + 1):
        for col in ("AD", "AE", "AF"):
            ws.cell(row=r, column=_ci(col)).value = None
    if "DiagnosisBlocks" in wb.sheetnames:
        del wb["DiagnosisBlocks"]
    wb.save(fixture)
    wb.close()

    params = {"io_list": {"path": fixture, "sheet": "NET SAFETY", "header_row": 1},
              "diag_bit_min": 0, "diag_bit_max": 62, "strike_handling": "error",
              "output_dir": os.path.join(tmp, "Output")}
    res = iolist_diag.populate(params, out_dir=os.path.join(tmp, "Output"), quiet=True)
    rs = {r.iorow.row: r for r in res.results}
    blocks = staging.load_diagnostic_blocks({"io_list": {"path": res.output_path}})

    def cab(rn):
        c = rs[rn].diag_cabinet
        return blocks.get(int(c), {}).get("fld", "") if c.lstrip("-").isdigit() else ""

    # object index sharing (KQ<->KI series, channel pairs, doors)
    check("KQ <-> KI series share one index",
          len({rs[118].index, rs[66].index, rs[67].index, rs[68].index, rs[69].index}) == 1
          and rs[118].index != INPUT_REQUIRED)
    check("KQ <-> primary KI share (r114<->r74, r115<->r75)",
          rs[114].index == rs[74].index and rs[115].index == rs[75].index)
    check("E1/2 <-> E2/2 share", rs[57].index == rs[61].index and rs[98].index == rs[102].index)
    check("door DI1/2 <-> DD <-> DI2/2 share", len({rs[106].index, rs[107].index, rs[110].index}) == 1)
    check("encoder N1/2 <-> N2/2 share", rs[80].index == rs[81].index)
    check("emergency areas share one index (Z1/Z2/Z3)", len({rs[41].index, rs[42].index, rs[43].index}) == 1)

    # diag block routing (through the FullName) + bit layout
    check("A -> node cabinet (+MS1.CC1)", cab(17).endswith("+MS1.CC1"))
    check("PW from top of node (bit 62)", rs[5].diag_bit == "62" and cab(5).endswith("+MC1.CC1"))
    check("KQ on its physical node (+PC1.CC1)", cab(114).endswith("+PC1.CC1"))
    check("E1/2 -> +EmergencyPushButtons", cab(57).endswith("+EmergencyPushButtons_1"))
    check("B1/2 -> +SafetyBreakers", cab(58).endswith("+SafetyBreakers_1"))
    check("F1/2 -> +FireAlarm", cab(60).endswith("+FireAlarm_1"))
    check("N1/2 bit 00, N2/2 bit 01 -> +SafetyEncoders",
          cab(80).endswith("+SafetyEncoders_1") and rs[80].diag_bit == "00" and rs[81].diag_bit == "01")
    check("door1 DI1/2 bit 01 + DD bit 02 -> +SafetyDoors",
          cab(106).endswith("+SafetyDoors_1") and rs[106].diag_bit == "01" and rs[107].diag_bit == "02")
    check("emergency areas Z1/Z2/Z3 -> bits 01/02/03 in +EmergencyAreas",
          cab(41).endswith("+EmergencyAreas_1") and (rs[41].diag_bit, rs[42].diag_bit, rs[43].diag_bit) == ("01", "02", "03"))
    check("lone field PA -> +FieldIODevices bit 62", cab(122).endswith("+FieldIODevices_1") and rs[122].diag_bit == "62")

    # DiagnosisBlocks structure
    wb = openpyxl.load_workbook(res.output_path, data_only=True)
    db = wb["DiagnosisBlocks"]
    hdr = [c.value for c in db[1]]
    check("DiagnosisBlocks header (8 cols, ID_Local first)",
          hdr[:6] == ["ID_Local", "ID_SWP", "Functional Unit", "Location", "FullName", "TemplateType"] and hdr[7] == "Count")
    body = [[db.cell(row=r, column=c).value for c in range(1, 9)] for r in range(2, db.max_row + 1)]
    check("ID_Local == ID_SWP every row", all(b[0] == b[1] for b in body))
    tts = [b[5] for b in body]
    check("Type-1 (FL) rows before Type-2 (generated)",
          tts == sorted(tts) and 1 in tts and 2 in tts)
    check("FullName materialized as text (staging-readable)", all(str(b[4]).startswith("=S1") for b in body))
    wb.close()

    # DR/DQ/DL are NOT auto-filled - they land in _UnresolvedIndex (§13.4)
    for rn in (111, 113, 116, 117, 82, 83):
        check(f"unresolved manual member r{rn} -> <input required>",
              rs[rn].index == INPUT_REQUIRED and rs[rn].unresolved)
    check("_UnresolvedIndex reported (>=6 entries)", res.unresolved >= 6)

    # the AC suggested-type array formula is preserved (never clobbered with literals)
    of = openpyxl.load_workbook(res.output_path)["NET SAFETY 50"]
    ac_col = cr.idx("suggested_type")
    ac_array = any(str(getattr(v, "ref", v)).startswith(cr.letter("suggested_type"))
                   for v in getattr(of, "array_formulae", {}).values())
    ac_literals = sum(1 for rr in range(2, of.max_row + 1)
                      if (lambda c: c is not None and not (isinstance(c, str) and c.startswith("=")))
                      (of.cell(row=rr, column=ac_col).value))
    check("AC array formula preserved, no literals written", ac_array and ac_literals == 0,
          f"array={ac_array} literals={ac_literals}")

    # §12 idempotency / §11 halt survive a re-run: feeding the output back must report the SAME
    # unresolved count (the <input required> markers don't masquerade as human-filled indices) and
    # must NOT renumber the already-written DiagnosisBlocks.
    res2 = iolist_diag.populate({"io_list": {"path": res.output_path, "sheet": "NET SAFETY", "header_row": 1},
                                 "diag_bit_min": 0, "diag_bit_max": 62, "strike_handling": "error",
                                 "output_dir": os.path.join(tmp, "Output2")},
                                out_dir=os.path.join(tmp, "Output2"), quiet=True)
    check("re-run reports the same unresolved count (sentinel not mistaken for human value)",
          res2.unresolved == res.unresolved, f"{res.unresolved} then {res2.unresolved}")
    b2 = staging.load_diagnostic_blocks({"io_list": {"path": res2.output_path}})
    o2 = openpyxl.load_workbook(res2.output_path, data_only=True)["NET SAFETY 50"]
    ae57 = str(o2.cell(row=57, column=cr.idx("diag_cabinet")).value)
    check("re-run keeps the AE<->block join intact",
          ae57.lstrip("-").isdigit() and b2.get(int(ae57), {}).get("fld", "").endswith("+EmergencyPushButtons_1"))

    shutil.rmtree(tmp, ignore_errors=True)
    print("ALL CHECKS PASS" if failures == 0 else f"{failures} CHECK(S) FAILED")
    raise SystemExit(0 if failures == 0 else 1)


if __name__ == "__main__":
    main()
