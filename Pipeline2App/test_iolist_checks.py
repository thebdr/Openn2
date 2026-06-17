#!/usr/bin/env python3
"""Tests for Phase 0A (pipeline2/core/iolist_checks.py) + the validate() phase interface.

Builds a synthetic I/O List workbook exercising every ported VBA check (a passing and a
failing case each), then asserts the resulting LogEntry findings. Also checks the phase
gating: the default validate() runs Phase 0A before Phase 1 (5 banners), and the designer
selection {"0A","1","2"} skips Diagnosis.

Run: python test_iolist_checks.py
"""
import os
import sys
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from pipeline2.core import config, iolist_checks, validation

failures = 0


def check(name, ok, detail=""):
    global failures
    print(("PASS  " if ok else "FAIL  ") + name + (f"  [{detail}]" if detail else ""))
    if not ok:
        failures += 1


COLS = config.load_iolist_columns()      # the 26 canonical names, in order


def _write_headers(ws, extra=None):
    for i, name in enumerate(COLS, 1):
        ws.cell(row=1, column=i, value=name)
    for col_idx, value in (extra or {}).items():
        ws.cell(row=1, column=col_idx, value=value)


def _set(ws, ref_row, **byletter):
    for col, val in byletter.items():
        ws[f"{col}{ref_row}"] = val


def build_wb(path):
    wb = Workbook()
    # --- main safety sheet: clean headers + a known-extra + unexpected + duplicate column ---
    ws = wb.active
    ws.title = "NET SAFETY 50"                       # contains SAFETY -> processed
    _write_headers(ws, extra={27: "Skip Reason",     # known pipeline column -> NO warning
                              28: "Foo Bar",          # unknown -> [WR] unexpected
                              29: "Device"})          # duplicate of a standard name -> [ER] + [WR]
    # row 2: clean alarm -> NO findings
    _set(ws, 2, R="A", K="EMERGENCY PUSH BUTTON PRESSED", G="I5.0", U="TS_001")
    # row 3: empty permanent part
    _set(ws, 3, R="A", K="", U="TS_003")
    # row 4: unknown permanent part + missing TS ref
    _set(ws, 4, R="W", K="NONSENSE PART", U="")
    # row 5: bad address (3 parts: I.2.5)
    _set(ws, 5, R="A", K="SAFETY DOOR OPEN", G="I.2.5", U="TS_005")
    # row 6: G and F both filled
    _set(ws, 6, R="A", K="SAFETY DOOR OPEN", G="I20.0", F="X", U="TS_006")
    # row 7: bad 4-part address (X.5.0.0)
    _set(ws, 7, R="A", K="SAFETY DOOR OPEN", G="X.5.0.0", U="TS_007")
    # row 8: PA node, all good
    _set(ws, 8, R="PA", O="S1", P="+C1", Q="-N1", V="192.168.50.10", W="n50-pc1", U="TS_008")
    # row 9: PA node, DUPLICATE ip (not .1) -> dup IP
    _set(ws, 9, R="PA", O="S1", P="+C1", Q="-N2", V="192.168.50.10", W="n50-pc2", U="TS_009")
    # row 10: PA node, missing IP (no dot) + bad profinet name + missing TS
    _set(ws, 10, R="PA", O="S1", P="+C1", Q="-N3", V="", W="bad_name", U="")
    # row 11: node DUPLICATE FU+LOC+DEV of row 8 (=S1/+C1/-N1), different IP
    _set(ws, 11, R="PA", O="S1", P="+C1", Q="-N1", V="192.168.50.20", W="n50-pc4", U="TS_011")
    # row 12: #REF error in the IP cell
    _set(ws, 12, R="PA", O="S1", P="+C1", Q="-N5", V="#REF!", W="n50-pc5", U="TS_012")
    # rows 13/14: two .1 IPs (duplicate but EXEMPT) -> no dup IP error
    _set(ws, 13, R="PA", O="S1", P="+C1", Q="-N6", V="192.168.50.1", W="n50-pc6", U="TS_013")
    _set(ws, 14, R="PA", O="S1", P="+C1", Q="-N7", V="192.168.50.1", W="n50-pc7", U="TS_014")

    # --- a 2-dot sheet with a typo in the position-17 header (Device) ---
    ws2 = wb.create_sheet("R1.S1.C2")
    _write_headers(ws2)
    ws2.cell(row=1, column=17, value="Devvice")       # wrong name at column 17

    # --- a 2-dot sheet with >64 (and >126) nodes ---
    ws3 = wb.create_sheet("R9.S1.C1")
    _write_headers(ws3)
    for i in range(130):                                # distinct IPs + distinct devices -> node count only
        r = i + 2
        ws3[f"V{r}"] = f"192.168.51.{i + 10}"
        ws3[f"O{r}"], ws3[f"P{r}"], ws3[f"Q{r}"] = "S9", "+C1", f"-N{i}"

    # --- a non-processed sheet (no dots, no SAFETY) ---
    wb.create_sheet("COVER")
    wb.save(path)


def run_0a(path, lang="en"):
    log = []
    iolist_checks.check_iolist_standalone({"io_list": {"path": path}}, log, lang)
    return log


def at(log, location):
    return [e for e in log if e.location == location]


def msg_at(log, location, level):
    return [e for e in at(log, location) if e.level == level]


def main():
    tmp = tempfile.mkdtemp(prefix="iolchk_")
    path = os.path.join(tmp, "io.xlsx")
    build_wb(path)
    log = run_0a(path)

    # every finding is a real LogEntry (location + level), never a raw string
    check("findings are LogEntry objects with a level",
          log and all(isinstance(e, validation.LogEntry) and e.level for e in log))

    # --- sheet selection ---
    check("processed sheet logged", any(e.message and "NET SAFETY 50" in e.message
                                        and e.level == "INFO" for e in log))
    check("COVER skipped", any(e.location == "COVER" and e.level == "INFO" for e in log))
    check("no 'no sheet found'", not any("No sheet" in e.message for e in log))

    # --- column checks ---
    check("col 28 'Foo Bar' unexpected (WARN)", msg_at(log, "NET SAFETY 50!AB1", "WARN"))
    check("col 27 'Skip Reason' NOT warned (known pipeline column)",
          not msg_at(log, "NET SAFETY 50!AA1", "WARN"))
    check("col 29 'Device' duplicated (FAIL)", msg_at(log, "NET SAFETY 50!AC1", "FAIL"))
    check("position-17 typo 'Devvice' -> wrong column FAIL",
          any(e.location == "R1.S1.C2!Q1" and e.level == "FAIL" for e in log))
    check("correct headers do NOT raise wrong-column on the main sheet",
          not any(e.location.startswith("NET SAFETY 50!") and e.location.endswith("1")
                  and e.level == "FAIL" and "column 17" in e.message for e in log))

    # --- per-row content ---
    check("row3 empty permanent part (FAIL @K3)", msg_at(log, "NET SAFETY 50!K3", "FAIL"))
    check("row2 clean row -> no K finding", not at(log, "NET SAFETY 50!K2"))
    check("row4 unknown permanent part (FAIL @K4)",
          any("NONSENSE PART" in e.message for e in msg_at(log, "NET SAFETY 50!K4", "FAIL")))
    check("row4 missing TS ref (FAIL @U4)", msg_at(log, "NET SAFETY 50!U4", "FAIL"))
    check("row5 bad address I.2.5 (FAIL @G5)", msg_at(log, "NET SAFETY 50!G5", "FAIL"))
    check("row6 G+F both filled (FAIL @V6)", msg_at(log, "NET SAFETY 50!V6", "FAIL"))
    check("row7 bad 4-part address X.5.0.0 (FAIL @G7)", msg_at(log, "NET SAFETY 50!G7", "FAIL"))
    check("row8 clean PA node -> no IP finding", not msg_at(log, "NET SAFETY 50!V8", "FAIL"))
    check("row9 duplicate IP (FAIL @V9)",
          any("192.168.50.10" in e.message for e in msg_at(log, "NET SAFETY 50!V9", "FAIL")))
    check("row10 IP missing (FAIL @V10)", msg_at(log, "NET SAFETY 50!V10", "FAIL"))
    check("row10 bad profinet name (FAIL @W10)", msg_at(log, "NET SAFETY 50!W10", "FAIL"))
    check("row11 duplicate FU+LOC+DEV (FAIL @W11)", msg_at(log, "NET SAFETY 50!W11", "FAIL"))
    check("row12 #REF IP error (FAIL @V12)", msg_at(log, "NET SAFETY 50!V12", "FAIL"))
    check("rows13/14 .1 IPs are EXEMPT from dup-IP",
          not msg_at(log, "NET SAFETY 50!V13", "FAIL") and not msg_at(log, "NET SAFETY 50!V14", "FAIL"))

    # --- per-row info block populated (same format as other phases) ---
    e11 = msg_at(log, "NET SAFETY 50!W11", "FAIL")[0]    # row 11 has FU+LOC+DEV
    check("finding carries the per-entry info block",
          "|" in e11.format() and "S1" in e11.info())

    # --- node counts (sheet R9.S1.C1, 130 nodes) ---
    check("nodes > 64 -> WARN", any(e.location == "R9.S1.C1" and e.level == "WARN" for e in log))
    check("nodes > 126 -> FAIL", any(e.location == "R9.S1.C1" and e.level == "FAIL" for e in log))

    # --- Italian language switch ---
    log_it = run_0a(path, lang="it")
    e_it = msg_at(log_it, "NET SAFETY 50!K3", "FAIL")[0]
    check("IT message differs from EN (permanent part empty)",
          "Parte permanente" in e_it.message)

    # --- phase gating via validate() ---
    params = {"io_list": {"path": path}}              # no ce -> phases 1/2 log+skip
    full = validation.validate(params, [])
    banners = [e for e in full if e.level == "PHASE"]
    titles = [b.message for b in banners]
    check("default validate -> 5 phase banners", len(banners) == 5, str(len(banners)))
    check("Phase 0A banner is first", "0A" in titles[0])
    i0a = next(i for i, e in enumerate(full) if e.level == "PHASE" and "0A" in e.message)
    i1 = next(i for i, e in enumerate(full) if e.level == "PHASE" and "PHASE 1" in e.message)
    check("Phase 0A precedes Phase 1", i0a < i1)
    check("0B placeholder present", any("0B" in t for t in titles)
          and any("no checks" in e.message.lower() for e in full))

    designer = validation.validate(params, [], phases={"0A", "1", "2"})
    dtitles = [e.message for e in designer if e.level == "PHASE"]
    check("designer phases {0A,1,2} -> no Diagnosis banner",
          not any("Diagnosis" in t for t in dtitles) and any("0A" in t for t in dtitles))

    print("ALL CHECKS PASS" if failures == 0 else f"{failures} CHECK(S) FAILED")
    raise SystemExit(0 if failures == 0 else 1)


if __name__ == "__main__":
    main()
