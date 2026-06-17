#!/usr/bin/env python3
"""Data-independent tests for the iolist_diag populator (pure functions + a synthetic round-trip).

Covers the §6 regex ladder + old->new map, FLD/object index sharing (E1/2<->E2/2, KQ<->KI,
DI1/2<->DD), in_diag gating + P-from-top/non-P-from-bottom bit layout, the Field inference, skip
rules, the <input required> path + _UnresolvedIndex, DiagnosisBlocks (read back via staging), and
idempotency. Uses openpyxl + temp workbooks only - no live project data.

Run:  python test_iolist_diag.py
"""
import os
import sys
import shutil
import tempfile

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

import openpyxl

from pipeline2.core import config, staging
from pipeline2 import iolist_diag
from pipeline2.iolist_diag import script_type as st, families as fam_mod, index_assign as ix, diag_alloc as da
from pipeline2.iolist_diag.columns import ColumnResolver
from pipeline2.iolist_diag.models import IoRow, RowResult, SignalTypeDef, INPUT_REQUIRED

failures = 0


def check(name, ok, detail=""):
    global failures
    print(("PASS  " if ok else "FAIL  ") + name + (f"  [{detail}]" if detail else ""))
    if not ok:
        failures += 1


def _row(**kw):
    base = dict(sheet="S", row=10, functional_unit="S1", location="+N1", device="-D1",
                desc_l1="", desc_l1b="", id_node="", addr="", type_hw="", mnemonic="",
                skip_reason="0", struck=False, ex_script_type="", ex_suggested=None,
                ex_index="", ex_diag_cabinet="", ex_diag_bit="")
    base.update(kw)
    return IoRow(**base)


# --------------------------------------------------------------------------- #
# 1. script_type port (§6) + old->new canonical map (§7.1)                    #
# --------------------------------------------------------------------------- #
def test_script_type():
    cases = [
        # (desc_l1, desc_l1b, addr, type_hw, id_node, expected_AC, expected_AB)
        ("CIRCUIT BREAKER TRIPPED", "LINE", "I1.0", "A", "", "A", "A"),
        ("EMERGENCY PUSH BUTTON PRESSED", "CH1 + MS1", "I20.0", "A", "", "E1/2", "E1/2"),
        ("EMERGENCY PUSH BUTTON PRESSED", "CH2 + MS1", "I20.4", "A", "", "E2/2", "E2/2"),
        ("SAFETY DOOR OPEN", "CH1", "I120.0", "", "", "DI1/2", "DI1/2"),
        ("SAFETY DOOR OPEN", "", "I120.1", "", "", "DD", "DD"),
        ("", "PRIMARY DOOR OPEN RESET", "I120.5", "", "", "DR", "DR"),
        ("SAFETY ENCODER FAILURE", "PHOTOCELL 1", "I101.6", "", "", "ENC1/2", "N1/2"),
        ("SAFETY ENCODER FAILURE", "PHOTOCELL 2", "I101.7", "", "", "ENC2/2", "N2/2"),
        ("SWITCH DISCONNECTOR IS OPEN", "CH1", "I20.1", "A", "", "B1/2", "B1/2"),
        ("EMERGENCY RESET", "", "I20.2", "", "", "RES", "R*"),
        ("EMERGENCY RESET", "AREA 2", "I20.6", "", "", "RES", "R2"),
        ("FEEDBACK OF SAFETY RELAY FAULT", "", "I100.0", "A", "", "KI", "KI"),
        ("", "EMERGENCY AREA 1", "Q0.0", "", "", "FA1", "Z1"),         # area named in desc_l1b (col L)
        ("SAFETY RELAY - ENABLE", "DRIVE UNIT", "Q130.0", "", "", "KQ", "KQ"),
        ("SAFETY DOOR OPEN", "SAFETY RELAY", "Q130.2", "", "", "DQ", "DQ"),
        ("PROFINET SWITCH", "", "", "PW", "10", "PW", "PW"),           # Node -> type_hw
        ("FIRE ALARM", "CH1", "I20.3", "", "", "???", INPUT_REQUIRED),  # unmatched In -> ??? -> marker
    ]
    for d1, d2, addr, thw, node, exp_ac, exp_ab in cases:
        r = _row(desc_l1=d1, desc_l1b=d2, addr=addr, type_hw=thw, id_node=node)
        ac = st.suggested_type(r)
        ab = st.to_canonical(ac, r)
        check(f"script_type {d1[:18]!r}/{d2[:8]!r} AC", str(ac) == exp_ac, f"got {ac!r} want {exp_ac!r}")
        check(f"script_type {d1[:18]!r}/{d2[:8]!r} AB", ab == exp_ab, f"got {ab!r} want {exp_ab!r}")
    # the formula's dangling case: not In/Out/Node -> boolean False
    check("script_type unclassified -> False", st.suggested_type(_row(addr="10000")) is False)


# --------------------------------------------------------------------------- #
# synthetic workbook helpers                                                  #
# --------------------------------------------------------------------------- #
_SPECS = [
    # functional_unit, location, device, desc_l1, desc_l1b, bit(addr), type_hw, id_node, mnemonic
    dict(functional_unit="S1", location="+N1", device="-F1", desc_l1="CIRCUIT BREAKER TRIPPED", bit="I1.0", type_hw="A"),
    dict(functional_unit="S1", location="+N1", device="-K1", desc_l1="PROFINET SWITCH", type_hw="PW", id_node="10"),
    dict(functional_unit="S1", location="+N2", device="-S1", desc_l1="EMERGENCY PUSH BUTTON PRESSED", desc_l1b="CH1", bit="I2.0"),
    dict(functional_unit="S1", location="+N2", device="-S1", desc_l1="EMERGENCY PUSH BUTTON PRESSED", desc_l1b="CH2", bit="I2.1"),
    dict(functional_unit="S1", location="+SG1", device="-B1", desc_l1="SAFETY DOOR OPEN", desc_l1b="CH1", bit="I3.0"),
    dict(functional_unit="S1", location="+SG1", device="-B1", desc_l1="SAFETY DOOR OPEN", bit="I3.1"),
    dict(functional_unit="S1", location="+N3", device="-Q1", desc_l1="SAFETY RELAY - ENABLE", bit="Q1.0"),
    dict(functional_unit="S1", location="+N3", device="-Q1", desc_l1="FEEDBACK OF SAFETY RELAY FAULT", bit="I4.0"),
    dict(functional_unit="S1", location="+N4", device="-X1", desc_l1b="PRIMARY DOOR OPEN RESET", bit="I5.0"),
    dict(functional_unit="S1", location="+N5", device="-Z1", desc_l1="SPARE", bit="I6.0", mnemonic="-"),
    dict(functional_unit="S1", location="+SM1", device="-X9", desc_l1="FIELD BUS FAILURE", type_hw="PA", id_node="7"),
]


def _build_wb(path):
    cr = ColumnResolver("IoList")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    ws["A1"] = "header"
    cr.set(ws, 2, "id_node", "50")               # node-meta row -> dropped by the reader
    for i, spec in enumerate(_SPECS):
        r = 3 + i
        for canon, val in spec.items():
            cr.set(ws, r, canon, val)
    wb.save(path)
    wb.close()
    return cr


def _params(path, out):
    return {"io_list": {"path": path, "sheet": "NET SAFETY", "header_row": 1},
            "diag_bit_min": 0, "diag_bit_max": 62, "strike_handling": "error",
            "output_dir": out}


def _by_fld(results):
    return {(r.iorow.location, r.iorow.device, r.iorow.desc_l1b): r for r in results}


# --------------------------------------------------------------------------- #
# 2. full populate round-trip on the synthetic workbook                       #
# --------------------------------------------------------------------------- #
def test_populate():
    tmp = tempfile.mkdtemp(prefix="idiagtest_")
    src = os.path.join(tmp, "io.xlsx")
    cr = _build_wb(src)
    out = os.path.join(tmp, "Output")
    res = iolist_diag.populate(_params(src, out), out_dir=out, quiet=True)

    rs = {r.iorow.row: r for r in res.results}
    # object index sharing
    check("E1/2 <-> E2/2 share index", rs[5].index == rs[6].index and rs[5].index != "")
    check("DI1/2 <-> DD share index", rs[7].index == rs[8].index and rs[7].index != "")
    check("KQ <-> KI share index", rs[9].index == rs[10].index and rs[9].index != "")
    # in_diag gating + cabinet routing (via the regenerated DiagnosisBlocks, read like staging)
    blocks = staging.load_diagnostic_blocks({"io_list": {"path": res.output_path}})
    def fln(r):
        c = rs[r].diag_cabinet
        return blocks.get(int(c), {}).get("fld", "") if c.lstrip("-").isdigit() else ""
    check("A -> node cabinet, bit 00", fln(3) == "S1+N1" and rs[3].diag_bit == "00")
    check("PW -> same node, bit 62 (P from top)", fln(4) == "S1+N1" and rs[4].diag_bit == "62")
    check("E1/2 -> +EmergencyPushButtons, bit 00", fln(5).endswith("+EmergencyPushButtons_1") and rs[5].diag_bit == "00")
    check("E2/2 not diagnosed", rs[6].diag_cabinet == "")
    check("DI1/2 -> +SafetyDoors bit 01, DD bit 02",
          fln(7).endswith("+SafetyDoors_1") and rs[7].diag_bit == "01" and rs[8].diag_bit == "02")
    check("KQ -> node cabinet, KI not diagnosed", fln(9) == "S1+N3" and rs[10].diag_cabinet == "")
    check("lone PA -> +FieldIODevices bit 62", fln(13).endswith("+FieldIODevices_1") and rs[13].diag_bit == "62")
    # skip + unresolved
    check("mnemonic '-' row skipped", rs[12].skipped and rs[12].index == "")
    check("DR -> <input required> (unresolved)", rs[11].index == INPUT_REQUIRED and rs[11].unresolved)
    # written cells in the output workbook
    wb = openpyxl.load_workbook(res.output_path, data_only=True)
    ws = wb["NET SAFETY 50"]
    check("AB written for A row", cr.get(ws, 3, "script_type") == "A")
    check("AE/AF written for A row", str(cr.get(ws, 3, "diag_cabinet")) == "000" and str(cr.get(ws, 3, "diag_bit")) == "00")
    check("DiagnosisBlocks + _UnresolvedIndex sheets present",
          "DiagnosisBlocks" in wb.sheetnames and "_UnresolvedIndex" in wb.sheetnames)
    check("DiagnosisBlocks Type-1 before Type-2",
          [b for b in blocks.values()][0]["template_type"] in ("1", "01"))
    wb.close()

    # idempotency: re-running on the populated output changes nothing in AB/AD/AE/AF
    out2 = os.path.join(tmp, "Output2")
    p2 = _params(res.output_path, out2)
    res2 = iolist_diag.populate(p2, out_dir=out2, quiet=True)
    a = openpyxl.load_workbook(res.output_path, data_only=True)["NET SAFETY 50"]
    b = openpyxl.load_workbook(res2.output_path, data_only=True)["NET SAFETY 50"]
    same = all(cr.get(a, r, col) == cr.get(b, r, col)
               for r in range(3, 3 + len(_SPECS)) for col in ("script_type", "index", "diag_cabinet", "diag_bit"))
    check("idempotent re-run (AB/AD/AE/AF unchanged)", same)
    check("re-run reports the SAME unresolved count (DR sentinel not re-read as a human index)",
          res2.unresolved == res.unresolved and res.unresolved >= 1)
    shutil.rmtree(tmp, ignore_errors=True)


# --------------------------------------------------------------------------- #
# 4. review regressions: struck/strike_handling, unknown type, stray KI       #
# --------------------------------------------------------------------------- #
def test_review_regressions():
    from openpyxl.styles import Font
    from pipeline2.iolist_diag.report import needs_attention
    tmp = tempfile.mkdtemp(prefix="idiagreg_")
    cr = ColumnResolver("IoList")
    src = os.path.join(tmp, "io.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    ws["A1"] = "header"
    cr.set(ws, 2, "id_node", "50")
    for canon, val in dict(functional_unit="S1", location="+N1", device="-F1",
                           desc_l1="CIRCUIT BREAKER TRIPPED", bit="I1.0", type_hw="A").items():
        cr.set(ws, 3, canon, val)                       # row 3: a struck A row
    cr.cell(ws, 3, "desc_l1").font = Font(strike=True)
    for canon, val in dict(functional_unit="S1", location="+N2", device="-X1",
                           desc_l1="MYSTERY DEVICE", bit="I2.0", script_type="WAT").items():
        cr.set(ws, 4, canon, val)                       # row 4: a pre-filled non-catalogue type
    wb.save(src)
    wb.close()

    def run(strike):
        out = os.path.join(tmp, "out_" + strike)
        p = {"io_list": {"path": src, "sheet": "NET SAFETY", "header_row": 1},
             "diag_bit_min": 0, "diag_bit_max": 62, "strike_handling": strike, "output_dir": out}
        return iolist_diag.populate(p, out_dir=out, quiet=True)

    re = run("error")
    rs = {r.iorow.row: r for r in re.results}
    check("struck row KEPT + indexed when strike_handling != exclude (§5)",
          not rs[3].skipped and rs[3].index not in ("", INPUT_REQUIRED))
    check("unknown (non-catalogue) type is reported in _UnresolvedIndex",
          rs[4].unknown_type and needs_attention(rs[4]) == ("script_type", "unknown_type"))
    rx = run("exclude")
    rsx = {r.iorow.row: r for r in rx.results}
    check("struck row SKIPPED when strike_handling = exclude", rsx[3].skipped)

    # stray KIx/n with no matching primary -> <input required>, never inherits a previous group
    fams = fam_mod.load_families()
    types = config.load_signal_types()

    def kires(row, st, dev):
        io = _row(row=row, functional_unit="S1", location="+P", device=dev, addr=f"I{row}.0")
        r = RowResult(iorow=io)
        r.script_type = st
        rec = config.resolve_type(types, st)
        r.type_def = SignalTypeDef.from_record(rec) if rec else None
        r.family = fam_mod.family_for(st, fams)
        return r
    res = [kires(10, "KQ", "-Q1"), kires(11, "KI1/2", "-Q1"), kires(12, "KI2/2", "-Q1"), kires(20, "KI2/3", "-Q9")]
    ix.assign_indices(res, fams)
    check("KI1/2 inherits its KQ (same FLD)", res[1].index == res[0].index and res[0].index != "")
    check("stray KI2/3 (no primary) -> <input required>", res[3].index == INPUT_REQUIRED and res[3].unresolved)
    shutil.rmtree(tmp, ignore_errors=True)


# --------------------------------------------------------------------------- #
# 3. ColumnResolver + family registry                                         #
# --------------------------------------------------------------------------- #
def test_resolver_and_families():
    cr = ColumnResolver("IoList")
    check("ColumnResolver canonical->letter", cr.letter("script_type") == "AB" and cr.letter("index") == "AD")
    fams = fam_mod.load_families()
    check("object_families.csv loaded", len(fams) >= 7)
    check("family_for(E2/2)=emergency_pb", getattr(fam_mod.family_for("E2/2", fams), "family", None) == "emergency_pb")
    check("family_for(DR)=door", getattr(fam_mod.family_for("DR", fams), "family", None) == "door")
    check("family_for(A)=None (standalone)", fam_mod.family_for("A", fams) is None)


def main():
    test_script_type()
    test_resolver_and_families()
    test_populate()
    test_review_regressions()
    print("ALL CHECKS PASS" if failures == 0 else f"{failures} CHECK(S) FAILED")
    raise SystemExit(0 if failures == 0 else 1)


if __name__ == "__main__":
    main()
