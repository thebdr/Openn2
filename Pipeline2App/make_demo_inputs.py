#!/usr/bin/env python3
"""make_demo_inputs.py - synthesise demo I/O List + C&E workbooks that exercise EVERY
validation case, for regression tests and demonstrations.

The two generated files (under demo/) mirror the real input structure (same sheets, headers,
column positions) but their rows are crafted so the validation log shows one example of each
PASS / FAIL / WARN / SKIP across all three phases:

  Phase 1  C&E in IOList (forward)        - match, address-mismatch, device-not-found, empty key, AREA match/miss
  Phase 2  IOList in C&E (reverse)        - typed yes/warn present/absent, fld-only & addr-only partials,
                                            untyped+safety, untyped plain, ce_excluded skip, ce_always_excluded (ch2) skip
  Phase 3  Diagnosis Coherence Check      - unique alarm + warning, missing cabinet, missing bit, duplicate slot

Run `python make_demo_inputs.py` to (re)generate the files and print the validation result.
The companion `test_demo_validation.py` asserts every case fires.
"""
from __future__ import annotations
import os
import sys

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from pipeline2.core import config

DEMO_DIR = os.path.join(_HERE, "demo")
IO_PATH = os.path.join(DEMO_DIR, "8XXX_IOList_DEMO.xlsx")
CE_PATH = os.path.join(DEMO_DIR, "8XXX_CE_DEMO.xlsx")
IO_SHEET = "NET SAFETY 50"
CE_SHEET = "CAUSE&EFFECT MATRIX"


def _cols(document: str) -> dict:
    """canonical -> column letter for a document, from config_project/input_docs/column_map.csv."""
    return {m["canonical"]: m["column"] for m in config.load_column_map(document)}


# --- the crafted I/O List rows --------------------------------------------- #
# Each dict is one NET SAFETY 50 row (canonical column -> value). The trailing comment names
# the validation case(s) the row triggers. fu/loc/dev = the device key (FLD); bit = address.
FU = "=S1"
IO_ROWS = [
    # ---- Phase-1 forward PASS + Phase-2 reverse typed-PASS + Phase-3 diag PASS ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10001", bit="I0.0", normal_condition="1",
         script_type="E1/2", desc_l1="EMERGENCY STOP", desc_l1b="PUSHBUTTON", type_hw="DI",
         drawing="DWG-001", index="0001", diag_cabinet="001", diag_bit="00"),       # fwd PASS / rev PASS(yes) / diag 001.00 alarm

    # ---- Phase-2 reverse FAIL: ce_mandatory=yes, absent from C&E ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10002", bit="I0.1", normal_condition="1",
         script_type="B1/2", desc_l1="SAFETY GATE SWITCH", type_hw="DI",
         drawing="DWG-001", index="0002", diag_cabinet="001", diag_bit="01"),       # rev FAIL(yes absent) / diag 001.01 alarm

    # ---- Phase-2 reverse WARN: ce_mandatory=warn, absent from C&E ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10003", bit="I0.2",
         script_type="RES", desc_l1="EMERGENCY RESET", type_hw="DI",
         drawing="DWG-001", index="0003"),                                          # rev WARN(warn absent) (RES not in-diag)

    # ---- Phase-1 forward FAIL (address mismatch) + Phase-2 reverse FAIL (fld-only partial) ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10004", bit="I0.3", normal_condition="1",
         script_type="DI1/2", desc_l1="SAFETY DOOR", type_hw="DI",
         drawing="DWG-001", index="0004", diag_cabinet="001", diag_bit="03"),       # fwd addr-mismatch / rev FAIL fld-only / diag 001.03

    # ---- Phase-2 reverse FAIL (addr-only partial): this device's address sits under another device in C&E ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10005", bit="I0.4", normal_condition="1",
         script_type="B1/2", desc_l1="SAFETY GATE 2", type_hw="DI",
         drawing="DWG-001", index="0005", diag_cabinet="001", diag_bit="04"),       # rev FAIL addr-only / diag 001.04

    # ---- Phase-2 reverse FAIL: untyped + a safety word in the description ----
    dict(functional_unit=FU, location="+DM2.CC1", device="-X20001", bit="I1.0",
         desc_l1="EMERGENCY STOP", desc_l1b="FIELD BOX", drawing="DWG-002"),        # rev FAIL untyped+safety

    # ---- Phase-2 reverse WARN: untyped, no safety word ----
    dict(functional_unit=FU, location="+DM2.CC1", device="-X20002", bit="I1.1",
         desc_l1="CONVEYOR RUN FEEDBACK", drawing="DWG-002"),                       # rev WARN untyped plain

    # ---- Phase-2 reverse SKIP: untyped, matches ce_excluded_words ("power supply") ----
    dict(functional_unit=FU, location="+DM2.CC1", device="-T20003", bit="I1.2",
         desc_l1="POWER SUPPLY UNIT FAILURE", drawing="DWG-002"),                   # rev SKIP (ce_excluded_words)

    # ---- Phase-2 reverse SKIP: ce_always_excluded_words ("ch2") WINS over a mandatory type + word ----
    dict(functional_unit=FU, location="+DM2.CC1", device="-Q20004", bit="Q1.0",
         script_type="KQ2/2", desc_l1="CONTACTOR FEEDBACK", desc_l1b="CH2",
         drawing="DWG-002", index="0004"),                                          # rev SKIP (always-excluded ch2)

    # ---- Phase-1 forward PASS + Phase-2 reverse PASS: untyped device present in C&E ----
    dict(functional_unit=FU, location="+DM2.CC1", device="-X20010", bit="I2.0",
         desc_l1="FIELD DEVICE HEALTHY", drawing="DWG-002"),                        # fwd PASS / rev PASS untyped

    # ---- Phase-1 forward AREA PASS + Phase-2 reverse PASS (KQ via AREA sheet) ----
    dict(functional_unit=FU, location="+DM3.CC1", device="-Q30001", bit="Q5.0", normal_condition="1",
         script_type="KQ", desc_l1="SAFETY ENABLE", type_hw="DQ",
         drawing="DWG-003", index="0001", diag_cabinet="001", diag_bit="30"),       # fwd AREA PASS / rev PASS(yes via AREA) / diag 001.30

    # ---- Phase-3 diagnosis PASS: a warning-family slot (type ends 'W') - unique ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-H10020", bit="I3.0",
         script_type="W", desc_l1="OVER-TEMPERATURE WARNING", type_hw="W",
         drawing="DWG-001", index="0020", diag_cabinet="001", diag_bit="00"),       # diag 001.00 WARNING (unique)

    # ---- Phase-3 diagnosis FAIL: in-diagnosis but missing Diag Cabinet ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10010", bit="I0.6",
         script_type="A", desc_l1="GENERAL ALARM", type_hw="A",
         drawing="DWG-001", index="0010", diag_cabinet="", diag_bit="10"),          # diag FAIL missing cabinet

    # ---- Phase-3 diagnosis FAIL: in-diagnosis but missing Diag Bit ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10011", bit="I0.7",
         script_type="A", desc_l1="GENERAL ALARM 2", type_hw="A",
         drawing="DWG-001", index="0011", diag_cabinet="001", diag_bit=""),         # diag FAIL missing bit

    # ---- Phase-3 diagnosis FAIL: duplicate cabinet/bit/family (both flagged) ----
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10012", bit="I1.5",
         script_type="A", desc_l1="DUPLICATE SLOT ALARM A", type_hw="A",
         drawing="DWG-001", index="0012", diag_cabinet="001", diag_bit="20"),       # diag FAIL duplicate 001.20
    dict(functional_unit=FU, location="+DM1.CC1", device="-S10013", bit="I1.6",
         script_type="A", desc_l1="DUPLICATE SLOT ALARM B", type_hw="A",
         drawing="DWG-001", index="0013", diag_cabinet="001", diag_bit="20"),       # diag FAIL duplicate 001.20
]

# DiagnosticBlocks rows: (ID_SWP, FullName, TemplateType)
DIAG_BLOCKS = [(1, "=S1+DM1.CC1", 1), (3, "=S1+DM3.CC1", 1)]

# C&E matrix rows (functional_unit, location, device, address). Blank device => empty-key case.
CE_MATRIX = [
    (FU, "+DM1.CC1", "-S10001", "I0.0"),    # matches row 1            -> fwd PASS
    (FU, "+DM1.CC1", "-S10004", "I9.9"),    # row-4 device, diff addr  -> fwd address mismatch
    (FU, "+DM1.CC1", "-S19999", "I0.4"),    # device not in I/O list   -> fwd not-found (and gives addr I0.4 for the addr-only reverse case)
    (FU, "+DM2.CC1", "-X20010", "I2.0"),    # matches the untyped row  -> fwd PASS / rev untyped PASS
    ("",  "",        "",        "I8.8"),    # address but no device    -> fwd empty device key
]
# AREA 1 rows (concat_id, address)
CE_AREA = [
    ("=S1+DM3.CC1-Q30001", "Q5.0"),         # matches the KQ row       -> fwd AREA PASS / rev KQ PASS
    ("=S1+DM3.CC1-Q39999", "Q5.9"),         # not in I/O list          -> fwd AREA not-found
]


def build(io_path: str = IO_PATH, ce_path: str = CE_PATH) -> tuple[str, str]:
    """Write the two demo workbooks; returns their paths."""
    os.makedirs(os.path.dirname(io_path), exist_ok=True)

    # ---- I/O List ----
    iocols = _cols("IoList")
    wb = Workbook()
    ws = wb.active
    ws.title = IO_SHEET
    for m in config.load_column_map("IoList"):                 # header row 1 = the expected headers
        ws[f"{m['column']}1"] = m["expected_header"]
    for i, row in enumerate(IO_ROWS, start=2):
        _text(ws, f"{iocols['part_no']}{i}", "DEMO")            # required, non-empty so the row isn't blank
        for canon, val in row.items():
            _text(ws, f"{iocols[canon]}{i}", val)              # _text keeps '=S1...' a string, not a formula
    db = wb.create_sheet("DiagnosticBlocks")
    for c, h in enumerate(["ID_Local", "ID_SWP", "Functional Unit", "Location",
                           "FullName", "TemplateType", "Notes"], start=1):
        db.cell(row=1, column=c, value=h)
    for r, (swp, full, tt) in enumerate(DIAG_BLOCKS, start=2):
        db[f"A{r}"], db[f"B{r}"], db[f"F{r}"] = swp, swp, tt
        _text(db, f"E{r}", full)                               # FullName '=S1+...' as string
    wb.save(io_path)

    # ---- C&E ----
    cecols, acols = _cols("CE"), _cols("AREA")
    wb2 = Workbook()
    m = wb2.active
    m.title = CE_SHEET
    m[f"{cecols['functional_unit']}2"] = "FUNCTIONAL UNIT"
    m[f"{cecols['location']}2"] = "LOCATION"
    m[f"{cecols['device']}2"] = "DEVICE"
    m[f"{cecols['address']}2"] = "BIT (ADDRESS)"
    for r, (fu, loc, dev, addr) in enumerate(CE_MATRIX, start=5):   # data from matrix_data_row 5
        _text(m, f"{cecols['functional_unit']}{r}", fu)
        _text(m, f"{cecols['location']}{r}", loc)
        _text(m, f"{cecols['device']}{r}", dev)
        m[f"{cecols['address']}{r}"] = addr
    a = wb2.create_sheet("AREA 1")
    a[f"{acols['concat_id']}3"] = "SIGLA CONTATTORE"
    a[f"{acols['address']}3"] = "DIGITAL OUTPUT"
    for r, (cid, addr) in enumerate(CE_AREA, start=4):             # data from area_data_row 4
        _text(a, f"{acols['concat_id']}{r}", cid)
        a[f"{acols['address']}{r}"] = addr
    wb2.save(ce_path)
    return io_path, ce_path


def _text(ws, cell, value):
    """Write a literal string ('=...' would otherwise be read as a formula)."""
    ws[cell] = value
    if isinstance(value, str) and value.startswith("="):
        ws[cell].data_type = "s"


def demo_params(io_path: str = IO_PATH, ce_path: str = CE_PATH, out_dir: str | None = None) -> dict:
    """A params dict pointing at the demo files, with the ce_* knobs set so every case shows
    (ce_full_check off -> ce_excluded_words active; ce_full_print on -> passes/skips logged)."""
    return {
        "io_list": {"path": io_path, "sheet": IO_SHEET, "header_row": 1},
        "ce": {"path": ce_path, "matrix_sheet": CE_SHEET, "matrix_header_row": 2,
               "matrix_data_row": 5, "area_header_row": 3, "area_data_row": 4},
        "output_dir": out_dir or os.path.join(DEMO_DIR, "Output"),
        "ce_fuzzy_chars": 2,
        "ce_mandatory_words": ["emergency", "safety", "relay", "contactor", "enable"],
        "ce_excluded_words": ["circuit breaker", "power supply", "profinet switch", "coupler", "door"],
        "ce_always_excluded_words": ["ch2"],
        "ce_full_check": False,
        "ce_full_print": True,
        "diag_bit_min": 0, "diag_bit_max": 62,
    }


def main():
    from pipeline2.core import staging, validation
    io_path, ce_path = build()
    params = demo_params()
    rows, warns = staging.load_io_list(params, config.load_signal_types())
    log = validation.validate(params, rows)
    out_dir = params["output_dir"]
    passed, failed, warned = validation.write_log(log, out_dir)
    skipped = sum(1 for e in log if e.level == "SKIP")
    print(f"demo I/O List : {io_path}")
    print(f"demo C&E      : {ce_path}")
    print(f"validation    : {passed} passed, {failed} failed, {warned} warning(s), {skipped} skipped")
    print(f"  -> {os.path.join(out_dir, 'Reports', 'documents_validation_report.txt')}  +  .html")
    print("\nFAIL / WARN / SKIP cases produced:")
    for e in log:
        if e.level in ("FAIL", "WARN", "SKIP"):
            print(f"  [{e.level:4}] {e.location}  {e.message}")


if __name__ == "__main__":
    main()
