"""ColumnResolver - resolve I/O-list columns by canonical name, never a hardcoded letter.

Built from config.load_column_map(document) (column_map.csv), it maps canonical <-> column
letter <-> 1-based index and reads/writes worksheet cells by canonical name. The whole stage
goes through this so a column move in column_map.csv needs no code change (spec §3).
"""
from __future__ import annotations
from openpyxl.utils import column_index_from_string

from pipeline2.core import config


class ColumnResolver:
    def __init__(self, document: str = "IoList"):
        self._letter: dict[str, str] = {}
        self._required: dict[str, bool] = {}
        self._expected: dict[str, str] = {}
        for m in config.load_column_map(document):
            self._letter[m["canonical"]] = m["column"]
            self._required[m["canonical"]] = m["required"]
            self._expected[m["canonical"]] = m["expected_header"]

    def letter(self, canonical: str) -> str:
        return self._letter[canonical]

    def idx(self, canonical: str) -> int:
        return column_index_from_string(self._letter[canonical])

    def get(self, ws, row: int, canonical: str):
        return ws.cell(row=row, column=self.idx(canonical)).value

    def set(self, ws, row: int, canonical: str, value) -> None:
        ws.cell(row=row, column=self.idx(canonical)).value = value

    def cell(self, ws, row: int, canonical: str):
        return ws.cell(row=row, column=self.idx(canonical))

    def expected_header(self, canonical: str) -> str:
        return self._expected.get(canonical, "")

    def __contains__(self, canonical: str) -> bool:
        return canonical in self._letter
