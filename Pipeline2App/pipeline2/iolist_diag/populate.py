"""The iolist_diag stage entry: populate a fresh I/O List, before staging (spec §0, §11).

`populate(params)` copies the input workbook (never mutating it in place - the copy preserves the
COVER sheet, the suggested-type array formula, the `_TRO_TRAILING` LAMBDA and all formatting), then
fills script_type (AB) + suggested_type (AC) + index (AD) + diag_cabinet (AE) + diag_bit (AF),
(re)generates the DiagnosisBlocks sheet and writes the `_UnresolvedIndex` report. It is idempotent
and non-destructive: every column is written ONLY into an empty cell, and each pre-filled cell is
preserved and audit-logged. Returns a PopulateResult; the caller (run.py / run_populate.py) halts
the pipeline when `unresolved > 0` (§11).

Two loads of the copy are used: data_only=True to read the cached cell VALUES (Skip Reason and the
suggested-type column are formulas), and the default formula-preserving load to write + save.
"""
from __future__ import annotations
from dataclasses import dataclass, field
import os
import shutil

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from pipeline2.core import config, i18n
from pipeline2.iolist_diag import blocks as blocks_mod
from pipeline2.iolist_diag import diag_alloc, families, index_assign, reader, report, script_type
from pipeline2.iolist_diag.columns import ColumnResolver
from pipeline2.iolist_diag.models import DASH, INPUT_REQUIRED, RowResult, SignalTypeDef


@dataclass
class PopulateResult:
    output_path: str
    sheets: list
    processed: int = 0
    indexed: int = 0
    diagnosed: int = 0
    unresolved: int = 0
    unknown: int = 0
    skipped: int = 0
    blocks: int = 0
    reported: int = 0          # rows in _UnresolvedIndex (unresolved + unmatched + unknown)
    audit: list = field(default_factory=list)
    results: list = field(default_factory=list)

    @property
    def counts(self) -> dict:
        return {"processed": self.processed, "indexed": self.indexed, "diag": self.diagnosed,
                "unresolved": self.unresolved, "unknown": self.unknown, "skipped": self.skipped}


def _skip_reason(io, type_def, params) -> str:
    strike_exclude = str(params.get("strike_handling", "exclude")).lower() == "exclude"
    sk = io.skip_reason.strip()
    if io.mnemonic.strip() == DASH:
        return "mnemonic"
    if io.struck and strike_exclude:
        return "struck"
    # Skip Reason (AA) mirrors the strike + mnemonic-"-" causes (both already handled above). Honor it
    # as a forward-compatible catch-all ONLY for a NON-struck row, so a struck row is excluded purely by
    # the strike rule (which respects strike_handling) and is NOT silently re-skipped via AA when
    # strike_handling != exclude (§5: struck rows are kept unless strike_handling = exclude).
    if sk and sk not in ("0", "0.0") and not io.struck:
        return "skip_reason"
    if type_def is not None and not io.description and not io.has_device:
        return "spare"            # a typed placeholder/SPARE row (skip + log, never unresolved)
    return ""


def _classify(io, types: dict, fams: list, params: dict) -> RowResult:
    """Compute script_type (Mode 1 compute / Mode 2 preserve+audit), resolve the type + family,
    and classify the §5 skip. Index + diag are assigned later over the non-skipped rows."""
    r = RowResult(iorow=io)
    sug = script_type.suggested_type(io)
    r.suggested_type = sug
    computed = script_type.to_canonical(sug, io)
    r.computed_script_type = computed
    ex = io.ex_script_type.strip()
    if ex:                                              # Mode 2: keep the human value, audit it
        r.script_type = ex
        r.prefilled = True
        r.audit_mismatch = bool(computed) and computed not in (ex, INPUT_REQUIRED)
    else:                                               # Mode 1: write the computed value
        r.script_type = computed
    rec = config.resolve_type(types, r.script_type) if r.script_type not in ("", INPUT_REQUIRED) else None
    r.type_def = SignalTypeDef.from_record(rec) if rec else None
    r.family = families.family_for(r.script_type, fams)
    r.unknown_type = bool(r.script_type) and r.script_type != INPUT_REQUIRED and r.type_def is None and r.family is None
    r.skip_reason = _skip_reason(io, r.type_def, params)
    r.skipped = bool(r.skip_reason)
    return r


def _ac_locked(wb, sheets: list, cr: ColumnResolver) -> bool:
    """True when the suggested_type (AC) column is governed by a (dynamic-)array formula on any I/O
    sheet. Then the spill owns the column: openpyxl reads its member cells as empty (data_only None),
    so writing into them would drop literal constants inside the array range -> Excel 'repaired
    records' / #SPILL!. We leave the live formula untouched (it already computes AC)."""
    if "suggested_type" not in cr:
        return False
    target = cr.idx("suggested_type")
    for sheet in sheets:
        for ref in getattr(wb[sheet], "array_formulae", {}).values():
            ref_str = getattr(ref, "ref", ref)              # str or ArrayFormula
            try:
                min_col, _, max_col, _ = range_boundaries(str(ref_str))
            except (TypeError, ValueError):
                continue
            if min_col is not None and min_col <= target <= max_col:
                return True
    return False


def _write_row(ws, cr: ColumnResolver, r: RowResult, ac_locked: bool) -> None:
    """Write AB/AC/AD/AE/AF into empty cells only (non-destructive, §12)."""
    io = r.iorow
    if not io.ex_script_type.strip() and r.script_type != "":
        cr.set(ws, io.row, "script_type", r.script_type)
    if not ac_locked and io.ex_suggested in (None, "") and r.suggested_type != "" and "suggested_type" in cr:
        cr.set(ws, io.row, "suggested_type", r.suggested_type)      # may be the formula's bool False
    if "index" in cr and not io.ex_index.strip() and r.index != "":
        cr.set(ws, io.row, "index", r.index)
    if "diag_cabinet" in cr and not io.ex_diag_cabinet.strip() and r.diag_cabinet:
        cr.set(ws, io.row, "diag_cabinet", r.diag_cabinet)
    if "diag_bit" in cr and not io.ex_diag_bit.strip() and r.diag_bit:
        cr.set(ws, io.row, "diag_bit", r.diag_bit)


def populate(params: dict, sheets_override=None, lang: str | None = None, out_dir: str | None = None,
             quiet: bool = False) -> PopulateResult:
    lang = lang or params.get("language") or i18n.DEFAULT_LANG
    src = params.get("io_list", {}).get("path")
    if not src or not os.path.isfile(src):
        raise SystemExit(f"ERROR iolist_diag: I/O List not found: {src!r}")

    out_root = out_dir or config.output_root(params)
    dest_dir = config.out_path(out_root, "populated_iolist")
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, os.path.basename(src))
    shutil.copy2(src, dest)                              # never mutate the input in place (§11)

    cr = ColumnResolver("IoList")
    wb_vals = load_workbook(dest, data_only=True)        # cached values (AA / AC are formulas)
    rows, sheets = reader.read_iolist(wb_vals, params, cr, sheets_override)
    wb_vals.close()

    types = config.load_signal_types()
    fams = families.load_families()
    results = [_classify(io, types, fams, params) for io in rows]
    index_assign.assign_indices(results, fams)
    diagblocks = diag_alloc.allocate(results, params)

    wb = load_workbook(dest)                             # formula-preserving (default) load to edit + save
    ac_locked = _ac_locked(wb, sheets, cr)
    for r in results:
        _write_row(wb[r.iorow.sheet], cr, r, ac_locked)
    # (Re)generate DiagnosisBlocks unless the input is already fully diagnosed (an existing sheet AND no
    # fresh AE was written): regenerating with fresh sequential IDs would mismatch the preserved human
    # cabinet numbers and break the AE<->block join. A fresh/empty list always regenerates.
    fresh_diag = any(r.diag_cabinet and not r.iorow.ex_diag_cabinet.strip() for r in results)
    has_blocks = any(s.strip().lower() in ("diagnosisblocks", "diagnosticblocks") for s in wb.sheetnames)
    blocks_written = fresh_diag or not has_blocks
    if blocks_written:
        blocks_mod.write_diagnosis_blocks(wb, diagblocks, results)
    n_report = report.write_unresolved_sheet(wb, results, cr, lang)
    wb.save(dest)

    # The §11 halt fires on genuine blockers only - an `<input required>` script_type or index.
    # Unknown-but-typed rows are reported (clickable) but do not by themselves stop the run.
    halt = sum(1 for r in results
               if not r.skipped and (r.script_type == INPUT_REQUIRED or r.unresolved))
    res = PopulateResult(
        output_path=dest, sheets=sheets, processed=len(results),
        indexed=sum(1 for r in results if r.index and r.index != INPUT_REQUIRED),
        diagnosed=sum(1 for r in results if r.diag_cabinet),
        unresolved=halt,
        unknown=sum(1 for r in results if r.unknown_type and not r.skipped),
        skipped=sum(1 for r in results if r.skipped),
        blocks=len(diagblocks) if blocks_written else 0,
        reported=n_report,
        audit=report.audit_lines(results, lang), results=results,
    )
    if not quiet:
        print(f"[{i18n.tr('idiag_phase', lang)}]")
        print("  " + report.summary_line(res.counts, lang))
        print("  " + i18n.tr("idiag_wrote", lang, path=dest))
        if n_report:
            print("  " + i18n.tr("idiag_unresolved_warn", lang, n=n_report, sheet="_UnresolvedIndex"))
        else:
            print("  " + i18n.tr("idiag_all_resolved", lang))
    return res
