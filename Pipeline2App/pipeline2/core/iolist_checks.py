"""Phase 0A - I/O List standalone validation.

A faithful Python port of the electrical designers' Excel/VBA tool
(`IOList check_v18.xlsm`, module `CheckIOListModule.CheckIOList`). It re-opens the I/O List
workbook and checks each safety sheet on its own - column headers, addresses, Profinet
nodes/IPs, the safety "permanent part" against the ported glossary, T.S. refs - WITHOUT any
Cause&Effect cross-reference (that is phases 1-2). Unlike `staging.load_io_list`, it works on
the RAW rows (no strike/skip exclusion) so the designer sees everything before sharing.

Each VBA `trace` becomes a first-class `LogEntry` (so it renders in the same
`[LEVEL] location  info  message` format as every other phase): `[ER]`->FAIL, `[WR]`->WARN,
`[I]`->INFO. Messages come from `i18n` (EN/IT). The future `matrix_checks.py` (phase 0B) is
this module's sibling.

Faithful-port notes (vs the VBA):
- sheet is processed iff its name has exactly 2 dots OR contains "SAFETY" (case-sensitive);
- `_clean` = strip spaces + CR/LF/FF, case-PRESERVING (ports `cleanString`);
- the positional column-name check is exact-after-`_clean` against the 26-name glossary;
- columns 27+ that match our pipeline's own extra columns (Skip Reason, Script Type, ...) do
  NOT raise the "unexpected column" warning (the only deliberate adaptation);
- duplicate-IP / duplicate-FU+LOC+DEV are evaluated only for rows that carry a Profinet IP,
  the node list is global across sheets, and IPs ending in ".1" are exempt from the dup check;
- node-count thresholds (>64 WARN, >126 ERROR) are per sheet.
"""
from __future__ import annotations
import os

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from . import config, i18n
from .validation import LogEntry

# Fixed I/O List columns (match config_project/input_docs/column_map.csv positions; the VBA hard-codes the same).
_FU, _LOC, _DEV, _TYPE = "O", "P", "Q", "R"
_ID, _BIT, _PP, _TSREF, _IP, _PNAME = "F", "G", "K", "U", "V", "W"
_DESC_L1B, _DRAWING, _SCRIPT, _INDEX = "L", "S", "AB", "AD"

_STD_COLS = 26                      # the glossary defines 26 positional columns
_ERROR_LITERALS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!", "#SPILL!"}
_FORBIDDEN_PNAME = ("_", ".", "/", "=", "*")


def _clean(s: str) -> str:
    """Ports VBA cleanString: remove spaces + CR(13)/FF(12)/LF(10), keep case."""
    return s.replace(" ", "").replace("\r", "").replace("\f", "").replace("\n", "")


def _s(value) -> str:
    """Cell value as the VBA would see it in a string context. Integral floats lose the
    trailing '.0' (so a numeric bit doesn't look like a dotted address)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _is_error(value) -> bool:
    return isinstance(value, str) and (value in _ERROR_LITERALS or
                                       (value.startswith("#") and value.endswith(("!", "?"))))


def _loc(sheet: str, col: str, row: int) -> str:
    return f"{sheet}!{col}{row}"


def _known_extra_headers() -> set:
    """Cleaned headers of our pipeline's own columns beyond the standard 26 (Skip Reason,
    Script Type, Suggested Type, Index, Diag Cabinet, Diag Bit, Hardware Parameters). These
    are expected in our template (empty for designers) and must NOT warn as 'unexpected'."""
    out = set()
    for m in config.load_column_map("IoList"):
        if column_index_from_string(m["column"]) > _STD_COLS and m.get("expected_header"):
            out.add(_clean(m["expected_header"]))
    return out


def check_iolist_standalone(params: dict, log: list, lang: str = "en") -> None:
    """Append Phase 0A findings (LogEntry) for the I/O List in `params`. Missing/unreadable
    workbook -> a single FAIL entry and return (the caller's banner already framed the phase)."""
    io = params.get("io_list") or {}
    path = io.get("path") or ""
    if not path or not os.path.exists(path):
        log.append(LogEntry("FAIL", "I/O List", "", "", i18n.tr("iolist_missing", lang, path=path or "(empty)"),
                            path=path))
        return
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001 - report, never raise (designer sees the problem)
        log.append(LogEntry("FAIL", "I/O List", "", "", i18n.tr("iolist_open_error", lang, error=e), path=path))
        return

    cols = config.load_iolist_columns()
    parts_clean = {_clean(p) for p in config.load_permanent_parts()}
    known_extra = _known_extra_headers()
    nodes: list[dict] = []             # global across sheets, like the VBA NodeList
    processed = 0

    for ws in wb.worksheets:
        name = ws.title
        # process only sheets named like "R1.S1.C1" (exactly 2 dots) or containing "SAFETY"
        if not (name.count(".") == 2 or "SAFETY" in name):
            log.append(LogEntry("INFO", name, "", "", i18n.tr("skip_sheet", lang, sheet=name), path=path))
            continue
        processed += 1
        log.append(LogEntry("INFO", name, "", "", i18n.tr("proc_sheet", lang, sheet=name), path=path))
        _check_columns(ws, name, cols, known_extra, path, log, lang)
        _check_rows(ws, name, parts_clean, nodes, path, log, lang)

    if processed == 0:
        log.append(LogEntry("FAIL", "I/O List", "", "", i18n.tr("no_sheet", lang), path=path))


def _check_columns(ws, sheet, cols, known_extra, path, log, lang) -> None:
    # positional check: header at column ll (1..26) must equal the glossary name (after _clean)
    for ll in range(1, _STD_COLS + 1):
        actual = _s(ws.cell(row=1, column=ll).value)
        expected = cols[ll - 1]
        if _clean(actual) != _clean(expected):
            letter = get_column_letter(ll)
            log.append(LogEntry("FAIL", _loc(sheet, letter, 1), "", "",
                                i18n.tr("col_wrong", lang, actual=actual, n=ll, expected=expected), path=path))
    # columns beyond the standard set: unexpected (unless a known pipeline column), and a
    # second copy of a standard name is a duplicate (exact match, as the VBA does)
    for zz in range(_STD_COLS + 1, _STD_COLS + 101):
        header = _s(ws.cell(row=1, column=zz).value)
        if not header:
            continue
        letter = get_column_letter(zz)
        if _clean(header) not in known_extra:
            log.append(LogEntry("WARN", _loc(sheet, letter, 1), "", "",
                                i18n.tr("col_unexpected", lang, n=zz, name=header), path=path))
        if any(header == c for c in cols):
            log.append(LogEntry("FAIL", _loc(sheet, letter, 1), "", "",
                                i18n.tr("col_duplicated", lang, n=zz, name=header), path=path))


def _check_rows(ws, sheet, parts_clean, nodes, path, log, lang) -> None:
    number_of_nodes = 0                # per-sheet count of IP-bearing rows (VBA NumberOfNodes)
    for qq in range(2, (ws.max_row or 1) + 1):
        def cell(col):
            return ws[f"{col}{qq}"].value

        fu, location, dev = _s(cell(_FU)), _s(cell(_LOC)), _s(cell(_DEV))
        type_row = _s(cell(_TYPE)).strip()
        g, f = _s(cell(_BIT)), _s(cell(_ID))
        ip_raw = cell(_IP)
        # per-row info block, so 0A entries render like every other phase
        row = {"functional_unit": fu, "location": location, "device": dev,
               "bit": g, "desc_l1": _s(cell(_PP)), "desc_l1b": _s(cell(_DESC_L1B)),
               "drawing": _s(cell(_DRAWING)), "script_type": _s(cell(_SCRIPT)), "index": _s(cell(_INDEX))}
        fld = "".join((fu, location, dev))

        def add(level, col, msgkey, **fmt):
            log.append(LogEntry(level, _loc(sheet, col, qq), fld, "", i18n.tr(msgkey, lang, **fmt),
                                path=path, row=row))

        # --- IP / node dedup (only rows carrying a Profinet IP) ---
        if _is_error(ip_raw):
            add("FAIL", _IP, "ip_error")
        else:
            ip = _s(ip_raw)
            if ip:
                number_of_nodes += 1
                for n in nodes:
                    if n["ip"] == ip and not ip.endswith(".1"):
                        add("FAIL", _IP, "ip_duplicated", ip=ip, first=n["label"])
                    if n["fu"] == fu and n["loc"] == location and n["dev"] == dev:
                        add("FAIL", _PNAME, "dup_fld", first=n["label"])
                if not any(n["ip"] == ip for n in nodes):
                    nodes.append({"ip": ip, "fu": fu, "loc": location, "dev": dev,
                                  "label": _loc(sheet, _BIT, qq)})

        # --- column G (address) implies column F (ID) empty ---
        if g and f:
            add("FAIL", _IP, "fg_exclusion")

        # --- address format (only dotted, non-Rockwell addresses) ---
        if "." in g and ":" not in g:
            arr = g.split(".")
            ub = len(arr) - 1
            if ub == 3 and arr[0] not in ("I", "Q", "O"):
                add("FAIL", _BIT, "address_wrong")
            elif ub == 1 and ("I" in arr[0] and "Q" in arr[0] and "O" in arr[0]):
                add("FAIL", _BIT, "address_wrong")
            elif ub not in (1, 3):
                add("FAIL", _BIT, "address_wrong")

        # --- permanent part (alarm/warning rows) ---
        if type_row in ("A", "W"):
            pp = _s(cell(_PP))
            if pp == "":
                add("FAIL", _PP, "pp_empty")
            elif _clean(pp) not in parts_clean:
                add("FAIL", _PP, "pp_unknown", value=pp)

        # --- T.S. ref (alarm/warning/profinet rows) ---
        if type_row in ("A", "W", "PA", "PW"):
            if "TS_" not in _s(cell(_TSREF)):
                add("FAIL", _TSREF, "ts_missing")

        # --- Profinet node identity (PA/PW) ---
        if type_row in ("PA", "PW"):
            if not _is_error(ip_raw) and "." not in _s(cell(_IP)):
                add("FAIL", _IP, "ip_missing")
            name = _s(cell(_PNAME))
            cleaned = name
            for ch in _FORBIDDEN_PNAME:
                cleaned = cleaned.replace(ch, "")
            if cleaned != name or not cleaned.startswith("n"):
                add("FAIL", _PNAME, "pname_wrong")

    # --- node count for this sheet (VBA: >126 ERROR, >64 WARN) ---
    if number_of_nodes > 126:
        log.append(LogEntry("FAIL", sheet, "", "", i18n.tr("too_many_nodes", lang, n=number_of_nodes), path=path))
    if number_of_nodes > 64:
        log.append(LogEntry("WARN", sheet, "", "", i18n.tr("nodes_over_64", lang, n=number_of_nodes), path=path))
