"""Staging: read the I/O List into canonical rows.

Reads the source workbook with openpyxl (NOT read_only, so cell.font.strike is
available). Columns are read by POSITION (column letter) per the column_map -
the document has duplicate headers ("Description language 1" twice), so a
header-name map is unsafe; instead the header at each position is verified and a
mismatch is reported as a warning (catches reordered/inserted columns).

A row is excluded when any of its cells is struck through (predisposition, when
strike_handling = "exclude") or it carries a non-empty Skip Reason. Each kept
row gets its resolved signal-type record (from Script Type) and its source row.
"""
from __future__ import annotations
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from . import config
from . import matrix
from . import outputs  # leaf module; provides the DB member name for name_in_db


class StagedRow(dict):
    """Canonical row; attribute-ish access via keys, plus _source_row / _type."""


def _is_struck(ws, row_index: int, last_col: int) -> bool:
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row_index, column=c)
        if cell.value not in (None, "") and cell.font and cell.font.strike:
            return True
    return False


def _norm(value):
    if value is None:
        return ""
    s = str(value).strip()
    return s


def _skip_reason_present(value) -> bool:
    """Skip Reason: empty or 0 means 'not skipped'."""
    s = _norm(value)
    return s not in ("", "0", "0.0")


def _hkey(s):
    # collapse internal whitespace: source headers wrap with embedded newlines,
    # e.g. "Functional \nunit"
    return " ".join(_norm(s).split()).lower()


def _load_sheet(ws, sheet, header_row, colmap, strike_exclude, signal_types):
    """Read one sheet -> (rows, [summary, *header-warnings]). Each row carries its
    source sheet + row so the GUI can jump back to the right cell."""
    header_warnings = []
    fields = []  # (col_index, canonical)
    for m in colmap:
        ci = column_index_from_string(m["column"])
        actual = _norm(ws.cell(row=header_row, column=ci).value)
        # prefix match: real headers carry verbose suffixes, e.g.
        # "Description language 1 (1° part) (permanent part)"; binding is positional
        if not _hkey(actual).startswith(_hkey(m["expected_header"])):
            msg = (f"IoList[{sheet}] column {m['column']}: header '{actual}' != expected "
                   f"'{m['expected_header']}' (mapping to {m['canonical']})")
            if m["required"]:
                raise SystemExit("ERROR " + msg)
            header_warnings.append(msg)
        fields.append((ci, m["canonical"]))

    last_col = ws.max_column
    rows: list[StagedRow] = []
    struck = skip = 0
    for r in range(header_row + 1, ws.max_row + 1):
        if all(ws.cell(row=r, column=ci).value in (None, "") for ci, _ in fields):
            continue  # entirely empty row
        row = StagedRow((canon, _norm(ws.cell(row=r, column=ci).value)) for ci, canon in fields)
        row["_source_row"] = r
        row["_source_sheet"] = sheet
        if _skip_reason_present(row.get("skip_reason")):
            skip += 1
            continue
        if strike_exclude and _is_struck(ws, r, last_col):
            struck += 1
            continue
        row["_type"] = config.resolve_type(signal_types, row.get("script_type"))
        rows.append(row)

    summary = f"IoList[{sheet}]: {len(rows)} rows kept, {struck} struck, {skip} skip-reason"
    return rows, [summary] + header_warnings


def load_io_list(params: dict, signal_types: dict) -> tuple[list[StagedRow], list[str]]:
    """Read the I/O List. `io_list.sheet` may be a single name or a list of names;
    rows from every sheet are concatenated (each tagged with its source sheet)."""
    spec = params["io_list"]
    header_row = spec["header_row"]
    colmap = config.load_column_map("IoList")
    patterns = config.as_sheet_list(spec.get("sheet"))
    if not patterns:
        raise SystemExit(f"no I/O List sheet configured ({spec['path']})")
    strike_exclude = str(params.get("strike_handling", "exclude")).lower() == "exclude"

    wb = load_workbook(spec["path"], data_only=True)
    # `sheet` entries are regex patterns matched against the workbook's sheet names
    sheets = config.resolve_sheets(patterns, wb.sheetnames)
    if not sheets:
        names = wb.sheetnames
        wb.close()
        raise SystemExit(f"no sheet matched {patterns} in {spec['path']} (have {names})")

    rows: list[StagedRow] = []
    warnings: list[str] = []
    for sheet in sheets:
        sheet_rows, sheet_warnings = _load_sheet(
            wb[sheet], sheet, header_row, colmap, strike_exclude, signal_types)
        rows.extend(sheet_rows)
        warnings.extend(sheet_warnings)
    wb.close()
    # enrich each row with its safety AREAs from the C&E workbook (best-effort:
    # '' when the C&E doc is absent). Paired channels share their device's areas.
    matrix.annotate_areas(params, rows)
    # diagnosis-cabinet blocks (DiagnosticBlocks sheet), keyed by cabinet id, for the per-row
    # diag_block_name (FullName, col E) / diag_block_template (TemplateType, col F)
    diag_blocks = load_diagnostic_blocks(params)
    # functional-unit column letter (for the per-row source_cell link back into the I/O List)
    fu_col = next((m["column"] for m in colmap if m["canonical"] == "functional_unit"), "O")
    for row in rows:
        # source_cell = where this row lives in the I/O List ('Sheet!<FU col><row>'), so the
        # CentralDatabase row links straight back to the source (validation uses the same shape)
        if not row.get("source_cell") and row.get("_source_row"):
            row["source_cell"] = f"{row.get('_source_sheet', '')}!{fu_col}{row['_source_row']}"
        # name_in_db = the DB member name (the type's db_element template); '' when the
        # type isn't DB-backed
        row["name_in_db"] = outputs.member_name(row)
        # name_in_tagtable = the PLC I/O tag name (the type's tag_name template), tagtable =
        # the tag-table (Path); both '' when not a taggable I/O point or the type is untagged
        if outputs._is_io_signal(row) and outputs.tag_name(row):
            row["name_in_tagtable"] = outputs.tag_name(row)
            row["tagtable"] = outputs._tagtable(row)
        else:
            row["name_in_tagtable"] = ""
            row["tagtable"] = ""
        # diag_desc = the diagnosis alarm/warning description (the type's diag_desc template)
        row["diag_desc"] = outputs.diag_desc(row)
        # diag_block_name / diag_block_template = the FullName + TemplateType of the cabinet
        # this signal is assigned to (DiagnosticBlocks cols E/F), looked up by Diag Cabinet
        dc = str(row.get("diag_cabinet", "")).strip()
        blk = diag_blocks.get(int(dc)) if dc.lstrip("-").isdigit() else None
        row["diag_block_name"] = blk["fld"] if blk else ""
        row["diag_block_template"] = blk["template_type"] if blk else ""
        # datablocks = the DB(s) this row is a member of ('|'-joined), '' when not DB-backed
        t = row.get("_type") or {}
        row["datablocks"] = ("|".join(t.get("db_names") or [])
                             if t.get("db_kind") in ("db", "safe_db") and row["name_in_db"] else "")
        # subnet_name = "Subnet" + 3rd octet of the node IP (e.g. 192.168.50.x -> Subnet50)
        octets = str(row.get("profinet_ip", "")).split(".")
        row["subnet_name"] = f"Subnet{octets[2]}" if len(octets) == 4 and octets[2] else ""
    _add_node_address_ranges(rows)
    # IsSorterArea: the row is in an area listed in params.yaml 'sorter_areas'
    sorter = config.sorter_areas(params)
    for row in rows:
        areas = {a for a in str(row.get("matrix_areas", "")).split("|") if a}
        row["IsSorterArea"] = "yes" if (areas & sorter) else ""
    return rows, warnings


def _addr_byte(bit):
    """('I'|'Q', byte) from an address cell like 'I20.0'/'Q130.1', else None."""
    b = str(bit or "").strip().upper()
    if b[:1] in ("I", "Q") and "." in b:
        try:
            return b[0], int(b[1:b.index(".")])
        except ValueError:
            return None
    return None


def _add_node_address_ranges(rows) -> None:
    """On each node row (one carrying a profinet_name) set I_/Q_ start/end byte from the
    addresses used at that node's location; '' on the others. A signal's node is then
    the node whose matching I/Q range contains the signal's address byte."""
    by_loc_I, by_loc_Q = {}, {}
    for r in rows:
        ab = _addr_byte(r.get("bit"))
        if ab:
            (by_loc_I if ab[0] == "I" else by_loc_Q).setdefault(r.get("location", ""), []).append(ab[1])
    for r in rows:
        if r.get("profinet_name"):
            ib = by_loc_I.get(r.get("location", ""), [])
            qb = by_loc_Q.get(r.get("location", ""), [])
            r["I_startByte"] = min(ib) if ib else ""
            r["I_endByte"] = max(ib) if ib else ""
            r["Q_startByte"] = min(qb) if qb else ""
            r["Q_endByte"] = max(qb) if qb else ""
        else:
            r["I_startByte"] = r["I_endByte"] = r["Q_startByte"] = r["Q_endByte"] = ""


def load_diagnostic_blocks(params: dict) -> dict:
    """Read the I/O List 'DiagnosticBlocks' sheet -> {cabinet_id:int -> {index, fld,
    template_type}}, keyed by ID_SWP. `index` = the 3-digit cabinet string ('000','001'…);
    `fld` = FullName (the cabinet's Functional unit+Location); `template_type` = the
    per-cabinet variant (e.g. '01'-'04'), '' when that column is absent/blank. Columns are
    found by header name; {} if the sheet (or ID_SWP column) is missing."""
    wb = load_workbook(params["io_list"]["path"], data_only=True)
    sheet = next((s for s in wb.sheetnames if s.strip().lower() == "diagnosticblocks"), None)
    if sheet is None:
        wb.close()
        return {}
    grid = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    if not grid:
        return {}
    header = [_hkey(c) for c in grid[0]]

    def col(name):
        key = _hkey(name)
        return next((i for i, h in enumerate(header) if h == key), None)

    c_id, c_fld, c_tt = col("ID_SWP"), col("FullName"), col("TemplateType")
    if c_id is None:
        return {}
    out = {}
    for r in grid[1:]:
        raw = _norm(r[c_id]) if c_id < len(r) else ""
        if not raw or not raw.lstrip("-").isdigit():
            continue
        cid = int(raw)
        out[cid] = {
            "index": f"{cid:03d}",
            "fld": _norm(r[c_fld]) if c_fld is not None and c_fld < len(r) else "",
            "template_type": _norm(r[c_tt]) if c_tt is not None and c_tt < len(r) else "",
        }
    return out
