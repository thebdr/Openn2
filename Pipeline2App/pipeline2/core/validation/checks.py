"""The phase checks: the forward C&E cross-check (Phase 1), the reverse C&E cross-check
(Phase 2), the diagnosis bit map (Phase 3), and the `validate` orchestrator."""
from __future__ import annotations
import json
import os
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .. import config, i18n
from .model import (
    ALL_PHASES, CE_MANDATORY_WORDS, DOC_IO, DOC_CE, Cmp, LogEntry,
    _key, _addr, _norm, _raw, _join_set, _banner, _is_warning, _io_loc, _dev_label,
    _row_text, _first_match, _matches_words,
)
from .indexes import build_io_index, build_ce_index


def _check_reference(sheet, key_cells, key, addr_cell, address, io_index, info_row, log,
                     io_path="", raw_key="", raw_addr="", lang="en"):
    """One C&E/AREA reference vs the I/O List; appends a PASS or FAIL entry carrying the
    comparison (`cmp`). The C&E cell ($ side) is the primary link; on an address-mismatch PARTIAL
    match the matched I/O row links too (both sides). `raw_key`/`raw_addr` are the C&E cell text."""
    location = f"{sheet}!{addr_cell}"
    if not key:
        log.append(LogEntry("FAIL", location, key, address, i18n.tr("xc_p1_empty", lang), row=info_row))
        return
    entry = io_index.get(key)
    # Phase 1: caller = C&E (raw_key/raw_addr), other = I/O List (looked up); other_doc = "IOList".
    if entry is None:                                  # device not declared in the I/O List
        cmp = Cmp(raw_addr, "", raw_key, "", False, False, DOC_IO)
        log.append(LogEntry("FAIL", location, key, address, i18n.tr("xc_p1_missing", lang),
                            row=info_row, cmp=cmp))
    elif address not in entry["addrs"]:                # partial: device present, address differs
        io_cell = (info_row or {}).get("source_cell", "")
        cmp = Cmp(raw_addr, _join_set(entry["raw_addr"].values()), raw_key, entry["raw_fld"],
                  False, True, DOC_IO)
        log.append(LogEntry("FAIL", location, key, address, i18n.tr("xc_p1_fld", lang),
                            row=info_row, cmp=cmp,
                            location2=io_cell if io_path else "", path2=io_path if io_cell else ""))
    else:                                              # full match
        cmp = Cmp(raw_addr, entry["raw_addr"].get(address, raw_addr), raw_key, entry["raw_fld"],
                  True, True, DOC_IO)
        log.append(LogEntry("PASS", location, key, address, i18n.tr("xc_p1_pass", lang),
                            row=info_row, cmp=cmp))


def check_ce_matrix(params: dict, io_index: dict, io_by_key: dict, log: list, lang: str = "en") -> None:
    spec = params["ce"]
    cols = {m["canonical"]: m["column"] for m in config.load_column_map("CE")}
    wb = load_workbook(spec["path"], data_only=True, read_only=True)
    sheet = config.resolve_sheet(spec["matrix_sheet"], wb.sheetnames)   # matrix_sheet is a regex
    if not sheet:
        log.append(LogEntry("INFO", str(spec["matrix_sheet"]), "", "", "matrix sheet not found - skipped"))
        wb.close()
        return
    ws = wb[sheet]
    fu_c = column_index_from_string(cols["functional_unit"])
    loc_c = column_index_from_string(cols["location"])
    dev_c = column_index_from_string(cols["device"])
    addr_c = column_index_from_string(cols["address"])

    checked = 0
    for r in range(spec["matrix_data_row"], ws.max_row + 1):
        fu = ws.cell(r, fu_c).value
        loc = ws.cell(r, loc_c).value
        dev = ws.cell(r, dev_c).value
        addr_val = ws.cell(r, addr_c).value
        address = _addr(addr_val)
        if not any(_norm(v) for v in (fu, loc, dev)) and not address:
            continue  # empty row
        key = _key(fu, loc, dev)
        key_cells = f"{cols['functional_unit']}{r}:{cols['device']}{r}"
        ref_row = {"functional_unit": _norm(fu), "location": _norm(loc), "device": _norm(dev), "bit": address}
        info_row = io_by_key.get(key) or ref_row     # prefer the matched I/O row (desc/drawing/type)
        _check_reference(sheet, key_cells, key, f"{cols['address']}{r}", address, io_index, info_row, log,
                         io_path=(params.get("io_list") or {}).get("path", ""),
                         raw_key=_raw(fu) + _raw(loc) + _raw(dev), raw_addr=_raw(addr_val), lang=lang)
        checked += 1
    wb.close()
    log.append(LogEntry("INFO", sheet, "", "", f"{checked} reference(s) checked"))


def check_area_sheets(params: dict, io_index: dict, io_by_key: dict, log: list, lang: str = "en") -> None:
    spec = params["ce"]
    cols = {m["canonical"]: m["column"] for m in config.load_column_map("AREA")}
    wb = load_workbook(spec["path"], data_only=True, read_only=True)
    key_c = column_index_from_string(cols["concat_id"])
    addr_c = column_index_from_string(cols["address"])

    area_sheets = [s for s in wb.sheetnames if s.strip().upper().startswith("AREA")
                   and any(ch.isdigit() for ch in s)]
    for sheet in area_sheets:
        ws = wb[sheet]
        checked = 0
        for r in range(spec["area_data_row"], ws.max_row + 1):
            key_val = ws.cell(r, key_c).value
            addr_val = ws.cell(r, addr_c).value
            address = _addr(addr_val)
            if not _norm(key_val) and not address:
                continue
            key = _key(key_val)
            ref_row = {"functional_unit": _norm(key_val), "location": "", "device": "", "bit": address}
            info_row = io_by_key.get(key) or ref_row
            _check_reference(sheet, f"{cols['concat_id']}{r}", key,
                             f"{cols['address']}{r}", address, io_index, info_row, log,
                             io_path=(params.get("io_list") or {}).get("path", ""),
                             raw_key=_raw(key_val), raw_addr=_raw(addr_val), lang=lang)
            checked += 1
        log.append(LogEntry("INFO", sheet, "", "", f"{checked} reference(s) checked"))
    wb.close()


def check_diagnosis_bits(params: dict, io_rows: list, log: list) -> None:
    """Validate the diagnosis bit map of the in-diagnosis signals: each must carry a numeric
    Diag Cabinet + Diag Bit, and (cabinet, bit) must be unique within its alarm/warning
    family. Missing/invalid -> FAIL; collision -> FAIL on every colliding row; otherwise PASS.
    Entries link back to the I/O List (Diag Cabinet col AE / Diag Bit col AF)."""
    try:
        cols = {m["canonical"]: m["column"] for m in config.load_column_map("IoList")}
    except Exception:  # noqa: BLE001 - fall back to the documented positions
        cols = {}
    cab_col, bit_col = cols.get("diag_cabinet", "AE"), cols.get("diag_bit", "AF")
    io_path = (params.get("io_list") or {}).get("path", "")

    diag_rows = [r for r in io_rows if (r.get("_type") or {}).get("in_diagnosis")]
    by_slot: dict[tuple, list] = {}
    for r in diag_rows:
        cab = str(r.get("diag_cabinet", "")).strip()
        bit = str(r.get("diag_bit", "")).strip()
        cab_ok, bit_ok = cab.lstrip("-").isdigit(), bit.lstrip("-").isdigit()
        if not cab_ok or not bit_ok:
            missing = ([f"Diag Cabinet ('{cab}')"] if not cab_ok else []) + \
                      ([f"Diag Bit ('{bit}')"] if not bit_ok else [])
            loc = _io_loc(r, cab_col if not cab_ok else bit_col)
            log.append(LogEntry("FAIL", loc, _dev_label(r), f"{cab}.{bit}",
                                f"in-diagnosis signal missing/invalid {', '.join(missing)}", io_path, row=r))
            continue
        fam = "WARNING" if _is_warning(r) else "ALARM"
        by_slot.setdefault((int(cab), int(bit), fam), []).append(r)

    for (cab, bit, fam), members in sorted(by_slot.items()):
        slot = f"{cab:03d}.{bit:02d} {fam}"
        if len(members) == 1:
            r = members[0]
            log.append(LogEntry("PASS", _io_loc(r, cab_col), _dev_label(r), slot,
                                "diagnosis slot unique", io_path, row=r))
        else:
            for r in members:
                others = ", ".join(_io_loc(m, bit_col) for m in members if m is not r)
                log.append(LogEntry("FAIL", _io_loc(r, bit_col), _dev_label(r), slot,
                                    f"duplicate diagnosis slot {slot} - also at {others}", io_path, row=r))

    if diag_rows:
        log.append(LogEntry("INFO", "Diagnosis", "", "", f"{len(diag_rows)} in-diagnosis signal(s) checked"))


def _ce_match(key, addr, raw_fld_io, raw_addr_io, by_key, by_addr, loc_by_key, loc_by_addr,
              raw_fld_by_key, raw_addr_by_key, raw_fld_by_addr):
    """Match an I/O (key, addr) against the C&E indexes. Returns (status, ce_loc, cmp):
    'full'     - the device and (when it has one) its address both line up;
    'fld_only' - the device is in the C&E but at (a) different address(es);
    'addr_only'- the address is in the C&E but under a different device;
    'missing'  - neither is in the C&E.
    The I/O side of `cmp` comes from the row (raw_fld_io / raw_addr_io), the C&E side from
    the indexes; `ce_loc` is the C&E cell for a PARTIAL match (link the C&E side too), '' otherwise."""
    # Phase 2: caller = I/O List (raw_*_io), other = C&E (looked up); other_doc = "CEMatrix".
    ce_addrs = by_key.get(key)
    if ce_addrs is not None and (not addr or addr in ce_addrs or not (ce_addrs - {""})):
        ce_addr = raw_addr_by_key.get(key, {}).get(addr, raw_addr_io)   # matched addr (or device has none)
        cmp = Cmp(raw_addr_io, ce_addr, raw_fld_io, raw_fld_by_key.get(key, ""), True, True, DOC_CE)
        return "full", "", cmp
    if ce_addrs is not None:                                            # device present, address differs
        ce_addr = _join_set(raw_addr_by_key.get(key, {}).values())
        cmp = Cmp(raw_addr_io, ce_addr, raw_fld_io, raw_fld_by_key.get(key, ""), False, True, DOC_CE)
        return "fld_only", loc_by_key.get(key, ""), cmp
    if addr and addr in by_addr:                                       # address under (a) different device(s)
        cmp = Cmp(raw_addr_io, raw_addr_io, raw_fld_io, _join_set(raw_fld_by_addr.get(addr, set())),
                  True, False, DOC_CE)
        return "addr_only", loc_by_addr.get(addr, ""), cmp
    cmp = Cmp(raw_addr_io, "", raw_fld_io, "", False, False, DOC_CE)   # neither present
    return "missing", "", cmp


def check_ce_mandatory(params: dict, io_rows: list, log: list, lang: str = "en") -> None:
    """Reverse C&E: every I/O signal that must be in the C&E has to appear there, matched by
    BOTH its device key (FLD) and its address. Typed rows use their type's `ce_mandatory`
    (yes -> FAIL when not fully present, warn -> WARN, no/'' -> skip); untyped/unknown rows
    are checked when they carry a device designation (FLD) - or, device-less, a non-empty
    desc_l1/desc_l1b - plus an address, an absent one being a FAIL when its description names
    a safety concept else a WARN. A partial match (only the FLD or only the address differs)
    is logged at the same severity, printing both sides. Links to the functional-unit cell.

    Untyped rows whose description matches `ce_excluded_words` (params) are skipped entirely -
    BUT a `ce_mandatory_words` match wins (still checked), and typed rows are never excluded
    (signal-type rules prevail). `ce_full_check` (params, default false) ignores that list.
    `ce_always_excluded_words` (params, substring only - NO fuzzy) ALWAYS skips a row (typed
    or untyped), taking precedence over everything (not overridden by a mandatory match, not
    bypassed by ce_full_check) - e.g. 'ch2' to drop paired channel-2 rows. Every PASS/SKIP is
    written to the log (the .txt/.html hold the full record, with reasons); `ce_full_print`
    (params) controls only whether the live display (GUI/console) shows them (greyed)."""
    spec = params.get("ce") or {}
    if not spec.get("path") or not os.path.exists(spec["path"]):
        return
    (by_key, by_addr, loc_by_key, loc_by_addr,
     raw_fld_by_key, raw_addr_by_key, raw_fld_by_addr) = build_ce_index(params)
    io_path = (params.get("io_list") or {}).get("path", "")
    ce_path = spec["path"]                     # for the C&E-side link on a partial match
    words = params.get("ce_mandatory_words") or list(CE_MANDATORY_WORDS)
    fuzzy = int(params.get("ce_fuzzy_chars", 2) or 0)
    excluded = [] if params.get("ce_full_check") else (params.get("ce_excluded_words") or [])
    always = params.get("ce_always_excluded_words") or []       # substring only, no fuzzy, wins all
    try:
        cols = {m["canonical"]: m["column"] for m in config.load_column_map("IoList")}
    except Exception:  # noqa: BLE001
        cols = {}
    fu_col = cols.get("functional_unit", "O")

    # every PASS/SKIP is logged (the .txt/.html are the full record); the live display
    # filters them out unless params.yaml ce_full_print is set.
    def skip(r, key, addr, why):
        nonlocal skipped
        skipped += 1
        log.append(LogEntry("SKIP", _io_loc(r, fu_col), key, addr, f"skipped - {why}", io_path, row=r))

    checked = skipped = 0
    for r in io_rows:
        key = _key(r.get("functional_unit"), r.get("location"), r.get("device"))
        addr = _addr(r.get("bit"))
        raw_fld_io = _raw(r.get("functional_unit")) + _raw(r.get("location")) + _raw(r.get("device"))
        raw_addr_io = _raw(r.get("bit"))
        text = _row_text(r)
        aw = _first_match(text, always, 0)         # ALWAYS excluded (substring, no fuzzy) - wins all
        if aw:
            skip(r, key, addr, f"matches ce_always_excluded_words '{aw}'")
            continue
        typ = r.get("_type")
        if typ:                                    # known signal type -> ce_mandatory governs
            mand = str(typ.get("ce_mandatory", "")).strip().lower()
            if mand not in ("yes", "warn") or not key:
                continue
            checked += 1
            status, ce_loc, cmp = _ce_match(key, addr, raw_fld_io, raw_addr_io, by_key, by_addr,
                                            loc_by_key, loc_by_addr, raw_fld_by_key, raw_addr_by_key,
                                            raw_fld_by_addr)
            loc = _io_loc(r, fu_col)
            if status == "full":
                log.append(LogEntry("PASS", loc, key, addr, i18n.tr("xc_p2_pass", lang),
                                    io_path, row=r, cmp=cmp))
            else:
                fail = (mand == "yes")
                mkey = ("xc_p2_missing_mandatory" if fail else "xc_p2_missing_plain") \
                    if status == "missing" else ("xc_p2_fld" if status == "fld_only" else "xc_p2_addr")
                log.append(LogEntry("FAIL" if fail else "WARN", loc, key, addr, i18n.tr(mkey, lang),
                                    io_path, row=r, cmp=cmp,
                                    location2=ce_loc, path2=ce_path if ce_loc else ""))
        else:                                      # untyped / unknown
            # check a real device (FLD), or - device-less - a row that still carries a
            # description (desc_l1/desc_l1b); a bare spare channel (no device, no desc) is
            # ignored. An address is always required.
            has_desc = bool(_norm(r.get("desc_l1")) or _norm(r.get("desc_l1b")))
            if not key or not addr or not (_norm(r.get("device")) or has_desc):
                continue
            safety = _matches_words(text, words, fuzzy)
            ew = _first_match(text, excluded, fuzzy) if (not safety and excluded) else None
            if ew:
                skip(r, key, addr, f"matches ce_excluded_words '{ew}'")
                continue
            checked += 1
            status, ce_loc, cmp = _ce_match(key, addr, raw_fld_io, raw_addr_io, by_key, by_addr,
                                            loc_by_key, loc_by_addr, raw_fld_by_key, raw_addr_by_key,
                                            raw_fld_by_addr)
            loc = _io_loc(r, fu_col)
            if status == "full":
                log.append(LogEntry("PASS", loc, key, addr, i18n.tr("xc_p2_pass", lang),
                                    io_path, row=r, cmp=cmp))
                continue
            mkey = ("xc_p2_missing_safety" if safety else "xc_p2_missing_plain") \
                if status == "missing" else ("xc_p2_fld" if status == "fld_only" else "xc_p2_addr")
            log.append(LogEntry("FAIL" if safety else "WARN", loc, key, addr, i18n.tr(mkey, lang),
                                io_path, row=r, cmp=cmp,
                                location2=ce_loc, path2=ce_path if ce_loc else ""))
    if checked:
        log.append(LogEntry("INFO", "IOList in C&E", "", "", f"{checked} I/O signal(s) checked vs C&E"))
    if skipped:
        log.append(LogEntry("INFO", "IOList in C&E", "", "", f"{skipped} row(s) skipped (excluded words)"))


def validate(params: dict, io_rows: list, phases: set | None = None, lang: str = "en",
             out_dir: str | None = None) -> list:
    """The validation log: the effective params.yaml is logged first (config audit), then the
    requested phases (default ALL) of the stable 5-phase interface:
      0A. I/O List standalone      - format/content checks, ported from the VBA tool (no C&E)
      0B. C&E Matrix standalone    - placeholder (no checks implemented yet)
      1.  C&E in IOList            - every C&E / AREA reference exists in the I/O List
      2.  IOList in C&E            - every mandatory I/O signal appears in the C&E (reverse)
      3.  Diagnosis Coherence Check - the in-diagnosis cabinet/bit map is complete + unique.
    Missing documents degrade gracefully: a phase that needs an absent doc logs and ends.
    `phases` selects which to run (e.g. the designer tool passes {"0A","0B","1","2"})."""
    phases = phases if phases is not None else set(ALL_PHASES)
    io_index = build_io_index(io_rows)
    io_by_key: dict[str, dict] = {}                 # key -> a representative I/O row (for info)
    for r in io_rows:
        k = _key(r.get("functional_unit"), r.get("location"), r.get("device"))
        if k:
            io_by_key.setdefault(k, r)
    ce_path = (params.get("ce") or {}).get("path", "")
    ce_ok = bool(ce_path and os.path.exists(ce_path))
    log: list[LogEntry] = []

    # log the effective params.yaml at the start, so the validation log records the config
    # (paths, ce_* knobs) that produced it
    log.append(LogEntry("INFO", "params.yaml", "", "",
                        json.dumps(params, indent=2, default=str)))

    if "0A" in phases:
        _banner(log, i18n.tr("phase_0a", lang))
        from .. import iolist_checks            # lazy: iolist_checks imports LogEntry from here
        iolist_checks.check_iolist_standalone(params, log, lang)

    if "0B" in phases:
        _banner(log, i18n.tr("phase_0b", lang))
        log.append(LogEntry("INFO", "C&E Matrix", "", "", i18n.tr("no_checks_yet", lang)))

    if "1" in phases:
        _banner(log, "PHASE 1 - C&E in IOList  (every C&E / AREA reference exists in the I/O List)")
        log.append(LogEntry("INFO", "I/O List", "", "", f"{len(io_index)} distinct device key(s) indexed"))
        if ce_ok:
            check_ce_matrix(params, io_index, io_by_key, log, lang)
            check_area_sheets(params, io_index, io_by_key, log, lang)
        else:
            log.append(LogEntry("WARN", "C&E", "", "", f"C&E document not found ({ce_path}) - phase skipped"))

    if "2" in phases:
        _banner(log, "PHASE 2 - IOList in C&E  (every mandatory I/O signal appears in the C&E)")
        if ce_ok:
            check_ce_mandatory(params, io_rows, log, lang)
        else:
            log.append(LogEntry("WARN", "C&E", "", "", f"C&E document not found ({ce_path}) - phase skipped"))

    if "3" in phases:
        _banner(log, "PHASE 3 - Diagnosis Coherence Check  (in-diagnosis cabinet/bit map)")
        check_diagnosis_bits(params, io_rows, log)

    # error registry + per-error treatment (warn/skip/skip_type) the user set on a prior run.
    # `out_dir` given (the pipeline + GUIs) -> write/merge error_management.csv and apply rules;
    # tests that call validate() without out_dir keep the raw FAIL log untouched.
    if out_dir:
        from .. import error_management
        error_management.run(log)          # registry lives in user_input/ (persists between runs)
    return log
