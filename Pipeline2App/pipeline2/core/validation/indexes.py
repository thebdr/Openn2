"""Indexes built over the I/O List and the C&E document, feeding the cross-checks."""
from __future__ import annotations
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .. import config
from .model import _key, _addr, _raw


def build_io_index(io_rows: list) -> dict:
    """key -> {addrs: set of normalized addresses, raw_addr: {norm_addr: raw cell text},
    raw_fld: raw FU+LOC+DEV text}, for every I/O row that has a key. The raw_* values feed the
    I/O-List side of the Phase-1 comparison table; `addrs` is the match set."""
    index: dict[str, dict] = {}
    for r in io_rows:
        key = _key(r.get("functional_unit"), r.get("location"), r.get("device"))
        if not key:
            continue
        entry = index.setdefault(key, {"addrs": set(), "raw_addr": {}, "raw_fld": ""})
        norm_addr = _addr(r.get("bit"))
        entry["addrs"].add(norm_addr)
        entry["raw_addr"].setdefault(norm_addr, _raw(r.get("bit")))
        if not entry["raw_fld"]:
            entry["raw_fld"] = (_raw(r.get("functional_unit")) + _raw(r.get("location"))
                                + _raw(r.get("device")))
    return index


def build_ce_index(params: dict):
    """Indexes over the C&E document (CAUSE&EFFECT MATRIX rows + every AREA n sheet):
    by_key  = device key (FU+LOC+DEV) -> set of addresses referenced for it (a '' member
              means the C&E listed the device with no address),
    by_addr = address -> set of device keys referenced at it,
    loc_by_key / loc_by_addr = the C&E cell ("Sheet!Cell", first occurrence) for a key / address,
              so a partial match can link back into the C&E workbook,
    raw_fld_by_key = key -> raw FU+LOC+DEV cell text (first occurrence),
    raw_addr_by_key = key -> {norm_addr: raw addr cell text},
    raw_fld_by_addr = norm_addr -> set of raw device texts referenced at it,
              the three raw_* feeding the $(C&E) side of the Phase-2 comparison table.
    A signal is fully 'in the C&E' when both its key and its address line up here."""
    spec = params["ce"]
    by_key: dict[str, set] = {}
    by_addr: dict[str, set] = {}
    loc_by_key: dict[str, str] = {}
    loc_by_addr: dict[str, str] = {}
    raw_fld_by_key: dict[str, str] = {}
    raw_addr_by_key: dict[str, dict] = {}
    raw_fld_by_addr: dict[str, set] = {}

    def add(key, addr, raw_key, raw_addr, sheet, key_cell, addr_cell):
        if not key:
            return
        by_key.setdefault(key, set()).add(addr)
        loc_by_key.setdefault(key, f"{sheet}!{key_cell}")
        raw_fld_by_key.setdefault(key, raw_key)
        if addr:
            by_addr.setdefault(addr, set()).add(key)
            loc_by_addr.setdefault(addr, f"{sheet}!{addr_cell}")
            raw_addr_by_key.setdefault(key, {}).setdefault(addr, raw_addr)
            raw_fld_by_addr.setdefault(addr, set()).add(raw_key)

    wb = load_workbook(spec["path"], data_only=True, read_only=True)
    mcols = {m["canonical"]: m["column"] for m in config.load_column_map("CE")}
    sheet = config.resolve_sheet(spec["matrix_sheet"], wb.sheetnames)   # matrix_sheet is a regex
    if sheet:
        ws = wb[sheet]
        fu_c = column_index_from_string(mcols["functional_unit"])
        loc_c = column_index_from_string(mcols["location"])
        dev_c = column_index_from_string(mcols["device"])
        addr_c = column_index_from_string(mcols["address"])
        for r in range(spec["matrix_data_row"], ws.max_row + 1):
            fu, loc, dev = ws.cell(r, fu_c).value, ws.cell(r, loc_c).value, ws.cell(r, dev_c).value
            av = ws.cell(r, addr_c).value
            add(_key(fu, loc, dev), _addr(av), _raw(fu) + _raw(loc) + _raw(dev), _raw(av),
                sheet, f"{mcols['functional_unit']}{r}", f"{mcols['address']}{r}")
    acols = {m["canonical"]: m["column"] for m in config.load_column_map("AREA")}
    key_c = column_index_from_string(acols["concat_id"])
    addr_c = column_index_from_string(acols["address"])
    for s in wb.sheetnames:
        if s.strip().upper().startswith("AREA") and any(ch.isdigit() for ch in s):
            ws = wb[s]
            for r in range(spec["area_data_row"], ws.max_row + 1):
                kv, av = ws.cell(r, key_c).value, ws.cell(r, addr_c).value
                add(_key(kv), _addr(av), _raw(kv), _raw(av),
                    s, f"{acols['concat_id']}{r}", f"{acols['address']}{r}")
    wb.close()
    return (by_key, by_addr, loc_by_key, loc_by_addr,
            raw_fld_by_key, raw_addr_by_key, raw_fld_by_addr)
