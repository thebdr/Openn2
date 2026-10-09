"""Regenerate `assets/ButtonsLayout.xlsx` - the operator's one-page picture of the phase bar - from the phase
registry the app draws (`PHASES` in pipeline5/systems/plc_based/siemens_s7/safety/main.py).

The workbook keeps PL3's oracle design (its column widths, its theme, the styles below - transcribed from the
original sheet): the Run Pipeline master, one column per phase (its number in row 1, its name in row 3, the ▾
chevron in row 4), and below it the phase's sub-buttons exactly as the bar's dropdown shows them - registry
order, `<number>  <label>` in English, filled by kind (action plain, open light blue, special orange, a disabled
one grey) - joined by `↓` between two action steps (the phase button runs them in order) and `—` elsewhere. A
legend under the grid names the colours and the Run-all order. The styles are defined here, never read back from
the sheet, so a rerun gives the same workbook. `tests/unit/test_gui_phase_model.py` fails when the workbook and
the registry drift apart: rerun this script after changing the registry.

Run:   <python> scripts/gen_buttons_layout.py
"""
from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PL5_ROOT = os.path.dirname(HERE)
sys.path.insert(0, PL5_ROOT)

WORKBOOK = os.path.join(PL5_ROOT, "assets", "ButtonsLayout.xlsx")
FIRST_COLUMN = 4                  # D: the first phase column (B = Run Pipeline, then a phase every 2nd column)
HEADER_ROW, CHEVRON_ROW, FIRST_SUB_ROW = 3, 4, 5
HEADER_HEIGHT, SUB_HEIGHT, LINK_HEIGHT = 45.75, 60.0, 18.75
CHEVRON = "▾"                     # the bar's chevron glyph
PHASE_ARROW = "→"
LEGEND = (("action", "action - the phase runs these in order (↓)"), ("open", "open - a file or folder"),
          ("special", "special - a separate operator action"), ("disabled", "disabled - shown greyed in the bar"))


def styles() -> dict:
    """{kind -> cell attributes}: the original sheet's styles (Courier New buttons in thin boxes, the bold phase
    heads and the pink master in medium boxes, theme fills) + the bar's disabled grey (theme.BUTTON_FILLS)."""
    from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
    text, mono, narrow = Color(theme=1, tint=0.0), "Courier New", "Aptos Narrow"
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def box(style):
        side = Side(style=style)
        return Border(left=side, right=side, top=side, bottom=side)

    def fill(theme, tint):
        return PatternFill(fill_type="solid", fgColor=Color(theme=theme, tint=tint))

    button = dict(font=Font(name=mono, sz=11, color=text), alignment=center, border=box("thin"))
    return {
        "run": dict(font=Font(name=mono, sz=11, color=text), fill=fill(8, 0.7999816888943144), alignment=center,
                    border=box("medium")),
        "header": dict(font=Font(name=mono, sz=11, b=True, color=text), alignment=center, border=box("medium")),
        "arrow": dict(font=Font(name=narrow, sz=16, b=True, color=text),
                      alignment=Alignment(horizontal="center", vertical="center")),
        "chevron": dict(font=Font(name=mono, sz=11, color=text), fill=fill(0, -0.1499984740745262), alignment=center),
        "number": dict(font=Font(name=narrow, sz=11, color="FF00B050"), alignment=Alignment(horizontal="center")),
        "link": dict(font=Font(name=mono, sz=14, color=text), alignment=center),
        "action": button,
        "open": dict(button, fill=fill(3, 0.8999908444471572)),
        "special": dict(button, fill=fill(5, 0.7999816888943144)),
        "disabled": dict(button, font=Font(name=mono, sz=11, i=True, color="FF636E72"),
                         fill=PatternFill(fill_type="solid", fgColor="FFB2BEC3")),
        "label": dict(font=Font(name=narrow, sz=11, b=True), alignment=Alignment(vertical="center")),
        "note": dict(font=Font(name=narrow, sz=11)),
    }


def phase_columns(phases) -> dict:
    """{phase number -> its column (1-based)}: the phases with a number, left to right in bar order."""
    numbered = [p for p in phases if p.number]
    return {p.number: FIRST_COLUMN + 2 * i for i, p in enumerate(numbered)}


def sub_text(sub, lang: str = "en") -> str:
    """A sub-button's text as the dropdown shows it."""
    from pipeline5.language import i18n
    return f"{sub.number}  {i18n.tr(sub.label_key, lang)}"


def header_text(phase, lang: str = "en") -> str:
    """A phase header: its name on two lines at the first space (the bar wraps it the same way)."""
    from pipeline5.language import i18n
    return i18n.tr(phase.name_key, lang).replace(" ", "\n", 1)


def link(a, b) -> str:
    """`↓` between two action steps (the phase runs them in order), `—` around an open/special button."""
    return "↓" if a.kind == "action" and b.kind == "action" else "—"


def build(path: str = WORKBOOK) -> str:
    import openpyxl
    from pipeline5.language import i18n
    from pipeline5.systems.plc_based.siemens_s7.safety.main import PHASES

    wb = openpyxl.load_workbook(path)             # kept for its column widths, theme and Normal style
    ws = wb.worksheets[0]
    last = max(ws.max_row, 60)
    for row in ws.iter_rows(min_row=1, max_row=last):
        for cell in row:
            cell.value = None
            cell.style = "Normal"
    for r in range(1, last + 1):
        ws.row_dimensions[r].height = None
    look = styles()

    def put(row, col, value, kind):
        cell = ws.cell(row=row, column=col, value=value)
        for attribute, setting in look[kind].items():
            setattr(cell, attribute, setting)
        return cell

    columns = phase_columns(PHASES.phases)
    ws.row_dimensions[HEADER_ROW].height = HEADER_HEIGHT
    run = next(p for p in PHASES.phases if not p.number)
    put(HEADER_ROW, 2, i18n.tr(run.name_key, "en"), "run")
    deepest = CHEVRON_ROW
    for phase in (p for p in PHASES.phases if p.number):
        col = columns[phase.number]
        put(1, col, phase.number, "number")
        put(HEADER_ROW, col, header_text(phase), "header")
        if col + 2 in columns.values():
            put(HEADER_ROW, col + 1, PHASE_ARROW, "arrow")
        put(CHEVRON_ROW, col, CHEVRON, "chevron")
        subs = list(phase.subs or ())
        for i, sub in enumerate(subs):
            row = FIRST_SUB_ROW + 2 * i
            put(row, col, sub_text(sub), sub.kind if sub.enabled else "disabled")
            ws.row_dimensions[row].height = SUB_HEIGHT
            if i + 1 < len(subs):
                put(row + 1, col, link(sub, subs[i + 1]), "link")
                ws.row_dimensions[row + 1].height = LINK_HEIGHT
            deepest = max(deepest, row)

    legend = deepest + 3
    put(legend, 2, "Legend", "label")
    for offset, (kind, text) in enumerate(LEGEND):
        put(legend, FIRST_COLUMN + 2 * offset, text, kind)
    ws.row_dimensions[legend].height = SUB_HEIGHT
    order = " → ".join(str(n) for n in PHASES.run_plan)
    absent = [str(p.number) for p in PHASES.phases if p.number and p.number not in PHASES.run_plan]
    put(legend + 2, 2, f"Run Pipeline (Run-all) order: {order}"
        + (f" - not in Run-all: {', '.join(absent)} (its own button only)" if absent else ""), "note")
    put(legend + 3, 2, "Generated from PHASES in pipeline5/systems/plc_based/siemens_s7/safety/main.py by "
                       "scripts/gen_buttons_layout.py - change the registry, then rerun the script.", "note")
    wb.save(path)
    return path


if __name__ == "__main__":
    print("written:", build())
