"""ph200 / 210 - the in-place AB/AC fill (phases/fillout/document_fill) over a synthetic raw I/O
List. The fill takes the CALLER's staged database since step 5 - `_fill_210` mirrors the run-plan."""
import os
import tempfile

from openpyxl import Workbook, load_workbook

from _harness import run, eq, ok
from pipeline5 import config
from pipeline5.phases.fillout import document_fill as fill
from pipeline5.phases.staging import iolist as staging
from pipeline5.systems import catalog


def _fill_210(params):
    """The 210-only run-plan wiring: stage the source doc, hand the database to fill_script_type."""
    database, _ = staging.stage(params, system=catalog.by_id("siemens_s7_safety"))
    return fill.fill_script_type(database, params)


def _sandboxed(fn):
    """Run `fn` with `config.database_dir` pointed at a throwaway dir: the domain build/fill paths
    SAVE the SSOT, and un-redirected that overwrote the builtin Shared/Database (the parity
    environment) with this module's synthetic staging on every gate run."""
    def wrapped():
        original = config.database_dir
        with tempfile.TemporaryDirectory() as sandbox:
            config.database_dir = lambda: sandbox
            try:
                fn()
            finally:
                config.database_dir = original
    return wrapped


_SHEET = "NETSAFETY"


def _raw_iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET
    for col, head in (("F", "ID"), ("G", "Bit"), ("K", "Desc1"), ("L", "Desc2"), ("R", "Type"), ("X", "Mnem")):
        ws[f"{col}1"] = head                                   # AA..AG headers left BLANK (a raw doc)
    ws["G2"] = "I0.0"; ws["K2"] = "EMERGENCY PUSH-BUTTON PRESSED"; ws["L2"] = "CH1"      # -> E1/2 (Mode-1)
    ws["F3"] = "5"; ws["R3"] = "PLC"; ws["K3"] = "MAIN PLC NODE"                          # node -> PLC
    ws["G4"] = "I0.1"; ws["K4"] = "SOMETHING NOT RECOGNISED AT ALL"                       # -> <input required>
    ws["G5"] = "I0.2"; ws["K5"] = "DOOR OPEN"; ws["L5"] = "CH2"; ws["AB5"] = "MANUAL"     # Mode-2 (computed DI2/2)
    wb.save(path)


def _params(path):
    return {"iolist_path": path, "iolist_params": {"sheets": _SHEET, "header_row": 1}}


def test_fill_ab_ac_modes_and_unresolved():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        _raw_iolist(p)
        res = _fill_210(_params(p))
        wb = load_workbook(p)
        ws = wb[_SHEET]
        eq(ws["AB2"].value, "E1/2", "Mode-1: AB written from the rules")
        eq(ws["AC2"].value, "E1/2", "AC mirrors the computed value")
        eq(ws["AB3"].value, "PLC", "node AB = Type col verbatim")
        eq(ws["AB4"].value, "<input required>", "unresolved AB -> sentinel")
        eq(ws["AB5"].value, "MANUAL", "Mode-2: a human AB value is preserved")
        eq(ws["AC5"].value, "DI2/2", "AC still shows what the rules computed (Mode-2)")
        eq(ws["AB1"].value, "Script Type", "a blank output header is written")
        ok("_UnresolvedIndex" in wb.sheetnames, "the _UnresolvedIndex sheet is added")
        eq((res["filled"], res["mismatch"], res["unresolved"]), (2, 1, 1), "counts: 2 filled, 1 kept, 1 unresolved")
        eq([f.type for f in res["findings"] if f.severity == "FAIL"], ["fill_unresolved"], "one blocking FAIL")
        eq([f.type for f in res["findings"] if f.type == "fill_type_mismatch"][0:1], ["fill_type_mismatch"], "Mode-2 audit WARN")
        ok(res["backup"], "the raw doc changed -> the backup is KEPT")
        # the finding hand-off: doc-located findings carry the WORKBOOK (the log links + the Findings
        # jump need it) and the treatable ones are RECORDED (ph200 is doc-only, but findings are facts)
        fail = next(f for f in res["findings"] if f.severity == "FAIL")
        eq(fail.doc, "io.xlsx", "a Sheet!Cell finding carries its workbook")
        ok("!" in fail.location, "…and a doc-located location")
        issues_path = os.path.join(config.database_dir(), "validation_issues.csv")
        ok(os.path.exists(issues_path), "the fill findings are recorded to validation_issues")
        with open(issues_path, encoding="utf-8") as h:
            body = h.read()
        ok(fail.uid in body and "fill_unresolved" in body,
           "the [FAIL] log line's uid IS in the table - the Findings jump lands")


def test_fill_is_idempotent_noop_drops_backup():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        _raw_iolist(p)
        _fill_210(_params(p))                                  # first run fills it
        res2 = _fill_210(_params(p))                           # second run changes nothing
        eq(res2["backup"], "", "no-op re-run deletes its backup")
        eq(res2["filled"], 0, "nothing re-filled (AB already present)")


if __name__ == "__main__":
    import sys
    sys.exit(run("fillout_fill", [
        ("fill_ab_ac_modes_and_unresolved", _sandboxed(test_fill_ab_ac_modes_and_unresolved)),
        ("fill_is_idempotent_noop_drops_backup", _sandboxed(test_fill_is_idempotent_noop_drops_backup)),
    ]))
