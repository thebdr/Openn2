"""Phase 510b - the tag-table XML emitter (plctags_xml_emitter): one `SW.Tags.PlcTagTable` document per tag
table of the collected tags, well-formed, mirroring the workbook row for row (the same tags per table, the
Hmi flags as the External* attributes, the %-prefixed address), the comment item only where a comment
exists, hex ids unique per document, the contract-v1 header (name = the exact table name) parseable by the
OP5 rules, file-safe file names (OP5's rule) with a collision refused, and the writer wired into the 510
projection beside the workbook under the same duplicate-tag gate."""
import os
import tempfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from _harness import run, eq, ok, raises
from pipeline5.systems.plc_based.siemens_s7 import openn_header as header
from pipeline5.systems.plc_based.siemens_s7 import plctags_xlsx_writer as io_tags
from pipeline5.systems.plc_based.siemens_s7 import plctags_xml_emitter as tag_xml
from pipeline5.truth.database import Database
from pipeline5.truth.table import Table

_TAGS = [
    {"name": "PB STOP", "path": "EMERGENCY_PushButtons", "data_type": "Bool", "address": "%I1.0",
     "comment": "Emergency stop [ =S1-F1 ]"},
    {"name": "PB RESET", "path": "EMERGENCY_PushButtons", "data_type": "Bool", "address": "%I1.1", "comment": ""},
    {"name": "PNC_Q_S1 HEARTBEAT", "path": "IF_SORTER-01", "data_type": "Word", "address": "%Q10000.0",
     "comment": "heart <beat> & more"},
    {"name": "Door Closed [ =S1+SG1-B1 ]", "path": "Alarms/Warnings: A|B", "data_type": "Bool",
     "address": "%I2.0", "comment": ""},
]


def _parse(path: str) -> ET.Element:
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return ET.fromstring(handle.read())


def _tags_of(root: ET.Element) -> list:
    return root.findall("./SW.Tags.PlcTagTable/ObjectList/SW.Tags.PlcTag")


def _workbook_rows(xlsx: str) -> list:
    wb = load_workbook(xlsx, read_only=True)
    try:
        return [r for r in wb["PLC Tags"].iter_rows(values_only=True)][1:]
    finally:
        wb.close()                                    # read-only mode keeps the zip open (Windows: the temp dir)


def test_one_document_per_table_mirroring_the_workbook():
    with tempfile.TemporaryDirectory() as d:
        xlsx = io_tags.write_plc_tags(_TAGS, d)
        paths = tag_xml.write_tag_tables(_TAGS, d)
        eq([os.path.basename(p) for p in paths],
           ["Alarms_Warnings_ A_B.xml", "EMERGENCY_PushButtons.xml", "IF_SORTER-01.xml"],
           "one file per table, table-sorted, the name made file-safe (OP5's rule: invalid chars -> _)")
        rows = _workbook_rows(xlsx)
        per_table_in_workbook = {}
        for r in rows:
            per_table_in_workbook.setdefault(r[1], []).append(r)
        total = 0
        for path in paths:
            root = _parse(path)
            eq(root.tag, "Document")
            eq(root.find("Engineering").get("version"), "V18")
            table_el = root.find("SW.Tags.PlcTagTable")
            eq(table_el.get("ID"), "0")
            table = table_el.findtext("AttributeList/Name")
            tags = _tags_of(root)
            wb_rows = per_table_in_workbook[table]
            eq(len(tags), len(wb_rows), f"{table}: one PlcTag per workbook row of that table")
            for tag, row in zip(tags, wb_rows):
                attrs = tag.find("AttributeList")
                eq([c.tag for c in attrs], ["DataTypeName", "ExternalAccessible", "ExternalVisible",
                                            "ExternalWritable", "LogicalAddress", "Name"],
                   "the attribute order of a TIA export (alphabetical)")
                eq((attrs.findtext("Name"), attrs.findtext("DataTypeName"), attrs.findtext("LogicalAddress")),
                   (row[0], row[2], row[3]), "name, data type and the %-prefixed address as in the workbook")
                eq((attrs.findtext("ExternalVisible"), attrs.findtext("ExternalAccessible"), attrs.findtext("ExternalWritable")),
                   tuple(str(v).lower() for v in row[5:8]), "the Hmi flags as the External* attributes")
                comment = tag.find("ObjectList/MultilingualText[@CompositionName='Comment']")
                if row[4]:
                    ok(comment is not None, f"{row[0]}: a comment item where the workbook has a comment")
                    item = comment.find("ObjectList/MultilingualTextItem")
                    eq((item.get("CompositionName"), item.findtext("AttributeList/Culture"), item.findtext("AttributeList/Text")),
                       ("Items", "en-US", row[4]), "en-US item with the comment text (escaped on the way out)")
                else:
                    ok(comment is None and tag.find("ObjectList") is None, f"{row[0]}: no comment, no ObjectList")
            total += len(tags)
        eq(total, len(rows), "the documents hold exactly the workbook's tags")


def test_ids_unique_and_hex_bytes_bom_crlf():
    many = [{"name": f"T{i}", "path": "BIG", "data_type": "Bool", "address": f"%I{i}.0", "comment": f"c{i}"}
            for i in range(12)]
    text = tag_xml.tag_table_xml("BIG", many)
    ids = [el.get("ID") for el in ET.fromstring(text).iter() if el.get("ID") is not None]
    eq(len(ids), len(set(ids)), "ids unique within the document")
    eq(ids[:4], ["0", "1", "2", "3"], "the table is 0, then tag / comment / item")
    ok("A" in ids and "B" in ids and "10" in ids, f"hex, as TIA writes them: {ids[8:14]}")
    ok(text.startswith('<?xml version="1.0" encoding="utf-8"?>\r\n<Document>\r\n  <Engineering version="V18" />\r\n'))
    ok("\n" not in text.replace("\r\n", ""), "CRLF only")
    ok(text.endswith("</Document>") and "<ObjectList>" in text, "no trailing newline")
    empty = tag_xml.tag_table_xml("EMPTY", [])
    ok("<ObjectList>" not in empty and "<Name>EMPTY</Name>" in empty, "an empty table: no ObjectList at all")
    with tempfile.TemporaryDirectory() as d:
        path = tag_xml.write_tag_tables(many, d)[0]
        raw = open(path, "rb").read()
        ok(raw.startswith(b'\xef\xbb\xbf<?xml version="1.0" encoding="utf-8"?>\r\n<!--\r\n#!openn\r\n#! kind: sw/tag-table\r\n'),
           "BOM, declaration, then the contract-v1 header as the first comment")
        ok(raw.endswith(b"</Document>") and b"\r\n" in raw, "CRLF, no trailing newline")


def test_header_per_table_and_the_double_dash_rule():
    header.begin_run("tags-run-1")
    tags = _TAGS + [{"name": "X", "path": "A--B", "data_type": "Bool", "address": "%I3.0", "comment": ""}]
    with tempfile.TemporaryDirectory() as d:
        paths = tag_xml.write_tag_tables(tags, d, plc="n0001-mc1-cc1-k65501")
        for path in paths:
            h = header.read_header(path)
            eq(h.status, "ok", f"{os.path.basename(path)}: {h.problems}")
            eq((h.get("kind"), h.get("schema"), h.get("producer"), h.get("run"), h.get("plc"), h.get("target"), h.get("source")),
               ("sw/tag-table", "1", "Pipeline5 5.0 (phase 510)", "tags-run-1", "n0001-mc1-cc1-k65501", "PLC tags",
                "signals + interface_elements"))
            table = _parse(path).findtext("SW.Tags.PlcTagTable/AttributeList/Name")
            expected = table.replace("--", "- -")              # an XML comment cannot hold `--`
            eq(h.get("name"), expected, "`name` = the exact table name (the `--` rule applied)")
        ok(os.path.exists(os.path.join(d, "A--B.xml")), "the file keeps the real name; only the comment is rewritten")
        ok("<Name>A--B</Name>" in open(os.path.join(d, "A--B.xml"), encoding="utf-8-sig").read(),
           "...and so does the document's own <Name>")


def test_file_names_and_collisions():
    eq(tag_xml.safe_file_name('a<b>c:d"e/f\\g|h?i*j'), "a_b_c_d_e_f_g_h_i_j")
    eq(tag_xml.safe_file_name("IF_SORTER+DIAG-02"), "IF_SORTER+DIAG-02", "+ and - are fine")
    eq(tag_xml.safe_file_name("tab\tname"), "tab_name", "control characters too")
    tags = [{"name": "a", "path": "A/B", "data_type": "Bool", "address": "%I0.0", "comment": ""},
            {"name": "b", "path": "A_B", "data_type": "Bool", "address": "%I0.1", "comment": ""}]
    raises(ValueError, lambda: tag_xml.tag_table_files(tags))


def _db(signals):
    cols = ["uid", "script_type", "bit", "name_in_tagtable", "tagtable", "type", "source_cell", "is_io", "profinet_name"]
    t = Table("signals", columns=cols, json_columns=["type"], key_columns=["uid"])
    for i, s in enumerate(signals):
        row = {"uid": f"u{i}", "source_cell": f"S!A{i}", **s}
        t.add(**row)
    return Database([t])


def test_wired_into_the_510_projection_under_the_duplicate_gate():
    from pipeline5 import config
    base = {"script_type": "DI1/2", "tagtable": "EMERGENCY_PushButtons", "type": {"category": "Safety", "io_comment": ""}}
    db = _db([{**base, "bit": "I1.0", "name_in_tagtable": "PB"}, {**base, "bit": "I1.1", "name_in_tagtable": "PB2"}])
    with tempfile.TemporaryDirectory() as d:
        res = io_tags.project(db, out_dir=d)
        eq(res["tables"], ["EMERGENCY_PushButtons"])
        eq([os.path.basename(p) for p in res["xml_files"]], ["EMERGENCY_PushButtons.xml"],
           "the 510 projection writes the tag-table XML beside the workbook")
        eq(len(_tags_of(_parse(res["xml_files"][0]))), 2)
    dup = _db([{**base, "bit": "I1.0", "name_in_tagtable": "PB"}, {**base, "bit": "I1.1", "name_in_tagtable": "pb"}])
    orig_params, orig_dbdir = config.load_params, config.database_dir
    with tempfile.TemporaryDirectory() as d:
        config.load_params = lambda *a, **k: {"iolist_path": os.path.join(d, "IOList.xlsx")}
        config.database_dir = lambda: d
        try:
            res = io_tags.project(dup, out_dir=d)
        finally:
            config.load_params, config.database_dir = orig_params, orig_dbdir
        eq((res["path"], res["xml_files"]), ("", []), "a duplicate tag: no workbook and no XML (the raw-FAIL guard)")
        eq([f for f in os.listdir(d) if f.endswith(".xml")], [], "nothing on disk")


if __name__ == "__main__":
    import sys
    sys.exit(run("plctags_xml", [
        ("one_document_per_table_mirroring_the_workbook", test_one_document_per_table_mirroring_the_workbook),
        ("ids_unique_and_hex_bytes_bom_crlf", test_ids_unique_and_hex_bytes_bom_crlf),
        ("header_per_table_and_the_double_dash_rule", test_header_per_table_and_the_double_dash_rule),
        ("file_names_and_collisions", test_file_names_and_collisions),
        ("wired_into_the_510_projection_under_the_duplicate_gate", test_wired_into_the_510_projection_under_the_duplicate_gate),
    ]))
