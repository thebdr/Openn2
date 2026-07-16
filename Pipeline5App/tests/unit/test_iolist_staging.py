"""Staging (domain.identity + domain.staging): the FLD/tag identity + the IoList read into rows."""
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font

from _harness import run, eq, ok
from pipeline5.truth import identity
from pipeline5.phases.staging import iolist as staging
from pipeline5.truth.signals import signals_table
from pipeline5.documents import xlsx_reader as workbook


def test_identity_flds_and_tag():
    row = {"functional_unit": "S1", "location": "+SG1", "device": "-B1", "script_type": "DI1/2", "bit": "I0.0",
           "type": {"category": "Safety", "tag_name": "Door [ {$combined_FLD} ]", "tagtable_name": "SAFETY_Doors"}}
    row["combined_FLD"] = identity.combined_fld(row)
    eq(identity.fld(row), "S1+SG1-B1", "iol_FLD = FU+LOC+DEV")
    eq(identity.combined_fld(row), "S1+SG1-B1", "combined == iol when there is no C&E FLD")
    ok(identity.is_io_signal(row), "an I/Q bit on a non-interface type is an I/O signal")
    eq(identity.tag_name(row), "Door [ S1+SG1-B1 ]", "tag_name template interpolated")
    eq(identity.tagtable(row), "SAFETY_Doors", "tagtable from the type")


def test_interp_delegates_to_expr_render():
    # empty mode: a missing field -> "" ; outer whitespace trimmed; numeric spec coerced; blank stays blank.
    eq(identity.interp("  {$a}-{$b}  ", {"a": "x", "b": "y"}), "x-y", "resolves + outer-trims")
    eq(identity.interp("{$a}-{$missing}", {"a": "x"}), "x-", "a missing field renders empty")
    eq(identity.interp("CAB{$diag_cabinet:03d}", {"diag_cabinet": "1"}), "CAB001", "numeric spec coerces")
    eq(identity.interp("CAB{$diag_cabinet:03d}", {"diag_cabinet": ""}), "CAB", "a blank value is not padded")
    eq(identity.interp("{$n:02d}", {"n": "0"}), "00", "a literal '0' coerces, not blank")
    eq(identity.interp(None, {}), "", "None template -> ''")


def test_combined_fld_appends_ce_when_different():
    row = {"functional_unit": "S1", "location": "+SG1", "device": "-B1",
           "ce_functional_unit": "S1", "ce_location": "+CE", "ce_device": "-B1"}
    eq(identity.combined_fld(row), "S1+SG1-B1 S1+CE-B1", "the C&E-side FLD is appended only when it differs")


_COLMAP = [{"column": c, "canonical": n} for c, n in
           [("A", "functional_unit"), ("B", "location"), ("C", "device"),
            ("D", "script_type"), ("E", "bit"), ("F", "skip_reason")]]
_TYPES = {"DI1/2": {"type_id": "DI1/2", "category": "Safety",
                    "tag_name": "Door [ {$combined_FLD} ]", "tagtable_name": "SAFETY_Doors"}}


def _iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    for col, head in zip("ABCDEF", ["FU", "Loc", "Dev", "Type", "Bit", "Skip"]):
        ws[f"{col}1"] = head
    ws["A2"] = "S1"; ws["B2"] = "+SG1"; ws["C2"] = "-B1"; ws["D2"] = "DI1/2"; ws["E2"] = "I0.0"
    ws["A3"] = "S2"; ws["B3"] = "+SG2"; ws["C3"] = "-B2"; ws["D3"] = "DI1/2"; ws["E3"] = "I0.1"; ws["F3"] = "skip me"
    ws["A4"] = "S3"; ws["B4"] = "+SG3"; ws["C4"] = "-B3"; ws["D4"] = "DI1/2"; ws["E4"] = "I0.2"
    ws["A4"].font = Font(strike=True)
    wb.save(path)


def test_read_view_drops_skip_and_struck():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx"); _iolist(p)
        v = workbook.open_sheet(p, "NET SAFETY 50", 1)
        rows = staging._read_view(v, _COLMAP, True, _TYPES)        # strike_exclude=True
        eq(len(rows), 1, "the Skip-Reason row + the struck row are dropped -> 1 kept")
        r = rows[0]
        eq(r["functional_unit"], "S1"); eq(r["script_type"], "DI1/2")
        eq(r["type"]["type_id"], "DI1/2", "the resolved type is attached")
        eq(r["source_sheet"], "NET SAFETY 50"); eq(r["source_row"], 2)
        v.close()


def test_read_view_keeps_struck_when_not_excluding():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx"); _iolist(p)
        v = workbook.open_sheet(p, "NET SAFETY 50", 1)
        rows = staging._read_view(v, _COLMAP, False, _TYPES)       # strike_exclude=False
        eq(len(rows), 2, "the struck row is KEPT when not excluding; the Skip-Reason row is still dropped")
        v.close()


def test_is_sorter_area_number_to_name():
    names = staging._sorter_area_names({"matrix_params": {"sorter_areas": [1]}})
    eq(names, {"AREA 1"}, "sorter number 1 -> area name 'AREA 1'")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 1"]}, names), "yes", "matches its sorter area")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 2", "AREA 1"]}, names), "yes", "yes when ANY area is a sorter")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 3"]}, names), "", "a non-sorter area -> ''")
    eq(staging._is_sorter_area({"matrix_areas": []}, names), "", "no areas -> ''")
    eq(staging._is_sorter_area({}, names), "", "missing matrix_areas -> ''")
    eq(staging._sorter_area_names({}), set(), "no sorter_areas configured -> no names")


def test_node_address_ranges_positional():
    rows = [
        {"source_sheet": "S1", "profinet_name": "n1", "bit": ""},      # node head (owns no address of its own)
        {"source_sheet": "S1", "bit": "I0.0"},
        {"source_sheet": "S1", "bit": "I20.7"},
        {"source_sheet": "S1", "bit": "Q0.0"},
        {"source_sheet": "S1", "profinet_name": "n2", "bit": "Q5.0"},  # next head opens a new span + owns its own Q5
        {"source_sheet": "S1", "bit": "I8.1"},
        {"source_sheet": "S2", "bit": "I3.0"},                          # past the sheet boundary -> owned by no node
    ]
    staging._add_node_address_ranges(rows)
    eq((rows[0]["I_startByte"], rows[0]["I_endByte"]), (0, 20), "n1 I range = min/max byte beneath it")
    eq((rows[0]["Q_startByte"], rows[0]["Q_endByte"]), (0, 0), "n1 owns the Q0 row in its span too")
    eq((rows[4]["I_startByte"], rows[4]["I_endByte"]), (8, 8), "n2 owns the I8 row beneath it on the same sheet")
    eq((rows[4]["Q_startByte"], rows[4]["Q_endByte"]), (5, 5), "n2 owns its own Q5 bit")
    eq(rows[1]["I_startByte"], "", "a non-node row NEVER carries a range")
    eq(rows[6]["I_startByte"], "", "the sheet boundary ends n2's span -> the S2 row is unowned")


def test_node_address_range_empty_when_no_addressed_rows():
    rows = [{"source_sheet": "S", "profinet_name": "n", "bit": ""}]
    staging._add_node_address_ranges(rows)
    eq(rows[0]["I_startByte"], "", "a node owning no addressed row -> empty range (like PL3's 7 empty nodes)")
    eq(rows[0]["Q_endByte"], "")


def test_dup_signal_uid_findings():
    """The post-stage duplicate-uid check (moved from the GUI into stage) is a stg_dup_signal_uid WARN,
    one per shared content-hash uid; the clean case emits nothing."""
    cols = ["combined_FLD", "script_type", "source_cell"]      # the stable key -> the content-hash uid
    t = signals_table(cols)
    for _ in range(2):                                          # two rows share the key -> the SAME uid
        t.add_row({"combined_FLD": "S1+SG1-B1", "script_type": "DI1/2", "source_cell": "NET!O2"})
    t.add_row({"combined_FLD": "S2+SG2-B2", "script_type": "DI1/2", "source_cell": "NET!O3"})   # unique
    fs = staging._dup_findings(t)
    eq(len(fs), 1, "one WARN per shared uid (the unique row produces none)")
    eq((fs[0].phase, fs[0].type, fs[0].severity), (300, "stg_dup_signal_uid", "WARN"), "the finding container")
    ok(fs[0].location in t.duplicate_uids(), "the finding locates the shared uid")
    eq(staging._dup_findings(signals_table(cols)), [], "no rows -> no findings (the clean path)")


def test_finalize_identity_ce_overwrite():
    """_finalize_identity is idempotent + C&E-overwriting: the no-C&E pass yields combined_FLD == iol_FLD
    and IsSorterArea '', and a SECOND pass once the matrix fields are present lifts combined_FLD to include
    the C&E side + sets IsSorterArea. This idempotence is what makes the 310/320 split byte-exact."""
    params = {"matrix_params": {"sorter_areas": [1]}}
    row = {"functional_unit": "S1", "location": "+SG1", "device": "-B1",
           "script_type": "DI1/2", "bit": "I0.0", "type": {"category": "Safety"},
           "source_sheet": "NET", "source_row": 2}
    staging._finalize_identity(params, [row])                       # the 310 pass: no C&E yet
    eq(row["iol_FLD"], "S1+SG1-B1")
    eq(row["combined_FLD"], "S1+SG1-B1", "no C&E -> combined == iol")
    eq(row["IsSorterArea"], "", "no matrix areas -> not a sorter area")
    eq(row["source_cell"], "NET!O2", "source_cell derived from the node sheet+row")
    row["ce_functional_unit"] = "S1"; row["ce_location"] = "+CE"; row["ce_device"] = "-B1"
    row["matrix_areas"] = ["AREA 1"]                                # the 320 pass: matrix.annotate has run
    staging._finalize_identity(params, [row])
    eq(row["combined_FLD"], "S1+SG1-B1 S1+CE-B1", "the C&E side is appended once it differs")
    eq(row["IsSorterArea"], "yes", "AREA 1 is a configured sorter area")


def test_annotate_cematrix_records_and_restamps():
    """annotate_cematrix enriches the staged signals in place, re-stamps each uid from the (now C&E)
    combined_FLD, and records the duplicate-uid findings to validation_issues. With no C&E document the
    combined_FLD is unchanged (uid stable) and a shared key still surfaces one stg_dup_signal_uid WARN."""
    from pipeline5.truth.database import Database
    from pipeline5.truth.diagnosis import diagnosis_cabinets_table
    table = signals_table(["functional_unit", "location", "device", "script_type", "bit"])
    for _ in range(2):                                             # two rows share the stable key -> one uid
        table.add_row({"functional_unit": "S1", "location": "+SG1", "device": "-B1", "script_type": "DI1/2",
                       "bit": "I0.0", "type": {"category": "Safety"}, "source_cell": "NET!O2"})
    table.add_row({"functional_unit": "S2", "location": "+SG2", "device": "-B2", "script_type": "DI1/2",
                   "bit": "I0.1", "type": {"category": "Safety"}, "source_cell": "NET!O3"})
    db = Database([table, diagnosis_cabinets_table()])
    _db, findings = staging.annotate_cematrix(db, {"matrix_params": {}}, save=False)   # no matrix_path -> empty C&E
    eq(len(findings), 1, "the shared-key rows surface one stg_dup_signal_uid WARN")
    eq(findings[0].type, "stg_dup_signal_uid", "the dup finding container")
    ok("validation_issues" in db, "the findings are recorded to validation_issues")
    eq(db["signals"].rows[0]["combined_FLD"], "S1+SG1-B1", "no C&E doc -> combined_FLD == iol_FLD (uid stable)")


def test_load_io_list_no_match_returns_empty():
    """When no sheet matches the pattern, load_io_list returns ([], []) - the predicate stage() turns into the
    blocking stg_no_io_sheet FAIL (replacing the old raise SystemExit). The FAIL emission + no-write is verified
    end-to-end by the data-dependent parity smoke; this pins the load-bearing no-match branch."""
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        wb = Workbook(); wb.active.title = "NotAnIoSheet"; wb.active["A1"] = "x"; wb.save(p)
        params = {"iolist_params": {"sheets": "^__no_match__$", "header_row": 1}}
        rows, matched = staging.load_io_list(params, {}, p)
        eq(matched, [], "no sheet matched the pattern -> empty match list")
        eq(rows, [], "no rows read")


def test_dup_type_index_findings():
    """A (script_type, index) pair MUST be unique - a duplicate is a stg_dup_type_index FAIL (one per
    occurrence). A distinct type OR a distinct index is fine; a blank index is skipped."""
    cols = ["combined_FLD", "script_type", "source_cell", "index"]
    t = signals_table(cols)
    t.add_row({"combined_FLD": "A", "script_type": "KQ", "index": "5", "source_cell": "NET!O11"})
    t.add_row({"combined_FLD": "B", "script_type": "KQ", "index": "5", "source_cell": "NET!O25"})   # dup (KQ,5)
    t.add_row({"combined_FLD": "C", "script_type": "KI", "index": "5", "source_cell": "NET!I3"})    # KI,5 - distinct type
    t.add_row({"combined_FLD": "D", "script_type": "KQ", "index": "6", "source_cell": "NET!O30"})   # KQ,6 - distinct index
    t.add_row({"combined_FLD": "E", "script_type": "KQ", "index": "", "source_cell": "NET!O40"})    # blank index - skipped
    fs = staging._dup_type_index_findings(t)
    eq(len(fs), 2, "one FAIL per occurrence of the (KQ,5) duplicate; distinct/blank rows are clean")
    eq((fs[0].phase, fs[0].type, fs[0].severity), (300, "stg_dup_type_index", "FAIL"), "the finding container")
    eq(sorted(f.location for f in fs), ["NET!O11", "NET!O25"], "each duplicate row located at its own cell")
    ok("KQ index 5" in fs[0].detail and "also at" in fs[0].detail, "loud: names the type/index + the sibling cell")
    eq(staging._dup_type_index_findings(signals_table(cols)), [], "no rows -> no findings")


if __name__ == "__main__":
    import sys
    sys.exit(run("staging", [
        ("identity_flds_and_tag", test_identity_flds_and_tag),
        ("interp_delegates_to_expr_render", test_interp_delegates_to_expr_render),
        ("combined_fld_appends_ce_when_different", test_combined_fld_appends_ce_when_different),
        ("read_view_drops_skip_and_struck", test_read_view_drops_skip_and_struck),
        ("read_view_keeps_struck_when_not_excluding", test_read_view_keeps_struck_when_not_excluding),
        ("is_sorter_area_number_to_name", test_is_sorter_area_number_to_name),
        ("node_address_ranges_positional", test_node_address_ranges_positional),
        ("node_address_range_empty_when_no_addressed_rows", test_node_address_range_empty_when_no_addressed_rows),
        ("dup_signal_uid_findings", test_dup_signal_uid_findings),
        ("dup_type_index_findings", test_dup_type_index_findings),
        ("finalize_identity_ce_overwrite", test_finalize_identity_ce_overwrite),
        ("annotate_cematrix_records_and_restamps", test_annotate_cematrix_records_and_restamps),
        ("load_io_list_no_match_returns_empty", test_load_io_list_no_match_returns_empty),
    ]))
