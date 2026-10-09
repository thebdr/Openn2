"""Phase 400c/400e - the IF_*.xlsx projection + the lossless I/O List insertion (domain.interface_xlsx;
the ZIP/XML-level graft of [[C-034]] - every existing part of the I/O List byte-identical)."""
import os
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table as XlTable
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo, TableFormula

from _harness import run, eq, ok
from pipeline5.systems.plc_based.siemens_s7 import interface_xlsx_writer as interface_xlsx
from pipeline5.phases.interfaces import builder as interfaces
from pipeline5.documents import xlsx_surgical_writer as xlsx_edit


def test_coerce_int_and_safe_name():
    eq(interface_xlsx._coerce_int("10"), 10)
    eq(interface_xlsx._coerce_int("10.0"), 10, "a float-read int string coerces")
    eq(interface_xlsx._coerce_int(""), None)
    eq(interface_xlsx._coerce_int("None"), None, "a stringified None -> None (the WORD bit)")
    eq(interface_xlsx._safe_name('SORTER+DIAG-02'), "SORTER+DIAG-02")
    eq(interface_xlsx._safe_name('A/B:C'), "A_B_C", "filesystem-unsafe chars replaced")


def _synthetic_sheet():
    wb = Workbook()
    ws = wb.active
    headers = ["Category", "Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit",
               "Signal Name Side 1", "Expression Side 1"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    ws.cell(2, 1, "TEMPLATE"); ws.cell(2, 3, ">"); ws.cell(2, 4, 0)     # one template data row
    ws.add_table(XlTable(displayName="T", ref="A1:G2"))
    return wb, ws


def _col(ws, header):
    return {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)}[header]


def test_append_custom_rows_bool_block_and_padding():
    wb, ws = _synthetic_sheet()
    elems = [{"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 0,
              "signal_name": "PNC_a", "expression": '"DB"."a"'},
             {"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 1,
              "signal_name": "PNC_b", "expression": '"DB"."b"'}]
    interface_xlsx._append_custom_rows(ws, elems)
    cat, off, bit, s1, e1 = (_col(ws, "Category"), _col(ws, "I/O Offset Byte"), _col(ws, "I/O Bit"),
                             _col(ws, "Signal Name Side 1"), _col(ws, "Expression Side 1"))
    # the 2 real signals on rows 3-4 (low bits of byte 10), then 14 blank padding rows -> a full 2-byte block
    eq((ws.cell(3, cat).value, ws.cell(3, off).value, ws.cell(3, bit).value), ("X", 10, 0))
    eq((ws.cell(3, s1).value, ws.cell(3, e1).value), ("PNC_a", '"DB"."a"'))
    eq((ws.cell(4, off).value, ws.cell(4, bit).value), (10, 1), "second signal, bit 1")
    eq(ws.cell(5, e1).value in (None, ""), True, "row 5 is blank padding (addressed, no signal)")
    eq((ws.cell(18, cat).value, ws.cell(18, off).value, ws.cell(18, bit).value), ("X", 11, 7),
       "the block runs a full 16 bits (byte 10 + 11), last padding row at 11.7")
    eq(ws.tables["T"].ref, "A1:G18", "the table ref extends to the last written row")


def test_append_custom_rows_skips_template_source():
    # 400f: a source="template" element is already in the copied sheet -> the mirror append must SKIP it.
    wb, ws = _synthetic_sheet()
    elems = [{"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 0,
              "signal_name": "PNC_native", "expression": '"Clock"', "source": "template"},
             {"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 0,
              "signal_name": "PNC_mirror", "expression": '"DB"."m"', "source": "mirror"}]
    interface_xlsx._append_custom_rows(ws, elems)
    s1 = _col(ws, "Signal Name Side 1")
    written = [ws.cell(r, s1).value for r in range(3, ws.max_row + 1) if ws.cell(r, s1).value]
    ok("PNC_mirror" in written, "the mirror element is appended")
    ok("PNC_native" not in written, "the template-native element is NOT re-appended (already in the sheet)")


def test_append_custom_rows_word_row_and_separator():
    wb, ws = _synthetic_sheet()
    elems = [{"category": "SPD", "direction": "Q", "data_type": "WORD", "offset_byte": 12, "bit": None,
              "signal_name": "PNC_speed", "expression": '"DB"."speed"'}]
    interface_xlsx._append_custom_rows(ws, elems)
    dt, off, bit = _col(ws, "Data Type"), _col(ws, "I/O Offset Byte"), _col(ws, "I/O Bit")
    eq((ws.cell(3, dt).value, ws.cell(3, off).value, ws.cell(3, bit).value), ("WORD", 12, None),
       "a WORD is one row, even byte, no bit")
    ok(ws.tables["T"].ref in ("A1:G3", "A1:G4"), "one word row (+ a trailing separator)")


# =================================================================================================== #
# Phase 400e - the lossless IF_ sheet insertion + the Excel-independent address-cache seeding
# =================================================================================================== #
def test_io_address_mirror():
    f = interfaces.io_address_side1
    isy = "I<base+offset>/.<bit>"
    eq(f(isy, 10000, ">", "BOOL", 0, 0), "Q10000.0", "BOOL output -> Q with bit")
    eq(f(isy, 10000, "<", "BOOL", 5, 3), "I10005.3", "BOOL input -> I, base+offset.bit")
    eq(f(isy, 10000, ">", "WORD", 26, None), "Q10026", "WORD -> byte only (TEXTBEFORE '/'), no bit")
    eq(f(isy, 10000, "", "BOOL", 0, 0), "", "no direction -> no address")
    eq(f(isy, None, ">", "BOOL", 0, 0), "", "no base -> no address")
    eq(f(isy, 10000, ">", "BOOL", None, 0), "", "no offset -> no address")
    # the SSOT table stores direction as I/Q (not </>): io_address_side1 accepts both spellings
    eq(f(isy, 10000, "Q", "BOOL", 0, 0), "Q10000.0", "table 'Q' == output '>'")
    eq(f(isy, 10000, "I", "BOOL", 5, 3), "I10005.3", "table 'I' == input '<'")


def test_resolve_num_chain():
    wb = Workbook(); ws = wb.active
    ws["C2"], ws["D2"] = 1, 0
    ws["C3"], ws["D3"] = "=C2", "=D2+1"        # the interface table's two offset/bit chain shapes
    ws["C4"], ws["D4"] = "=C3", "=D3+1"
    ws["C5"], ws["C6"] = "12.0", "x"
    memo = {}
    eq(interface_xlsx._resolve_num(ws, "C4", memo), 1, "=C3 -> =C2 -> 1")
    eq(interface_xlsx._resolve_num(ws, "D4", memo), 2, "=D3+1 -> (=D2+1)+1 -> 2")
    eq(interface_xlsx._resolve_num(ws, "C5", memo), 12, "float-string literal -> int")
    ok(interface_xlsx._resolve_num(ws, "C6", memo) is None, "non-numeric -> None")


def _make_if_addr_sheet(path, base=10000):
    """A minimal inserted-IF_ sheet: the address-relevant headers + an Input Format / Base Address on
    row 2, a (placeholder) LET-formula address cell per data row, and an offset/bit chain (=C3 / =D3+1)."""
    wb = Workbook(); ws = wb.active; ws.title = "IF_T"
    cols = ["Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit", "I/O Address Side 1",
            "Base Address", "Input Format", "Signal Name Side 1"]     # A..H ; offset=C bit=D addr=E
    for i, h in enumerate(cols, start=1):
        ws.cell(1, i, h)
    LET = "=$A$1"                                                     # any formula -> data_type 'f'

    def setrow(r, dt, d, off, bit, sig):
        ws.cell(r, 1, dt); ws.cell(r, 2, d); ws.cell(r, 3, off); ws.cell(r, 4, bit)
        ws.cell(r, 5, LET); ws.cell(r, 8, sig)
    setrow(2, "BOOL", ">", 0, 0, "SIG1")            # Q<base>.0
    setrow(3, "BOOL", ">", 1, 0, "SIG2")            # Q<base+1>.0
    setrow(4, "BOOL", ">", "=C3", "=D3+1", "SIG3")  # CHAIN -> off 1, bit 1 -> Q<base+1>.1
    setrow(5, "WORD", ">", 4, None, "SIG4")         # WORD -> Q<base+4>, no bit
    setrow(6, "BOOL", "<", 0, 0, "SIG5")            # I<base>.0
    ws.cell(7, 5, LET)                              # blank separator: LET present, no direction -> ''
    ws.cell(2, 6, base); ws.cell(2, 7, "I<base+offset>/.<bit>")
    wb.save(path)


def test_interface_address_caches():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(p)
        cache = interface_xlsx._interface_address_caches(p)
        eq(cache.get("E2"), ("str", "Q10000.0"), "BOOL output bit 0")
        eq(cache.get("E3"), ("str", "Q10001.0"), "BOOL output byte+1")
        eq(cache.get("E4"), ("str", "Q10001.1"), "offset/bit CHAIN resolved (=C3 / =D3+1)")
        eq(cache.get("E5"), ("str", "Q10004"), "WORD: byte only, no bit")
        eq(cache.get("E6"), ("str", "I10000.0"), "BOOL input -> I")
        ok("E7" not in cache, "blank separator row (no direction) is not seeded")


def test_insert_seeds_interface_address_cache():
    # the phase-510 bug: the inserted IF_ address LET has no Excel cache, so a data_only reader sees
    # None and skips. insert SEEDS a computed cache that data_only sees (no Excel), keeping the formula.
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(ifp)
        iol = os.path.join(d, "iol.xlsx")
        wb = Workbook(); wb.active.title = "NET SAFETY 50"; wb.active["A1"] = "src"; wb.save(iol)
        acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ok(any("seeded" in a and "IF_SORTER-01" in a for a in acts), "seeding is reported")
        ws = load_workbook(iol, data_only=True)["IF_SORTER-01"]
        eq(ws["E2"].value, "Q10000.0", "data_only reader sees the seeded address (no Excel needed)")
        eq(ws["E4"].value, "Q10001.1", "the chained-offset address is seeded too")
        live = load_workbook(iol)["IF_SORTER-01"]["E2"].value                # formula kept for Excel
        ok(isinstance(live, str) and live.startswith("="), "the LET formula survives for an engineer")


def _make_if_with_table(path):
    """A minimal IF_ sheet whose table carries a source dataDxfId + a calculatedColumnFormula (both must
    be stripped on copy, else Excel drops the table), plus a tiny array-free body."""
    wb = Workbook(); ws = wb.active; ws.title = "SORTER"
    hdr = ["Category", "Description", "Data Type", "Direction </>", "I/O Offset Byte",
           "I/O Bit", "I/O Address Side 1", "Signal Name Side 1"]
    for i, h in enumerate(hdr, start=1):
        ws.cell(1, i, h)
    ws.cell(2, 3, "BOOL"); ws.cell(2, 4, ">"); ws.cell(2, 5, 0); ws.cell(2, 6, 0)
    ws.cell(2, 7, "=DT[[#This Row],[I/O Offset Byte]]")
    cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(hdr)]
    cols[6].dataDxfId = 99                                                    # out-of-range dxf ref
    cols[6].calculatedColumnFormula = TableFormula(attr_text="DT[[#This Row],[I/O Offset Byte]]")
    t = Table(displayName="DT", ref="A1:H2"); t.tableColumns = cols
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def test_insert_lossless_idempotent_and_table_clean():
    import warnings as _w
    _w.simplefilter("ignore")
    from openpyxl.worksheet.formula import ArrayFormula
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_SORTER-01.xlsx"); _make_if_with_table(ifp)
        iol = os.path.join(d, "iolist.xlsx")
        wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
        ws["A1"] = "src"; ws["B1"] = "=A1"          # a real formula (its cache must survive)
        ws["C2"] = ArrayFormula("C2:C4", "=A1")     # an array on an untouched sheet -> must survive AS IS
        wb.save(iol)
        nspart = xlsx_edit._sheet_name_to_part(open(iol, "rb").read())["NET SAFETY 50"]
        with zipfile.ZipFile(iol) as z:
            before = z.read(nspart)

        acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ok(any("inserted" in a and "IF_SORTER-01" in a for a in acts), "sheet inserted")
        wb2 = load_workbook(iol)
        eq(wb2.sheetnames, ["NET SAFETY 50", "IF_SORTER-01"], "IF_ added as the LAST sheet, original kept")
        eq(list(wb2["IF_SORTER-01"].tables), ["DT_IF_SORTER_01"], "table renamed per-instance")
        eq(wb2["IF_SORTER-01"]["G2"].value, "=DT_IF_SORTER_01[[#This Row],[I/O Offset Byte]]",
           "the cell formula follows the renamed table")
        eq(wb2["NET SAFETY 50"]["B1"].value, "=A1", "the existing formula survives (lossless)")
        wb2.close()
        with zipfile.ZipFile(iol) as z:
            tx = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/tables/")]
            eq(z.read(nspart), before, "the untouched sheet - array formula and all - is byte-identical")
            ok('t="array"' in z.read(nspart).decode("utf-8"), "the array formula is still there (never frozen)")
        ok(tx and not any("DxfId" in x for x in tx), "source dataDxfId stripped (else Excel drops the table)")
        ok(not any("calculatedColumnFormula" in x for x in tx), "stale calc-column formula stripped")

        size = os.path.getsize(iol)
        acts2 = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])   # idempotent
        eq(acts2, ["IF_SORTER-01: already present in the I/O List - skipped"], "second run skips the sheet")
        acts3 = interface_xlsx.insert_sheets_into_iolist(iol, [("if_sorter-01", ifp)])
        eq(acts3, ["if_sorter-01: already present in the I/O List - skipped"],
           "the presence check is Excel's - case-insensitive (two such names would need a repair)")
        eq(os.path.getsize(iol), size, "a skip writes nothing")
        eq(load_workbook(iol).sheetnames.count("IF_SORTER-01"), 1, "no duplicate sheet")
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_NEW", ifp), ("IF_SORTER-01", ifp)]),
           ["IF_NEW: inserted into the I/O List", "IF_SORTER-01: already present in the I/O List - skipped"],
           "the actions follow the order the sheets were given")


# =================================================================================================== #
# [[C-034]] - the insert on a MODERN workbook: every existing part byte-identical (the graft)
# =================================================================================================== #
_CT_MAIN = "application/vnd.openxmlformats-officedocument.spreadsheetml."
_NS = ('xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
_PKG = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
_BOOKKEEPING = ("xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/styles.xml")


def _zip(path, parts):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content if isinstance(content, bytes) else content.encode("utf-8"))


def _devmode() -> bytes:
    """A minimal well-formed DEVMODEW (220 bytes: a device name, spec 0x0401, size 220, orientation + paper size
    set - landscape A4). Excel parses the printer-settings part and REFUSES a workbook whose blob is not one
    (checked 2026-10-09), so the fixture carries a real-shaped one."""
    import struct
    return ("Fixture Printer".encode("utf-16-le").ljust(64, b"\0") + struct.pack("<HHHHI", 0x0401, 0, 220, 0, 0x3)
            + struct.pack("<8h", 2, 9, 0, 0, 100, 1, 0, 0) + struct.pack("<5h", 0, 0, 0, 0, 0) + b"\0" * 64
            + struct.pack("<HIII", 0, 0, 0, 0) + struct.pack("<II", 0, 0) + b"\0" * 32)


def _modern_iolist(path, defined_names=(), table="T_Types", crowded=False):
    """An I/O List as a modern Excel writes it - with what openpyxl cannot carry: a DYNAMIC ARRAY (cm +
    xl/metadata.xml), a THREADED comment (threadedComments + persons + its legacy placeholder), an add-in
    binding (webextensions), a featurePropertyBag, printer settings, an EMPTY-TEXT cell, a calcChain, a
    cached formula, a table (T_Types, id 1), no <numFmts> and a self-closing <dxfs/>. Excel 16 opens it
    (checked 2026-10-09: the array spills 1 / 2 / 3, one threaded comment, landscape). `crowded`: the
    styles shaped like FVT's - its own custom number formats (a date cell E2 on 164), more fills /
    borders, two dxfs - so a source index lands on a DIFFERENT target index."""
    dn = "".join(f'<definedName name="{n}">FamilyCheck!$A$1</definedName>' for n in defined_names)
    ms = "http://schemas.microsoft.com/office"
    parts = {
        "[Content_Types].xml": (
            _DECL + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            f'<Default Extension="bin" ContentType="{_CT_MAIN}printerSettings"/>'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/xl/workbook.xml" ContentType="{_CT_MAIN}sheet.main+xml"/>'
            f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{_CT_MAIN}worksheet+xml"/>'
            f'<Override PartName="/xl/worksheets/sheet2.xml" ContentType="{_CT_MAIN}worksheet+xml"/>'
            f'<Override PartName="/xl/styles.xml" ContentType="{_CT_MAIN}styles+xml"/>'
            f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT_MAIN}sharedStrings+xml"/>'
            f'<Override PartName="/xl/metadata.xml" ContentType="{_CT_MAIN}sheetMetadata+xml"/>'
            f'<Override PartName="/xl/comments1.xml" ContentType="{_CT_MAIN}comments+xml"/>'
            f'<Override PartName="/xl/tables/table1.xml" ContentType="{_CT_MAIN}table+xml"/>'
            '<Override PartName="/xl/threadedComments/threadedComment1.xml" '
            'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
            '<Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/>'
            '<Override PartName="/xl/featurePropertyBag/featurePropertyBag.xml" '
            'ContentType="application/vnd.ms-excel.featurepropertybag+xml"/>'
            f'<Override PartName="/xl/calcChain.xml" ContentType="{_CT_MAIN}calcChain+xml"/>'
            '<Override PartName="/xl/webextensions/taskpanes.xml" '
            'ContentType="application/vnd.ms-office.webextensiontaskpanes+xml"/>'
            '<Override PartName="/xl/webextensions/webextension1.xml" '
            'ContentType="application/vnd.ms-office.webextension+xml"/>'
            '<Override PartName="/docProps/core.xml" '
            'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
            '<Override PartName="/docProps/app.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/></Types>'),
        "_rels/.rels": (
            _DECL + _PKG +
            f'<Relationship Id="rId4" Type="{ms}/2011/relationships/webextensiontaskpanes" '
            'Target="xl/webextensions/taskpanes.xml"/>'
            f'<Relationship Id="rId3" Type="{_R}/extended-properties" Target="docProps/app.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/'
            'core-properties" Target="docProps/core.xml"/>'
            f'<Relationship Id="rId1" Type="{_R}/officeDocument" Target="xl/workbook.xml"/></Relationships>'),
        "xl/workbook.xml": (
            _DECL + f'<workbook {_NS}><workbookPr codeName="ThisWorkbook"/><bookViews><workbookView activeTab="0"/>'
            '</bookViews><sheets><sheet name="NET SAFETY 50" sheetId="1" r:id="rId1"/>'
            '<sheet name="FamilyCheck" sheetId="7" r:id="rId2"/></sheets><definedNames>'
            '<definedName name="_xlnm._FilterDatabase" localSheetId="0" hidden="1">'
            f"'NET SAFETY 50'!$A$1:$D$4</definedName>{dn}</definedNames><calcPr calcId=\"191029\"/></workbook>"),
        "xl/_rels/workbook.xml.rels": (
            _DECL + _PKG +
            f'<Relationship Id="rId8" Type="{_R}/sharedStrings" Target="sharedStrings.xml"/>'
            f'<Relationship Id="rId1" Type="{_R}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{_R}/worksheet" Target="worksheets/sheet2.xml"/>'
            f'<Relationship Id="rId7" Type="{_R}/styles" Target="styles.xml"/>'
            f'<Relationship Id="rId12" Type="{_R}/calcChain" Target="calcChain.xml"/>'
            f'<Relationship Id="rId11" Type="{ms}/2022/11/relationships/FeaturePropertyBag" '
            'Target="featurePropertyBag/featurePropertyBag.xml"/>'
            f'<Relationship Id="rId10" Type="{ms}/2017/10/relationships/person" Target="persons/person.xml"/>'
            f'<Relationship Id="rId9" Type="{_R}/sheetMetadata" Target="metadata.xml"/></Relationships>'),
        "xl/worksheets/sheet1.xml": (
            _DECL + f'<worksheet {_NS}><dimension ref="A1:D4"/><sheetViews><sheetView tabSelected="1" '
            'workbookViewId="0"/></sheetViews><sheetFormatPr defaultRowHeight="15"/><sheetData>'
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" s="1" t="s"><v>2</v></c></row>'
            '<row r="2"><c r="A2"><v>10</v></c><c r="B2" cm="1"><f t="array" ref="B2:B4">_xlfn.SEQUENCE(3)</f>'
            '<v>1</v></c><c r="C2" t="s"><v>1</v></c><c r="D2"><f>A2*2</f><v>20</v></c></row>'
            '<row r="3"><c r="B3"><v>2</v></c></row><row r="4"><c r="B4"><v>3</v></c></row></sheetData>'
            '<autoFilter ref="A1:D4"/>'
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
            '<pageSetup paperSize="9" orientation="landscape" r:id="rId1"/><legacyDrawing r:id="rId3"/>'
            '</worksheet>'),
        "xl/worksheets/_rels/sheet1.xml.rels": (
            _DECL + _PKG +
            f'<Relationship Id="rId1" Type="{_R}/printerSettings" Target="../printerSettings/printerSettings1.bin"/>'
            f'<Relationship Id="rId2" Type="{_R}/comments" Target="../comments1.xml"/>'
            f'<Relationship Id="rId3" Type="{_R}/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
            f'<Relationship Id="rId4" Type="{ms}/2017/10/relationships/threadedComment" '
            'Target="../threadedComments/threadedComment1.xml"/></Relationships>'),
        "xl/worksheets/sheet2.xml": (
            _DECL + f'<worksheet {_NS}><dimension ref="A1:A2"/><sheetViews><sheetView workbookViewId="0"/>'
            '</sheetViews><sheetFormatPr defaultRowHeight="15"/><sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>type_id</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>KQ</t></is></c></row></sheetData>'
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
            '<tableParts count="1"><tablePart r:id="rId1"/></tableParts></worksheet>'),
        "xl/worksheets/_rels/sheet2.xml.rels": (
            _DECL + _PKG +
            f'<Relationship Id="rId1" Type="{_R}/table" Target="../tables/table1.xml"/></Relationships>'),
        "xl/tables/table1.xml": (
            _DECL + '<table xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" id="1" '
            f'name="{table}" displayName="{table}" ref="A1:A2" totalsRowShown="0"><autoFilter ref="A1:A2"/>'
            '<tableColumns count="1"><tableColumn id="1" name="type_id"/></tableColumns>'
            '<tableStyleInfo name="TableStyleMedium2" showFirstColumn="0" showLastColumn="0" '
            'showRowStripes="1" showColumnStripes="0"/></table>'),
        "xl/styles.xml": (
            _DECL + '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x14ac" '
            'xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac">'
            '<fonts count="2" x14ac:knownFonts="1"><font><sz val="10"/><name val="Arial"/><family val="2"/></font>'
            '<font><b/><sz val="10"/><name val="Arial"/><family val="2"/></font></fonts>'
            '<fills count="2"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
            '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/></cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles><dxfs count="0"/>'
            '<tableStyles count="0" defaultTableStyle="TableStyleMedium2" defaultPivotStyle="PivotStyleLight16"/>'
            '</styleSheet>'),
        "xl/sharedStrings.xml": (
            _DECL + '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" '
            'uniqueCount="3"><si><t>Tag</t></si><si><t/></si><si><t>Spill</t></si></sst>'),
        "xl/metadata.xml": (
            _DECL + '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:xda="{ms}/spreadsheetml/2017/dynamicarray"><metadataTypes count="1">'
            '<metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" '
            'merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1" '
            'cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst>'
            '<ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}"><xda:dynamicArrayProperties fDynamic="1" '
            'fCollapsed="0"/></ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/>'
            '</bk></cellMetadata></metadata>'),
        "xl/comments1.xml": (
            _DECL + '<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><authors>'
            '<author>tc={11111111-1111-4111-8111-111111111111}</author></authors><commentList>'
            '<comment ref="A1" authorId="0"><text><t>[Threaded comment] check the tag</t></text></comment>'
            '</commentList></comments>'),
        "xl/drawings/vmlDrawing1.vml": (
            '<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:x="urn:schemas-microsoft-com:office:excel"><v:shape id="_x0000_s1025" type="#_x0000_t202">'
            '<x:ClientData ObjectType="Note"><x:Row>0</x:Row><x:Column>0</x:Column></x:ClientData></v:shape></xml>'),
        "xl/threadedComments/threadedComment1.xml": (
            _DECL + f'<ThreadedComments xmlns="{ms}/spreadsheetml/2018/threadedcomments" '
            'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><threadedComment ref="A1" '
            'dT="2026-10-09T10:00:00.00" personId="{22222222-2222-4222-8222-222222222222}" '
            'id="{11111111-1111-4111-8111-111111111111}"><text>check the tag</text></threadedComment>'
            '</ThreadedComments>'),
        "xl/persons/person.xml": (
            _DECL + f'<personList xmlns="{ms}/spreadsheetml/2018/threadedcomments" '
            'xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><person displayName="Engineer" '
            'id="{22222222-2222-4222-8222-222222222222}" userId="engineer" providerId="None"/></personList>'),
        "xl/featurePropertyBag/featurePropertyBag.xml": (
            _DECL + f'<FeaturePropertyBags xmlns="{ms}/spreadsheetml/2022/featurepropertybag">'
            '<bag type="Checkbox"/></FeaturePropertyBags>'),
        "xl/printerSettings/printerSettings1.bin": _devmode(),
        "xl/calcChain.xml": (
            _DECL + '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<c r="B2" i="1" a="1"/><c r="D2"/></calcChain>'),
        "xl/webextensions/taskpanes.xml": (
            _DECL + f'<wetp:taskpanes xmlns:wetp="{ms}/webextensions/taskpanes/2010/11">'
            f'<wetp:taskpane dockstate="right" visibility="0" width="350" row="1"><wetp:webextensionref '
            f'xmlns:r="{_R}" r:id="rId1"/></wetp:taskpane></wetp:taskpanes>'),
        "xl/webextensions/_rels/taskpanes.xml.rels": (
            _DECL + _PKG + f'<Relationship Id="rId1" Type="{ms}/2011/relationships/webextension" '
            'Target="webextension1.xml"/></Relationships>'),
        "xl/webextensions/webextension1.xml": (
            _DECL + f'<we:webextension xmlns:we="{ms}/webextensions/webextension/2010/11" '
            'id="{33333333-3333-4333-8333-333333333333}"><we:reference id="wa000000000" version="1.0.0.0" '
            'store="en-US" storeType="OMEX"/><we:alternateReferences/><we:properties/><we:bindings/>'
            f'<we:snapshot xmlns:r="{_R}"/></we:webextension>'),
        "docProps/core.xml": (
            _DECL + '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/'
            'core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:creator>engineer</dc:creator>'
            '</cp:coreProperties>'),
        "docProps/app.xml": (
            _DECL + '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
            '<Application>Microsoft Excel</Application></Properties>'),
    }
    if crowded:   # FVT's shape: custom number formats 164 / 165 (one in use), more fills / borders, two dxfs
        parts["xl/styles.xml"] = parts["xl/styles.xml"].replace(
            '<fonts count="2"',
            '<numFmts count="2"><numFmt numFmtId="164" formatCode="[$-410]d\-mmm\-yyyy;@"/>'
            '<numFmt numFmtId="165" formatCode="0.0%"/></numFmts><fonts count="2"').replace(
            '<fills count="2">', '<fills count="4">').replace(
            '<fill><patternFill patternType="gray125"/></fill></fills>',
            '<fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor '
            'rgb="FF00FF00"/><bgColor indexed="64"/></patternFill></fill><fill><patternFill patternType="solid">'
            '<fgColor rgb="FFFF0000"/><bgColor indexed="64"/></patternFill></fill></fills>').replace(
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>',
            '<borders count="2"><border><left/><right/><top/><bottom/><diagonal/></border><border><left/><right/>'
            '<top/><bottom style="double"><color rgb="FFFF0000"/></bottom><diagonal/></border></borders>').replace(
            '<cellXfs count="2">', '<cellXfs count="3">').replace(
            '</cellXfs>', '<xf numFmtId="164" fontId="0" fillId="2" borderId="1" xfId="0" applyNumberFormat="1" '
            'applyFill="1" applyBorder="1"/></cellXfs>').replace(
            '<dxfs count="0"/>', '<dxfs count="2"><dxf><font><b/></font></dxf><dxf><fill><patternFill><bgColor '
            'rgb="FF00B0F0"/></patternFill></fill></dxf></dxfs>')
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
            '<c r="D2"><f>A2*2</f><v>20</v></c>', '<c r="D2"><f>A2*2</f><v>20</v></c><c r="E2" s="2"><v>45000</v></c>')
    _zip(path, parts)


def _make_rich_if(path, font="Arial Narrow", hyperlink=False, totals=False):
    """A generated IF_ sheet with what the graft must carry: a styled header (font / solid fill / border /
    alignment), a CUSTOM number format, a conditional format (a dxf), a validation, a merge, a column width,
    a text that starts with '=', a styled column (B) and row (6), the address headers the seeding reads (offset
    D / bit E / address F / base G / format H), a formula naming the table beside a string literal (K3),
    optionally an EXTERNAL hyperlink
    (I2), and a table carrying a source dxf, a calculated-column formula and cell-style names (all three
    point into the source workbook). The sheet is tabSelected (the FVTGENERIC template is) - the graft must
    not keep that."""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws = wb.active; ws.title = "SORTER"
    hdr = ["Category", "Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit", "I/O Address Side 1",
           "Base Address", "Input Format", "Signal Name Side 1"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(1, i, h)
        c.font = Font(name=font, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor="FF7030A0")
        c.border = Border(bottom=Side(style="thick", color="FF000000"))
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    rows = [("SAFE", "BOOL", ">", 0, 0, "PNC_A"), ("SAFE", "BOOL", "<", 1, 2, "=TRIB-CA01"),
            ("SPD", "WORD", ">", 4, None, "PNC_SPEED")]
    for r, (cat, dt, d, off, bit, sig) in enumerate(rows, start=2):
        ws.cell(r, 1, cat); ws.cell(r, 2, dt); ws.cell(r, 3, d); ws.cell(r, 4, off); ws.cell(r, 5, bit)
        ws.cell(r, 6, "=DT[[#This Row],[I/O Offset Byte]]+$G$2")
        ws.cell(r, 9, sig)
    ws["I3"].data_type = "s"                                   # an FLD-like text, never a formula
    ws["G2"] = 10000; ws["G2"].number_format = "#,##0.000"
    ws["H2"] = "I<base+offset>/.<bit>"
    ws["K1"] = "NOTE"; ws.merge_cells("K1:L1")
    ws["K3"] = '=COUNTA(DT[Category])&" DT"'
    if hyperlink:
        ws["I2"].hyperlink = "https://example.com/interfaces"
    ws.column_dimensions["A"].width = 20.5
    ws.column_dimensions["B"].font = Font(name="Consolas", italic=True)        # a styled column  (<col style>)
    ws.row_dimensions[6].fill = PatternFill("solid", fgColor="FFFFFF00")       # a styled row     (<row s>)
    ws.sheet_view.tabSelected = True                          # as FVT's FVTGENERIC sheets carry it
    ws.conditional_formatting.add("A2:A4", CellIsRule(operator="equal", formula=['"SPD"'],
                                                      fill=PatternFill("solid", bgColor="FFFFC7CE")))
    dv = DataValidation(type="list", formula1='"<,>"'); dv.add("C2:C4"); ws.add_data_validation(dv)
    cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(hdr)]
    cols[5].dataDxfId = 99
    cols[5].calculatedColumnFormula = TableFormula(attr_text="DT[[#This Row],[I/O Offset Byte]]+$G$2")
    cols[0].dataCellStyle = "Source Style"
    t = Table(displayName="DT", ref="A1:I5" if totals else "A1:I4"); t.tableColumns = cols
    t.headerRowCellStyle, t.dataCellStyle = "Source Heading", "Source Style"
    if totals:          # the total row SHOWN (Excel: Table Design > Total Row): a label, a count, a custom formula
        from openpyxl.worksheet.filters import AutoFilter
        t.totalsRowCount, t.autoFilter = 1, AutoFilter(ref="A1:I4")
        cols[0].totalsRowLabel, cols[1].totalsRowFunction, cols[3].totalsRowFunction = "Total", "count", "custom"
        cols[3].totalsRowFormula = TableFormula(attr_text="SUM(DT[I/O Offset Byte])")
        ws["A5"], ws["B5"], ws["D5"] = "Total", "=SUBTOTAL(103,DT[Data Type])", "=SUM(DT[I/O Offset Byte])"
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def _parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _only_inserts(a: bytes, b: bytes, collections: bool = False) -> bool:
    """`b` is `a` with text INSERTED only - nothing removed or rewritten. `collections`: a styles part, whose
    collections' `count` attributes change and whose empty collection (`<dxfs count="0"/>`) opens up when
    appended to - both normalised away first (every self-closing tag written open on both sides)."""
    import difflib
    import re as _re
    a, b = a.decode("utf-8"), b.decode("utf-8")
    if collections:
        a, b = ([_re.sub(r"<([\w:.-]+)([^<>]*?)\s*/>", r"<\1\2></\1>",
                         _re.sub(r'(?<![\w:])count="\d+"', 'count="#"', x)) for x in (a, b)])
    ops = difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()
    return all(op[0] in ("equal", "insert") for op in ops)


def test_insert_into_a_modern_workbook_keeps_every_part():
    import warnings as _w
    import xml.etree.ElementTree as ET
    from pipeline5.documents.xlsx_sheet_graft import _ser
    _w.simplefilter("ignore")
    M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_A.xlsx"); _make_rich_if(ifp)
        with zipfile.ZipFile(ifp) as z:
            ok(b'tabSelected="1"' in z.read("xl/worksheets/sheet1.xml"), "precondition: the source is tab-selected")
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        before = _parts(iol)

        acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp), ("IF_SORTER-02", ifp)])
        eq(acts, ["IF_SORTER-01: inserted into the I/O List",
                  "IF_SORTER-01: seeded 3 I/O Address Side 1 values (Excel-independent)",
                  "IF_SORTER-02: inserted into the I/O List",
                  "IF_SORTER-02: seeded 3 I/O Address Side 1 values (Excel-independent)"], "the exact actions")
        after = _parts(iol)
        with zipfile.ZipFile(iol) as z:
            eq([n for n in z.namelist() if n in before], list(before), "every original part kept, in its order")
        for name, data in before.items():
            if name not in _BOOKKEEPING:
                eq(after[name], data, f"{name} byte-identical")
        ok(b'cm="1"' in after["xl/worksheets/sheet1.xml"] and b't="array"' in after["xl/worksheets/sheet1.xml"],
           "the dynamic array is still a dynamic array (metadata.xml, its index, kept too)")
        ok(b'<c r="C2" t="s"><v>1</v></c>' in after["xl/worksheets/sheet1.xml"], "the empty-text cell stays text")
        eq(sorted(set(after) - set(before)),
           ["xl/tables/table2.xml", "xl/tables/table3.xml", "xl/worksheets/_rels/sheet3.xml.rels",
            "xl/worksheets/_rels/sheet4.xml.rels", "xl/worksheets/sheet3.xml", "xl/worksheets/sheet4.xml"],
           "exactly the grafted sheets' parts are added (at free part names)")

        # the four bookkeeping parts: appended to, never rewritten
        for name in ("xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            ok(_only_inserts(before[name], after[name]), f"{name}: insertions only")
        ok(_only_inserts(before["xl/styles.xml"], after["xl/styles.xml"], collections=True),
           "xl/styles.xml: insertions only, as TEXT (its root, namespaces, mc:Ignorable untouched) - counts aside")
        ok(not _only_inserts(b'<dxfs count="0"/><x/>', b'<dxfs count="1"><dxf/></dxfs>', collections=True),
           "(the styles check does see a removal)")
        for tp in ("xl/tables/table2.xml", "xl/tables/table3.xml"):
            for gone in (b"DxfId", b"CellStyle", b"calculatedColumnFormula"):
                ok(gone not in after[tp], f"{tp}: no {gone.decode()} (it points into the source workbook)")
        wbx = after["xl/workbook.xml"].decode("utf-8")
        ok('<sheet name="IF_SORTER-01" sheetId="8" r:id="rId13"/><sheet name="IF_SORTER-02" sheetId="9" '
           'r:id="rId14"/></sheets>' in wbx, "the sheets appended LAST (localSheetId stays valid), fresh ids")
        ok('<calcPr calcId="191029" fullCalcOnLoad="1"/>' in wbx, "Excel computes the uncached formulas on open")
        ctx = after["[Content_Types].xml"].decode("utf-8")
        for part in ("/xl/worksheets/sheet3.xml", "/xl/worksheets/sheet4.xml", "/xl/tables/table2.xml",
                     "/xl/tables/table3.xml"):
            eq(ctx.count(f'PartName="{part}"'), 1, f"one content-type Override for {part}")
        so, sn = ET.fromstring(before["xl/styles.xml"]), ET.fromstring(after["xl/styles.xml"])
        eq([e.tag for e in sn], [M + "numFmts"] + [e.tag for e in so],
           "the styleSheet keeps its elements in order (numFmts created as the first)")
        for eo in so:
            en = sn.find(eo.tag)
            kids_o, kids_n = [_ser(c) for c in eo], [_ser(c) for c in en]
            eq(kids_n[:len(kids_o)], kids_o, f"styles {eo.tag}: every original entry kept in place")
            eq({k: v for k, v in en.attrib.items() if k != "count"}, {k: v for k, v in eo.attrib.items() if k != "count"},
               f"styles {eo.tag}: no attribute but count changes")
            if en.get("count") is not None:
                eq(en.get("count"), str(len(kids_n)), f"styles {eo.tag}: count = its entries")
        eq([(n.get("numFmtId"), n.get("formatCode")) for n in sn.find(M + "numFmts")], [("164", "#,##0.000")],
           "the custom number format, at the first custom id")
        eq(len(sn.find(M + "dxfs")), 1, "the conditional format's dxf appended (the table's dxfs dropped)")
        eq(len(sn.find(M + "cellXfs")), 2 + 4, "the header's, the number format's, the column's and the row's cell "
           "formats appended (the source's default one maps to the I/O List's default)")

        # one more graft of the SAME template: every style already there -> styles.xml unchanged
        ifp2 = os.path.join(d, "IF_B.xlsx"); _make_rich_if(ifp2)
        interface_xlsx.insert_sheets_into_iolist(iol, [("IF_FVTGENERIC-03", ifp2)])
        eq(_parts(iol)["xl/styles.xml"], after["xl/styles.xml"], "an equal style is reused, never re-appended")

        # read back (openpyxl): values, styles, seeded addresses, tables, the sheet facts
        wb = load_workbook(iol)
        eq(wb.sheetnames, ["NET SAFETY 50", "FamilyCheck", "IF_SORTER-01", "IF_SORTER-02", "IF_FVTGENERIC-03"],
           "the order of the sheets")
        ws = wb["IF_SORTER-01"]
        eq([ws.cell(1, c).value for c in (1, 3, 6)], ["Category", "Direction </>", "I/O Address Side 1"], "headers")
        eq((ws["A2"].value, ws["D3"].value, ws["E3"].value, ws["G2"].value), ("SAFE", 1, 2, 10000), "values")
        eq((ws["I3"].value, ws["I3"].data_type), ("=TRIB-CA01", "s"), "a leading '=' stays text")
        eq(ws["F2"].value, "=DT_IF_SORTER_01[[#This Row],[I/O Offset Byte]]+$G$2", "the formula follows its table")
        eq(ws["K3"].value, '=COUNTA(DT_IF_SORTER_01[Category])&" DT"', "the table renamed, the literal kept")
        h = ws["A1"]
        eq((h.font.name, h.font.b, h.font.color.rgb, h.fill.fgColor.rgb, h.border.bottom.style,
            h.alignment.horizontal, h.alignment.wrap_text),
           ("Arial Narrow", True, "FFFFFFFF", "FF7030A0", "thick", "center", True), "the header's style")
        eq(ws["G2"].number_format, "#,##0.000", "the custom number format")
        eq((ws["A2"].font.name, ws["A2"].style_id), ("Arial", 0), "an unstyled cell takes the I/O List's default")
        eq([str(r) for r in ws.merged_cells.ranges], ["K1:L1"], "the merge")
        eq(ws.column_dimensions["A"].width, 20.5, "the column width")
        eq((ws.column_dimensions["B"].font.name, ws.column_dimensions["B"].font.i), ("Consolas", True),
           "the column's style (remapped)")
        eq(ws.row_dimensions[6].fill.fgColor.rgb, "FFFFFF00", "the row's style (remapped)")
        eq([(str(v.sqref), v.formula1) for v in ws.data_validations.dataValidation], [("C2:C4", '"<,>"')],
           "the validation")
        cf = [(str(rng.sqref), r.formula, r.dxf.fill.bgColor.rgb)
              for rng in ws.conditional_formatting for r in rng.rules]
        eq(cf, [("A2:A4", ['"SPD"'], "FFFFC7CE")], "the conditional format and its (remapped) dxf")
        ok(not ws.sheet_view.tabSelected, "the grafted sheet is not tab-selected (one active tab)")
        eq([(t.name, t.ref) for t in ws.tables.values()], [("DT_IF_SORTER_01", "A1:I4")], "the renamed table")
        eq(sorted(t.id for s in wb.worksheets for t in s.tables.values()), [1, 2, 3, 4], "unique table ids")
        wd = load_workbook(iol, data_only=True)
        eq([wd["IF_SORTER-02"][r].value for r in ("F2", "F3", "F4")], ["Q10000.0", "I10001.2", "Q10004"],
           "the seeded I/O Address Side 1 values (data_only, no Excel)")
        eq(wd["IF_SORTER-02"]["F2"].data_type, "s", "seeded as a STRING result")
        import re as _re
        f2 = _re.search(r'<c r="F2"[^>]*>', after["xl/worksheets/sheet4.xml"].decode("utf-8")).group(0)
        ok(' t="str"' in f2, f"the seeded cell is a formula with a string result: {f2}")
        eq((wd["NET SAFETY 50"]["D2"].value, wd["NET SAFETY 50"]["B3"].value), (20, 2), "the I/O List's caches")


def test_graft_remaps_into_a_crowded_styles_part():
    """On a target whose styles are shaped like FVT's (its own custom number formats 164 / 165, more fills,
    borders and dxfs than the source has), every source index lands on a DIFFERENT target index: the inserted
    sheet still shows its own formats, the I/O List's own cells keep theirs (a reused id would restyle them)."""
    import warnings as _w
    import xml.etree.ElementTree as ET
    _w.simplefilter("ignore")
    M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF.xlsx"); _make_rich_if(ifp)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol, crowded=True)
        before = _parts(iol)
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])[0],
           "IF_SORTER-01: inserted into the I/O List", "inserted")
        after = _parts(iol)
        ok(_only_inserts(before["xl/styles.xml"], after["xl/styles.xml"], collections=True), "styles: insertions only")
        sn = ET.fromstring(after["xl/styles.xml"])
        eq([(n.get("numFmtId"), n.get("formatCode")) for n in sn.find(M + "numFmts")],
           [("164", "[$-410]d\\-mmm\\-yyyy;@"), ("165", "0.0%"), ("166", "#,##0.000")],
           "the new number format takes the id AFTER the I/O List's own")
        wb = load_workbook(iol)
        e2 = wb["NET SAFETY 50"]["E2"]
        eq((e2.number_format, e2.fill.fgColor.rgb, e2.border.bottom.style),
           ("[$-410]d\\-mmm\\-yyyy;@", "FF00FF00", "double"), "the I/O List's own date cell keeps its format")
        ws = wb["IF_SORTER-01"]
        h = ws["A1"]
        eq((h.font.name, h.fill.fgColor.rgb, h.border.bottom.style), ("Arial Narrow", "FF7030A0", "thick"),
           "the header's font / fill / border, remapped past the I/O List's own")
        eq(ws["G2"].number_format, "#,##0.000", "the custom number format, remapped to 166")
        cf = [(str(rng.sqref), r.dxf.fill.bgColor.rgb, r.dxf.font) for rng in ws.conditional_formatting for r in rng.rules]
        eq(cf, [("A2:A4", "FFFFC7CE", None)], "the conditional format's own dxf (not the I/O List's bold one)")
        eq((ws.column_dimensions["B"].font.name, ws.row_dimensions[6].fill.fgColor.rgb), ("Consolas", "FFFFFF00"),
           "the column's and the row's styles")


def _sst_source(path):
    """A source written with SHARED strings (a rich-text run, an empty string) - not openpyxl's inline ones."""
    _zip(path, {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="/xl/workbook.xml" ContentType="{_CT_MAIN}sheet.main+xml"/>'
            f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{_CT_MAIN}worksheet+xml"/>'
            f'<Override PartName="/xl/styles.xml" ContentType="{_CT_MAIN}styles+xml"/>'
            f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT_MAIN}sharedStrings+xml"/></Types>'),
        "_rels/.rels": _PKG + f'<Relationship Id="rId1" Type="{_R}/officeDocument" Target="xl/workbook.xml"/>'
                              '</Relationships>',
        "xl/workbook.xml": f'<workbook {_NS}><sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": (
            _PKG + f'<Relationship Id="rId1" Type="{_R}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{_R}/styles" Target="styles.xml"/>'
            f'<Relationship Id="rId3" Type="{_R}/sharedStrings" Target="sharedStrings.xml"/></Relationships>'),
        "xl/worksheets/sheet1.xml": (
            f'<worksheet {_NS}><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>'
            '<c r="C1" t="s"><v>2</v></c><c r="D1"><v>7</v></c></row></sheetData></worksheet>'),
        "xl/styles.xml": (
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="1">'
            '<font><sz val="11"/><name val="Calibri"/></font></fonts><fills count="1"><fill><patternFill/></fill>'
            '</fills><borders count="1"><border/></borders><cellXfs count="1"><xf numFmtId="0" fontId="0" '
            'fillId="0" borderId="0"/></cellXfs></styleSheet>'),
        "xl/sharedStrings.xml": (
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" uniqueCount="3">'
            '<si><t>plain</t></si><si><r><rPr><b/></rPr><t>bold</t></r><r><t xml:space="preserve"> rest</t></r>'
            '</si><si><t/></si></sst>'),
    })


def test_graft_shared_strings_become_inline():
    """The grafted cells become inline strings with their runs intact - the target's sharedStrings.xml is
    never rewritten (its indices would shift under every existing cell)."""
    with tempfile.TemporaryDirectory() as d:
        src = os.path.join(d, "src.xlsx"); _sst_source(src)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        before = _parts(iol)
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_S", src)]), ["IF_S: inserted into the I/O List"],
           "inserted (no address headers -> nothing seeded)")
        after = _parts(iol)
        eq(after["xl/sharedStrings.xml"], before["xl/sharedStrings.xml"], "the target's shared strings untouched")
        eq(after["xl/styles.xml"], before["xl/styles.xml"], "the source's default format appends no style")
        sheet = after["xl/worksheets/sheet3.xml"].decode("utf-8")
        ok('t="s"' not in sheet, "no shared-string reference left in the grafted sheet")
        ok('<c r="B1" t="inlineStr"><is><r><rPr><b/></rPr><t>bold</t></r><r><t xml:space="preserve"> rest</t></r>'
           '</is></c>' in sheet, "the rich-text runs are kept in the inline string")
        ok('<c r="C1" t="inlineStr"><is><t/></is></c>' in sheet, "the empty string stays an (empty) text cell")
        ws = load_workbook(iol)["IF_S"]
        eq([ws["A1"].value, ws["D1"].value], ["plain", 7], "the values read back")
        eq(str(ws["B1"].value), "bold rest", "the rich text reads back whole")


def _variant(src, dst, old: bytes, new: bytes, part="xl/worksheets/sheet1.xml"):
    """`src` with ONE byte-level edit in one part, written to `dst`."""
    p = _parts(src)
    ok(old in p[part], f"fixture: {old!r} in {part}")
    p[part] = p[part].replace(old, new, 1)
    _zip(dst, p)


def _deflate_damaged(src, dst, part="xl/worksheets/sheet1.xml"):
    """`src` with ONE byte of `part`'s compressed stream flipped - the first flip that makes inflating it fail
    (zlib.error: a damaged deflate stream, the case a CRC check never reaches)."""
    import struct
    import zlib
    with open(src, "rb") as fh:
        data = bytearray(fh.read())
    with zipfile.ZipFile(src) as z:
        info = z.getinfo(part)
    n, m = struct.unpack("<HH", bytes(data[info.header_offset + 26:info.header_offset + 30]))
    start = info.header_offset + 30 + n + m
    for k in range(info.compress_size):
        bad = bytearray(data)
        bad[start + k] ^= 0xFF
        try:
            zlib.decompress(bytes(bad[start:start + info.compress_size]), -15)
        except zlib.error:
            with open(dst, "wb") as fh:
                fh.write(bad)
            return
    raise AssertionError("fixture: no single flipped byte broke the deflate stream")


def test_graft_refuses_what_it_cannot_carry():
    """A sheet the graft cannot read or carry faithfully is refused - a WARN, nothing of it left behind (not
    a style it appended before the refusal, not a part name, not an id) - while the other sheets still
    graft; when none grafts, the I/O List is not written at all."""
    import warnings as _w
    from openpyxl.comments import Comment
    from pipeline5.documents import xlsx_sheet_graft as graft
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        good = os.path.join(d, "IF_GOOD.xlsx"); _make_rich_if(good)
        rich = os.path.join(d, "rich.xlsx"); _make_rich_if(rich, font="Comic Sans MS")
        bad = os.path.join(d, "IF_BAD.xlsx"); _make_rich_if(bad, font="Comic Sans MS")
        wb = load_workbook(bad); wb.active["A2"].comment = Comment("a note", "eng"); wb.save(bad)
        cm = os.path.join(d, "IF_CM.xlsx"); _variant(rich, cm, b'<c r="A2"', b'<c r="A2" cm="1"')
        vm = os.path.join(d, "IF_VM.xlsx"); _variant(rich, vm, b'<c r="A3"', b'<c r="A3" vm="1"')
        sty = os.path.join(d, "IF_STY.xlsx"); _variant(rich, sty, b'<c r="A4"', b'<c r="A4" s="99"')
        s0 = os.path.join(d, "s0.xlsx"); _sst_source(s0)
        sst = os.path.join(d, "IF_SST.xlsx"); _variant(s0, sst, b'<c r="C1" t="s"><v>2</v>', b'<c r="C1" t="s"><v>99</v>')
        empty = os.path.join(d, "IF_EMPTY.xlsx"); open(empty, "wb").close()
        typ = os.path.join(d, "IF_TYP.xlsx")                   # a malformed attribute (openpyxl: TypeError)
        _variant(rich, typ, b'fontId="1"', b'fontId="one"', part="xl/styles.xml")
        flip = os.path.join(d, "IF_FLIP.xlsx"); _deflate_damaged(rich, flip)   # zlib.error, not BadZipFile

        # the graft's own guards, exactly (cm / vm / a dangling style come AFTER the header's Comic Sans styles
        # were appended - the snapshot must take them back)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        before = _parts(iol)
        res = graft.graft_sheets(iol, [{"title": t, "source": p} for t, p in
                                       (("IF_CM", cm), ("IF_VM", vm), ("IF_STY", sty), ("IF_SST", sst),
                                        ("IF_TYP", typ), ("IF_FLIP", flip))])
        eq([(r["title"], r["error"]) for r in res[:4]], [
            ("IF_CM", "cell A2 carries cell metadata (a dynamic array / rich value) the graft does not copy"),
            ("IF_VM", "cell A3 carries cell metadata (a dynamic array / rich value) the graft does not copy"),
            ("IF_STY", "a cell names style 99, the source has 5"),
            ("IF_SST", "cell C1 names shared string 99, the source has 3")], "each refusal and its reason")
        eq(res[4]["error"], "ValueError: invalid literal for int() with base 10: 'one'", "a malformed index")
        ok(res[5]["error"].startswith("error: Error -3 while decompressing data"), f"a damaged part: {res[5]}")
        eq(_parts(iol), before, "nothing grafted -> the I/O List is not written")

        # through the 400e entry: every refusal is a WARN (whichever layer catches it), the I/O List untouched
        named = [("IF_BAD", bad), ("IF_CM", cm), ("IF_VM", vm), ("IF_STY", sty), ("IF_SST", sst),
                 ("IF_EMPTY", empty), ("IF_TYP", typ), ("IF_FLIP", flip)]
        acts = interface_xlsx.insert_sheets_into_iolist(iol, named)
        eq([a.split(":", 1)[0] for a in acts], [f"[WARN] {t}" for t, _ in named], "one WARN per refused sheet")
        ok(all(": not inserted - " in a for a in acts), f"each says so: {acts}")
        ok(acts[0].endswith(" part the graft does not copy"), f"a comment's part is refused: {acts[0]}")
        eq(acts[5], "[WARN] IF_EMPTY: not inserted - its workbook cannot be read (BadZipFile: File is not a zip "
                    "file)", "an unreadable IF_ workbook refuses itself - it never aborts the insert")
        ok(acts[6].startswith("[WARN] IF_TYP: not inserted - its workbook cannot be read (TypeError: "), acts[6])
        ok(acts[7].startswith("[WARN] IF_FLIP: not inserted - its workbook cannot be read ("), acts[7])
        eq(_parts(iol), before, "nothing grafted -> the I/O List is not written")

        # mixed, the good sheet LAST and FIRST: the refused ones leave NOTHING - byte for byte the good one alone
        ref = os.path.join(d, "ref.xlsx"); _modern_iolist(ref)
        interface_xlsx.insert_sheets_into_iolist(ref, [("IF_GOOD", good)])
        acts = interface_xlsx.insert_sheets_into_iolist(iol, named + [("IF_GOOD", good)])
        eq(acts[8:], ["IF_GOOD: inserted into the I/O List",
                      "IF_GOOD: seeded 3 I/O Address Side 1 values (Excel-independent)"], "the good one grafted")
        eq(_parts(iol), _parts(ref), "the refused sheets left nothing behind (styles, part names, ids)")
        ok(b"Comic Sans MS" not in _parts(iol)["xl/styles.xml"], "no refused style")
        iol2 = os.path.join(d, "iol2.xlsx"); _modern_iolist(iol2)
        acts = interface_xlsx.insert_sheets_into_iolist(iol2, [("IF_GOOD", good)] + named)
        eq(acts[:2], ["IF_GOOD: inserted into the I/O List",
                      "IF_GOOD: seeded 3 I/O Address Side 1 values (Excel-independent)"], "good first: grafted")
        eq(_parts(iol2), _parts(ref), "a refusal AFTER a graft rolls back only itself")
        iol3 = os.path.join(d, "iol3.xlsx"); _modern_iolist(iol3)
        interface_xlsx.insert_sheets_into_iolist(iol3, [("IF_X-1", cm), ("IF_X_1", good)])
        eq(list(load_workbook(iol3)["IF_X_1"].tables), ["DT_IF_X_1"], "a refused sheet's table name is not taken")


def test_graft_bookkeeping_edges():
    """No calcPr -> one is added (fullCalcOnLoad) after the defined names; one title twice in a run -> the
    second skipped like an already-present sheet (Excel's case-insensitive rule); a title Excel would refuse
    -> that sheet refused, the workbook untouched."""
    import warnings as _w
    from pipeline5.documents import xlsx_sheet_graft as graft
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF.xlsx"); _make_rich_if(ifp)
        base = os.path.join(d, "base.xlsx"); _modern_iolist(base)
        iol = os.path.join(d, "iol.xlsx")
        _variant(base, iol, b'<calcPr calcId="191029"/>', b"", part="xl/workbook.xml")
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_A", ifp), ("if_a", ifp)]),
           ["IF_A: inserted into the I/O List", "IF_A: seeded 3 I/O Address Side 1 values (Excel-independent)",
            "if_a: already present in the I/O List - skipped"], "the second spelling of one title is skipped")
        ok('</definedNames><calcPr fullCalcOnLoad="1"/></workbook>' in _parts(iol)["xl/workbook.xml"].decode(),
           "no calcPr -> one added in its schema place")
        before = _parts(iol)
        res = graft.graft_sheets(iol, [{"title": t, "source": ifp} for t in ("bad/name", "x" * 32, "'quoted'", "",
                                                                             "if_A")])
        eq([r["error"] for r in res], ["'bad/name' is no valid sheet name", f"'{'x' * 32}' is no valid sheet name",
                                       "\"'quoted'\" is no valid sheet name", "'' is no valid sheet name",
                                       "a sheet of that name is already in the workbook"], "each title refused")
        eq(_parts(iol), before, "nothing grafted -> nothing written")


def test_graft_table_names_unique_and_formula_rename():
    """The grafted table takes `<name>_<title>`, unique against the workbook's tables, its defined names and
    the tables grafted before it in the same run (`_2` on a clash); the formulas follow - a whole NAME only:
    never in a string literal, a quoted sheet name or a structured reference's [column] specifiers."""
    import warnings as _w
    from pipeline5.documents import xlsx_sheet_graft as graft
    rx = graft._table_renamer({"DT": "DT_X", "Data": "Data_X", "Side": "Side_X"})
    for formula, want in (
            ('DT[[#This Row],[A]]&"DT"&dt[B]&DT2[C]&SUM(DT)', 'DT_X[[#This Row],[A]]&"DT"&DT_X[B]&DT2[C]&SUM(DT_X)'),
            ("Data[@[Data Type]]", "Data_X[@[Data Type]]"),
            ("Side[[#This Row],[I/O Address Side 1]]", "Side_X[[#This Row],[I/O Address Side 1]]"),
            ("'Data'!A1+Data[x]", "'Data'!A1+Data_X[x]"),
            ("Data[[#Headers],[a'[b']]]&Data", "Data_X[[#Headers],[a'[b']]]&Data_X"),
            ('"say ""DT"" "&DT[A]', '"say ""DT"" "&DT_X[A]'),
            ("X.DT+DT.Y+[1]Sheet!DT", "X.DT+DT.Y+[1]Sheet!DT_X"),
            ("Data!A1+SUM(Data!A1:A3)+Data[x]", "Data!A1+SUM(Data!A1:A3)+Data_X[x]")):
        eq(graft._rename_in_formula(formula, rx), want, formula)
    eq(graft._table_renamer({}), None, "nothing to rename -> no pattern")
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF.xlsx"); _make_rich_if(ifp)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol, defined_names=("DT_IF_SORTER_01",))
        interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ws = load_workbook(iol)["IF_SORTER-01"]
        eq(list(ws.tables), ["DT_IF_SORTER_01_2"], "a defined name of that spelling -> _2")
        eq(ws["F3"].value, "=DT_IF_SORTER_01_2[[#This Row],[I/O Offset Byte]]+$G$2", "the formula follows")

        iol2 = os.path.join(d, "iol2.xlsx"); _modern_iolist(iol2, table="DT_IF_SORTER_01")
        interface_xlsx.insert_sheets_into_iolist(iol2, [("IF_SORTER-01", ifp)])
        eq(list(load_workbook(iol2)["IF_SORTER-01"].tables), ["DT_IF_SORTER_01_2"],
           "an existing TABLE of that spelling (a renamed IF_ sheet keeps its tables) -> _2")

        iol3 = os.path.join(d, "iol3.xlsx"); _modern_iolist(iol3)
        interface_xlsx.insert_sheets_into_iolist(iol3, [("IF_A-1", ifp), ("IF_A_1", ifp)])
        wb = load_workbook(iol3)
        eq((list(wb["IF_A-1"].tables), list(wb["IF_A_1"].tables)), (["DT_IF_A_1"], ["DT_IF_A_1_2"]),
           "two titles that clean up alike in one run -> the second _2")
        eq(wb["IF_A_1"]["F2"].value, "=DT_IF_A_1_2[[#This Row],[I/O Offset Byte]]+$G$2", "each follows its own")


def test_graft_total_row_and_workbook_names():
    """A table whose total row is SHOWN keeps it - each column's function, label and custom formula (its table
    renamed): without them Excel repairs the whole I/O List (refute round 3, measured in Excel). A sheet whose
    formulas use a name defined in its OWN workbook is refused (the name would not resolve in the I/O List, or
    resolve to the I/O List's own); an unused name, and Excel's own `_xl...` names, are no reason."""
    import warnings as _w
    from openpyxl.workbook.defined_name import DefinedName
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        tot = os.path.join(d, "IF_TOT.xlsx"); _make_rich_if(tot, totals=True)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_TOT", tot)])[0], "IF_TOT: inserted into the I/O List",
           "a total-row table is inserted")
        tx = _parts(iol)["xl/tables/table2.xml"].decode("utf-8")
        for piece in ('ref="A1:I5" headerRowCount="1" totalsRowCount="1"', '<autoFilter ref="A1:I4"/>',
                      '<tableColumn id="1" name="Category" totalsRowLabel="Total"/>',
                      '<tableColumn id="2" name="Data Type" totalsRowFunction="count"/>',
                      '<tableColumn id="4" name="I/O Offset Byte" totalsRowFunction="custom"><totalsRowFormula>'
                      'SUM(DT_IF_TOT[I/O Offset Byte])</totalsRowFormula></tableColumn>'):
            ok(piece in tx, f"the total row kept: {piece}")
        ws = load_workbook(iol)["IF_TOT"]
        eq((ws.tables["DT_IF_TOT"].totalsRowCount, ws["A5"].value, ws["B5"].value),
           (1, "Total", "=SUBTOTAL(103,DT_IF_TOT[Data Type])"), "the total row's cells follow the table")

        def named(path, formula=None):
            _make_rich_if(path)
            wb = load_workbook(path)
            wb.defined_names["SPARE_BYTES"] = DefinedName("SPARE_BYTES", attr_text="4")
            if formula:
                wb.active["L3"] = formula
            wb.save(path)
        used = os.path.join(d, "IF_USED.xlsx"); named(used, "=SPARE_BYTES*8")
        lit = os.path.join(d, "IF_LIT.xlsx"); named(lit, '="SPARE_BYTES"&COUNTA(DT[Category])')
        unused = os.path.join(d, "IF_UNUSED.xlsx"); named(unused)
        xl0 = os.path.join(d, "xl0.xlsx"); _make_rich_if(xl0)
        wb = load_workbook(xl0); wb.active["L4"] = "=_xlfn.SINGLE(A2)"; wb.save(xl0)
        xlfn = os.path.join(d, "IF_XLFN.xlsx")
        old = b"<definedNames />" if b"<definedNames />" in _parts(xl0)["xl/workbook.xml"] else b"<definedNames/>"
        _variant(xl0, xlfn, old, b'<definedNames><definedName name="_xlfn.SINGLE" hidden="1">#NAME?</definedName>'
                                 b'</definedNames>', part="xl/workbook.xml")
        iol2 = os.path.join(d, "iol2.xlsx"); _modern_iolist(iol2)
        acts = interface_xlsx.insert_sheets_into_iolist(iol2, [("IF_USED", used), ("IF_LIT", lit),
                                                               ("IF_UNUSED", unused), ("IF_XLFN", xlfn)])
        eq(acts[0], "[WARN] IF_USED: not inserted - its formulas use 'SPARE_BYTES', defined in its own workbook - "
                    "the graft does not copy workbook names", "a used workbook name refuses the sheet")
        eq([a for a in acts[1:] if "inserted" in a], ["IF_LIT: inserted into the I/O List",
                                                     "IF_UNUSED: inserted into the I/O List",
                                                     "IF_XLFN: inserted into the I/O List"],
           "a name only in a literal, an unused name, Excel's own _xlfn. name: inserted")


def test_graft_carries_external_hyperlinks():
    import re as _re
    import warnings as _w
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF.xlsx"); _make_rich_if(ifp, hyperlink=True)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        eq(interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])[0],
           "IF_SORTER-01: inserted into the I/O List", "a sheet with an external hyperlink is inserted")
        p = _parts(iol)
        rid = _re.search(r'<hyperlink\b[^>]*\br:id="([^"]+)"', p["xl/worksheets/sheet3.xml"].decode("utf-8")).group(1)
        rel = _re.search(r'<Relationship Id="' + rid + r'"[^>]*/>', p["xl/worksheets/_rels/sheet3.xml.rels"].decode())
        ok(rel and 'TargetMode="External"' in rel.group(0) and 'Target="https://example.com/interfaces"' in rel.group(0),
           f"the hyperlink's relationship comes along under its own id ({rel and rel.group(0)})")
        eq(load_workbook(iol)["IF_SORTER-01"]["I2"].hyperlink.target, "https://example.com/interfaces", "read back")


def test_insert_is_one_atomic_write():
    """The grafted workbook is written to a temp file and swapped in: a swap that fails (the I/O List open in
    Excel) leaves the I/O List byte-identical and no temp file behind."""
    import warnings as _w
    from pipeline5.documents import xlsx_sheet_graft as graft
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF.xlsx"); _make_rich_if(ifp)
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        with open(iol, "rb") as fh:
            original = fh.read()

        def locked(src, dst):
            raise PermissionError("the I/O List is open in Excel")
        real, graft.os.replace = graft.os.replace, locked
        try:
            acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        finally:
            graft.os.replace = real
        eq(acts, ["[WARN] save failed (the I/O List is open in Excel) - I/O List left unchanged (restore from the "
                  ".bak if needed)"], "the failure is reported as a WARN")
        with open(iol, "rb") as fh:
            eq(fh.read(), original, "the I/O List is byte-identical")
        eq(sorted(os.listdir(d)), ["IF.xlsx", "iol.xlsx"], "no temp file left behind")


def test_insert_interface_sheets_backs_up_first():
    import warnings as _w
    _w.simplefilter("ignore")
    with tempfile.TemporaryDirectory() as d:
        os.makedirs(os.path.join(d, "out"))
        _make_rich_if(os.path.join(d, "out", "IF_SORTER-01.xlsx"))
        iol = os.path.join(d, "iol.xlsx"); _modern_iolist(iol)
        with open(iol, "rb") as fh:
            original = fh.read()
        db = {"interfaces": [{"instance": "SORTER-01"}, {"instance": "SORTER-02"}]}   # 02 never projected
        acts = interface_xlsx.insert_interface_sheets(db, iolist_path=iol, out_dir=os.path.join(d, "out"))
        ok(acts[0].startswith("backed up the I/O List -> iol.xlsx.bak_") and acts[0].endswith(".xlsx"), acts[0])
        eq(acts[1:], ["IF_SORTER-01: inserted into the I/O List",
                      "IF_SORTER-01: seeded 3 I/O Address Side 1 values (Excel-independent)"],
           "only the projected one")
        baks = [f for f in os.listdir(d) if f.startswith("iol.xlsx.bak_")]
        eq(len(baks), 1, "one timestamped backup")
        with open(os.path.join(d, baks[0]), "rb") as fh:
            eq(fh.read(), original, "the backup is the I/O List as it was")
        eq(sorted(os.listdir(d)), sorted(["out", "iol.xlsx"] + baks), "no temp file left behind")


if __name__ == "__main__":
    import sys
    sys.exit(run("interface_xlsx", [
        ("coerce_int_and_safe_name", test_coerce_int_and_safe_name),
        ("append_custom_rows_bool_block_and_padding", test_append_custom_rows_bool_block_and_padding),
        ("append_custom_rows_skips_template_source", test_append_custom_rows_skips_template_source),
        ("append_custom_rows_word_row_and_separator", test_append_custom_rows_word_row_and_separator),
        ("io_address_mirror", test_io_address_mirror),
        ("resolve_num_chain", test_resolve_num_chain),
        ("interface_address_caches", test_interface_address_caches),
        ("insert_seeds_interface_address_cache", test_insert_seeds_interface_address_cache),
        ("insert_lossless_idempotent_and_table_clean", test_insert_lossless_idempotent_and_table_clean),
        ("insert_into_a_modern_workbook_keeps_every_part", test_insert_into_a_modern_workbook_keeps_every_part),
        ("graft_remaps_into_a_crowded_styles_part", test_graft_remaps_into_a_crowded_styles_part),
        ("graft_shared_strings_become_inline", test_graft_shared_strings_become_inline),
        ("graft_refuses_what_it_cannot_carry", test_graft_refuses_what_it_cannot_carry),
        ("graft_bookkeeping_edges", test_graft_bookkeeping_edges),
        ("graft_table_names_unique_and_formula_rename", test_graft_table_names_unique_and_formula_rename),
        ("graft_total_row_and_workbook_names", test_graft_total_row_and_workbook_names),
        ("graft_carries_external_hyperlinks", test_graft_carries_external_hyperlinks),
        ("insert_is_one_atomic_write", test_insert_is_one_atomic_write),
        ("insert_interface_sheets_backs_up_first", test_insert_interface_sheets_backs_up_first),
    ]))
