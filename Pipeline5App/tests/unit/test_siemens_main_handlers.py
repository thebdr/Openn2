"""The Siemens phase handlers (safety/main.py), run HEADLESS against the real fixture documents -
the phase buttons without Tk.

This is the step-5 payoff of the PhaseContext seam: every handler reports only through
`ctx.emit/status/gate/render/records/halt`, so the whole Run-all chain executes here with the
App's REAL gate semantics (treatments applied, halt on effective severity) - exactly the dispatch
the App performs, minus the window. It catches the bug class the byte-parity oracle CANNOT see:
the oracle drives the chain functions directly with hand-passed arguments, while THIS runs the
handlers' own wiring (the 900 handler shipped for a whole step calling `engine.build(database)`
without its system - a guaranteed click-time crash the oracle never noticed).

Sandboxed like the oracle (module-attr overrides for the Database/Output roots), the treatments
registry copied aside (run_validation reconciles it - the builtin file must not be rewritten by a
test), and - also like the oracle - the document-MUTATING legs stay off: the 200 fill / 245 risky
handlers are not in the run_plan, and the sandbox forces `insert_interface_sheets` OFF (the builtin
fixture config ships it ON for real operator runs; a test must not edit the repo's fixture doc).
The source document's mtime is asserted unchanged after the whole chain.
"""
import os
import shutil
import tempfile

from _harness import run, eq, ok
from pipeline5 import config
from pipeline5.config import paths as p5paths
from pipeline5.findings import treatments
from pipeline5.systems.plc_based.siemens_s7.safety import main
from pipeline5.systems.plc_based.siemens_s7.safety.system import SYSTEM
from pipeline5.systems.system_contract import PhaseContext


class _Host:
    """The App's seam wiring, as collectors: what `App._ctx()` hands a handler, minus Tk/queue.
    gate/render mirror App._gate/_render - the treatments registry is APPLIED and the halt flag
    follows the EFFECTIVE severity (the severity contract), not the raw one."""

    def __init__(self, run: str = ""):
        self.run = run             # the generation id (the App mints one per button press, contract v1 `run`)
        self.lines = []            # every (level, message) a handler emitted
        self.rendered = []         # every findings/records batch handed to gate/render/records
        self.halted = False

    def ctx(self):
        return PhaseContext(
            system=SYSTEM,
            emit=lambda level, msg: self.lines.append((level, msg)),
            status=lambda _text: None,
            gate=self._gate,
            render=self._render,
            records=lambda recs: self.rendered.append(list(recs)),
            halt=self._halt,
            lang="en",
            run=self.run,
        )

    def _apply(self, findings):
        self.rendered.append(list(findings))
        return treatments.apply(findings, treatments.load())

    def _gate(self, findings, label=""):
        if treatments.should_halt(self._apply(findings)):
            self.halted = True
            return False
        return True

    def _render(self, findings, label=""):
        if treatments.should_halt(self._apply(findings)):
            self.halted = True

    def _halt(self):
        self.halted = True

    def levels(self):
        return [level for level, _ in self.lines]


def _sandboxed(fn):
    """The oracle's isolation, harness-side: Database/Output into temp dirs (module-attr overrides),
    the treatments registry copied into the sandbox (run_validation's reconcile must not rewrite the
    builtin file), and `insert_interface_sheets` forced OFF (the one doc-writing leg inside a
    run_plan handler - the oracle skips it too; the repo's fixture doc stays untouched)."""
    def wrapped():
        orig_db, orig_out = p5paths._BUILTIN_DATABASE, p5paths._BUILTIN_OUTPUT
        orig_registry = treatments.registry_path
        orig_load_params = config.load_params
        with tempfile.TemporaryDirectory() as db_dir, tempfile.TemporaryDirectory() as out_dir:
            registry_copy = os.path.join(db_dir, "finding_treatments.csv")
            if os.path.isfile(orig_registry()):
                shutil.copy2(orig_registry(), registry_copy)

            def load_params_no_insert(path=None):
                params = orig_load_params(path)
                section = params.get("iolist_params")
                if isinstance(section, dict):
                    section["insert_interface_sheets"] = False
                return params

            p5paths._BUILTIN_DATABASE, p5paths._BUILTIN_OUTPUT = db_dir, out_dir
            treatments.registry_path = lambda: registry_copy
            config.load_params = load_params_no_insert
            try:
                fn()
            finally:
                p5paths._BUILTIN_DATABASE, p5paths._BUILTIN_OUTPUT = orig_db, orig_out
                treatments.registry_path = orig_registry
                config.load_params = orig_load_params
    return wrapped


def test_run_plan_executes_end_to_end():
    """Every run_plan phase through HANDLERS[phase.handler](ctx), each with a fresh host (handlers
    are self-contained) - exactly the dispatch the App performs. Each handler must complete (a
    crash = red), report a RSLT line (the success marker - a silent no-op handler is a bug), and
    NEVER halt on the clean fixture.

    RE-PINNED 2026-07-17 (the fire-alarm ruling): the two historic `iotag_duplicate` FAILs were a
    CONFIG defect - the F1/2+F2/2 rows shared one tag rule over one shared device - not a fixture
    defect; the rule now disambiguates with {$script_type}, so phase 500's 510 leg WRITES PLCTags
    (asserted below). A halt anywhere in the chain is a regression."""
    doc_path = config.load_params()["iolist_path"]
    doc_mtime = os.path.getmtime(doc_path)
    for number in SYSTEM.phases.run_order():
        phase = SYSTEM.phases.by_number(number)
        host = _Host()
        SYSTEM.handlers[phase.handler](host.ctx())
        ok("RSLT" in host.levels(), f"phase {number} ({phase.handler}) emitted a RSLT line "
                                    f"(got {host.levels()})")
        eq(host.halted, False, f"phase {number} must not halt on the clean fixture")
        if number == 500:
            ok(any("510:" in msg for lvl, msg in host.lines if lvl == "RSLT"),
               "the 510 leg reports its written tag surface (PLCTags is a real output again)")
    eq(os.path.getmtime(doc_path), doc_mtime,
       "the source I/O List was NOT modified by the run_plan chain")


def test_staging_fires_reaction_hooks():
    """The C-024 run-plan wiring: run_staging fires `before_300` (no database yet) then `after_300`
    (the freshly staged database) through the REAL rule loaders - deleting or typo'ing a hook call
    in safety/main.py dies HERE (the parity oracle drives staging directly and cannot see it).
    With the shipped EMPTY rules the fire is a no-op, so the staged output stays untouched."""
    from pipeline5.phases.chain_reactions import engine as reactions
    calls = []
    real_fire = reactions.fire

    seen = {}

    def recording_fire(hook, database, **kw):
        calls.append((hook, database is not None and "signals" in database))
        seen[hook] = None if database is None else {name: database[name].columns for name in database.names()}
        return real_fire(hook, database, **kw)

    host = _Host()
    reactions.fire = recording_fire
    try:
        SYSTEM.handlers["staging"](host.ctx())
    finally:
        reactions.fire = real_fire
    eq(calls, [("before_300", False), ("after_300", True)],
       "before fires with no database, after with the staged one - in that order")
    # P-012: the DECLARED hooks (System.reaction_hooks - what the template builder previews against)
    # are the fired ones, in order, each loading exactly the tables (+ declared columns) it saw
    eq(list(SYSTEM.reaction_hooks), [hook for hook, _db in calls], "the declared hooks = the fired ones")
    for hook, loader in SYSTEM.reaction_hooks.items():
        loaded = None if loader is None else loader(SYSTEM)
        eq(None if loaded is None else {name: loaded[name].columns for name in loaded.names()}, seen[hook],
           f"{hook}: the declared Database = the one the fire saw")
    ok("RSLT" in host.levels(), "staging completed with the hooks live")
    eq(host.halted, False, "the dark engine never perturbs the run")


def _reaction_project(project, params):
    """A PROJECT (tier 1 of the 4-tier walk) carrying its own chain-reaction rule pair: a before_300
    header, a before_300 add_rows that can only fail (no database yet), an after_300 line drilling
    the real params, an after_300 rule on a typo'd table, a rule on a hook the run-plan never fires,
    an after_300 rule whose condition does not compile (a decimal slice - it used to crash every
    hook), and an after_300 loop over the STAGED signals reading a declared column staging leaves
    unfilled (the loop-schema check used to reject it)."""
    import csv
    from ruamel.yaml import YAML
    shared = os.path.join(project, "config_project", "shared")
    rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
    os.makedirs(shared)
    os.makedirs(rx_dir)
    with open(os.path.join(shared, "project_params.yaml"), "w", encoding="utf-8") as handle:
        YAML(typ="safe").dump(params, handle)
    rules = [
        ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
        ["probe_hdr", "before_300", "", "", "file", "rx/probe.txt", "hdr_txt", "a header before staging"],
        ["probe_early", "before_300", "", "", "add_rows", "signals", "rows_tpl", "no database exists yet"],
        ["probe_file", "after_300", "", "", "file", "rx/probe.txt", "probe_txt", "one line per run"],
        ["probe_bad", "after_300", "no_such_table", "", "file", "rx/bad.txt", "probe_txt", "a typo'd table"],
        ["probe_cold", "after_520", "", "", "file", "rx/cold.txt", "probe_txt", "a hook nobody fires"],
        ["probe_mal", "after_300", "signals", 'extract($mnemonic, /(\\d+)/, 1.3) = "1"', "file", "rx/mal.txt",
         "probe_txt", "1.3 typed for the slice 1:3"],
        ["probe_cols", "after_300", "", "", "file", "rx/cols.txt", "cols_txt", "a declared, unfilled column"],
        ["probe_top", "after_300", "signals", "", "file", "rx/top.txt", "top_txt", "the same, as a MATCHED row"],
        ["probe_count", "after_300", "", "", "file", "rx/count.txt", "count_txt", "a predicate in a strict hole"],
    ]
    with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rules)
    with open(os.path.join(rx_dir, "templates.yaml"), "w", encoding="utf-8") as handle:
        handle.write("hdr_txt: |-\n  HEADER {$_rule.hook}\n"
                     "probe_txt: |-\n  [{$_params.project_code}] {$_rule.hook} {$_rule.name}\n"
                     "cols_txt: |-\n  @for $r in signals: [{$r.name_in_db}]\n"
                     "top_txt: |-\n  [{$name_in_db}]\n"
                     "count_txt: |-\n  PEC {count(signals, $script_type = \"PEC\")}\n"
                     "rows_tpl:\n  - label: \"x\"\n")


def _read_csv(path):
    import csv
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _rx_rendered(host):
    return {(f.type, f.location) for batch in host.rendered for f in batch
            if getattr(f, "type", "").startswith("rx_")}


def test_configured_rules_fire_through_the_real_run_plan():
    """E1 + E3 (refuter round 6): a PROJECT carrying its own rule pair (tier 1 of the 4-tier walk)
    through the REAL run_staging - the real loaders, the real params / templates / output-root
    defaults, the before_300 Deferred settled into the staged record, the run-plan's declared hooks,
    and every failure RENDERED. The unit tests inject all of these by hand; the empty shipped CSV
    makes the other smoke a no-op - so dropping a loader, a default, the settle, the hooks, or a
    ctx.render call dies only HERE."""
    params = config.load_params()                   # the builtin fixture params (doc paths absolute)
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _reaction_project(project, params)
        host = _Host()
        config.use_project(project)
        try:
            SYSTEM.handlers["staging"](host.ctx())
            out_root, db_dir = config.output_root(), config.database_dir()
            ok("RSLT" in host.levels(), "staging completed with configured rules live")
            eq(host.halted, False, "ERRR reactions never halt the run")
            with open(os.path.join(out_root, "rx", "probe.txt"), encoding="utf-8") as handle:
                eq(handle.read(), f"HEADER before_300\n[{params['project_code']}] after_300 probe_file\n",
                   "the real loaders fired both hooks: the before_300 header, then the after_300 line "
                   "drilling the REAL project params - written under the REAL output root")
            ok(not os.path.exists(os.path.join(out_root, "rx", "bad.txt")), "the failing rule wrote nothing")
            ok(not os.path.exists(os.path.join(out_root, "rx", "cold.txt")), "the unfired-hook rule never ran")
            ok(not os.path.exists(os.path.join(out_root, "rx", "mal.txt")), "the uncompilable rule never ran")
            staged = _read_csv(os.path.join(db_dir, "signals.csv"))
            with open(os.path.join(out_root, "rx", "cols.txt"), encoding="utf-8") as handle:
                cols = handle.read().splitlines()
            ok(len(staged) > 0 and len(cols) == len(staged) and all(c[:1] == "[" and c[-1:] == "]" for c in cols),
               f"one line per STAGED signal reading the declared-but-unfilled column ({len(cols)} vs {len(staged)})")
            with open(os.path.join(out_root, "rx", "top.txt"), encoding="utf-8") as handle:
                top = handle.read().splitlines()
            eq(len(top), len(staged), "…and as a MATCHED row's top-level field: one line per staged signal")
            pec = sum(1 for row in staged if row.get("script_type") == "PEC")
            with open(os.path.join(out_root, "rx", "count.txt"), encoding="utf-8") as handle:
                eq(handle.read(), f"PEC {float(pec)}\n", "a data-function predicate in a strict hole counts the staged rows")
            log = sorted((r["hook"], r["rule"], r["outcome"])
                         for r in _read_csv(os.path.join(db_dir, "chain_reactions_log.csv")))
            eq(log, [("after_300", "probe_bad", "rx_unknown_table"), ("after_300", "probe_cols", "ok"),
                     ("after_300", "probe_count", "ok"), ("after_300", "probe_file", "ok"),
                     ("after_300", "probe_top", "ok"),
                     ("before_300", "probe_early", "rx_unknown_table"), ("before_300", "probe_hdr", "ok")],
               "every firing audited - the database-less before_300 ones SETTLED into the staged record")
            recorded = [(r["type"], r["location"])
                        for r in _read_csv(os.path.join(db_dir, "validation_issues.csv"))
                        if r["type"].startswith("rx_")]
            eq(sorted(recorded), [("rx_bad_condition", "probe_mal"), ("rx_unfired_hook", "probe_cold"),
                                  ("rx_unknown_table", "probe_bad"), ("rx_unknown_table", "probe_early")],
               "every reaction finding recorded ONCE (the index-wide ones surfaced at both hooks)")
        finally:
            config.use_project(previous_project)
    eq(_rx_rendered(host), {("rx_unknown_table", "probe_early"), ("rx_unknown_table", "probe_bad"),
                            ("rx_unfired_hook", "probe_cold"), ("rx_bad_condition", "probe_mal")},
       "each finding was RENDERED through ctx.render - never silent")
    eq(sum(1 for batch in host.rendered for f in batch if getattr(f, "type", "") == "rx_unfired_hook"), 2,
       "the unfired-hook finding rendered at BOTH fires - each passes the run-plan's hooks (refuter round 11)")


def test_halted_staging_lists_the_deferred_trail_and_records_nothing():
    """Round-6 implications: when staging HALTS (a blocking FAIL - the dup type/index case saves its
    database, the no-I/O-sheet case writes NOTHING), the before_300 Deferred must not be settled
    into a half-committed record - and must not vanish either: its findings were rendered, and its
    firings are listed in the log (INFO, never hidden). after_300 never fires."""
    from pipeline5.findings.finding import Finding
    from pipeline5.phases.staging import iolist
    params = config.load_params()
    previous_project = config.active_project()
    real_stage = iolist.stage

    def halting_stage(*args, **kwargs):
        database, findings = real_stage(*args, **kwargs)
        return database, findings + [Finding(phase=300, type="probe_forced_fail", severity="FAIL",
                                             detail="forced halt", location="probe")]
    with tempfile.TemporaryDirectory() as project:
        _reaction_project(project, params)
        host = _Host()
        config.use_project(project)
        iolist.stage = halting_stage
        try:
            SYSTEM.handlers["staging"](host.ctx())
            out_root, db_dir = config.output_root(), config.database_dir()
            eq(host.halted, True, "the forced FAIL halted staging")
            with open(os.path.join(out_root, "rx", "probe.txt"), encoding="utf-8") as handle:
                eq(handle.read(), "HEADER before_300\n",
                   "before_300 ran (its append-only side effect stays); after_300 never fired")
            eq(_read_csv(os.path.join(db_dir, "chain_reactions_log.csv")), [],
               "no reaction audit was committed into the halted run's record")
            ok(not any(r["type"].startswith("rx_")
                       for r in _read_csv(os.path.join(db_dir, "validation_issues.csv"))),
               "...and no reaction finding either")
        finally:
            iolist.stage = real_stage
            config.use_project(previous_project)
    listed = [msg for level, msg in host.lines if level == "INFO" and "not recorded, staging halted" in msg]
    eq(len(listed), 2, "both deferred before_300 firings are LISTED in the log")
    ok(any("probe_hdr: ok (1 created)" in msg for msg in listed), "...the header's success")
    ok(any("probe_early: rx_unknown_table" in msg for msg in listed), "...and the failure's outcome")
    eq(_rx_rendered(host), {("rx_unknown_table", "probe_early"), ("rx_unfired_hook", "probe_cold"),
                            ("rx_bad_condition", "probe_mal")},
       "before_300 - the ONLY fire of a halted run - rendered its own finding, the unfired hook (it "
       "passes the run-plan's hooks), and after_300's malformed rule (another hook's rule still "
       "surfaces): refuter round 7 E1/E2")


def test_a_corrupt_reaction_record_is_reported_not_crashed_or_overwritten():
    """Refuter round 8 R4 through the real run-plan: a hand-edited chain_reactions_log.csv gone ragged
    used to crash run_staging (after staging had saved and before_300 had appended) - no RSLT, the
    trail lost. It is a RENDERED finding now, the run completes, and the ragged file is untouched."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _reaction_project(project, params)
        host = _Host()
        config.use_project(project)
        try:
            db_dir = config.database_dir()
            os.makedirs(db_dir, exist_ok=True)
            ragged = os.path.join(db_dir, "chain_reactions_log.csv")
            with open(ragged, "w", encoding="utf-8", newline="") as handle:
                handle.write("uid,hook,rule,action,target,matches,created,outcome\n"
                             "u1,after_500,x,file,t,1,1,ok,EXTRA-CELL\n")
            with open(ragged, "rb") as handle:
                before = handle.read()
            SYSTEM.handlers["staging"](host.ctx())
            ok("RSLT" in host.levels(), "staging completed - the corrupt record did not crash it")
            eq(host.halted, False, "an ERRR never halts")
            with open(ragged, "rb") as handle:
                eq(handle.read(), before, "the ragged record was left untouched")
        finally:
            config.use_project(previous_project)
    eq(sum(1 for batch in host.rendered for f in batch if getattr(f, "type", "") == "rx_record_unreadable"), 2,
       "RENDERED twice - by the before_300 settle AND the after_300 fire: neither path is silent")


def _spawn_project(project, params, header: bool = False):
    """A project whose after_300 rules CASCADE within one hook: the templates.yaml manual's own
    example (spawn a DIAG signal per `A` row), then a rule that counts and lists those spawns.
    `header` adds the run-plan's business AROUND after_300: a before_300 file rule that succeeds and
    one that fails (no database yet), a rule on a hook the run-plan never fires (a rule-INDEX
    problem - every hook's business), and an after_300 rule listing the audit + the reaction
    findings - the before_300 firing is SETTLED into the staged database before after_300 reads it."""
    import csv
    from ruamel.yaml import YAML
    shared = os.path.join(project, "config_project", "shared")
    rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
    os.makedirs(shared)
    os.makedirs(rx_dir)
    with open(os.path.join(shared, "project_params.yaml"), "w", encoding="utf-8") as handle:
        YAML(typ="safe").dump(params, handle)
    with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows([
            ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
            ["probe_spawn", "after_300", "signals", '$script_type = "A"', "add_rows", "signals", "diag_rows", ""],
            ["probe_read", "after_300", "", "", "file", "rx/diag.txt", "diag_txt", "reads the spawns"],
        ] + ([["probe_hdr", "before_300", "", "", "file", "rx/hdr.txt", "hdr_txt", "settled"],
               ["probe_early", "before_300", "signals", "", "file", "rx/early.txt", "hdr_txt", "fails"],
               ["probe_never", "after_999", "", "", "file", "rx/never.txt", "hdr_txt", "never fired"],
               ["probe_log", "after_300", "", "", "file", "rx/log.txt", "log_txt", "reads the record"]]
              if header else []))
    with open(os.path.join(rx_dir, "templates.yaml"), "w", encoding="utf-8") as handle:
        handle.write("diag_rows:\n  - script_type: \"DIAG\"\n    mnemonic: \"DIAG-{$mnemonic}\"\n"
                     "diag_txt: |-\n  DIAG rows: {count(signals, $script_type = \"DIAG\")}\n"
                     "  @for $s in signals where $script_type = \"DIAG\": {$s.spawned_by}\n"
                     "hdr_txt: |-\n  HEADER {$_rule.hook}\n"
                     "log_txt: |-\n  @for $l in chain_reactions_log: LOG {$l.hook} {$l.rule} {$l.outcome}\n"
                     "  @for $i in validation_issues where $type ~ /^rx_/: ISSUE {$i.type} {$i.location}\n"
                     "  END\n")


def _staged_count(host):
    """The RSLT line's `staged N signals` N (what the run REPORTS it saved)."""
    import re
    for level, msg in host.lines:
        match = re.search(r"staged (\d+) signals", msg) if level == "RSLT" else None
        if match:
            return int(match.group(1))
    return None


def test_a_cascade_within_one_hook_through_the_real_run_plan():
    """Refuter round 9 (medium) through run_staging: the manual's own example spawned 24 DIAG signals,
    yet the next rule in the same hook counted 0.0 and listed none - every rule audited `ok`, nothing
    rendered: a SILENT wrong answer in the engine's core use (a chain reaction within one hook)."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _spawn_project(project, params)
        host = _Host()
        config.use_project(project)
        try:
            SYSTEM.handlers["staging"](host.ctx())
            saved = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            sources = sum(1 for row in saved if row.get("script_type") == "A")
            spawned = sum(1 for row in saved if row.get("script_type") == "DIAG")
            ok(sources > 0 and spawned == sources, f"one DIAG spawn per A row, saved ({spawned} vs {sources})")
            with open(os.path.join(config.output_root(), "rx", "diag.txt"), encoding="utf-8") as handle:
                eq(handle.read(), f"DIAG rows: {float(sources)}\n" + "probe_spawn\n" * sources,
                   "the later rule COUNTED and LISTED the earlier rule's spawns")
            eq(_staged_count(host), len(saved), "the RSLT count matches the saved signals")
        finally:
            config.use_project(previous_project)


def test_the_builder_models_what_the_after_300_fire_gets():
    """C-025 refute round 1 + round 2 (#1): the builder previewed against the Database/ folder (a
    previous run's spawns, later phases' write-backs), then modelled the settle as "attach the on-disk
    record" - not this run's before_300 audit and findings. Built BEFORE each run, as a user opens
    the builder (a failing before_300 rule and a rule-index problem among the rules), the model is
    compared with what the REAL after_300 fire gets: every table's rows, in order - the audit and
    validation_issues included - and the preview of a rule listing them is the text the fire
    appends. Run 2 re-stages from the documents: run 1's saved spawns are not in its input."""
    from pipeline5.phases.chain_reactions import engine as reactions
    from pipeline5.truth.database import Database
    from pipeline5.workbench import template_doc
    params = config.load_params()
    previous_project = config.active_project()
    seen, real_fire = [], reactions.fire

    def rows_of(database):
        return {name: [dict(row) for row in database[name]] for name in database.names()}

    def recording_fire(hook, database, **kw):
        if hook == "after_300":                       # the fire's input, before it runs
            seen.append(rows_of(database))
        return real_fire(hook, database, **kw)

    def model():
        """The builder, opened before the run: the first after_300 rule's Database (no spawns before
        it) and probe_log's preview - built in memory (Database.save counted: nothing may be saved)."""
        real_save, saves = Database.save, []
        Database.save = lambda self, directory: saves.append(directory)
        try:
            session = template_doc.Session(None)
            with open(session.active, encoding="utf-8") as handle:
                doc = template_doc.parse(handle.read())
            first = next(r for r in session.rules if r.fire_when == "after_300")
            view = rows_of(session.view(first, doc.templates)[0])
            shown = session.preview_text(doc, next(r for r in session.rules if r.name == "probe_log"), 0)
        finally:
            Database.save = real_save
        return view, shown, saves

    with tempfile.TemporaryDirectory() as project:
        _spawn_project(project, params, header=True)
        config.use_project(project)
        try:
            out = config.output_root()
            log_path = os.path.join(out, "rx", "log.txt")
            for number in (1, 2):
                view, shown, saves = model()
                eq(saves, [], f"run {number}: building the model saved nothing")
                if number == 1:
                    ok(not os.path.exists(os.path.join(out, "rx")),
                       "…and wrote no output (its dry before_300 fire writes nothing)")
                else:
                    ok("spawned_by" in _read_csv(os.path.join(config.database_dir(), "signals.csv"))[0],
                       "the saved Database carries run 1's spawns - what the round-1 builder previewed against")
                earlier = ""                                  # the file APPENDS: run 2 adds to run 1's text
                if os.path.exists(log_path):
                    with open(log_path, encoding="utf-8") as handle:
                        earlier = handle.read()
                reactions.fire = recording_fire
                try:
                    SYSTEM.handlers["staging"](_Host().ctx())
                finally:
                    reactions.fire = real_fire
                fired = seen[-1]
                eq(sorted(view), sorted(fired), f"run {number}: the same tables - the settled record included")
                for name in fired:
                    eq(view[name], fired[name], f"run {number}: {name} - the same rows, in order")
                ok("spawned_by" not in fired["signals"][0], f"run {number}: the fire's input has no spawn column")
                ok(any(r["rule"] == "probe_early" and r["outcome"] == "rx_unknown_table"
                       for r in fired["chain_reactions_log"]), "…THIS run's before_300 audit, the failure included")
                with open(log_path, encoding="utf-8") as handle:
                    appended = handle.read()[len(earlier):]
                lines = shown.split("\n")
                body = appended.split("\n")[:-1]
                ok(lines[0].startswith("APPEND to ") and os.path.normcase(os.path.abspath(lines[0][10:-1]))
                   == os.path.normcase(os.path.abspath(log_path)), f"run {number}: the previewed path ({lines[0]})")
                eq(lines[1:1 + len(body)], body, f"run {number}: the preview is the text the fire appends")
                ok(all(line.startswith("note: ") for line in lines[1 + len(body):]), "…then notes only")
                ok(any(line.startswith("ISSUE rx_unknown_table probe_early") for line in body)
                   and any(line.startswith("ISSUE rx_unfired_hook probe_never") for line in body),
                   f"…which lists this run's before_300 findings: {body}")
        finally:
            config.use_project(previous_project)


def test_the_builder_lints_every_leg_that_fires_after_300():
    """C-025 refute round 2 (#2): the 310 button (Stage I/O List) fires after_300 over the I/O List
    ALONE - no C&E pass, so no validation_issues table. The builder modelled only the full staging:
    a rule listing validation_issues linted clean while the 310 fire rejects it. Each rule is linted
    on every leg that fires its hook now; what only a leg rejects names the leg (a problem every leg
    has is reported once)."""
    from pipeline5.workbench import template_doc
    rules = [["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
             ["issues", "after_300", "", "", "file", "rx/issues.txt", "issues_txt", ""]]
    templates = ("issues_txt: |-\n  @for $i in validation_issues: ISSUE {$i.type}\n  END {$nme}\n"
                 "hdr_txt: |-\n  HEADER\n")
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:           # + a before_300 rule: its settle attaches the
        _rules_project(project, params, rules + [["hdr", "before_300", "", "", "file", "rx/hdr.txt", "hdr_txt",
                                                  ""]], templates)  # record - validation_issues included
        config.use_project(project)
        try:
            session = template_doc.Session(None)
            doc = template_doc.parse(templates)
            leg = [p for p in session.check(doc, session.rules[0]) if p.message.startswith("on the ")]
            eq(leg, [], "after a before_300 settle the 310 leg HAS validation_issues (the record attached)")
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx(), only=310)
            fired = [f.detail for batch in host.rendered for f in batch
                     if getattr(f, "type", "") == "rx_bad_template" and f.location == "issues"]
            ok(fired and "validation_issues" not in fired[0] and "$nme" in fired[0],
               f"…and the real 310 fire loops it (it stops at the missing field): {fired}")
        finally:
            config.use_project(previous_project)
    with tempfile.TemporaryDirectory() as project:
        _rules_project(project, params, rules, templates)
        config.use_project(project)
        try:
            session = template_doc.Session(None)
            doc = template_doc.parse(templates)
            rule = session.rules[0]
            placed = session.check(doc, rule)
            leg = [p for p in placed if p.message.startswith("on the 310 Stage I/O List leg: ")]
            eq([(p.severity, doc.text[p.start:p.end]) for p in leg], [("error", "validation_issues")],
               f"with no settle the 310 leg rejects the loop table, located ({[p.message for p in placed]})")
            eq([(p.severity, doc.text[p.start:p.end]) for p in placed if not p.message.startswith("on the ")],
               [("error", "$nme")], "the full staging accepts the loop - the missing field (every leg's) once")
            ok("the 310 Stage I/O List leg also fires after_300" in session.preview_text(doc, rule, 0),
               "the preview says another leg fires the hook")
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx(), only=310)
            fired = [f for batch in host.rendered for f in batch
                     if getattr(f, "type", "") == "rx_bad_template" and f.location == "issues"]
            ok(fired and "validation_issues" in fired[0].detail, f"the real 310 fire rejects it: {fired}")
        finally:
            config.use_project(previous_project)


def test_the_builder_gates_as_the_run_plan_does():
    """C-025 refute round 2 (#3) + the implications check: the builder halted on a staging finding's RAW
    severity, the run-plan on its severity after the treatments registry - and the run-plan then let
    the reactions fire on a raw FAIL the operator DOWNGRADED (every generation phase refuses one; with
    no I/O sheet it saved an empty Database over the record). Now: a WARN escalated to 'fail' halts
    the run; a FAIL downgraded to 'warn' lets staging proceed but the reactions never run on it; an
    untreated WARN fires. The builder agrees in each case - and a stop is said as a stop, not as the
    rx_unknown_table errors of a fire that never runs; the 310 leg (no second staging pass) fires in
    all three, and is checked."""
    from pipeline5.findings.finding import Finding
    from pipeline5.phases.staging import iolist
    from pipeline5.workbench import template_doc
    rules = [["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
             ["per_signal", "after_300", "signals", "", "file", "rx/sig.txt", "sig_txt", ""],
             ["typo", "after_300", "signals", "", "file", "rx/typo.txt", "typo_txt", ""]]
    templates = "sig_txt: |-\n  SIG {$mnemonic}\ntypo_txt: |-\n  T {$mnemonc}\n"
    params = config.load_params()
    previous_project = config.active_project()
    real_dups = iolist._dup_type_index_findings        # (the full staging's second pass only - not on 310)
    dup = Finding(phase=300, type="stg_dup_type_index", severity="FAIL", detail="duplicated PEC index 7",
                  location="IoList!AD12")
    uid = Finding(phase=300, type="stg_dup_signal_uid", severity="WARN", detail="duplicated signal uid",
                  location="IoList!AD13")
    cases = [(dup, "warn", False, "staging (300) has a blocking stg_dup_type_index the treatments downgraded"),
             (uid, "fail", True, "staging (300) halts on stg_dup_signal_uid"),
             (uid, "", False, None)]
    try:
        for finding, treatment, halts, stop in cases:
            iolist._dup_type_index_findings = lambda table, _f=finding: real_dups(table) + [_f]
            treatments.write(treatments.registry_path(),
                             {finding.uid: treatments.Treatment(uid=finding.uid, treatment=treatment, type=finding.type)})
            label = f"a {finding.severity} treated {treatment!r}"
            with tempfile.TemporaryDirectory() as project:
                _rules_project(project, params, rules, templates)
                config.use_project(project)
                try:
                    session = template_doc.Session(None)
                    doc = template_doc.parse(templates)
                    rule, typo = session.rules
                    placed = session.check(doc, rule)
                    shown = session.preview_text(doc, rule, 0)
                    typo_placed = session.check(doc, typo)
                    context = session.context(typo, doc.templates)
                    host = _Host()
                    SYSTEM.handlers["staging"](host.ctx())
                    path = os.path.join(config.output_root(), "rx", "sig.txt")
                    written = open(path, encoding="utf-8").read().splitlines() if os.path.exists(path) else []
                    host310 = _Host()
                    SYSTEM.handlers["staging"](host310.ctx(), only=310)
                finally:
                    config.use_project(previous_project)
            eq(host.halted, halts, f"{label}: the real run {'halts' if halts else 'runs'}")
            eq(bool(written), stop is None, f"{label}: the reactions {'fire' if stop is None else 'do not fire'}")
            eq(session.halt("after_300") is None, stop is None, f"{label}: the builder agrees")
            skipped = [m for level, m in host.lines if level == "WARN" and "never run on a raw FAIL" in m]
            eq(bool(skipped), not halts and stop is not None, f"{label}: the run says the reactions were skipped")
            fired310 = [f.detail for batch in host310.rendered for f in batch
                        if getattr(f, "type", "") == "rx_bad_template" and f.location == "typo"]
            ok(not host310.halted and fired310 and "$mnemonc" in fired310[0],
               f"{label}: the 310 leg fires all the same - and rejects the typo ({fired310})")
            if stop is not None:
                eq([(p.severity, p.label) for p in placed], [("warning", "rule")], f"{label}: one note, no errors")
                ok(placed[0].message.startswith(f"the run stops before after_300 fires - {stop}"), placed[0].message)
                ok(shown.startswith("the run stops before after_300 fires") and "APPEND" not in shown, shown)
                ok("note: the 310 Stage I/O List leg fires after_300 all the same" in shown, "…the 310 leg still fires")
                eq([(p.severity, p.message.split(" - ")[0]) for p in typo_placed][1:],
                   [("error", "on the 310 Stage I/O List leg: missing field $mnemonc")],
                   f"{label}: the stop note, then the 310 leg's own check ({[p.message for p in typo_placed]})")
                ok("mnemonic" in (context.fields or ()) and "signals" in (context.tables or {}),
                   f"{label}: completions come from the 310 leg that fires (C-025 refute round 3, F7)")
            else:
                eq(placed, [], f"{label}: no false errors")
                eq(shown.split("\n")[1], written[0], f"{label}: the preview is the fire's first line")
    finally:
        iolist._dup_type_index_findings = real_dups


def test_a_downgraded_no_io_sheet_never_saves_over_the_record():
    """The implications check (item 1), the data-loss case - and C-024 refute round 20 (G1: pinned on the 300
    leg only, a guard skipped on 310 / 320 survived; G3: the dark engine warned). With no I/O sheet matched,
    staging returns an EMPTY Database and a raw FAIL. Downgraded in the registry, the gate let the run on -
    and the after_300 fire (any rule) saved that empty Database over signals.csv. On EVERY staging leg
    (300, 310, 320) now: downgraded, the reactions are skipped (a WARN; the before_300 firing listed, not
    recorded); untreated, the run halts (listed as halted) - either way signals.csv, the audit and the
    rule's file are untouched, and the builder says the run stops there. With no rules at all (the engine
    dark, as shipped) the skip says nothing."""
    from pipeline5.phases.staging import iolist
    from pipeline5.workbench import template_doc
    rules = [["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
             ["per_signal", "after_300", "signals", "", "file", "rx/sig.txt", "sig_txt", ""],
             ["hdr", "before_300", "", "", "file", "rx/hdr.txt", "sig_txt", ""]]
    templates = "sig_txt: |-\n  SIG {$_rule.name}\n"
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _rules_project(project, params, rules, templates)
        config.use_project(project)
        try:
            SYSTEM.handlers["staging"](_Host().ctx())                  # a good run: signals.csv saved
            saved = os.path.join(config.database_dir(), "signals.csv")
            before = len(_read_csv(saved))
            ok(before > 0, f"the good run saved {before} signals")
            audit = os.path.join(config.database_dir(), "chain_reactions_log.csv")
            sig = os.path.join(config.output_root(), "rx", "sig.txt")

            def snapshot():
                with open(audit, "rb") as a, open(sig, "rb") as s:
                    return len(_read_csv(saved)), a.read(), s.read()
            recorded = snapshot()
            broken = dict(params, iolist_params=dict(params["iolist_params"], sheets=["NoSuchSheet"]))
            _database, findings = iolist.stage_iolist(broken, save=False)
            eq([f.type for f in findings], ["stg_no_io_sheet"], "the probe's staging finds no I/O sheet")
            real_load = config.load_params
            config.load_params = lambda path=None: broken
            try:
                for treatment, halts in (("warn", False), ("", True)):
                    treatments.write(treatments.registry_path(), {findings[0].uid: treatments.Treatment(
                        uid=findings[0].uid, treatment=treatment, type=findings[0].type)})
                    session = template_doc.Session(None)
                    stops = (session.halt("after_300"), session.halt("after_300", "310 Stage I/O List"))
                    why = "halts on stg_no_io_sheet" if halts else "has a blocking stg_no_io_sheet the treatments downgraded"
                    ok(all(stop and why in stop for stop in stops), f"treated {treatment!r}: the builder, both legs: {stops}")
                    for only in (None, 310, 320):
                        label = f"leg {only or 300}, treated {treatment!r}"
                        host = _Host()
                        SYSTEM.handlers["staging"](host.ctx(), only=only)
                        eq(host.halted, halts, f"{label}: the run {'halts' if halts else 'goes on'}")
                        eq(snapshot(), recorded, f"{label}: signals.csv, the audit and the rule's file untouched")
                        listed = "not recorded, staging halted" if halts else "not recorded, a raw FAIL"
                        ok(any(level == "INFO" and f"before_300 reaction hdr: ok (1 created) - {listed}" in m
                               for level, m in host.lines), f"{label}: the before_300 firing is listed")
                        warned = [m for level, m in host.lines if level == "WARN" and "never run on a raw FAIL" in m]
                        eq(len(warned), 0 if halts else 1, f"{label}: the skip {'is a halt' if halts else 'is said'}")
                treatments.write(treatments.registry_path(), {findings[0].uid: treatments.Treatment(
                    uid=findings[0].uid, treatment="warn", type=findings[0].type)})
                rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
                with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
                    handle.write(",".join(rules[0]) + "\n")            # the engine DARK, as shipped
                host = _Host()
                SYSTEM.handlers["staging"](host.ctx())
                eq([m for _level, m in host.lines if "reaction" in m], [], "the dark engine: the skip says nothing")
                ok("RSLT" in host.levels(), "…and the staging reports as ever")
                with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
                    handle.write(",".join(rules[0]) + "\n" + ",".join(rules[2]) + "\n")   # before_300 only
                host = _Host()
                SYSTEM.handlers["staging"](host.ctx())
                eq(len([m for level, m in host.lines if level == "WARN" and "never run on a raw FAIL" in m]), 1,
                   "only a before_300 rule dropped: the skip is said (C-024 refute round 21, N1)")
            finally:
                config.load_params = real_load
        finally:
            config.use_project(previous_project)


def test_text_utf8_cannot_store_never_crashes_the_run():
    """C-024 refute round 21 (F1) through the REAL run-plan: an add_rows value holding a JSON-style emoji
    escape (`"\\uD83D\\uDE00"` - YAML decodes it to two lone surrogates) crashed the staging handler at the
    save ("surrogates not allowed"; no RSLT, the spawns and findings lost). The value is refused before
    the add now: the run completes, the finding is rendered AND recorded, signals.csv holds the staged
    rows."""
    rules = [["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
             ["emoji", "after_300", "signals", '$script_type = "A"', "add_rows", "signals", "emoji_rows", ""]]
    templates = 'emoji_rows:\n  - script_type: "E"\n    mnemonic: "E-{$mnemonic}-\\uD83D\\uDE00"\n'
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _rules_project(project, params, rules, templates)
        config.use_project(project)
        try:
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx())
            ok("RSLT" in host.levels(), "the run completes")
            eq(_rx_rendered(host), {("rx_bad_template", "emoji")}, "the finding is rendered")
            recorded = _read_csv(os.path.join(config.database_dir(), "validation_issues.csv"))
            ok(any(r["type"] == "rx_bad_template" and r["location"] == "emoji" for r in recorded), "…and recorded")
            saved = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            eq((_staged_count(host), sum(1 for r in saved if r.get("script_type") == "E")), (len(saved), 0),
               "signals.csv holds the staged rows - no spawn")
        finally:
            config.use_project(previous_project)


def test_a_field_name_or_a_quoted_text_utf8_cannot_store_never_wipes_the_record():
    """C-024 refute round 22 (F-A, F-B) + C-025 refute round 5 (#1), through the REAL run-plan. A row template's
    FIELD NAME holding the escape pair (`"tag\\uD83D\\uDE00"` - a YAML key: two lone surrogates) was added, then
    the save failed on the CSV header AFTER opening signals.csv for writing: the staged table emptied (0 bytes)
    while RSLT reported it - and the builder approved the rule. A finding QUOTING such text (templates.yaml
    naming a template twice, the second value holding the escape) crashed the run at the record (its uid
    hashes UTF-8). Now the field name is the rule's rx_bad_template on BOTH legs that fire after_300 - nothing
    added, the lint and the preview saying the same - signals.csv stays byte for byte what staging saved, the
    audit reaches the disk, and a quoted text is recorded escaped, on either hook."""
    from pipeline5.phases.chain_reactions import engine as reactions
    from pipeline5.workbench import template_doc
    header = ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"]
    escape = "\\uD83D\\uDE00"                                   # the YAML escapes, as an author types them
    params = config.load_params()
    previous_project = config.active_project()
    real_fire, staged = reactions.fire, {}

    def spy(hook, database, **kw):                            # signals.csv as staging saved it, before the fire
        if hook == "after_300":
            with open(os.path.join(config.database_dir(), "signals.csv"), "rb") as handle:
                staged["signals"] = handle.read()
        return real_fire(hook, database, **kw)

    templates = 'name_rows:\n  - script_type: "E"\n    "tag' + escape + '": "E-{$mnemonic}"\n'
    rule = ["rx_name", "after_300", "signals", '$script_type = "A"', "add_rows", "signals", "name_rows", ""]
    for only in (None, 310):
        leg = only or 300
        with tempfile.TemporaryDirectory() as project:
            _rules_project(project, params, [header, rule], templates)
            config.use_project(project)
            reactions.fire = spy
            try:
                if only is None:                                  # the builder, opened before the run
                    session = template_doc.Session(None)
                    doc = template_doc.parse(templates)
                    placed = session.check(doc, session.rules[0])
                    shown = session.preview_text(doc, session.rules[0], 0).splitlines()
                    ok(any(p.severity == "error" and "the field name 'tag\\ud83d\\ude00'" in p.message for p in placed),
                       f"the lint names the field: {[p.message for p in placed]}")
                host = _Host()
                SYSTEM.handlers["staging"](host.ctx(), only=only)
                db_dir = config.database_dir()
                found = [f for batch in host.rendered for f in batch if getattr(f, "type", "") == "rx_bad_template"]
                eq([(f.location, "entry 1: the field name 'tag\\ud83d\\ude00'" in f.detail) for f in found],
                   [("rx_name", True)], f"leg {leg}: the rule's rx_bad_template, naming the entry and the field")
                if only is None:
                    eq(shown[0], f"rx_bad_template: {found[0].detail}", "…the preview said the fire's finding")
                with open(os.path.join(db_dir, "signals.csv"), "rb") as handle:
                    eq(handle.read(), staged.pop("signals"), f"leg {leg}: signals.csv byte for byte as staged")
                audit = _read_csv(os.path.join(db_dir, "chain_reactions_log.csv"))
                eq([(r["rule"], r["created"], r["outcome"]) for r in audit], [("rx_name", "0", "rx_bad_template")],
                   f"leg {leg}: the audit on disk - nothing created")
                eq(_staged_count(host), len(_read_csv(os.path.join(db_dir, "signals.csv"))), f"leg {leg}: RSLT = the saved rows")
            finally:
                reactions.fire = real_fire
                config.use_project(previous_project)
    duplicated = 'hdr: "HEADER"\nhdr: "HEADER ' + escape + '"\n'   # the YAML error quotes the decoded value
    for hook in ("after_300", "before_300"):
        with tempfile.TemporaryDirectory() as project:
            _rules_project(project, params, [header, ["rx_dup", hook, "", "", "file", "rx/d.txt", "hdr", ""]], duplicated)
            config.use_project(project)
            try:
                host = _Host()
                SYSTEM.handlers["staging"](host.ctx())
                ok("RSLT" in host.levels(), f"{hook}: the run completes")
                eq(_rx_rendered(host), {("rx_templates_unreadable", "chain_reactions/templates.yaml")}, f"{hook}: rendered")
                recorded = [r for r in _read_csv(os.path.join(config.database_dir(), "validation_issues.csv"))
                            if r["type"] == "rx_templates_unreadable"]
                eq(len(recorded), 1, f"{hook}: …and recorded")
                ok("\\ud83d" in recorded[0]["detail"] and "\ud83d" not in recorded[0]["detail"],
                   f"{hook}: the quoted text recorded escaped - {ascii(recorded[0]['detail'])[:160]}")
            finally:
                config.use_project(previous_project)


def test_a_device_name_target_is_refused_through_the_run_plan():
    """C-024 refute round 23 (F1) through the REAL run-plan: an absolute target naming `NUL:` (a ':' after the device
    name) reached the null device - the ':' was read on the NORMALIZED path (`\\\\.\\NUL`, all of it a drive) - audited
    ok on both hooks, nothing on disk. Refused before the write now, rendered and recorded, on both hooks."""
    header = ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"]
    params = dict(config.load_params())
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project, tempfile.TemporaryDirectory() as elsewhere:
        params["rx_dir"] = elsewhere
        _rules_project(project, params, [header,
                                         ["nul_before", "before_300", "", "", "file", elsewhere.replace("\\", "/") + "/nul:",
                                          "hdr", ""],
                                         ["nul_after", "after_300", "", "", "file", "{$_params.rx_dir}/NUL:", "hdr", ""]],
                       'hdr: "HEADER"\n')
        config.use_project(project)
        try:
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx())
            ok("RSLT" in host.levels(), "the run completes")
            eq(_rx_rendered(host), {("rx_file_write", "nul_before"), ("rx_file_write", "nul_after")},
               "both refused - rendered")
            audit = _read_csv(os.path.join(config.database_dir(), "chain_reactions_log.csv"))
            eq(sorted((r["rule"], r["outcome"]) for r in audit),
               [("nul_after", "rx_file_write"), ("nul_before", "rx_file_write")], "…audited so, never ok")
            recorded = {(r["type"], r["location"]) for r in _read_csv(os.path.join(config.database_dir(), "validation_issues.csv"))}
            ok({("rx_file_write", "nul_before"), ("rx_file_write", "nul_after")} <= recorded, "…and recorded")
            eq(os.listdir(elsewhere), [], "nothing written")
        finally:
            config.use_project(previous_project)


def test_spawns_are_saved_beside_a_ragged_record():
    """Refuter round 9: with a ragged chain_reactions_log.csv the engine skipped the WHOLE save - the
    after_300 spawns silently never reached signals.csv while RSLT reported them. The rest of the
    Database is saved now; only the unreadable record is left untouched (and reported)."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _spawn_project(project, params)
        host = _Host()
        config.use_project(project)
        try:
            db_dir = config.database_dir()
            os.makedirs(db_dir, exist_ok=True)
            ragged = os.path.join(db_dir, "chain_reactions_log.csv")
            with open(ragged, "w", encoding="utf-8", newline="") as handle:
                handle.write("uid,hook,rule,action,target,matches,created,outcome\nu1,after_500,x,f,t,1,1,ok,EXTRA\n")
            with open(ragged, "rb") as handle:
                before = handle.read()
            SYSTEM.handlers["staging"](host.ctx())
            saved = _read_csv(os.path.join(db_dir, "signals.csv"))
            ok(any(row.get("script_type") == "DIAG" for row in saved), "the spawns were SAVED")
            eq(_staged_count(host), len(saved), "the RSLT count matches the saved signals")
            with open(ragged, "rb") as handle:
                eq(handle.read(), before, "the ragged record is untouched")
        finally:
            config.use_project(previous_project)
    ok(any(getattr(f, "type", "") == "rx_record_unreadable" for batch in host.rendered for f in batch),
       "the unreadable record was rendered")


def test_a_hidden_record_file_is_saved_whole_through_the_run_plan():
    """C-024 refute round 24 (gap 1) through the REAL run-plan: run 2 with the audit HIDDEN rendered
    rx_record_unwritable ("a file held open") while signals.csv held run 2's spawns and validation_issues run 2's
    rx_file_write - the audit still run 1's, saying the refused rule ran ok: a record half one run's, half the last
    (the save opened each file to append - Windows grants that on a hidden file - then re-created it, which Windows
    refuses). The save writes in place now: run 2's record whole, the audit still hidden."""
    if os.name != "nt":
        return                                                # (a Windows attribute)
    import ctypes
    from ruamel.yaml import YAML
    header = ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"]
    templates = ('diag:\n  - script_type: "DIAG"\n    mnemonic: "D-{$mnemonic}-run{$_params.rx_run}"\n'
                 'line: |-\n  run {$_params.rx_run}\n')
    rules = [header, ["spawn", "after_300", "signals", '$script_type = "A"', "add_rows", "signals", "diag", ""],
             ["out", "after_300", "", "", "file", "rx/{$_params.rx_name}", "line", ""]]
    params = dict(config.load_params())
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        params.update(rx_run="1", rx_name="ok.txt")
        _rules_project(project, params, rules, templates)
        config.use_project(project)
        try:
            SYSTEM.handlers["staging"](_Host().ctx())
            db_dir = config.database_dir()
            audit = os.path.join(db_dir, "chain_reactions_log.csv")
            eq([(r["rule"], r["outcome"]) for r in _read_csv(audit)], [("spawn", "ok"), ("out", "ok")], "(run 1)")
            params.update(rx_run="2", rx_name="NUL")                  # run 2: other spawns, a refused target
            with open(os.path.join(project, "config_project", "shared", "project_params.yaml"), "w",
                      encoding="utf-8") as handle:
                YAML(typ="safe").dump(params, handle)
            ctypes.windll.kernel32.SetFileAttributesW(audit, 0x2)
            try:
                host = _Host()
                SYSTEM.handlers["staging"](host.ctx())
                eq(ctypes.windll.kernel32.GetFileAttributesW(audit) & 0x2, 0x2, "(the audit is still hidden)")
            finally:
                ctypes.windll.kernel32.SetFileAttributesW(audit, 0x80)
            eq(_rx_rendered(host), {("rx_file_write", "out")}, "run 2: the refused target - no rx_record_unwritable")
            spawned = [r for r in _read_csv(os.path.join(db_dir, "signals.csv")) if r.get("script_type") == "DIAG"]
            ok(spawned and all(r["mnemonic"].endswith("-run2") for r in spawned), "signals.csv: run 2's spawns")
            eq([(r["rule"], r["created"], r["outcome"]) for r in _read_csv(audit)],
               [("spawn", str(len(spawned)), "ok"), ("out", "0", "rx_file_write")], "the hidden audit: run 2's, whole")
            eq([(r["type"], r["location"]) for r in _read_csv(os.path.join(db_dir, "validation_issues.csv"))
                if r["type"].startswith("rx_")], [("rx_file_write", "out")], "…and the findings it audits")
        finally:
            config.use_project(previous_project)


def _rules_project(project, params, rules, templates_text):
    """A minimal project carrying exactly `rules` (reactions.csv rows, header first) + `templates_text`."""
    import csv
    from ruamel.yaml import YAML
    shared = os.path.join(project, "config_project", "shared")
    rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
    os.makedirs(shared)
    os.makedirs(rx_dir)
    with open(os.path.join(shared, "project_params.yaml"), "w", encoding="utf-8") as handle:
        YAML(typ="safe").dump(params, handle)
    with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rules)
    with open(os.path.join(rx_dir, "templates.yaml"), "w", encoding="utf-8") as handle:
        handle.write(templates_text)


def test_a_blank_range_end_invents_nothing_on_the_real_fixture():
    """Refuter round 10 + the user's decision (2026-09-25: a BLANK range end is an EMPTY range):
    a byte map over the staged node rows used to invent a `QB0` line for every input-only node (its Q
    start/end are blank - they read as 0), audited ok. The file must hold exactly the REAL Q bytes."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _rules_project(project, params, [
            ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
            ["byte_map", "after_300", "signals", "present($I_startByte) or present($Q_startByte)", "file",
             "rx/bytes.txt", "bytes_txt", "the refuter's byte map"],
        ], "bytes_txt: |-\n  @for $b in {$Q_startByte}..{$Q_endByte}: QB{$b}\n")
        host = _Host()
        config.use_project(project)
        try:
            SYSTEM.handlers["staging"](host.ctx())
            staged = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            with open(os.path.join(config.output_root(), "rx", "bytes.txt"), encoding="utf-8") as handle:
                got = handle.read()
        finally:
            config.use_project(previous_project)
    expected, blank_q = "", 0
    for row in staged:
        start, stop = (row.get("Q_startByte") or "").strip(), (row.get("Q_endByte") or "").strip()
        if not ((row.get("I_startByte") or "").strip() or start):
            continue                                                   # not matched by the condition
        if start and stop:
            expected += "\n".join(f"QB{b}" for b in range(int(float(start)), int(float(stop)) + 1)) + "\n"
        else:
            expected += "\n"                                           # blank Q: an EMPTY range
            blank_q += 1
    ok(blank_q > 0, f"the fixture has input-only nodes with blank Q bytes ({blank_q})")
    eq(got, expected, "exactly the real Q bytes - no invented QB0 for a blank end")


def test_310_leg_appends_reaction_findings_to_the_existing_record():
    """Refuter round 7 B3 through the real run-plan: the 310 leg (Stage I/O List) stages a Database
    WITHOUT validation_issues; a reaction finding used to create an empty table and save it OVER the
    record an earlier phase wrote. It is appended to the on-disk record now."""
    from pipeline5.findings.finding import Finding, record_standalone
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _reaction_project(project, params)
        host = _Host()
        config.use_project(project)
        try:
            db_dir = config.database_dir()
            record_standalone([Finding(phase=110, type="probe_prior", severity="FAIL",
                                       detail="from phase 110", location="X1")], directory=db_dir)
            SYSTEM.handlers["staging"](host.ctx(), only=310)
            recorded = {(r["type"], r["location"]) for r in _read_csv(os.path.join(db_dir, "validation_issues.csv"))}
            ok(("probe_prior", "X1") in recorded, "the earlier phase's record SURVIVED the 310 leg")
            ok({("rx_unknown_table", "probe_bad"), ("rx_unknown_table", "probe_early")} <= recorded,
               "...with the reaction findings appended")
        finally:
            config.use_project(previous_project)


def test_sub_phase_dispatch():
    """A dropdown action: the parent handler with only=<sub> - the 310 stage-I/O-List leg, plus a
    validation sub (110), through the same dispatch the App's _sub_worker performs."""
    host = _Host()
    phase = SYSTEM.phases.phase_of_sub(310)
    SYSTEM.handlers[phase.handler](host.ctx(), only=310)
    ok(any("310" in msg for _lvl, msg in host.lines), "the 310 leg announces itself")
    ok("RSLT" in host.levels(), "the 310 leg reports a RSLT")
    host2 = _Host()
    phase = SYSTEM.phases.phase_of_sub(110)
    SYSTEM.handlers[phase.handler](host2.ctx(), only=110)
    ok("RSLT" in host2.levels(), "the 110 validator reports a RSLT")


def _qbad_project(project, params):
    """A PROJECT whose after_300 rule spawns, for every contactor output, its QBAD signal at the input of the
    same offset, named as the 05 Output Feedback block references it (the FVT convention, [[C-030]])."""
    import csv
    from ruamel.yaml import YAML
    shared = os.path.join(project, "config_project", "shared")
    rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
    os.makedirs(shared)
    os.makedirs(rx_dir)
    with open(os.path.join(shared, "project_params.yaml"), "w", encoding="utf-8") as handle:
        YAML(typ="safe").dump(params, handle)
    with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows([
            ["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"],
            ["qbad", "after_300", "signals", "$script_type ~ /^KQ/ and $bit ~ /^Q/", "add_rows", "signals",
             "qbad_rows", "the QBAD of every contactor output"]])
    with open(os.path.join(rx_dir, "templates.yaml"), "w", encoding="utf-8") as handle:
        handle.write("qbad_rows:\n"
                     "  - script_type: \"KB\"\n"
                     "    bit: \"{regex_replace($bit, /^Q/, 'I')}\"\n"
                     "    functional_unit: \"{$functional_unit}\"\n"
                     "    location: \"{$location}\"\n"
                     "    device: \"{$device}\"\n"
                     "    name_in_tagtable: \"QBAD_{$name_in_tagtable}\"\n")


def _plc_tag_names(out_root):
    import glob
    from openpyxl import load_workbook
    names = {}
    for path in glob.glob(os.path.join(out_root, "**", "PLCTags*.xlsx"), recursive=True):
        wb = load_workbook(path, read_only=True)
        for row in wb["PLC Tags"].iter_rows(min_row=2, values_only=True):
            names[row[0]] = row[3]
        wb.close()
    return names


def test_generated_signals_reach_generation():
    """C-030 through the REAL run-plan: a project rule spawns a QBAD per contactor output; 300 stages them as
    first-class signals (saved, provenance, unique uids); a LONE 510 button - no 300 before it in this host -
    re-fires the rule in memory and emits their tags at the input of the same offset; coverage counts them
    as generated and none is an ORPHAN; validation (100) judges the documents only."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as project:
        _qbad_project(project, params)
        config.use_project(project)
        try:
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx())
            sig = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            kq = {r["uid"]: r for r in sig if r["script_type"].startswith("KQ") and r["bit"].startswith("Q")}
            gen = [r for r in sig if r.get("spawned_by") == "qbad"]
            ok(kq and len(gen) == len(kq), f"one QBAD per contactor output ({len(gen)} / {len(kq)})")
            g = gen[0]
            k = kq[g["source_uid"]]
            eq(g["bit"], "I" + k["bit"][1:], "the input of the same offset")
            eq(g["name_in_tagtable"], "QBAD_" + k["name_in_tagtable"], "the name the 05 block references")
            ok(g["source_cell"].startswith("qbad:") and "!" not in g["source_cell"], g["source_cell"])
            eq(len({r["uid"] for r in sig}), len(sig), "every uid unique")
            ok(any("generated" in m for lv, m in host.lines if lv == "RSLT"), "the RSLT line counts them")

            lone = _Host()
            SYSTEM.handlers["data_blocks"](lone.ctx(), only=510)
            tags = _plc_tag_names(config.output_root())
            eq(tags.get(g["name_in_tagtable"]), "%" + g["bit"], "a lone 510 emits the generated tag")
            ok(all(("QBAD_" + r["name_in_tagtable"]) in tags for r in kq.values()), "every QBAD tag emitted")

            cov = _Host()
            SYSTEM.handlers["reporting"](cov.ctx())
            report = [os.path.join(dp, f) for dp, _d, fs in os.walk(config.output_root()) for f in fs
                      if f == "io_project_coverage_report.txt"]
            text = open(report[0], encoding="utf-8").read()
            ok(f"generated {len(gen)}," in text, "coverage counts the generated signals")
            orphan_part = text.split("ORPHAN", 1)[1].split("UNPLACED", 1)[0]
            ok("qbad:" not in orphan_part, "no generated signal is an ORPHAN")

            val = _Host()
            SYSTEM.handlers["validation"](val.ctx())
            located = [f for batch in val.rendered for f in batch
                       if str(getattr(f, "location", "")).startswith("qbad:")]
            eq(located, [], "validation judges the documents only")
        finally:
            config.use_project(previous_project)


def _absorb_project(project, params, rules, templates):
    """A PROJECT with its own after_300 rules (`rules` = data rows of reactions.csv) and templates.yaml text."""
    import csv
    from ruamel.yaml import YAML
    shared = os.path.join(project, "config_project", "shared")
    rx_dir = os.path.join(project, "config_project", "systems", SYSTEM.id, "chain_reactions")
    os.makedirs(shared)
    os.makedirs(rx_dir)
    with open(os.path.join(shared, "project_params.yaml"), "w", encoding="utf-8") as handle:
        YAML(typ="safe").dump(params, handle)
    with open(os.path.join(rx_dir, "reactions.csv"), "w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(
            [["name", "fire_when", "source_table", "condition", "action", "target", "template", "comment"]] + rules)
    with open(os.path.join(rx_dir, "templates.yaml"), "w", encoding="utf-8") as handle:
        handle.write(templates)


_QBAD_RULE = ["qbad", "after_300", "signals", "$script_type ~ /^KQ/ and $bit ~ /^Q/", "add_rows", "signals",
              "qbad_rows", "the QBAD of every contactor output"]
_QBAD_TPL = ("qbad_rows:\n  - script_type: \"KB\"\n    bit: \"{regex_replace($bit, /^Q/, 'I')}\"\n"
             "    functional_unit: \"{$functional_unit}\"\n    location: \"{$location}\"\n    device: \"{$device}\"\n"
             "    name_in_tagtable: \"QBAD_{$name_in_tagtable}\"\n")


def _run_all(host):
    for number in SYSTEM.phases.run_order():
        SYSTEM.handlers[SYSTEM.phases.by_number(number).handler](host.ctx())


def _builder_data(out_root):
    """{relative path: content} of the BuilderData tree - .xlsx by cell values (zip containers are not
    byte-stable), everything else by bytes."""
    from openpyxl import load_workbook
    tree = {}
    for dirpath, _dirs, files in os.walk(out_root):
        for name in files:
            path = os.path.join(dirpath, name)
            rel = os.path.relpath(path, out_root).replace(os.sep, "/")
            if "BuilderData" not in rel or name.endswith(".openn"):
                continue                 # the workspace config / a header sidecar: per generation by design
            if name.endswith(".xlsx"):
                wb = load_workbook(path, read_only=True)
                tree[rel] = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
                wb.close()
            else:
                tree[rel] = _unstamped(open(path, "rb").read())
    return tree


def _unstamped(raw: bytes) -> bytes:
    """The bytes without the contract-v1 `#!openn` header: its `run` and `generated` differ per
    generation by design (contract §6.2); the content below it must not."""
    from pipeline5.systems.plc_based.siemens_s7 import openn_header
    try:
        return openn_header.strip_header(raw.decode("utf-8")).encode("utf-8")
    except UnicodeDecodeError:
        return raw


def test_run_all_writes_a_ready_op5_workspace():
    """Contract v1, the producer side (Shared/PL5_OP5_contract.md §3, §4, §6): after Run-all the BuilderData
    tree is the VCI-shaped workspace - the config names the run, every file carries a valid header with THAT
    run (one id for the whole chain), the PLC folder is the Plc station of Stations.csv, the hardware sits
    at the workspace root, the software under `<PLC>/Program blocks`, the tag workbook under `<PLC>/PLC
    tags` with its sidecar, the template copies under `Templates/`, no legacy folder; a lone 700 afterwards
    mints a new run - the config follows it and the other phases' files are stale (their run differs) -
    nothing is swept."""
    from pipeline5.systems.plc_based.siemens_s7 import openn_header as header
    plc = "n0001-mc1-cc1-k65501"
    host = _Host(run="run-all-1")
    _run_all(host)
    eq(host.halted, False, "the clean fixture does not halt")
    root = os.path.join(config.output_root(), "TiaPortalProjectInterface", "BuilderData")
    cfg_path = os.path.join(root, ".openn", "workspace.openn.config")
    cfg = header.read_header(cfg_path, header.WORKSPACE_REQUIRED)
    eq((cfg.status, cfg.get("run"), cfg.get("contract"), cfg.get("plcs"), cfg.get("templates")),
       ("ok", "run-all-1", "1", plc, "Templates"), f"the workspace config: {cfg.problems}")
    files = sorted(os.path.relpath(os.path.join(dp, f), root).replace(os.sep, "/")
                   for dp, _d, fs in os.walk(root) for f in fs)
    ok("Devices & networks/Stations.csv" in files and "Devices & networks/Modules.csv" in files, files)
    ok(f"{plc}/PLC tags/PLCTags.xlsx" in files and f"{plc}/PLC tags/PLCTags.xlsx.openn" in files, files)
    tag_tables = [f for f in files if f.startswith(f"{plc}/PLC tags/") and f.endswith(".xml")]
    ok(len(tag_tables) >= 2, f"one sw/tag-table XML per tag table beside the workbook (contract 9.1): {tag_tables}")
    blocks = [f for f in files if f.startswith(f"{plc}/Program blocks/")]
    ok(any(f.endswith(".xml") for f in blocks) and any(f.endswith(".csv") for f in blocks)
       and any(f.endswith(".scl") for f in blocks) and f"{plc}/Program blocks/InstanceDBs.csv" in files, blocks)
    ok(any(f.startswith("Templates/TEMPLATE--") for f in files), files)
    ok(not any(f.startswith(("HardwareConfiguration/", "SoftwareBlocks/", "PlcTags/")) for f in files),
       "the legacy layout is not written any more")
    bad = []
    for rel in files:
        if rel.startswith(".openn/") or rel.endswith(".openn"):
            continue
        h = header.read_header(os.path.join(root, rel))
        top = rel.split("/")[0]
        if h.status != "ok" or h.get("run") != "run-all-1":
            bad.append((rel, h.status, h.get("run"), h.problems))
        elif top == "Devices & networks" and (h.get("plc") or h.get("target") != "Devices & networks"):
            bad.append((rel, "hardware placement", h.get("plc"), h.get("target")))
        elif top == "Templates" and h.get("kind") != "sw/block-template":
            bad.append((rel, "template kind", h.get("kind")))
        elif top == plc and (h.get("plc") != plc or h.get("target") != rel.split("/")[1]):
            bad.append((rel, "plc/target vs location", h.get("plc"), h.get("target")))
    eq(bad, [], "every file: a valid header, the chain's run, plc/target agreeing with its location")
    kinds = {rel: header.read_header(os.path.join(root, rel)).get("kind") for rel in files if not rel.endswith(".openn")}
    eq(kinds["Devices & networks/Stations.csv"], "hw/stations")
    eq(kinds["Devices & networks/Modules.csv"], "hw/modules")
    eq(kinds[f"{plc}/Program blocks/InstanceDBs.csv"], "sw/instance-db")
    eq({kinds[f] for f in tag_tables}, {"sw/tag-table"}, "every tag-table XML declares sw/tag-table")
    eq(kinds[f"{plc}/Program blocks/02_COM.xml"], "sw/data-block")
    eq(kinds[f"{plc}/Program blocks/03_Zone Cumulative.xml"], "sw/code-block")
    eq(kinds[f"{plc}/Program blocks/04_ESTOP.csv"], "sw/block-gen")
    eq(kinds[f"{plc}/Program blocks/Diagnostic_for_OPC.scl"], "sw/source")
    eq(header.read_header(os.path.join(root, plc, "Program blocks", "Diagnostic_for_OPC.scl")).get("name"),
       "06_Diagnostic for OPC", "the FUNCTION's name, which differs from the file's")
    # a lone phase button afterwards: ITS run - the others' files are stale for OP5, nothing is deleted
    before = set(files)
    SYSTEM.handlers["hardware"](_Host(run="lone-700").ctx())
    eq(header.read_header(cfg_path, header.WORKSPACE_REQUIRED).get("run"), "lone-700", "the config follows the lone run")
    eq(header.read_header(os.path.join(root, "Devices & networks", "Stations.csv")).get("run"), "lone-700")
    eq(header.read_header(os.path.join(root, plc, "PLC tags", "PLCTags.xlsx")).get("run"), "run-all-1",
       "untouched by the lone phase = stale for OP5 (its run differs from the workspace's)")
    after = set(os.path.relpath(os.path.join(dp, f), root).replace(os.sep, "/") for dp, _d, fs in os.walk(root) for f in fs)
    eq(after, before, "a generation never sweeps: stale files stay for OP5 to list")


def test_generated_signals_change_only_their_own_outputs():
    """C-030 refute round 1, at the clause's altitude: Run-all WITH the QBAD rule vs WITHOUT it - every
    BuilderData file but the PLC tag tables is byte-identical (the hardware included: a generated signal sits
    under no station), the tags gain exactly the QBADs, and after Run-all (its last step, 100, judges the
    documents) the saved signals still hold every generated row."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as on, tempfile.TemporaryDirectory() as off:
        _absorb_project(on, params, [_QBAD_RULE], _QBAD_TPL)
        _absorb_project(off, params, [], "")
        try:
            config.use_project(off)
            _run_all(_Host())
            base = _builder_data(config.output_root())
            config.use_project(on)
            _run_all(_Host())
            with_rule = _builder_data(config.output_root())
            sig = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            log = _read_csv(os.path.join(config.database_dir(), "chain_reactions_log.csv"))
        finally:
            config.use_project(previous_project)
    changed = sorted(k for k in set(base) | set(with_rule) if base.get(k) != with_rule.get(k))
    ok(changed and all("PLC tags" in k for k in changed), f"only the PLC tag tables change: {changed}")
    gen = [r for r in sig if r.get("spawned_by") == "qbad"]
    created = sum(int(r["created"]) for r in log if r["rule"] == "qbad")
    ok(created > 0 and len(gen) == created, f"after Run-all the saved signals hold the {created} generated rows ({len(gen)})")


def test_lone_buttons_in_a_fresh_project():
    """C-030 refute round 1: in FRESH projects (nothing ran before) a lone 510 emits the generated tags and a lone
    700 the same hardware as Run-all, and neither writes the reaction record (phase 300 owns it)."""
    params = config.load_params()
    previous_project = config.active_project()
    with tempfile.TemporaryDirectory() as ref, tempfile.TemporaryDirectory() as a, \
            tempfile.TemporaryDirectory() as b:
        for p in (ref, a, b):
            _absorb_project(p, params, [_QBAD_RULE], _QBAD_TPL)
        try:
            config.use_project(ref)
            _run_all(_Host())
            ref_tree = _builder_data(config.output_root())
            config.use_project(a)
            SYSTEM.handlers["data_blocks"](_Host().ctx(), only=510)
            tags = _builder_data(config.output_root())
            no_log_a = not os.path.exists(os.path.join(config.database_dir(), "chain_reactions_log.csv"))
            config.use_project(b)
            SYSTEM.handlers["hardware"](_Host().ctx())
            hw = _builder_data(config.output_root())
            no_log_b = not os.path.exists(os.path.join(config.database_dir(), "chain_reactions_log.csv"))
        finally:
            config.use_project(previous_project)
    tag_files = [k for k in tags if "PLC tags" in k]
    ok(tag_files and all(tags[k] == ref_tree.get(k) for k in tag_files), "a lone 510 = Run-all's tags (the QBADs incl.)")
    hw_files = [k for k in hw if "Devices & networks" in k]
    ok(hw_files and all(hw[k] == ref_tree.get(k) for k in hw_files), "a lone 700 = Run-all's hardware")
    ok(no_log_a and no_log_b, "a generation button writes no reaction record")


def test_cascade_orphan_and_validation_see_absorbed_rows():
    """C-030 refute round 1: within ONE hook a later rule reads an earlier rule's spawns ABSORBED - a condition on
    a derived field matches them and its source_uid is their saved uid; a generated row landing in no output is
    an ORPHAN; a spawn the 140 check would flag (a C&E-mandatory type unknown to the C&E) stays unseen by 100."""
    params = config.load_params()
    previous_project = config.active_project()
    rules = [_QBAD_RULE,
             ["echo", "after_300", "signals", "$script_type = 'KB' and $tagtable = 'SAFETY_Contactors'", "add_rows",
              "signals", "echo_rows", "reads the absorbed QBADs"],
             ["lost", "after_300", "", "", "add_rows", "signals", "lost_rows", "a signal nothing consumes"]]
    tpl = (_QBAD_TPL + "echo_rows:\n  - script_type: \"ZZ_ECHO\"\n    functional_unit: \"{$functional_unit}\"\n"
           "    device: \"{$device}\"\n    desc_l1: \"echo\"\n"
           "lost_rows:\n  - script_type: \"KQ\"\n    functional_unit: \"=NOWHERE\"\n    device: \"-K999\"\n"
           "    bit: \"Q999.7\"\n    desc_l1: \"EMERGENCY\"\n")
    with tempfile.TemporaryDirectory() as project:
        _absorb_project(project, params, rules, tpl)
        try:
            config.use_project(project)
            host = _Host()
            SYSTEM.handlers["staging"](host.ctx())
            sig = _read_csv(os.path.join(config.database_dir(), "signals.csv"))
            qbad = {r["uid"] for r in sig if r.get("spawned_by") == "qbad"}
            echo = [r for r in sig if r.get("spawned_by") == "echo"]
            ok(qbad and len(echo) == len(qbad), f"the echo matched every absorbed QBAD ({len(echo)} / {len(qbad)})")
            ok(all(r["source_uid"] in qbad for r in echo), "its source_uid is a saved QBAD uid")
            eq(len({r["source_cell"] for r in echo}), len(echo), "one provenance per source - no #n chain")
            SYSTEM.handlers["reporting"](_Host().ctx())
            cov = _read_csv(os.path.join(config.database_dir(), "coverage.csv"))
            lost = [r for r in cov if r.get("source_cell", "").startswith("lost:")]
            ok(lost and lost[0].get("kind") == "generated", "the lost spawn is a generated row in coverage")
            val = _Host()
            SYSTEM.handlers["validation"](val.ctx())
            seen = [f for batch in val.rendered for f in batch
                    if any(str(getattr(f, k, "")).startswith(("lost:", "qbad:", "echo:")) for k in ("location", "location2"))
                    or "=NOWHERE" in str(getattr(f, "detail", ""))]
            eq(seen, [], "validation judges the documents only - the lost KQ is not checked against the C&E")
        finally:
            config.use_project(previous_project)


def test_a_transfer_area_fail_halts_the_interface_outputs():
    """C-031 through the real handlers, on the fixture's two IOC interfaces under the coupler n0006-mc1-cc1-k66201
    (SORTER-01 at 10000, SORTER+DIAG-02 at 20000): with the shipped database 700 writes the coupler's four areas
    into Modules.csv; a project database whose OUT area holds 2 bytes makes 400 halt before the IF_ workbooks and
    the SCL, and the 510 leg before PLCTags - while 700 still writes the areas (the overflow is the layout's)."""
    import csv
    import glob
    from pipeline5.config.paths import DEVICE_TYPES_DB_DEFAULT
    out_root = p5paths._BUILTIN_OUTPUT

    def areas():
        path = glob.glob(os.path.join(out_root, "**", "Modules.csv"), recursive=True)
        ok(len(path) == 1, path)
        text = _unstamped(open(path[0], "rb").read()).decode("utf-8")
        rows = [r for r in csv.reader(line for line in text.splitlines() if line and not line.startswith("#"))]
        return [(r[0], r[1], r[2], r[3], r[4], r[5]) for r in rows if r[3].startswith("TransferArea")]

    expected = [("n0006-mc1-cc1-k66201", "1", "SORTER-01_IN", "TransferArea-IN", "10000", ""),
                ("n0006-mc1-cc1-k66201", "2", "SORTER-01_OUT", "TransferArea-OUT", "", "10000"),
                ("n0006-mc1-cc1-k66201", "3", "SORTER-02_IN", "TransferArea-IN", "20000", ""),
                ("n0006-mc1-cc1-k66201", "4", "SORTER-02_OUT", "TransferArea-OUT", "", "20000")]
    host = _Host()
    SYSTEM.handlers["hardware"](host.ctx())
    eq(host.halted, False)
    eq(areas(), expected, "the shipped database: two areas per interface, I / Q Addr = the base")

    small = os.path.join(p5paths._BUILTIN_DATABASE, "SmallAreas.csv")
    with open(DEVICE_TYPES_DB_DEFAULT, encoding="utf-8") as src, open(small, "w", encoding="utf-8") as dst:
        dst.write(src.read().replace("),LocalToPartnerLength=128,,", "),LocalToPartnerLength=2,,"))
    original = config.load_params
    config.load_params = lambda path=None: dict(original(path), device_types_db=small)
    try:
        host = _Host()
        SYSTEM.handlers["interfaces"](host.ctx())
        eq(host.halted, True, "400 halts on the overflow")
        overflow = [f for batch in host.rendered for f in batch if f.type == "if_area_overflow"]
        eq(sorted(f.detail.split(":")[0] for f in overflow), ["interface 'SORTER+DIAG-02'", "interface 'SORTER-01'"])
        eq(glob.glob(os.path.join(out_root, "**", "IF_*.xlsx"), recursive=True), [], "no IF_ workbook")
        eq(glob.glob(os.path.join(out_root, "**", "*.scl"), recursive=True), [], "no MachineInterfaces SCL")
        ok("RSLT" not in host.levels(), host.lines)

        host = _Host()
        SYSTEM.handlers["data_blocks"](host.ctx(), only=510)
        eq(host.halted, True, "the 510 leg halts on the same interfaces")
        eq(glob.glob(os.path.join(out_root, "**", "PLCTags*.xlsx"), recursive=True), [], "no PLCTags")

        host = _Host()
        SYSTEM.handlers["hardware"](host.ctx())
        eq(host.halted, False, "700 is not the layout's judge")
        eq(areas(), expected, "the areas are still written")
    finally:
        config.load_params = original


if __name__ == "__main__":
    import sys
    sys.exit(run("siemens_main_handlers", [
        ("run_plan_executes_end_to_end", _sandboxed(test_run_plan_executes_end_to_end)),
        ("staging_fires_reaction_hooks", _sandboxed(test_staging_fires_reaction_hooks)),
        ("configured_rules_fire_through_the_real_run_plan",
         _sandboxed(test_configured_rules_fire_through_the_real_run_plan)),
        ("halted_staging_lists_the_deferred_trail_and_records_nothing",
         _sandboxed(test_halted_staging_lists_the_deferred_trail_and_records_nothing)),
        ("a_corrupt_reaction_record_is_reported_not_crashed_or_overwritten",
         _sandboxed(test_a_corrupt_reaction_record_is_reported_not_crashed_or_overwritten)),
        ("a_cascade_within_one_hook_through_the_real_run_plan",
         _sandboxed(test_a_cascade_within_one_hook_through_the_real_run_plan)),
        ("the_builder_models_what_the_after_300_fire_gets", _sandboxed(test_the_builder_models_what_the_after_300_fire_gets)),
        ("the_builder_lints_every_leg_that_fires_after_300",
         _sandboxed(test_the_builder_lints_every_leg_that_fires_after_300)),
        ("the_builder_gates_as_the_run_plan_does", _sandboxed(test_the_builder_gates_as_the_run_plan_does)),
        ("a_downgraded_no_io_sheet_never_saves_over_the_record",
         _sandboxed(test_a_downgraded_no_io_sheet_never_saves_over_the_record)),
        ("text_utf8_cannot_store_never_crashes_the_run", _sandboxed(test_text_utf8_cannot_store_never_crashes_the_run)),
        ("a_field_name_or_a_quoted_text_utf8_cannot_store_never_wipes_the_record",
         _sandboxed(test_a_field_name_or_a_quoted_text_utf8_cannot_store_never_wipes_the_record)),
        ("a_device_name_target_is_refused_through_the_run_plan",
         _sandboxed(test_a_device_name_target_is_refused_through_the_run_plan)),
        ("spawns_are_saved_beside_a_ragged_record", _sandboxed(test_spawns_are_saved_beside_a_ragged_record)),
        ("a_hidden_record_file_is_saved_whole_through_the_run_plan",
         _sandboxed(test_a_hidden_record_file_is_saved_whole_through_the_run_plan)),
        ("a_blank_range_end_invents_nothing_on_the_real_fixture",
         _sandboxed(test_a_blank_range_end_invents_nothing_on_the_real_fixture)),
        ("generated_signals_reach_generation", _sandboxed(test_generated_signals_reach_generation)),
        ("generated_signals_change_only_their_own_outputs", _sandboxed(test_generated_signals_change_only_their_own_outputs)),
        ("lone_buttons_in_a_fresh_project", _sandboxed(test_lone_buttons_in_a_fresh_project)),
        ("cascade_orphan_and_validation_see_absorbed_rows", _sandboxed(test_cascade_orphan_and_validation_see_absorbed_rows)),
        ("310_leg_appends_reaction_findings_to_the_existing_record",
         _sandboxed(test_310_leg_appends_reaction_findings_to_the_existing_record)),
        ("sub_phase_dispatch", _sandboxed(test_sub_phase_dispatch)),
        ("run_all_writes_a_ready_op5_workspace", _sandboxed(test_run_all_writes_a_ready_op5_workspace)),
        ("a_transfer_area_fail_halts_the_interface_outputs",
         _sandboxed(test_a_transfer_area_fail_halts_the_interface_outputs)),
    ]))
