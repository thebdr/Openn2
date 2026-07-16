"""Phase 510 - project the resolved I/O signals + the interface_elements to the PLCTags.xlsx BuilderData
surface (what OP4 imports). A PURE projection of the `signals` + `interface_elements` SSOT tables - no
I/O List read-back: the interface addresses are the `io_address_side1` column phase 400 already STORED
(equal to the value 400e seeds into the inserted IF_ sheets - verified). Clean-room port of PL3's
signals.generate_io_tags / build_io_tags / interface_tags / write_plc_tags.

Two tag sources, mixed into one workbook (sorted by Path = tag table):
  (a) resolved I/O signals  - Name=name_in_tagtable, Path=tagtable, Address=%-bit, Comment=io_comment,
      Data Type=Bool (PL3 hard-codes Bool for direct I/O points);
  (b) interface tags        - one per interface_elements row: Name=signal_name, Path=IF_<instance>,
      Address=%io_address_side1, Comment=description, Data Type=Bool/Word.

The DUPLICATE-TAG GATE: two tags landing on the same (tag table, name) [case-insensitive] break the TIA
import, so each 2nd+ occurrence is an `iotag_duplicate` FAIL linking BOTH producing I/O-List rows
(location = the duplicate's row, location2 = the first occurrence's) and the workbook is NOT written
(the raw-FAIL guard). The same name on TWO tables stays legal: one signal mirrors into several IF_
tables by design.

Output -> `config.io_tags_dir()`/PLCTags.xlsx (sheets "PLC Tags" + "TagTable Properties"; all values text).
"""
from __future__ import annotations

import os
from collections import Counter

from openpyxl import Workbook

from pipeline5 import config
from pipeline5.findings import gate as run
from pipeline5.truth.database import Database
from pipeline5.findings.finding import Finding, record_standalone
from pipeline5.truth import identity
from pipeline5.truth.signals import signals_table


from pipeline5.phases.io_tags.collector import (
    IF_PREFIX,
    PROP_COLUMNS,
    TAG_COLUMNS,
    TAG_TABLE_FILE,
    duplicate_findings,
    interface_tags,
    io_signal_tags,
)

def _append_text_row(ws, values) -> None:
    """Append a row; force data_type='s' on any cell starting with =/+/- so a tag/FLD/address isn't read
    as a formula (e.g. a Name like `... [ =S1+MS1.CC1-F09001 ]`)."""
    ws.append(values)
    r = ws.max_row
    for c, v in enumerate(values, 1):
        if isinstance(v, str) and v[:1] in ("=", "+", "-"):
            ws.cell(r, c).data_type = "s"


def _group_tables(tags) -> dict:
    """{path: [tag, ...]} preserving insertion order within each table."""
    out: dict = {}
    for t in tags:
        out.setdefault(t["path"], []).append(t)
    return out


def write_plc_tags(tags, out_dir) -> str:
    """Write PLCTags.xlsx: a "PLC Tags" sheet (one row per tag, sorted by Path) + a "TagTable Properties"
    sheet (one row per distinct Path). All values text; Hmi flags the literal 'True'. Returns the path."""
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, TAG_TABLE_FILE)
    tables = _group_tables(tags)
    wb = Workbook()
    ws = wb.active
    ws.title = "PLC Tags"
    ws.append(TAG_COLUMNS)
    for table in sorted(tables):
        for t in tables[table]:
            _append_text_row(ws, [t["name"], table, t["data_type"], t["address"], t["comment"],
                                  "True", "True", "True", "", ""])
    props = wb.create_sheet("TagTable Properties")
    props.append(PROP_COLUMNS)
    for table in sorted(tables):
        _append_text_row(props, [table, "", ""])
    wb.save(path)
    wb.close()
    return path


def project(database: Database | None = None, out_dir: str | None = None) -> dict:
    """Project the two SSOT sources to `out_dir`/PLCTags.xlsx (out_dir defaults to `config.io_tags_dir()`).
    Returns {'path', 'total', 'io_count', 'iface_count', 'tables', 'findings'} - `findings` = the
    `iotag_no_address` WARNs + the `iotag_duplicate` FAILs (the caller renders/gates them). On a duplicate
    tag the workbook is NOT written (`run.has_blocking` raw-FAIL guard - never hand OP4 a broken import
    surface; 'path' comes back '') and the FAILs alone are RECORDED to `validation_issues.csv`
    (`record_standalone`, the ph200 seam - the log's [FAIL]->Findings jump needs the row); the WARNs stay
    render-only as before."""
    if database is None:
        colmap = config.load_column_map("IoList")
        database = Database([signals_table([m["canonical"] for m in colmap])]).load(config.database_dir())
    out_dir = out_dir or config.io_tags_dir()
    io = io_signal_tags(database)
    iface, findings = interface_tags(database)
    tags = io + iface
    findings += duplicate_findings(tags)
    if run.has_blocking(findings):                    # duplicate tags -> record the FAILs, write NOTHING
        record_standalone([f for f in findings if f.type == "iotag_duplicate"])
        path = ""
    else:
        path = write_plc_tags(tags, out_dir)
    return {"path": path, "total": len(tags), "io_count": len(io), "iface_count": len(iface),
            "tables": sorted({t["path"] for t in tags}), "findings": findings}
