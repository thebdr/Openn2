"""Annotate staged rows with their C&E data: `matrix_areas` + the C&E-side FLD
(`ce_functional_unit`/`ce_location`/`ce_device`) + `numerazione_linea` + `areas_description`. Derived
from the Cause & Effect workbook (best-effort: blanks when the C&E doc is absent). Clean-room port of
PL3's matrix, reading the nested `matrix_params` and storing the multi-valued fields as REAL list cells.

  * input signal  (I-address): matched by address on the CAUSE&EFFECT MATRIX -> the area columns marked
    "X" (area columns = headers matching the area regex) + the matrix-side FLD (cols J/K/L).
  * output signal (Q-address): matched by address on the AREA n sheets -> those sheet names +
    `numerazione_linea`.
Paired channels of one device share the union of their C&E data (the matrix lists the device once).
"""
from __future__ import annotations

import os
import re

from openpyxl.utils import column_index_from_string as _ci

from pipeline5.core import config
from pipeline5.io import xlsx_reader as workbook


def _norm(value) -> str:
    """Address/marker key: whitespace removed, upper-cased."""
    return "".join(str(value or "").split()).upper()


def _s(value) -> str:
    return str(value if value is not None else "").strip()


def _area_regex(params):
    pattern = config.get_param(params, "matrix_params.area_sheets.name")
    if not pattern:
        return None
    try:
        return re.compile(config.js_to_re(pattern), re.IGNORECASE)
    except re.error:
        return None


def area_lookup(params: dict) -> tuple:
    """(inputs, outputs, area_desc):
      inputs:  norm(I-addr) -> {areas:[names], ce_functional_unit, ce_location, ce_device}
      outputs: norm(Q-addr) -> {areas:[sheet names], numerazione_linea}
      area_desc: area name -> its CONCEPT description
    Empty when the C&E document is missing."""
    matrix_path = params.get("matrix_path")
    if not matrix_path or not os.path.exists(matrix_path):
        return {}, {}, {}
    wb = workbook.open_workbook(matrix_path, data_only=True, read_only=True)
    inputs, outputs, area_desc = {}, {}, {}
    cemap = {m["canonical"]: m["column"] for m in config.load_column_map("CE")}
    areamap = {m["canonical"]: m["column"] for m in config.load_column_map("AREA")}
    area_re = _area_regex(params)

    # --- CAUSE&EFFECT MATRIX: input address -> areas marked X + the matrix-side FLD ---
    area_cols = []   # ordered (col_index, area_name) of the matrix EFFECT/area columns
    msheet = config.resolve_sheet(config.get_param(params, "matrix_params.ce_sheet.name"), wb.sheetnames)
    if msheet:
        ws = wb[msheet]
        hr = int(config.get_param(params, "matrix_params.ce_sheet.header_row", 2) or 2)
        dr = int(config.get_param(params, "matrix_params.ce_sheet.data_row", 5) or 5)
        addr_c = _ci(cemap["address"])
        fu_c, loc_c, dev_c = _ci(cemap["functional_unit"]), _ci(cemap["location"]), _ci(cemap["device"])
        for c in range(1, ws.max_column + 1):
            name = _s(ws.cell(hr, c).value)
            if name and (area_re is None or area_re.search(name)):
                area_cols.append((c, name))
        for r in range(dr, ws.max_row + 1):
            addr = _norm(ws.cell(r, addr_c).value)
            if not addr:
                continue
            inputs[addr] = {
                "areas": [name for c, name in area_cols if _s(ws.cell(r, c).value).upper() == "X"],
                "ce_functional_unit": _s(ws.cell(r, fu_c).value),
                "ce_location": _s(ws.cell(r, loc_c).value),
                "ce_device": _s(ws.cell(r, dev_c).value),
            }

    # --- AREA n sheets: output Q address -> sheet names + numerazione_linea ---
    area_sheets = config.resolve_sheets(config.get_param(params, "matrix_params.area_sheets.name"), wb.sheetnames)
    out_addr_c = _ci(areamap["address"])
    num_c = _ci(areamap["numerazione_linea"]) if "numerazione_linea" in areamap else None
    adr = int(config.get_param(params, "matrix_params.area_sheets.data_row", 4) or 4)
    for sheet in area_sheets:
        a = wb[sheet]
        for r in range(adr, a.max_row + 1):
            q = _norm(a.cell(r, out_addr_c).value)
            if not q:
                continue
            entry = outputs.setdefault(q, {"areas": [], "numerazione_linea": ""})
            entry["areas"].append(sheet.strip())
            if not entry["numerazione_linea"] and num_c:
                entry["numerazione_linea"] = _s(a.cell(r, num_c).value)

    # --- CONCEPT sheet: one description per area, positionally aligned with the matrix area columns ---
    concept_name = config.get_param(params, "matrix_params.concept_sheet.name")
    csheet = config.resolve_sheet(concept_name, wb.sheetnames) if concept_name else None
    if csheet and area_cols:
        cs = wb[csheet]
        crow = int(config.get_param(params, "matrix_params.concept_sheet.header_row", 2) or 2)
        descs = [" ".join(_s(cs.cell(crow, c).value).split())
                 for c in range(1, cs.max_column + 1) if _s(cs.cell(crow, c).value)]
        for name, desc in zip([nm for _, nm in area_cols], descs):
            area_desc[name] = desc
    wb.close()
    return inputs, outputs, area_desc


_CE_FIELDS = ("ce_functional_unit", "ce_location", "ce_device", "numerazione_linea")


def annotate(params: dict, rows: list) -> list:
    """Set matrix_areas (list) + ce_FLD + numerazione_linea + areas_description (list) on every staged
    row by its I/Q address. Paired channels (same pair_key + FLD) share the union of areas + each
    other's C&E fields."""
    inputs, outputs, area_desc = area_lookup(params)
    found = {}
    for row in rows:
        bit = _norm(row.get("bit"))
        head = bit[:1]
        if head == "I":
            found[id(row)] = dict(inputs.get(bit) or {})
        elif head == "Q":
            found[id(row)] = dict(outputs.get(bit) or {})
        else:
            found[id(row)] = {}

    # pair inheritance: union the areas + share the ce fields across a device's channels
    groups: dict = {}
    for row in rows:
        pair_key = ((row.get("type") or {}).get("pair_key") or "").strip().upper()
        if not pair_key:
            continue
        key = (pair_key, _norm(row.get("functional_unit")), _norm(row.get("location")), _norm(row.get("device")))
        groups.setdefault(key, []).append(row)
    for members in groups.values():
        union, merged = [], {}
        for m in members:
            d = found[id(m)]
            for area in d.get("areas", []):
                if area not in union:
                    union.append(area)
            for field in _CE_FIELDS:
                if not merged.get(field) and d.get(field):
                    merged[field] = d[field]
        for m in members:
            if union:
                found[id(m)]["areas"] = union
            for field in _CE_FIELDS:
                if not found[id(m)].get(field) and merged.get(field):
                    found[id(m)][field] = merged[field]

    for row in rows:
        d = found[id(row)]
        areas = d.get("areas", [])
        row["matrix_areas"] = list(areas)                                          # a real list cell
        row["areas_description"] = [area_desc[a] for a in areas if area_desc.get(a)]
        for field in _CE_FIELDS:
            row[field] = d.get(field, "")
    return rows
