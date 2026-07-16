"""THE one shared workbook reader. Every phase opens sheets through here (ported from PL3, verbatim
behavior - it is architecture-independent: PL4 reads the source documents exactly as PL3 did).

It owns the openpyxl load-mode policy, regex sheet resolution, by-position cell access, strike
detection, and the header check (which normalizes newline->space, then prefix-matches - real I/O List
headers wrap across lines and carry suffixes like "(1 part)").

Load modes:
- read (default): data_only=True, read_only=False -> cached values AND cell fonts (strike) available.
- write: load_for_write(path) -> data_only=False, read_only=False -> formulas preserved for an
  edit+save path (the populator).
"""
from __future__ import annotations

import os
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from pipeline5.core import config


class FileLockedError(Exception):
    """A workbook could not be opened because it is locked (open in Excel). Carries the offending path;
    the GUI turns it into a retry/abort prompt, the engine/CLI into a clear FAIL."""

    def __init__(self, path: str):
        self.path = path
        super().__init__(f"{os.path.basename(path)} may already be in use (open in Excel?)")


def _load(path: str, **kwargs):
    """load_workbook, but a PermissionError (file locked / open in Excel) becomes FileLockedError."""
    try:
        return load_workbook(path, **kwargs)
    except PermissionError as error:
        raise FileLockedError(path) from error


def _normalize_header(value) -> str:
    """Newline -> space, then whitespace-collapse + strip. The canonical header form for matching."""
    text = str(value if value is not None else "").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def _column_index(column) -> int:
    """Accept a 1-based int or a column letter ('A', 'AB'); return the 1-based index."""
    if isinstance(column, int):
        return column
    return column_index_from_string(str(column).strip())


def open_workbook(path: str, *, data_only: bool = True, read_only: bool = False):
    return _load(path, data_only=data_only, read_only=read_only)


def load_for_write(path: str):
    """Formula-preserving load for editing + saving (the populator). Not data_only, not read_only."""
    return _load(path, data_only=False, read_only=False)


def available_sheets(path: str) -> list:
    workbook = _load(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


class SheetView:
    """A positional view over one worksheet. Cells are addressed by (row, col) where col is a 1-based
    int or a column letter. Row/col are 1-based to match Excel + openpyxl."""

    def __init__(self, worksheet, doc_label: str = "", header_row: int = 1, first_data_row: int | None = None):
        self.ws = worksheet
        self.name = worksheet.title
        self.doc = doc_label
        self.header_row = header_row
        self.first_data_row = first_data_row if first_data_row is not None else header_row + 1

    # --- dimensions ------------------------------------------------------------------------------ #
    @property
    def max_row(self) -> int:
        return self.ws.max_row

    @property
    def max_col(self) -> int:
        return self.ws.max_column

    # --- values ---------------------------------------------------------------------------------- #
    def cell(self, row: int, column):
        """The cell value (by position). None when empty."""
        return self.ws.cell(row=row, column=_column_index(column)).value

    def text(self, row: int, column) -> str:
        """The cell value as a stripped string ('' when empty)."""
        value = self.cell(row, column)
        return "" if value is None else str(value).strip()

    # --- headers --------------------------------------------------------------------------------- #
    def header(self, column) -> str:
        """The normalized (newline->space, collapsed) header text at `column`."""
        return _normalize_header(self.cell(self.header_row, column))

    def header_matches(self, column, expected: str) -> bool:
        """True if the header at `column` starts with `expected` (both normalized, case-insensitive)."""
        return self.header(column).lower().startswith(_normalize_header(expected).lower())

    # --- strikethrough --------------------------------------------------------------------------- #
    def cell_struck(self, row: int, column) -> bool:
        font = self.ws.cell(row=row, column=_column_index(column)).font
        return bool(font and font.strike)

    def row_struck(self, row: int, columns=None) -> bool:
        """True if ANY cell in the row (over `columns`, default the used columns) is struck through."""
        columns = columns if columns is not None else range(1, self.max_col + 1)
        return any(self.cell_struck(row, column) for column in columns)

    # --- iteration + links ----------------------------------------------------------------------- #
    def data_rows(self):
        """Row indices from first_data_row through max_row (inclusive)."""
        return range(self.first_data_row, self.max_row + 1)

    def location(self, row: int, column) -> str:
        """A clickable 'sheet!cell' link string, e.g. 'NET SAFETY 50!G7'. The workbook name is NOT
        included (carried separately as `self.doc` for the GUI)."""
        return f"{self.name}!{get_column_letter(_column_index(column))}{row}"

    def close(self):
        try:
            self.ws.parent.close()
        except Exception:  # noqa: BLE001
            pass


def open_sheet(path: str, sheet_pattern, header_row: int = 1, *, doc_label: str | None = None,
               first_data_row: int | None = None, data_only: bool = True,
               read_only: bool = False) -> SheetView:
    """Open the first sheet of `path` matching `sheet_pattern` (regex via config.resolve_sheet).
    Raises ValueError when no sheet matches. The caller may .close() the view when done."""
    workbook = open_workbook(path, data_only=data_only, read_only=read_only)
    name = config.resolve_sheet(sheet_pattern, workbook.sheetnames)
    if name is None:
        workbook.close()
        raise ValueError(f"no sheet matching {sheet_pattern!r} in {os.path.basename(path)} "
                         f"(have: {workbook.sheetnames})")
    return SheetView(workbook[name], doc_label or os.path.basename(path), header_row, first_data_row)


def open_sheets(path: str, sheet_pattern, header_row: int = 1, *, doc_label: str | None = None,
                first_data_row: int | None = None, data_only: bool = True,
                read_only: bool = False) -> list:
    """Open ALL sheets matching `sheet_pattern` (the I/O List sheet may be a list/regex). Returns a list
    of SheetView (possibly empty). Each view shares the one underlying workbook."""
    workbook = open_workbook(path, data_only=data_only, read_only=read_only)
    names = config.resolve_sheets(sheet_pattern, workbook.sheetnames)
    label = doc_label or os.path.basename(path)
    return [SheetView(workbook[name], label, header_row, first_data_row) for name in names]
