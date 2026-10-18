"""Annotate staged rows with their C&E data: `matrix_areas` + the C&E-side FLD
(`ce_functional_unit`/`ce_location`/`ce_device`) + `numerazione_linea` + `areas_description`. Derived
from the Cause & Effect workbook (best-effort: blanks when the C&E doc is absent). Clean-room port of
PL3's matrix, reading the nested `matrix_params` and storing the multi-valued fields as REAL list cells.

  * input signal  (I-address): matched by address on the CAUSE&EFFECT MATRIX -> the area columns marked
    "X" (area columns = headers matching the area regex) + the matrix-side FLD (cols J/K/L).
  * output signal (Q-address): matched by address on the AREA n sheets -> those sheet names +
    `numerazione_linea`.
Paired channels of one device share the union of their C&E data (the matrix lists the device once).

The AREAS themselves ([[C-040]]): `area_hierarchy` numbers every area - a plain area its own (`AREA 3` -> 3), with
`matrix_params.sub_areas` on a sub-area `<parent>_<part>` its parent's x 1000 + its ordinal (`AREA 1_TRIB` ->
1001) - and `concept_descriptions` matches each CONCEPT cell to the area it NAMES (position only for a cell
naming none); a sub-area without its own cell takes its parent's description. Staging stores the result in
the `areas` table (`iolist._store_areas`).
"""
from __future__ import annotations

import os
import re

from openpyxl.utils import column_index_from_string as _ci

from pipeline5 import config
from pipeline5.documents import xlsx_reader as workbook
from pipeline5.truth.addresses import key as _addr_key   # [[C-022]]: the canonical spelling


def _norm(value) -> str:
    """Address/marker key: whitespace removed, upper-cased."""
    return "".join(str(value or "").split()).upper()


def _s(value) -> str:
    return str(value if value is not None else "").strip()


def _area_regex(params):
    pattern = config.get_param(params, "matrix_params.area_sheets.name")
    if not pattern:
        return None
    try:
        return re.compile(config.js_to_re(pattern), re.IGNORECASE)
    except re.error:
        return None


def _name_key(name) -> str:
    """An area name compared whitespace- and case-insensitively (a header 'AREA 1' == a CONCEPT 'area  1')."""
    return " ".join(str(name or "").split()).upper()


def _area_number(name: str, area_re) -> int | None:
    """A PLAIN area's own number: the digits of the area pattern's match when that match IS the whole name
    (`AREA 3` -> 3; `AREA 1_TRIB` is no plain area -> None); without a pattern, a name ending in its digits."""
    key = " ".join(str(name).split())
    if area_re is not None:
        m = area_re.search(key)
        if not m or m.group(0).strip().upper() != key.upper():
            return None
        digits = re.findall(r"\d+", m.group(0))
    else:
        m = re.fullmatch(r"\D*?(\d+)", key)
        digits = [m.group(1)] if m else []
    return int(digits[-1]) if digits else None


def area_hierarchy(names: list, area_re, sub_areas: bool) -> list:
    """The declared area hierarchy ([[C-040]]): one dict per area name, in the given order - `number` (a plain area
    its own; a SUB-AREA its parent's x 1000 + its ordinal among that parent's sub-areas, `AREA 1_TRIB` -> 1001;
    '' when neither), `parent` / `part` (a sub-area: `<parent>_<part>` with `<parent>` a plain area of `names`, the
    longest such) and `order` (1-based). `sub_areas` off = every name an independent area (no parent)."""
    plain = {n: _area_number(n, area_re) for n in names}
    parents = sorted((n for n in names if plain[n] is not None), key=len, reverse=True)
    out, ordinals = [], {}
    for i, name in enumerate(names, start=1):
        parent = ""
        if sub_areas and plain[name] is None:
            parent = next((p for p in parents if _name_key(name).startswith(_name_key(p) + "_")), "")
        if parent:
            ordinals[parent] = ordinals.get(parent, 0) + 1
            number = plain[parent] * 1000 + ordinals[parent]
            part = " ".join(name.split())[len(" ".join(parent.split())) + 1:]
        else:
            number, part = plain[name], ""
        out.append({"name": name, "number": "" if number is None else number, "parent": parent,
                    "part": part, "order": i})
    return out


def _names_area(text: str, name: str) -> bool:
    """A CONCEPT cell `text` NAMES `name` when it starts with it at a word boundary ('AREA 1 (TRIB+TRIH)' names
    AREA 1; 'AREA 12 ...' and 'AREA 1_TRIB ...' do not)."""
    t, n = _name_key(text), _name_key(name)
    if not t.startswith(n):
        return False
    nxt = t[len(n):len(n) + 1]
    return not nxt or not (nxt.isalnum() or nxt == "_")


def concept_descriptions(cells: list, column_areas: list, all_areas: list) -> dict:
    """{area name: CONCEPT description} ([[C-040]], every project): a cell describes the area it NAMES (the longest
    name it starts with, at a word boundary - first such cell wins); the cells naming no area fall back to
    position - the i-th of them to the i-th matrix area column still without one (a CONCEPT sheet whose cells
    name no area reads exactly as before)."""
    desc, unnamed = {}, []
    candidates = sorted(all_areas, key=lambda n: len(_name_key(n)), reverse=True)
    for text in cells:
        hit = next((n for n in candidates if _names_area(text, n)), None)
        if hit is None:
            unnamed.append(text)
        elif hit not in desc:
            desc[hit] = text
    for name, text in zip([n for n in column_areas if n not in desc], unnamed):
        desc[name] = text
    return desc


def area_lookup(params: dict) -> tuple:
    """(inputs, outputs, area_desc, areas):
      inputs:  norm(I-addr) -> {areas:[names], ce_functional_unit, ce_location, ce_device}
      outputs: norm(Q-addr) -> {areas:[sheet names], numerazione_linea}
      area_desc: area name -> its CONCEPT description (a sub-area without its own: its parent's)
      areas: the declared hierarchy (`area_hierarchy` + `description`), matrix columns first, then any area only
             an AREA sheet names
    Empty when the C&E document is missing."""
    matrix_path = params.get("matrix_path")
    if not matrix_path or not os.path.exists(matrix_path):
        return {}, {}, {}, []
    wb = workbook.open_workbook(matrix_path, data_only=True, read_only=True)
    inputs, outputs, area_desc = {}, {}, {}
    cemap = {m["canonical"]: m["column"] for m in config.load_column_map("CE")}
    areamap = {m["canonical"]: m["column"] for m in config.load_column_map("AREA")}
    area_re = _area_regex(params)

    # --- CAUSE&EFFECT MATRIX: input address -> areas marked X + the matrix-side FLD ---
    area_cols = []   # ordered (col_index, area_name) of the matrix EFFECT/area columns
    msheet = config.resolve_sheet(config.get_param(params, "matrix_params.ce_sheet.name"), wb.sheetnames)
    if msheet:
        ws = wb[msheet]
        hr = int(config.get_param(params, "matrix_params.ce_sheet.header_row", 2) or 2)
        dr = int(config.get_param(params, "matrix_params.ce_sheet.data_row", 5) or 5)
        addr_c = _ci(cemap["address"])
        fu_c, loc_c, dev_c = _ci(cemap["functional_unit"]), _ci(cemap["location"]), _ci(cemap["device"])
        for c in range(1, ws.max_column + 1):
            name = _s(ws.cell(hr, c).value)
            if name and (area_re is None or area_re.search(name)):
                area_cols.append((c, name))
        for r in range(dr, ws.max_row + 1):
            addr = _addr_key(ws.cell(r, addr_c).value)
            if not addr:
                continue
            inputs[addr] = {
                "areas": [name for c, name in area_cols if _s(ws.cell(r, c).value).upper() == "X"],
                "ce_functional_unit": _s(ws.cell(r, fu_c).value),
                "ce_location": _s(ws.cell(r, loc_c).value),
                "ce_device": _s(ws.cell(r, dev_c).value),
            }

    # --- AREA n sheets: output Q address -> sheet names + numerazione_linea ---
    area_sheets = config.resolve_sheets(config.get_param(params, "matrix_params.area_sheets.name"), wb.sheetnames)
    out_addr_c = _ci(areamap["address"])
    num_c = _ci(areamap["numerazione_linea"]) if "numerazione_linea" in areamap else None
    adr = int(config.get_param(params, "matrix_params.area_sheets.data_row", 4) or 4)
    for sheet in area_sheets:
        a = wb[sheet]
        for r in range(adr, a.max_row + 1):
            q = _addr_key(a.cell(r, out_addr_c).value)
            if not q:
                continue
            entry = outputs.setdefault(q, {"areas": [], "numerazione_linea": ""})
            entry["areas"].append(sheet.strip())
            if not entry["numerazione_linea"] and num_c:
                entry["numerazione_linea"] = _s(a.cell(r, num_c).value)

    # --- the areas: the matrix columns, then any area only an AREA sheet names; their hierarchy ([[C-040]]) ---
    names = [nm for _, nm in area_cols]
    for sheet in area_sheets:
        if sheet.strip() not in names:
            names.append(sheet.strip())
    areas = area_hierarchy(names, area_re, bool(config.get_param(params, "matrix_params.sub_areas", False)))

    # --- CONCEPT sheet: one description per area, matched by NAME (position for the cells naming none) ---
    concept_name = config.get_param(params, "matrix_params.concept_sheet.name")
    csheet = config.resolve_sheet(concept_name, wb.sheetnames) if concept_name else None
    if csheet and area_cols:
        cs = wb[csheet]
        crow = int(config.get_param(params, "matrix_params.concept_sheet.header_row", 2) or 2)
        descs = [" ".join(_s(cs.cell(crow, c).value).split())
                 for c in range(1, cs.max_column + 1) if _s(cs.cell(crow, c).value)]
        area_desc.update(concept_descriptions(descs, [nm for _, nm in area_cols], names))
    for a in areas:                                   # a sub-area without its own cell: its parent's description
        if a["parent"] and not area_desc.get(a["name"]) and area_desc.get(a["parent"]):
            area_desc[a["name"]] = area_desc[a["parent"]]
        a["description"] = area_desc.get(a["name"], "")
    wb.close()
    return inputs, outputs, area_desc, areas


_CE_FIELDS = ("ce_functional_unit", "ce_location", "ce_device", "numerazione_linea")


def annotate(params: dict, rows: list, areas: list | None = None) -> list:
    """Set matrix_areas (list) + ce_FLD + numerazione_linea + areas_description (list) on every staged
    row by its I/Q address. Paired channels (same pair_key + FLD) share the union of areas + each
    other's C&E fields. `areas`, when given, receives the declared area hierarchy ([[C-040]])."""
    inputs, outputs, area_desc, hierarchy = area_lookup(params)
    if areas is not None:
        areas[:] = hierarchy
    found = {}
    for row in rows:
        bit = _addr_key(row.get("bit"))
        head = bit[:1]
        if head == "I":
            found[id(row)] = dict(inputs.get(bit) or {})
        elif head == "Q":
            found[id(row)] = dict(outputs.get(bit) or {})
        else:
            found[id(row)] = {}

    # pair inheritance: union the areas + share the ce fields across a device's channels
    groups: dict = {}
    for row in rows:
        pair_key = ((row.get("type") or {}).get("pair_key") or "").strip().upper()
        if not pair_key:
            continue
        key = (pair_key, _norm(row.get("functional_unit")), _norm(row.get("location")), _norm(row.get("device")))
        groups.setdefault(key, []).append(row)
    for members in groups.values():
        union, merged = [], {}
        for m in members:
            d = found[id(m)]
            for area in d.get("areas", []):
                if area not in union:
                    union.append(area)
            for field in _CE_FIELDS:
                if not merged.get(field) and d.get(field):
                    merged[field] = d[field]
        for m in members:
            if union:
                found[id(m)]["areas"] = union
            for field in _CE_FIELDS:
                if not found[id(m)].get(field) and merged.get(field):
                    found[id(m)][field] = merged[field]

    for row in rows:
        d = found[id(row)]
        areas = d.get("areas", [])
        row["matrix_areas"] = list(areas)                                          # a real list cell
        row["areas_description"] = [area_desc[a] for a in areas if area_desc.get(a)]
        for field in _CE_FIELDS:
            row[field] = d.get(field, "")
    return rows
