"""Interface generation (phase 400): per-machine interface I/O tables from the staged database.

A staged row whose Script Type == "IOC" marks a machine interface; its Index has the form
MACHINETYPE-nn (e.g. SORTER-01). The machine type (SORTER) selects a sheet in the MachineInterfaces
template workbook (one sheet per type, plus a "<GENERIC>" fallback when the type has no dedicated
sheet); the trailing digits (01) are the instance index. For each IOC instance the template is
copied, every sheet but the chosen type sheet is removed, and the IOC row's base address + base node
are plugged into the Side-1 row, the index into the Index column, and in place of every "<index>"
token. Output: <interfaces_dir>/IF_<instance>.xlsx, one per instance, PRESERVED if it already exists.

interfaces_dir lives under ProjectDocumentation - this is documentation/intermediate, NOT the
Open2App BuilderData import surface. The template's LET/SUBSTITUTE formulas + Excel tables survive
the openpyxl round-trip; the engineer fills/extends Side 2 by hand.

Clean-room rebuild of Pipeline2's interfaces/interface_tool.py (reference for INTENT only). Two entry
points share one primitive (generate_one):
  - generate(rows, ...)   the 410 auto path - one interface per IOC row of the database;
  - generate_one(...)     the primitive the 430 "Generate Custom Interface..." popup will call (GUI
                          milestone): one interface from an explicit machine type / base address /
                          node side 1 / node side 2 / index.
"""
from __future__ import annotations
import datetime
import io
import os
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from copy import copy, deepcopy
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter
from openpyxl.worksheet.table import Table, TableColumn

_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

from pipeline3.core import config
from pipeline3.domain import identity

TRIGGER_TYPE = "IOC"
GENERIC_SHEET = "<GENERIC>"
INDEX_TOKEN = "<index>"
_DIAG_RE = re.compile(r"\+\s*DIAG", re.IGNORECASE)   # the +DIAG marker in an IOC Index
_DIR_TO_IO = {"<": "I", ">": "Q"}                    # template "Direction </>" -> I/Q

# Template header labels (row 1) the plug logic binds to (by name, whitespace-trimmed).
_H_SIDE = "Side"
_H_INDEX = "Index"
_H_BASE_NODE = "Base Node"
_H_BASE_ADDR = "Base Address"


def _detect_diag(index_field) -> bool:
    """True iff the Index carries a +DIAG marker (e.g. 'SORTER+DIAG-02' -> mirror all in_diag)."""
    return bool(_DIAG_RE.search(str(index_field or "")))


def parse_instance(index_field) -> tuple[str, str]:
    """'SORTER-01' -> ('SORTER', '01'): machine type = leading letters, index = trailing digits.
    A '+DIAG' marker is stripped first ('SORTER+DIAG-02' -> ('SORTER', '02'))."""
    s = _DIAG_RE.sub("", str(index_field or "")).strip()
    type_match = re.match(r"^[A-Za-z]+", s)
    digit_match = re.search(r"(\d+)\s*$", s)
    return (type_match.group(0) if type_match else "",
            digit_match.group(1) if digit_match else "")


def find_interfaces(rows) -> tuple[list[dict], list[str]]:
    """One record per IOC row of the staged database. Returns (records, warnings).
    Struck rows are already excluded by staging (strike_handling); a row with no Index is warned."""
    records, warnings = [], []
    for r in rows or []:
        if str(r.get("script_type", "") or "").strip().upper() != TRIGGER_TYPE:
            continue
        instance = str(r.get("index", "") or "").strip()
        loc = f"{r.get('_source_sheet', '')}!{r.get('_source_row', '')}"
        if not instance:
            warnings.append(f"IOC row {loc} has no Index - skipped")
            continue
        machine_type, index = parse_instance(instance)
        records.append({
            "instance": instance,                       # full name (keeps +DIAG) -> filename
            "machine_type": machine_type,
            "index": index,                             # the nn, used by interface_mapping
            "is_diag": _detect_diag(instance),
            "base": str(r.get("bit", "") or "").strip(),
            "base_node": str(r.get("id_node", "") or "").strip(),
            "device": str(r.get("device", "") or "").strip(),
            "ip": str(r.get("profinet_ip", "") or "").strip(),
            "source": loc,
        })
    return records, warnings


def template_sheet_names(template_path: str) -> list[str]:
    """The template's sheet names (the machine-type list; <GENERIC> is the fallback). Used by the
    430 popup's machine-type dropdown and by generate()/generate_one() to choose a sheet."""
    wb = load_workbook(template_path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def choose_sheet(sheet_names, machine_type: str) -> str:
    """Exact (case-insensitive) type match, else the generic fallback."""
    for name in sheet_names:
        if name.upper() == str(machine_type or "").upper():
            return name
    return GENERIC_SHEET


def _safe_name(text) -> str:
    bad = '\\/:*?"<>[]'
    return "".join("_" if ch in bad else ch for ch in str(text))[:31]


def _to_int(text):
    """Parse a cell to int, tolerating float-read strings ('10000.0') and floats - staging reads
    numeric cells with data_only=True, and openpyxl widens cached/exported numbers to float, so a
    Bit/ID byte can arrive as '10000.0'. Blank -> None (intentionally not plugged); same intent as
    domain/validation/diagnosis._int."""
    if text is None:
        return None
    s = str(text).strip()
    if s == "":
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return int(float(s))
        except ValueError:
            return None


def _header_index(ws, header_row: int = 1) -> dict:
    return {str(c.value).strip(): c.column for c in ws[header_row] if c.value not in (None, "")}


def _side_row(ws, side_c, value_c, side: int):
    """Row whose 'Side' cell == `side` (1 or 2). With no Side column only Side 1 is determinable
    (the first row carrying a value in value_c, matching the reference); Side 2 -> None."""
    if side_c:
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, side_c).value).strip() in (str(side), f"{side}.0"):
                return r
        return None
    if side == 1 and value_c:
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, value_c).value not in (None, ""):
                return r
    return None


def _plug(ws, base_address, node_side1, node_side2, index) -> None:
    """Plug the Side-1 base address + node, the Side-2 node (only when given), the Index column (on
    the Side-1 and Side-2 rows ONLY), and replace every <index> token. The Side-1 base address only
    is set (Side 2's address stays at the template default) - the auto path leaves Side 2 to be
    filled by hand."""
    headers = _header_index(ws, 1)
    side_c = headers.get(_H_SIDE)
    index_c = headers.get(_H_INDEX)
    base_c = headers.get(_H_BASE_ADDR)
    node_c = headers.get(_H_BASE_NODE)

    base_num = _to_int(base_address)
    n1 = _to_int(node_side1)
    n2 = _to_int(node_side2)

    if base_c and base_num is not None:
        r = _side_row(ws, side_c, base_c, 1)
        if r:
            ws.cell(r, base_c).value = base_num
    if node_c and n1 is not None:
        r = _side_row(ws, side_c, node_c, 1)
        if r:
            ws.cell(r, node_c).value = n1
    if node_c and n2 is not None:
        r = _side_row(ws, side_c, node_c, 2)
        if r:
            ws.cell(r, node_c).value = n2

    index = str(index or "").strip()
    if index_c and index:
        # the Index value belongs to the Side rows (AC2 = Side 1, AC3 = Side 2) - it is NOT scattered
        # across every populated Index cell; the signal rows below carry the instance via <index>.
        for side in (1, 2):
            r = _side_row(ws, side_c, index_c, side)
            if r:
                ws.cell(r, index_c).value = index
    if index:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and INDEX_TOKEN in cell.value:
                    cell.value = cell.value.replace(INDEX_TOKEN, index)


# --------------------------------------------------------------------------------------------- #
# Signal mirroring (the "Interface Mapping" feature): append project signals as a custom data
# block onto the generated interface, laid out by script_type one byte at a time.
# --------------------------------------------------------------------------------------------- #

@dataclass
class _Elem:
    """One element to mirror onto an interface. `mirror_name` is the binding/identity (-> Expression
    Side 1 + the dedup key); `signal_name` is the interface tag name (-> Signal Name Side 1);
    `description` -> Description. `direction` 'Q' (output, '>') or 'I' (input, '<'); offset_byte/bit
    filled by allocate_bytes. diag_* carried from the source signal."""
    script_type: str
    mirror_name: str
    signal_name: str = ""
    description: str = ""
    fu: str = ""
    loc: str = ""
    dev: str = ""
    data_type: str = "BOOL"
    direction: str = "Q"
    source: str = "mirror"
    diag_cabinet: str = ""
    swp_cabinet: str = ""
    diag_bit: str = ""
    src_row: dict = field(default_factory=dict)
    offset_byte: object = None
    bit: object = None


def _mirror_name(row) -> str:
    """The signal's interface representation: the TIA-qualified DB ref, else its tag (shared with
    diagnosis via identity.plc_binding). '' when the signal has neither."""
    return identity.plc_binding(row)


def _idx_key(value):
    """Canonical comparable interface index: the int value when numeric ('01' -> 1), else the
    stripped string, else None (blank). Used CONSISTENTLY by membership (_index_in), the uniqueness
    ERROR, and the dangling-mapping warning so '01' and '1' are the same interface everywhere."""
    n = _to_int(value)
    if n is not None:
        return n
    s = str(value or "").strip()
    return s or None


def _index_in(index, mapping) -> bool:
    """True iff the `|`-list `mapping` (e.g. '01|02') contains `index`, normalized via _idx_key so
    '01' == '1'."""
    want = _idx_key(index)
    if want is None:
        return False
    return any(_idx_key(tok) == want for tok in str(mapping or "").split("|") if tok.strip())


def _load_mirror_cfg() -> dict:
    """The 3 rule tables + the byte gap, loaded once per generate() run."""
    return {
        "db_rules": config.load_rules("datablock_elements_rules.csv"),
        "diag_rules": config.load_rules("diagnosis_logic_rules.csv", config.diagnosis_dir()),
        "if_rules": config.load_interface_elements_rules(),
        "gap": config.INTERFACE_CUSTOM_GAP,
    }


def _rule_tagname(rule, row, member, direction) -> str:
    """A rule element's interface tag name: the rule's interface_tagname template with {member} and
    {direction} bound (+ the source row's fields); falls back to the member when the rule defines none.
    {interface_name}/{interface_id} are kept for the generator (interp_keep)."""
    ctx = dict(row)
    ctx["member"] = member
    ctx["direction"] = direction
    return identity.interp_keep(rule.get("interface_tagname", ""), ctx) or member


def collect_mirror_set(rows, *, index, is_diag, db_rules, diag_rules, if_rules):
    """Ordered list of _Elem to mirror onto interface `index`, plus warnings.
    (a) signals whose interface_mapping lists this index (or every in_diag signal when is_diag) -> Q;
    (c) their datablock/diagnosis rule followers -> Q; (d) their interface_elements_rules elements ->
    the rule's direction (I/Q). Deduped on (direction, script_type, mirror_name)."""
    warnings = []
    elems, seen = [], set()

    def _add(e):
        if not e.mirror_name:
            return False
        key = (e.direction, e.script_type, e.mirror_name)
        if key in seen:
            return False
        seen.add(key)
        elems.append(e)
        return True

    direct = []
    for r in rows or []:
        st = str(r.get("script_type", "") or "").strip()
        if st.upper() == TRIGGER_TYPE:                  # never mirror the IOC rows themselves
            continue
        picked = _index_in(index, r.get("interface_mapping"))
        if not picked and is_diag and (r.get("_type") or {}).get("in_diagnosis"):
            picked = True
        if not picked:
            continue
        name = _mirror_name(r)
        if not name:
            warnings.append(f"{st} {identity.fld(r)} "
                            f"({r.get('_source_sheet', '')}!{r.get('_source_row', '')}): "
                            f"no db_element or tag - not mirrored")
            continue
        e = _Elem(script_type=st, mirror_name=name,
                  signal_name=str(r.get("interface_tagname", "") or ""),
                  description=identity.interp((r.get("_type") or {}).get("tag_name", ""), r),
                  fu=str(r.get("functional_unit", "") or ""), loc=str(r.get("location", "") or ""),
                  dev=str(r.get("device", "") or ""), direction="Q", source="mirror",
                  diag_cabinet=str(r.get("diag_cabinet", "") or ""),
                  swp_cabinet=str(r.get("swp_cabinet", "") or ""),
                  diag_bit=str(r.get("diag_bit", "") or ""), src_row=r)
        if _add(e):
            direct.append(e)

    # (c) datablock + diagnosis rule followers (matched on the source signal's script_type) -> Q.
    # These are DB members (the rule names the DB), so render them TIA-qualified "<db>"."<member>"
    # to match the mirror-name format.
    for e in direct:
        r = e.src_row
        for rule in list(db_rules) + list(diag_rules):
            if e.script_type in rule["required_types"]:
                member = identity.interp(rule["member"], r)
                if not member:
                    continue
                db = rule.get("db_name", "")
                nm = f'"{db}"."{member}"' if db else f'"{member}"'
                _add(_Elem(script_type=rule["name"], mirror_name=nm,
                           signal_name=_rule_tagname(rule, r, member, "Q"), description=member,
                           fu=e.fu, loc=e.loc, dev=e.dev, direction="Q", source="follow",
                           diag_cabinet=e.diag_cabinet, swp_cabinet=e.swp_cabinet,
                           diag_bit=e.diag_bit, src_row=r))
    # (d) interface_elements_rules -> the rule's own direction / data_type / group label
    for e in direct:
        r = e.src_row
        for rule in if_rules:
            if e.script_type in rule["required_types"]:
                member = identity.interp(rule["member"], r)
                if not member:
                    continue
                _add(_Elem(script_type=rule["script_type"], mirror_name=member,
                           signal_name=_rule_tagname(rule, r, member, rule["direction"]),
                           description=member,
                           fu=e.fu, loc=e.loc, dev=e.dev, data_type=rule["data_type"],
                           direction=rule["direction"], source="if_rule",
                           diag_cabinet=e.diag_cabinet, swp_cabinet=e.swp_cabinet,
                           diag_bit=e.diag_bit, src_row=r))
    return elems, warnings


def allocate_bytes(elems, last, gap) -> None:
    """Assign (offset_byte, bit) to each _Elem, per direction independently. The custom area starts
    2-byte-aligned at last[dir] + 1 + gap. Each BOOL script_type group reserves a FULL 2-byte block
    (padded up to a 16-bit multiple): its signals take the low bits in order (the layout fills the
    rest with blank addressed rows for the engineer). Each WORD element takes its own even-aligned
    2-byte slot (bit None). Never two script_types in one byte (each block is whole, 2-byte aligned)."""
    for direction in ("Q", "I"):
        de = [e for e in elems if e.direction == direction]
        if not de:
            continue
        byte = last.get(direction, -1) + 1 + gap
        if byte % 2:
            byte += 1
        order, groups = [], {}
        for e in de:
            if e.script_type not in groups:
                groups[e.script_type] = []
                order.append(e.script_type)
            groups[e.script_type].append(e)
        for st in order:
            g = groups[st]
            if g[0].data_type == "WORD":
                for e in g:
                    if byte % 2:
                        byte += 1
                    e.offset_byte, e.bit = byte, None
                    byte += 2
            else:
                if byte % 2:
                    byte += 1
                for i, e in enumerate(g):
                    e.offset_byte, e.bit = byte + i // 8, i % 8
                byte += ((len(g) + 15) // 16) * 2        # full 2-byte block(s)


def _find_data_table(ws):
    """The data table on `ws` (the one whose header row carries 'Category'), as
    (name, c1, r1, c2, r2, {header: col}). None if absent. Robust to table renames."""
    for name in list(ws.tables):
        t = ws.tables[name]
        c1, r1, c2, r2 = range_boundaries(t.ref)
        hdr = {str(ws.cell(r1, c).value).strip(): c
               for c in range(c1, c2 + 1) if ws.cell(r1, c).value not in (None, "")}
        if "Category" in hdr:
            return name, c1, r1, c2, r2, hdr
    return None


def _template_last_used_byte(template_path, sheet_name) -> dict:
    """Max I/O Offset Byte per direction among the template sheet's existing data rows ('<'=I,
    '>'=Q), read from cached values (data_only) since offset cells may be relative formulas. Computed
    fresh every generation so a hand-edited template is honored. {'I': int|-1, 'Q': int|-1}."""
    last = {"I": -1, "Q": -1}
    wb = load_workbook(template_path, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return last
        found = _find_data_table(wb[sheet_name])
        if not found:
            return last
        _name, _c1, r1, _c2, r2, hdr = found
        dcol, ocol = hdr.get("Direction </>"), hdr.get("I/O Offset Byte")
        if not dcol or not ocol:
            return last
        ws = wb[sheet_name]
        for r in range(r1 + 1, r2 + 1):
            d = _DIR_TO_IO.get(str(ws.cell(r, dcol).value or "").strip())
            o = ws.cell(r, ocol).value
            if d and isinstance(o, (int, float)):
                last[d] = max(last[d], int(o))
    finally:
        wb.close()
    return last


def _set_text(cell, value) -> None:
    """Write a string, forcing data_type='s' for a leading =/+/- so a tag/FLD isn't read as a
    formula (e.g. '=S1', '+MS1.CC1', '-K1')."""
    s = "" if value is None else str(value)
    cell.value = s
    if s[:1] in ("=", "+", "-"):
        cell.data_type = "s"


def _append_custom_rows(ws, elems) -> None:
    """Lay the mirrored elements onto `ws`'s data table (writing by column name): Signal Name Side 1
    = the interface tag name, Expression Side 1 = the binding, Description = the human name. A BOOL
    script_type group fills a FULL 2-byte block - its signals on the low bits, then blank addressed
    rows up to the 16-bit boundary (for the engineer to fill) - then a blank separator row. A WORD
    element is one row + a blank separator. The address formula is copied from an existing row; the
    table ref + autoFilter extend to the last content row."""
    found = _find_data_table(ws)
    if not found:
        return
    name, c1, r1, c2, r2, hdr = found
    table = ws.tables[name]
    s1c, s2c = hdr.get("I/O Address Side 1"), hdr.get("I/O Address Side 2")
    f1 = f2 = None
    for r in range(r1 + 1, r2 + 1):
        v = ws.cell(r, s1c).value if s1c else None
        if isinstance(v, str) and v.startswith("="):
            f1 = v
            f2 = ws.cell(r, s2c).value if s2c else None
            break

    def _put(r, col_key, value, text=False):
        col = hdr.get(col_key)
        if not col:
            return
        if text:
            _set_text(ws.cell(r, col), value)
        else:
            ws.cell(r, col).value = value

    def _write(r, *, category, dt, direction, offset, bit, name_, expr, desc, fu, loc, dev, dc, swp, db):
        _put(r, "Category", category, text=True)
        _put(r, "Description", desc, text=True)
        _put(r, "Functional Unit", fu, text=True)
        _put(r, "Location", loc, text=True)
        _put(r, "Device", dev, text=True)
        _put(r, "Data Type", dt)
        _put(r, "Direction </>", ">" if direction == "Q" else "<")
        if offset is not None:
            _put(r, "I/O Offset Byte", offset)
        if bit is not None:
            _put(r, "I/O Bit", bit)
        _put(r, "Signal Name Side 1", name_, text=True)
        _put(r, "Expression Side 1", expr, text=True)
        if s1c and f1 is not None:
            ws.cell(r, s1c).value = f1
        if s2c and f2 is not None:
            ws.cell(r, s2c).value = f2
        if dc:
            _put(r, "diag_cabinet", dc)
        if swp:
            _put(r, "swp_cabinet", swp)
        if db:
            _put(r, "diag_bit", db)

    # group by direction (Q then I) then script_type (first-seen) - matching the allocator
    blocks = []
    for direction in ("Q", "I"):
        seen = {}
        for e in elems:
            if e.direction != direction:
                continue
            if e.script_type not in seen:
                seen[e.script_type] = []
                blocks.append((direction, e.script_type, seen[e.script_type]))
            seen[e.script_type].append(e)

    row, last = r2, r2
    for direction, st, g in blocks:
        if g[0].data_type == "WORD":
            for e in g:                                   # one row + a blank separator per word
                row += 1; last = row
                _write(row, category=st, dt="WORD", direction=direction, offset=e.offset_byte,
                       bit=None, name_=e.signal_name, expr=e.mirror_name, desc=e.description,
                       fu=e.fu, loc=e.loc, dev=e.dev, dc=e.diag_cabinet, swp=e.swp_cabinet, db=e.diag_bit)
                row += 1
        else:                                             # full 2-byte BOOL block (used + blank) + separator
            start = g[0].offset_byte
            for b in range(((len(g) + 15) // 16) * 16):
                row += 1; last = row
                if b < len(g):
                    e = g[b]
                    _write(row, category=st, dt="BOOL", direction=direction, offset=start + b // 8,
                           bit=b % 8, name_=e.signal_name, expr=e.mirror_name, desc=e.description,
                           fu=e.fu, loc=e.loc, dev=e.dev, dc=e.diag_cabinet, swp=e.swp_cabinet, db=e.diag_bit)
                else:
                    _write(row, category=st, dt="BOOL", direction=direction, offset=start + b // 8,
                           bit=b % 8, name_="", expr="", desc="", fu="", loc="", dev="",
                           dc="", swp="", db="")
            row += 1

    if last > r2:
        newref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last}"
        table.ref = newref
        if table.autoFilter:
            table.autoFilter.ref = newref


def mirror_onto(ws, rows, *, index, machine_type, is_diag, template_path, sheet_name, cfg=None) -> list:
    """Append the mirrored custom rows for interface `index` onto `ws` (the chosen, retitled sheet).
    `sheet_name` is the ORIGINAL template sheet name (for last-used-byte lookup). Returns warnings."""
    cfg = cfg or _load_mirror_cfg()
    elems, warnings = collect_mirror_set(
        rows, index=index, is_diag=is_diag,
        db_rules=cfg["db_rules"], diag_rules=cfg["diag_rules"], if_rules=cfg["if_rules"])
    if not elems:
        return warnings
    # Fill the per-interface tokens the staged/rule interface_tagname kept ({interface_name} = the
    # machine type, {interface_id} = this interface's index).
    iname, iid = str(machine_type or ""), str(index or "")
    for e in elems:
        e.signal_name = e.signal_name.replace("{interface_name}", iname).replace("{interface_id}", iid)
    allocate_bytes(elems, _template_last_used_byte(template_path, sheet_name), cfg["gap"])
    _append_custom_rows(ws, elems)
    return warnings


def generate_one(template_path, out_dir, *, machine_type, base_address="", node_side1="",
                 node_side2="", index="", instance=None, template_sheets=None,
                 mirror_rows=None, is_diag=False, mirror_cfg=None,
                 mirror_warnings=None, overwrite=False) -> tuple[str, str]:
    """Generate ONE interface workbook from explicit inputs - the primitive the 430 popup calls and
    the 410 auto path reuses. Returns (path, status) with status in {'created', 'preserved'}.
    PRESERVES an existing IF_<instance>.xlsx unless `overwrite` (then it regenerates). `instance`
    defaults to '<machine_type>-<index>' (or just <machine_type> when no index). Pass `template_sheets`
    to avoid re-opening the template per call. When `mirror_rows` is given, the mirrored custom block
    is appended after the standard plug (see mirror_onto)."""
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"interface template not found: {template_path}")
    os.makedirs(out_dir, exist_ok=True)
    index = str(index or "").strip()
    if instance is None:
        instance = f"{machine_type}-{index}" if index else str(machine_type or "")
    target = os.path.join(out_dir, f"IF_{_safe_name(instance)}.xlsx")
    if os.path.exists(target) and not overwrite:
        return target, "preserved"

    sheets = template_sheets if template_sheets is not None else template_sheet_names(template_path)
    chosen = choose_sheet(sheets, machine_type)
    if chosen not in sheets:
        raise ValueError(f"template has no '{machine_type}' sheet and no {GENERIC_SHEET} fallback "
                         f"(sheets: {sheets})")
    shutil.copyfile(template_path, target)
    wb = load_workbook(target)
    for name in list(wb.sheetnames):
        if name != chosen:
            del wb[name]
    ws = wb[chosen]
    ws.title = _safe_name(instance)
    _plug(ws, base_address, node_side1, node_side2, index)
    if mirror_rows is not None:
        mw = mirror_onto(ws, mirror_rows, index=index, machine_type=machine_type, is_diag=is_diag,
                         template_path=template_path, sheet_name=chosen, cfg=mirror_cfg)
        if mirror_warnings is not None and mw:
            mirror_warnings.extend(f"{instance}: {w}" for w in mw)
    wb.save(target)
    wb.close()
    return target, "created"


def _copy_sheet(src, dst, title) -> None:
    """Copy a generated interface sheet `src` into workbook `dst` as a new sheet `title`: values +
    formulas (the per-sheet table names are made unique and rewritten in the formulas), cell styles,
    column widths, merged cells, and the tables (re-created)."""
    suffix = re.sub(r"[^A-Za-z0-9_]", "_", title)
    renames = {nm: f"{nm}_{suffix}" for nm in list(src.tables)}     # unique table names in the iolist
    new = dst.create_sheet(title)
    for row in src.iter_rows():
        for cell in row:
            v = cell.value
            if v is None and not cell.has_style:
                continue
            if isinstance(v, str) and v.startswith("="):
                for old, newn in renames.items():
                    v = v.replace(old, newn)
            nc = new.cell(row=cell.row, column=cell.column, value=v)
            if cell.data_type == "s" and isinstance(v, str) and v[:1] in ("=", "+", "-"):
                nc.data_type = "s"
            if cell.has_style:
                nc.font = copy(cell.font)
                nc.fill = copy(cell.fill)
                nc.border = copy(cell.border)
                nc.alignment = copy(cell.alignment)
                nc.number_format = cell.number_format
    for letter, dim in src.column_dimensions.items():
        if dim.width:
            new.column_dimensions[letter].width = dim.width
    for mc in list(src.merged_cells.ranges):
        new.merge_cells(str(mc))
    for nm in list(src.tables):
        st = src.tables[nm]
        nt = Table(displayName=renames[nm], ref=st.ref)     # constructor sets name == displayName
        nt.headerRowCount = st.headerRowCount
        nt.totalsRowCount = st.totalsRowCount
        # Fresh columns with id + name ONLY. The source columns carry dataDxfId (indices into the
        # SOURCE workbook's differential-format records - out of range in the I/O List) and
        # calculatedColumnFormula (referencing the OLD table name); either makes Excel drop the whole
        # table. The cells keep their own copied/rewritten formulas, so addresses still compute.
        nt.tableColumns = [TableColumn(id=i + 1, name=col.name) for i, col in enumerate(st.tableColumns)]
        if st.tableStyleInfo is not None:
            nt.tableStyleInfo = deepcopy(st.tableStyleInfo)
        new.add_table(nt)


def _sheet_name_to_part(zbytes) -> dict:
    """{sheet display name -> worksheet part path} from an xlsx's workbook.xml + rels."""
    out = {}
    with zipfile.ZipFile(io.BytesIO(zbytes)) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rid = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
        sheets_el = wb.find(_MAIN_NS + "sheets")
        for sh in (sheets_el if sheets_el is not None else []):
            t = rid.get(sh.get(_REL_NS + "id"), "")
            out[sh.get("name")] = (t.lstrip("/") if t.startswith("/")
                                   else t if t.startswith("xl/") else "xl/" + t)
    return out


def _capture_formula_caches(zbytes) -> dict:
    """{sheet name -> {cell ref -> (t_attr, cached value text)}} for every formula cell carrying a
    cached value. openpyxl drops these on re-save, so they are restored afterwards (else data_only
    readers - staging/validation, e.g. Profinet IP/name - would see None)."""
    caches = {}
    with zipfile.ZipFile(io.BytesIO(zbytes)) as z:
        present = set(z.namelist())
        for nm, part in _sheet_name_to_part(zbytes).items():
            if part not in present:
                continue
            cc = {}
            for c in ET.fromstring(z.read(part)).iter(_MAIN_NS + "c"):
                f, v = c.find(_MAIN_NS + "f"), c.find(_MAIN_NS + "v")
                if f is not None and v is not None and v.text is not None:
                    cc[c.get("r")] = (c.get("t"), v.text)
            if cc:
                caches[nm] = cc
    return caches


def _patch_formula_cache(xml: str, ref: str, t, value: str) -> str:
    """Restore a single formula cell's cached value (+ its result-type t) in an openpyxl-written
    worksheet XML string (openpyxl emits an empty <v/> for formula cells)."""
    esc = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    pat = re.compile(r'(<c r="' + re.escape(ref) + r'")([^>]*)(>)(.*?)(</c>)', re.DOTALL)

    def repl(m):
        attrs = re.sub(r'\s+t="[^"]*"', "", m.group(2))
        if t:
            attrs += f' t="{t}"'
        body = re.sub(r"<v\s*/>|<v>\s*</v>", f"<v>{esc}</v>", m.group(4))
        return m.group(1) + attrs + m.group(3) + body + m.group(5)

    return pat.sub(repl, xml, count=1)


def insert_sheets_into_iolist(iolist_path, sheets) -> list:
    """Insert each (title, if_path) interface sheet into the I/O List ONLY if a sheet of that name is
    not already present, LOSSLESSLY. openpyxl writes the combined workbook (keeping every existing
    formula), then the existing formula cells' CACHED VALUES - which openpyxl drops on re-save - are
    patched back from the original, so both the formulas and their values survive. ONE atomic save (the
    caller backs the I/O List up first). Returns per-interface action strings."""
    with open(iolist_path, "rb") as fh:
        caches = _capture_formula_caches(fh.read())

    dst = load_workbook(iolist_path)               # data_only=False -> keep the existing formulas
    actions, changed = [], False
    for title, if_path in sheets:
        if title in dst.sheetnames:
            actions.append(f"{title}: already present in the I/O List - skipped")
            continue
        src_wb = load_workbook(if_path)
        _copy_sheet(src_wb[src_wb.sheetnames[0]], dst, title)
        src_wb.close()
        actions.append(f"{title}: inserted into the I/O List")
        changed = True
    if not changed:
        dst.close()
        return actions

    tmp = iolist_path + ".tmp_insert.xlsx"
    try:
        dst.save(tmp)
        dst.close()
        with open(tmp, "rb") as fh:
            data = fh.read()
        tparts = _sheet_name_to_part(data)
        repl = {}
        for nm, cc in caches.items():
            part = tparts.get(nm)
            if not part:
                continue
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                xml = z.read(part).decode("utf-8")
            for ref, (t, v) in cc.items():
                xml = _patch_formula_cache(xml, ref, t, v)
            repl[part] = xml.encode("utf-8")
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                zout.writestr(it, repl.get(it.filename) or zin.read(it.filename))
        with open(tmp, "wb") as fh:
            fh.write(out.getvalue())
        os.replace(tmp, iolist_path)
    except (OSError, ET.ParseError, zipfile.BadZipFile) as e:
        try:
            dst.close()
        except Exception:  # noqa: BLE001
            pass
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        actions.append(f"save failed ({e}) - I/O List left unchanged (restore from the .bak if needed)")
    return actions


def generate(rows, out_dir, template_path, *, overwrite=False, iolist_path=None) -> dict:
    """410 auto path: one interface per IOC row of the staged database, each with its mirrored
    custom block. Returns {'created', 'preserved', 'warnings', 'fallback', 'errors', 'iolist'}.
    With `overwrite` an existing IF_*.xlsx is regenerated (not preserved). With `iolist_path` each
    generated interface sheet is inserted (named IF_<instance>) into that I/O List workbook if not
    already present (a timestamped '.bak' of the I/O List is taken first); the actions land in
    result['iolist']."""
    os.makedirs(out_dir, exist_ok=True)
    records, warnings = find_interfaces(rows)
    result = {"created": [], "preserved": [], "warnings": list(warnings), "fallback": [], "errors": [],
              "iolist": []}
    if not records:
        return result
    if not os.path.exists(template_path):
        result["warnings"].append(f"interface template not found: {template_path}")
        return result

    # Interface indices must be unique so an interface_mapping value resolves to exactly one
    # interface - a collision is an ERROR (files are still generated by their full names). Key by the
    # SAME normalization the resolver uses (_idx_key), so '02' and '2' collide as the duplicate they are.
    by_index = {}
    for itf in records:
        by_index.setdefault(_idx_key(itf["index"]), []).append(itf)
    for idx, group in by_index.items():
        if len(group) > 1:
            insts = ", ".join(f"{g['instance']} ({g['source']})" for g in group)
            result["errors"].append(f"duplicate interface index {idx!r} across IOC rows: {insts}")

    # interface_mapping values that point at no IOC interface (WARN, not fatal). Exclude None keys so
    # a genuinely dangling non-numeric token is still surfaced.
    ioc_idx = {k for itf in records if (k := _idx_key(itf["index"])) is not None}
    unknown = sorted({tok.strip() for r in (rows or [])
                      for tok in str(r.get("interface_mapping", "") or "").split("|")
                      if tok.strip() and _idx_key(tok) not in ioc_idx})
    if unknown:
        result["warnings"].append(f"interface_mapping references no IOC interface: {unknown}")

    cfg = _load_mirror_cfg()
    sheets = template_sheet_names(template_path)
    has_generic = GENERIC_SHEET in sheets
    seen = {}        # instance -> source, to flag within-run filename collisions
    to_insert = []   # (sheet title, IF path) for the I/O List insertion
    for itf in records:
        mt = itf["machine_type"]
        exact = bool(mt) and any(s.upper() == mt.upper() for s in sheets)
        # Degrade gracefully (per-row warning) when neither a type sheet nor the <GENERIC> fallback
        # exists - one bad/typo'd machine must not abort generation for every other valid one. (The
        # direct 430 primitive keeps its hard raise; this is the 410 batch path.)
        if not exact and not has_generic:
            result["warnings"].append(
                f"{itf['instance']} ({itf['source']}): no '{mt}' sheet and no {GENERIC_SHEET} "
                f"fallback - skipped")
            continue
        if mt and not exact:
            result["fallback"].append(itf["instance"])
        path, status = generate_one(
            template_path, out_dir, machine_type=mt, base_address=itf["base"],
            node_side1=itf["base_node"], node_side2="", index=itf["index"],
            instance=itf["instance"], template_sheets=sheets,
            mirror_rows=rows, is_diag=itf["is_diag"], mirror_cfg=cfg,
            mirror_warnings=result["warnings"], overwrite=overwrite)
        # A second IOC row with the same instance name hits the just-created file and reads as
        # "preserved"; that is a same-run collision (the 2nd row's data is dropped), not the
        # legitimate prior-run preserve - surface it instead of silently swallowing it.
        if status == "preserved" and itf["instance"] in seen:
            result["warnings"].append(
                f"duplicate IOC instance {itf['instance']} ({itf['source']} vs "
                f"{seen[itf['instance']]}) - second occurrence ignored")
        else:
            result[status].append(os.path.basename(path))
        seen.setdefault(itf["instance"], itf["source"])
        if iolist_path:
            # Inserted I/O List sheets are PREFIXED with IF_ (e.g. IF_SORTER-01): they read as
            # interfaces at a glance, stay clear of the NET SAFETY sheets staging/validation match,
            # and give phase 500 (510 I/O Tags) a stable prefix to find the interface tags by.
            to_insert.append((_safe_name(f"IF_{itf['instance']}"), path))

    if iolist_path and to_insert and os.path.exists(iolist_path):
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        bak = f"{iolist_path}.bak_{stamp}.xlsx"
        try:
            shutil.copyfile(iolist_path, bak)
        except OSError as e:
            result["iolist"].append(f"I/O List backup failed ({e}); sheet insertion skipped")
            return result
        result["iolist"].append(f"backed up the I/O List -> {os.path.basename(bak)}")
        result["iolist"].extend(insert_sheets_into_iolist(iolist_path, to_insert))
    return result
