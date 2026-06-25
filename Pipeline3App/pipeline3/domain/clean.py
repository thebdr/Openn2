"""Clean Addresses & Electrical Names (the manual phase-100 buttons 165 / 175).

Hand-authored I/O List + C&E documents carry stray apostrophes (') and spaces in the address +
electrical-name cells. The cross-check matcher (validation.model.vm.key/addr, matrix._norm) removes
whitespace + uppercases but NOT apostrophes, so e.g. `-S1'2` never reconciles with `-S12` -> phase
130/140 report spurious mismatches. This cleaner removes the configured SYMBOLS and collapses ALL
whitespace from those cells (case preserved; the matcher already uppercases), writing them back
SURGICALLY (io.xlsx_edit, no openpyxl round-trip) with a timestamped backup (dropped on a no-op).
Manual-only (never in the pipeline DAG). A FORMULA cell in a target column is left untouched.
"""
from __future__ import annotations
import os
import re
import shutil
from dataclasses import dataclass, field

from openpyxl import load_workbook

from pipeline3.core import config
from pipeline3.domain.iolist_diag.columns import ColumnResolver
from pipeline3.domain.iolist_diag.populate import _backup_path, _same_content
from pipeline3.io import xlsx_edit

# stray symbols removed from addresses & electrical names so the cross-check matches (extend as found)
CLEAN_SYMBOLS = ["'"]


@dataclass
class CleanResult:
    path: str = ""
    changed: int = 0
    sheets: list = field(default_factory=list)
    backup: str = ""
    skipped_formula: int = 0
    note: str = ""


def _clean_value(v):
    """The pure rule: a str -> every char in CLEAN_SYMBOLS removed + ALL whitespace collapsed (case
    kept); None / numeric / bool -> None (the caller skips it). A leading '='/'+'/'-' survives because
    the writer materializes the cell as an inline string (TEXT, not a formula)."""
    if not isinstance(v, str):
        return None
    s = v
    for ch in CLEAN_SYMBOLS:
        s = s.replace(ch, "")
    return re.sub(r"\s", "", s)


def _clean_workbook(path: str, make_jobs, emit) -> CleanResult:
    """Open `path` once (data_only=False, to see formula cells), let `make_jobs(sheetnames)` return the
    list of (sheet, data_start_row, ColumnResolver, [canonical_columns]) to scan, collect the cell edits,
    then write them back via io.xlsx_edit with a backup (dropped if nothing changed)."""
    res = CleanResult(path=path)
    backup = _backup_path(path)
    shutil.copy2(path, backup)

    cell_edits: dict = {}
    wb = load_workbook(path, data_only=False)
    try:
        for sheet, data_start, cr, cols in make_jobs(wb.sheetnames):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            emit(f"{sheet}: cleaning columns " + ", ".join(f"{cr.letter(c)}:{c}" for c in cols)
                 + f" (rows {data_start}..{ws.max_row})")
            edits = {}
            for row in range(data_start, ws.max_row + 1):
                for canon in cols:
                    cell = cr.cell(ws, row, canon)
                    if cell.data_type == "f":           # a formula -> leave it (don't flatten its cache)
                        res.skipped_formula += 1
                        continue
                    cleaned = _clean_value(cell.value)
                    if cleaned is None or cleaned == cell.value:
                        continue
                    ref = f"{cr.letter(canon)}{row}"
                    edits[ref] = cleaned
                    emit(f"{sheet}!{ref}: '{cell.value}' -> '{cleaned}'")
            if edits:
                cell_edits[sheet] = edits
                res.sheets.append(sheet)
    finally:
        wb.close()

    res.changed = sum(len(v) for v in cell_edits.values())
    if cell_edits:
        for sheet_name, master, rng in xlsx_edit.edit_workbook(path, cell_edits=cell_edits):
            emit(f"[WARN] froze array {sheet_name}!{master} (spill {rng}) to its cached values")
    if not cell_edits or _same_content(backup, path):
        if os.path.exists(backup):
            os.remove(backup)
        backup = ""
    res.backup = backup
    emit(f"cleaned {res.changed} cell(s) in {len(res.sheets)} sheet(s)"
         + (f"; skipped {res.skipped_formula} formula cell(s)" if res.skipped_formula else "")
         + (f"; backup {os.path.basename(backup)}" if backup else ""))
    return res


def clean_iolist(params: dict, emit=print) -> CleanResult:
    """165: clean the I/O List's address (bit) + electrical names (functional_unit/location/device) across
    every sheet matching io_list.sheet."""
    iol = params.get("io_list") or {}
    path = iol.get("path")
    if not path or not os.path.exists(path):
        raise ValueError(f"I/O List not found: {path}")
    cr = ColumnResolver("IoList")
    cols = ["bit", "functional_unit", "location", "device"]
    data_start = int(iol.get("header_row", 1) or 1) + 1

    def make_jobs(names):
        return [(s, data_start, cr, cols) for s in config.resolve_sheets(iol.get("sheet"), names)]

    return _clean_workbook(path, make_jobs, emit)


def clean_cematrix(params: dict, emit=print) -> CleanResult:
    """175: clean the C&E matrix sheet (address + functional_unit/location/device) AND the AREA sheets
    (contactor sigla + address). The C&E document being absent is a no-op (not an error)."""
    ce = params.get("ce") or {}
    path = ce.get("path")
    if not path or not os.path.exists(path):
        res = CleanResult(note="C&E Matrix not configured / not found - nothing to clean")
        emit(res.note)
        return res
    ce_cr, area_cr = ColumnResolver("CE"), ColumnResolver("AREA")
    ce_cols = ["address", "functional_unit", "location", "device"]
    area_cols = ["concat_id", "address"]
    m_start = int(ce.get("matrix_data_row", 5) or 5)
    a_start = int(ce.get("area_data_row", 4) or 4)

    def make_jobs(names):
        jobs = []
        msheet = config.resolve_sheet(ce.get("matrix_sheet"), names)
        if msheet:
            jobs.append((msheet, m_start, ce_cr, ce_cols))
        for s in config.resolve_sheets(params.get("areas"), names):
            jobs.append((s, a_start, area_cr, area_cols))
        return jobs

    return _clean_workbook(path, make_jobs, emit)
