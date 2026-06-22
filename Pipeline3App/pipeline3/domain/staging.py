"""Staging (phase 300): read the (populated) I/O List into canonical rows = the single database.

Reads every matched sheet through the ONE shared reader (io.workbook), by column position per
column_map (header verified by prefix, newline->space). Excludes struck rows only when
strike_handling == "exclude", and rows carrying a Skip Reason. Each kept row gets its resolved
signal type and is enriched with C&E data (matrix.annotate) + identity fields (identity.*) +
diagnosis-block lookup + node address ranges. `write_io_database` persists the single IODatabase.csv.
"""
from __future__ import annotations
import os

from openpyxl import load_workbook

from pipeline3.core import config
from pipeline3.io import workbook as wbk, csv_tables
from pipeline3.domain import matrix, identity
from pipeline3.domain.iolist_diag.models import DIAGBLOCKS_SHEET, DIAGBLOCKS_LEGACY, UNRESOLVED_SHEET

_GENERATED = {DIAGBLOCKS_SHEET.lower(), DIAGBLOCKS_LEGACY.lower(), UNRESOLVED_SHEET.lower()}

# IODatabase.csv = the IoList canonical columns + these derived/enriched columns.
_EXTRA = ["matrix_areas", "areas_description",
          "ce_functional_unit", "ce_location", "ce_device", "numerazione_linea",
          "iol_FLD", "ce_FLD", "combined_FLD",
          "IsSorterArea", "name_in_db", "name_in_tagtable", "tagtable", "datablocks", "diag_desc",
          "interface_tagname", "diag_block_name", "diag_block_template", "swp_cabinet", "subnet_name",
          "I_startByte", "I_endByte", "Q_startByte", "Q_endByte",
          "source_cell", "_source_sheet", "_source_row", "type_id_resolved", "type_category"]


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _skip_reason_present(v) -> bool:
    return _norm(v) not in ("", "0", "0.0")


def _read_view(v, colmap, strike_exclude, signal_types, warnings) -> list:
    # header verification (warning for optional, error for required). The pipeline-generated columns
    # (preliminary_check_exclude=yes: AA-AG) carry no header on an as-authored doc - the populator
    # writes them - so a missing header there is a warning even when the column is "required" (this
    # is what lets the validation-only/designer profile stage a raw, un-filled I/O List).
    for m in colmap:
        if not v.header_matches(m["column"], m["expected_header"]):
            msg = (f"IoList[{v.name}] column {m['column']}: header {v.header(m['column'])!r} != "
                   f"expected {m['expected_header']!r} (-> {m['canonical']})")
            if m["required"] and not m.get("preliminary_check_exclude"):
                raise SystemExit("ERROR " + msg)
            warnings.append(msg)
    rows = []
    for r in v.data_rows():
        if all(v.text(r, m["column"]) == "" for m in colmap):
            continue
        row = {m["canonical"]: v.text(r, m["column"]) for m in colmap}
        row["_source_row"] = r
        row["_source_sheet"] = v.name
        if _skip_reason_present(row.get("skip_reason")):
            continue
        if strike_exclude and v.row_struck(r):
            continue
        row["_type"] = config.resolve_type(signal_types, row.get("script_type"))
        rows.append(row)
    return rows


def load_io_list(params: dict, signal_types: dict, io_path: str) -> tuple:
    """Read the (matched) I/O sheets of `io_path` into enriched staged rows.
    Returns (rows, warnings, matched_sheets, skipped_sheets)."""
    colmap = config.load_column_map("IoList")
    header_row = int(params.get("io_list", {}).get("header_row", 1) or 1)
    strike_exclude = str(params.get("strike_handling", "exclude")).lower() == "exclude"
    sheet_pattern = params["io_list"]["sheet"]

    views = wbk.open_sheets(io_path, sheet_pattern, header_row, doc_label=os.path.basename(io_path))
    all_names = wbk.available_sheets(io_path)
    matched = [v.name for v in views]
    if not matched:
        raise SystemExit(f"no I/O sheet matched {sheet_pattern!r} in {io_path} (have {all_names})")
    skipped = [n for n in all_names if n not in matched and n.strip().lower() not in _GENERATED]

    rows, warnings = [], []
    for v in views:
        rows.extend(_read_view(v, colmap, strike_exclude, signal_types, warnings))
    views[0].close()

    matrix.annotate(params, rows)
    diag_blocks = load_diagnostic_blocks(io_path)
    fu_col = next((m["column"] for m in colmap if m["canonical"] == "functional_unit"), "O")
    for row in rows:
        if not row.get("source_cell") and row.get("_source_row"):
            row["source_cell"] = f"{row.get('_source_sheet', '')}!{fu_col}{row['_source_row']}"
        # combined device keys (must precede name_in_db/diag_desc - the templates reference them)
        row["iol_FLD"] = identity.fld(row)
        row["ce_FLD"] = identity.ce_fld(row)
        row["combined_FLD"] = identity.combined_fld(row)
        row["name_in_db"] = identity.member_name(row)
        if identity.is_io_signal(row) and identity.tag_name(row):
            row["name_in_tagtable"] = identity.tag_name(row)
            row["tagtable"] = identity.tagtable(row)
        else:
            row["name_in_tagtable"] = ""
            row["tagtable"] = ""
        row["diag_desc"] = identity.diag_desc(row)
        row["interface_tagname"] = identity.interface_tagname(row)
        dc = str(row.get("diag_cabinet", "")).strip()
        blk = diag_blocks.get(int(dc)) if dc.lstrip("-").isdigit() else None
        row["diag_block_name"] = blk["fld"] if blk else ""
        row["diag_block_template"] = blk["template_type"] if blk else ""
        row["swp_cabinet"] = blk["swp"] if blk else ""
        t = row.get("_type") or {}
        row["datablocks"] = ("|".join(t.get("db_names") or [])
                             if t.get("db_kind") in ("db", "safe_db") and row["name_in_db"] else "")
        octets = str(row.get("profinet_ip", "")).split(".")
        row["subnet_name"] = f"Subnet{octets[2]}" if len(octets) == 4 and octets[2] else ""
    _add_node_address_ranges(rows)
    sorter = config.sorter_areas(params)
    for row in rows:
        areas = {a for a in str(row.get("matrix_areas", "")).split("|") if a}
        row["IsSorterArea"] = "yes" if (areas & sorter) else ""
    return rows, warnings, matched, skipped


def _addr_byte(bit):
    b = str(bit or "").strip().upper()
    if b[:1] in ("I", "Q") and "." in b:
        try:
            return b[0], int(b[1:b.index(".")])
        except ValueError:
            return None
    return None


def _add_node_address_ranges(rows) -> None:
    by_loc_I, by_loc_Q = {}, {}
    for r in rows:
        ab = _addr_byte(r.get("bit"))
        if ab:
            (by_loc_I if ab[0] == "I" else by_loc_Q).setdefault(r.get("location", ""), []).append(ab[1])
    for r in rows:
        if r.get("profinet_name"):
            ib = by_loc_I.get(r.get("location", ""), [])
            qb = by_loc_Q.get(r.get("location", ""), [])
            r["I_startByte"] = min(ib) if ib else ""
            r["I_endByte"] = max(ib) if ib else ""
            r["Q_startByte"] = min(qb) if qb else ""
            r["Q_endByte"] = max(qb) if qb else ""
        else:
            r["I_startByte"] = r["I_endByte"] = r["Q_startByte"] = r["Q_endByte"] = ""


def load_diagnostic_blocks(io_path: str) -> dict:
    """Read the DiagnosisBlocks sheet -> {cabinet_id:int -> {index, fld, template_type, swp}} keyed
    by ID_Local (= the I/O List row's diag_cabinet, the local cabinet). `swp` is the ID_SWP (software
    cabinet; equals ID_Local on a generated sheet, may differ in general). {} if the sheet/column is
    missing. Accepts the legacy 'DiagnosticBlocks' spelling; falls back to ID_SWP as the key when
    ID_Local is absent (older sheets)."""
    wb = load_workbook(io_path, data_only=True)
    sheet = next((s for s in wb.sheetnames
                  if s.strip().lower() in (DIAGBLOCKS_SHEET.lower(), DIAGBLOCKS_LEGACY.lower())), None)
    if sheet is None:
        wb.close()
        return {}
    grid = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    if not grid:
        return {}
    header = [" ".join(str(c or "").split()).lower() for c in grid[0]]

    def col(name):
        key = " ".join(name.split()).lower()
        return next((i for i, h in enumerate(header) if h == key), None)

    c_local, c_swp, c_fld, c_tt = col("ID_Local"), col("ID_SWP"), col("FullName"), col("TemplateType")
    c_key = c_local if c_local is not None else c_swp
    if c_key is None:
        return {}
    out = {}
    for r in grid[1:]:
        raw = _norm(r[c_key]) if c_key < len(r) else ""
        if not raw or not raw.lstrip("-").isdigit():
            continue
        cid = int(raw)
        out[cid] = {
            "index": f"{cid:03d}",
            "fld": _norm(r[c_fld]) if c_fld is not None and c_fld < len(r) else "",
            "template_type": _norm(r[c_tt]) if c_tt is not None and c_tt < len(r) else "",
            "swp": _norm(r[c_swp]) if (c_swp is not None and c_swp < len(r)) else raw,
        }
    return out


def write_io_database(rows: list, out_root: str) -> str:
    """Write the single database to IODatabase.csv (canonical columns + derived/enriched)."""
    cols = [m["canonical"] for m in config.load_column_map("IoList")]
    path = config.out_path(out_root, "io_database")
    table = []
    for r in rows:
        t = r.get("_type") or {}
        vals = [r.get(c, "") for c in cols]
        for e in _EXTRA:
            if e == "type_id_resolved":
                vals.append(t.get("type_id", ""))
            elif e == "type_category":
                vals.append(t.get("category", ""))
            else:
                vals.append(r.get(e, ""))
        table.append(vals)
    return csv_tables.write_table(path, cols + _EXTRA, table)
