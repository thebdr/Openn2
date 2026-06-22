"""Pipeline Coverage Report (phase 900 / 910).

Traces each staged signal's life from validation (100) through staging (300) to every downstream
output, then flags two cross-check defects. **The trace reads the ACTUAL output files on disk** (not
an in-memory re-derivation) - phases 400-800 permit user inputs here and there (the editable .xlsm
block shells in override/keep mode, hand-filled interface Side-2 rows, error-management treatments),
so only the produced artifacts reflect the real project state. It reuses each staged row's identity
(name_in_db / name_in_tagtable / plc_binding / profinet_name - shared with the generators) to find
where that signal landed across the artifacts:

  interfaces (400)  -> the IF_*.xlsx Expression Side 1 (the mirror binding) + the IF_ instance files;
  I/O tags (510)    -> PlcTags/PLCTags.xlsx (the tag names);
  data blocks (520) -> ImportReady *.db / *.xml GlobalDB members;
  diagnosis (600)   -> DiagList_IO/Logic.csv PLC_Binding + Diagnostic_for_OPC.scl bindings;
  hardware (700)    -> HardwareConfiguration Stations.csv / Modules.csv station names;
  software (800)    -> CreationInfo/*.csv cells + ImportReady FC-XML component refs.

Findings:
  - **ORPHAN**  : a staged SIGNAL whose identity appears in NONE of the output artifacts (staged, but
                  nothing consumes it - usually an untyped / mis-typed signal no generator picked up);
  - **UNPLACED**: a qualified `"<db>"."<member>"` reference an output emits (a diagnosis binding / an
                  interface-mirror Expression) whose DB is one Pipeline3 generates (it has a produced
                  .db/.xml) but whose member is NOT in that DB's actual members - an output pointing at
                  a data-block member that was never created.

Two row kinds are never ORPHAN: raw I/O **channels** (untyped rows carrying an I/Q address - physical
module points, covered by hardware) and **structural** rows (headers / node-meta with no address).

Writes Reports/io_project_coverage_report.{csv,txt} - documentation, NOT a Open2App import surface.
It reads only artifacts already on disk (richer once 400-800 have run); a missing artifact degrades
to a noted gap. References to DBs Pipeline3 does NOT generate (05_EM_STATE, SPEED_STATE_REC, ... -
hand-authored in TIA) are out of scope - that is phase 920's job (the TIA project export).
"""
from __future__ import annotations
import csv as _csv
import glob
import os
import re

from openpyxl import load_workbook

from pipeline3.core import config
from pipeline3.domain import identity, signals, diagnosis, interfaces
from pipeline3.io import csv_tables

# the downstream outputs a signal can land in (the lifecycle stages after staging/validation).
STAGES = ("interfaces", "io_tags", "data_blocks", "diagnosis", "hardware", "software")
_STAGE_LABEL = {"interfaces": "400 interfaces", "io_tags": "510 I/O tags",
                "data_blocks": "520 data blocks", "diagnosis": "600 diagnosis",
                "hardware": "700 hardware", "software": "800 software"}

_QUALREF = re.compile(r'"([^"]+)"\."([^"]+)"')   # a TIA-qualified "<db>"."<member>" reference
_PADS = set(signals.DB_CONSTANTS)                 # Always FALSE / Always TRUE / No Operation


# --------------------------------------------------------------------------------------------- #
# row classification
# --------------------------------------------------------------------------------------------- #
def is_signal(row) -> bool:
    """A staged row is a SIGNAL (traceable, ORPHAN-able) when it resolved to a type OR carries a
    DB-member / PLC-tag identity. Untyped raw channels and structural rows are not signals."""
    if (row.get("_type") or {}).get("type_id"):
        return True
    return any(str(row.get(k, "") or "").strip() for k in ("name_in_db", "name_in_tagtable"))


def row_kind(row) -> str:
    """'signal' | 'channel' (untyped row with an I/Q address = a raw module point) | 'structural'."""
    if is_signal(row):
        return "signal"
    return "channel" if str(row.get("bit", "") or "").strip()[:1].upper() in ("I", "Q") else "structural"


# --------------------------------------------------------------------------------------------- #
# reading the on-disk output artifacts
# --------------------------------------------------------------------------------------------- #
def _xml_unescape(s: str) -> str:
    return (s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"'))


def _read_tags(out_root) -> tuple:
    """The tag names from PlcTags/PLCTags.xlsx ('PLC Tags' sheet, column Name). (set, found)."""
    path = config.out_path(out_root, "io_tags_dir", signals.TAG_TABLE_FILE)
    if not os.path.exists(path):
        return set(), False
    names = set()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["PLC Tags"] if "PLC Tags" in wb.sheetnames else wb[wb.sheetnames[0]]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row:
                continue
            if row[0] not in (None, ""):
                names.add(str(row[0]).strip())
    finally:
        wb.close()
    return names, True


def _read_db_members(out_root) -> tuple:
    """{db_name: set(member)} from ImportReady *.db (external source) + *.xml GlobalDB (safe F_DB).
    db_name = the file stem. FC XML (a software block, no <Member>) is skipped. (dict, found)."""
    import_dir = config.out_path(out_root, "blocks_import_dir")
    members, found = {}, False
    for p in glob.glob(os.path.join(import_dir, "*.db")):
        found = True
        stem = os.path.splitext(os.path.basename(p))[0]
        text = open(p, encoding="utf-8-sig").read()
        for m in re.finditer(r'"([^"]+)"\s*:\s*Bool', text):
            members.setdefault(stem, set()).add(m.group(1))
    for p in glob.glob(os.path.join(import_dir, "*.xml")):
        text = open(p, encoding="utf-8-sig").read()
        if "SW.Blocks.GlobalDB" not in text:                 # an FC / non-DB block - not a data block
            continue
        found = True
        stem = os.path.splitext(os.path.basename(p))[0]
        for m in re.finditer(r'<Member Name="([^"]+)"', text):
            members.setdefault(stem, set()).add(_xml_unescape(m.group(1)))
    return members, found


def _read_diagnosis(out_root) -> tuple:
    """The diagnosis bindings (the bound signals) + consumed refs, from DiagList_IO/Logic.csv
    (PLC_Binding column) and Diagnostic_for_OPC.scl (the IN_xx bindings). (bindings set, consumed
    [(origin, db, member)], found)."""
    diag_dir = config.out_path(out_root, "diagnosis_dir")
    import_dir = config.out_path(out_root, "blocks_import_dir")
    bindings, consumed, found = set(), [], False
    for fn in ("DiagList_IO.csv", "DiagList_Logic.csv"):
        p = os.path.join(diag_dir, fn)
        if not os.path.exists(p):
            continue
        found = True
        for r in csv_tables.read_rows(p):
            b = str(r.get("PLC_Binding", "") or "").strip()
            if b:
                bindings.add(b)
            for m in _QUALREF.finditer(b):
                consumed.append(("diagnosis-list", m.group(1), m.group(2)))
    scl = os.path.join(import_dir, diagnosis.DIAG_SCL_FILE)
    if os.path.exists(scl):
        found = True
        for m in _QUALREF.finditer(open(scl, encoding="utf-8-sig").read()):
            bindings.add(f'"{m.group(1)}"."{m.group(2)}"')
            consumed.append(("diagnosis-SCL", m.group(1), m.group(2)))
    return bindings, consumed, found


def _read_interfaces(out_root) -> tuple:
    """The interface mirror bindings (Expression Side 1) + the interface instance names (the IF_*.xlsx
    files) + consumed refs. (exprs set, instances set, consumed [(origin, db, member)], found)."""
    iface_dir = config.out_path(out_root, "interfaces_dir")
    exprs, instances, consumed, found = set(), set(), [], False
    for p in glob.glob(os.path.join(iface_dir, "IF_*.xlsx")):
        found = True
        stem = os.path.splitext(os.path.basename(p))[0]
        instances.add(stem[3:] if stem.startswith("IF_") else stem)     # IF_SORTER-01 -> SORTER-01
        wb = load_workbook(p, data_only=True)
        try:
            for sh in wb.sheetnames:
                ws = wb[sh]
                tbl = interfaces._find_data_table(ws)
                if not tbl:
                    continue
                _name, _c1, r1, _c2, r2, hdr = tbl
                ec = hdr.get("Expression Side 1")
                if not ec:
                    continue
                for r in range(r1 + 1, r2 + 1):
                    v = str(ws.cell(r, ec).value or "").strip()
                    if v:
                        exprs.add(v)
                        for m in _QUALREF.finditer(v):
                            consumed.append(("interface-mirror", m.group(1), m.group(2)))
        finally:
            wb.close()
    return exprs, instances, consumed, found


def _read_csv_positional(path, col) -> set:
    """A format-2 CSV (no header row, '#'/'#!' comment lines): the distinct non-empty values of the
    0-based column `col`."""
    out = set()
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in _csv.reader(f):
            if not row or row[0].startswith("#"):
                continue
            if len(row) > col and str(row[col]).strip():
                out.add(str(row[col]).strip())
    return out


def _read_hardware(out_root) -> tuple:
    """The generated station names (Stations.csv col 1 = Station Name, Modules.csv col 0 = Station
    Name; both format-2). (set, found)."""
    hw = config.out_path(out_root, "hardware_dir")
    stations, found = set(), False
    sp = os.path.join(hw, "Stations.csv")
    if os.path.exists(sp):
        found = True
        stations |= _read_csv_positional(sp, 1)
    mp = os.path.join(hw, "Modules.csv")
    if os.path.exists(mp):
        found = True
        stations |= _read_csv_positional(mp, 0)
    return stations, found


def _read_software(out_root) -> tuple:
    """{emitted value: set(block id)} from the CreationInfo $/#/%/@ CSVs (@-row cells, pads dropped)
    plus the ImportReady FC-XML <Component> refs (the direct-XML blocks, e.g. 03). (dict, found)."""
    creation_dir = config.out_path(out_root, "blocks_creation_dir")
    import_dir = config.out_path(out_root, "blocks_import_dir")
    refs, found = {}, False
    for p in glob.glob(os.path.join(creation_dir, "*.csv")):
        if os.path.basename(p) == "InstanceDBs.csv":
            continue
        found = True
        block = os.path.basename(p).split("_", 1)[0]
        with open(p, encoding="utf-8-sig", newline="") as f:
            for row in _csv.reader(f):
                if not row or row[0] != "@":
                    continue
                for c in row[1:]:
                    c = str(c).strip()
                    if c and c not in _PADS:
                        refs.setdefault(c, set()).add(block)
    for p in glob.glob(os.path.join(import_dir, "*.xml")):
        text = open(p, encoding="utf-8-sig").read()
        if "SW.Blocks.FC" not in text:                       # only the direct-XML software blocks (03)
            continue
        found = True
        block = os.path.basename(p).split("_", 1)[0]
        for m in re.finditer(r'<Component Name="([^"]+)"', text):
            v = _xml_unescape(m.group(1))
            if v and v not in _PADS:
                refs.setdefault(v, set()).add(block)
    return refs, found


def collect_outputs(out_root) -> dict:
    """Read every on-disk output artifact under `out_root` into the emitted-identity sets the trace
    attributes against. A missing artifact degrades to empty + found=False (noted in the report)."""
    tags, f_tags = _read_tags(out_root)
    dbm, f_db = _read_db_members(out_root)
    diagb, diag_consumed, f_diag = _read_diagnosis(out_root)
    ifx, ifi, iface_consumed, f_iface = _read_interfaces(out_root)
    stations, f_hw = _read_hardware(out_root)
    swrefs, f_sw = _read_software(out_root)
    return {
        "tags": tags, "db_members": dbm, "diag_bindings": diagb,
        "iface_exprs": ifx, "iface_instances": ifi, "stations": stations, "sw_refs": swrefs,
        "consumed": diag_consumed + iface_consumed,
        "found": {"interfaces": f_iface, "io_tags": f_tags, "data_blocks": f_db,
                  "diagnosis": f_diag, "hardware": f_hw, "software": f_sw},
    }


# --------------------------------------------------------------------------------------------- #
# UNPLACED
# --------------------------------------------------------------------------------------------- #
def find_unplaced(produced: dict, consumed: list) -> list:
    """The UNPLACED findings: each consumed `(origin, db, member)` whose db is generated by Pipeline3
    (db in `produced` = a DB with an actual .db/.xml on disk) but whose member is NOT among that DB's
    actual members. Refs to DBs Pipeline3 does NOT generate (external/TIA - 920 scope) are ignored.
    Deduped on (db, member), origins joined. Returns [{'db', 'member', 'referenced_by'}] sorted."""
    unplaced_map = {}
    for origin, d, m in consumed:
        if d in produced and m not in produced[d]:
            unplaced_map.setdefault((d, m), set()).add(origin)
    return [{"db": d, "member": m, "referenced_by": "|".join(sorted(origins))}
            for (d, m), origins in sorted(unplaced_map.items())]


# --------------------------------------------------------------------------------------------- #
# attribution (pure: rows + the read outputs -> per-row placement + findings)
# --------------------------------------------------------------------------------------------- #
def attribute(rows, outputs) -> dict:
    """Attribute each staged row to the output stages whose artifacts carry its identity, and compute
    the ORPHAN / UNPLACED findings. Pure over `rows` + `outputs` (collect_outputs' result), so it is
    unit-testable with a synthetic `outputs` and never touches disk."""
    rows = list(rows or [])
    tags = outputs["tags"]
    dbm = outputs["db_members"]
    diagb = outputs["diag_bindings"]
    ifx, ifi = outputs["iface_exprs"], outputs["iface_instances"]
    stations = outputs["stations"]
    swrefs = outputs["sw_refs"]

    records, orphans = [], []
    for r in rows:
        member = str(r.get("name_in_db", "") or "").strip()
        tag = str(r.get("name_in_tagtable", "") or "").strip()
        binding = identity.plc_binding(r)
        pname = str(r.get("profinet_name", "") or "").strip()
        place = {s: set() for s in STAGES}

        if tag and tag in tags:
            place["io_tags"].add(str(r.get("tagtable", "") or "").strip() or "yes")
        if member:
            place["data_blocks"].update(d for d, ms in dbm.items() if member in ms)
        if binding and binding in diagb:
            place["diagnosis"].add("yes")
        if str(r.get("script_type", "")).strip().upper() == interfaces.TRIGGER_TYPE \
                and str(r.get("index", "") or "").strip() in ifi:
            place["interfaces"].add("defines")
        if binding and binding in ifx:
            place["interfaces"].add("mirror")
        if pname and pname in stations:
            place["hardware"].add("station")
        elif diagnosis._addr_byte(r.get("bit")):
            node = diagnosis.node_of(rows, r)
            if node and str(node.get("profinet_name", "") or "").strip() in stations:
                place["hardware"].add("module")
        for ident in {member, tag, binding}:
            if ident and ident in swrefs:
                place["software"].update(swrefs[ident])

        landed = [s for s in STAGES if place[s]]
        kind = row_kind(r)
        rec = {
            "source": str(r.get("source_cell", "") or ""),
            "kind": kind,
            "script_type": str(r.get("script_type", "") or ""),
            "FLD": str(r.get("iol_FLD", "") or ""),
            "address": str(r.get("bit", "") or ""),
            "type": (r.get("_type") or {}).get("type_id", ""),
            "name_in_db": member,
            "tag": tag,
            "binding": binding,
        }
        for s in STAGES:
            rec[s] = "|".join(sorted(place[s]))
        rec["outputs"] = len(landed)
        rec["flag"] = "ORPHAN" if (kind == "signal" and not landed) else ""
        if rec["flag"] == "ORPHAN":
            orphans.append(rec)
        records.append(rec)

    unplaced = find_unplaced(outputs["db_members"], outputs["consumed"])

    stage_counts = {s: sum(1 for rec in records if rec[s]) for s in STAGES}
    kind_counts = {k: sum(1 for rec in records if rec["kind"] == k)
                   for k in ("signal", "channel", "structural")}
    stats = {"rows": len(records), "kinds": kind_counts, "stages": stage_counts,
             "orphans": len(orphans), "unplaced": len(unplaced),
             "generated_dbs": sorted(outputs["db_members"]),
             "missing": [s for s in STAGES if not outputs["found"].get(s)]}
    return {"records": records, "orphans": orphans, "unplaced": unplaced, "stats": stats}


def trace(rows, out_root) -> dict:
    """Read the on-disk outputs under `out_root` and attribute every staged row against them."""
    return attribute(rows, collect_outputs(out_root))


# --------------------------------------------------------------------------------------------- #
# report rendering
# --------------------------------------------------------------------------------------------- #
_CSV_COLUMNS = ["source", "kind", "script_type", "FLD", "address", "type",
                "name_in_db", "tag", "binding", *STAGES, "outputs", "flag"]


def render_csv(result) -> tuple:
    """(header, rows) for the per-signal coverage CSV (one row per staged row)."""
    rows = [[rec.get(c, "") for c in _CSV_COLUMNS] for rec in result["records"]]
    return _CSV_COLUMNS, rows


def render_txt(result) -> str:
    """The human coverage summary."""
    st = result["stats"]
    L = ["PIPELINE COVERAGE REPORT  (phase 910)",
         "=" * 52,
         f"Staged rows: {st['rows']}   "
         f"(signals {st['kinds']['signal']}, raw I/O channels {st['kinds']['channel']}, "
         f"structural {st['kinds']['structural']})",
         "Source: the on-disk output artifacts under the output root.",
         ""]
    if st["missing"]:
        L.append("MISSING outputs (run the phase to populate; stage shown empty below): "
                 + ", ".join(_STAGE_LABEL[s] for s in st["missing"]))
        L.append("")
    L.append("Stage coverage (rows landing in each output):")
    for s in STAGES:
        miss = "  (not generated)" if s in st["missing"] else ""
        L.append(f"  {_STAGE_LABEL[s]:<16}{st['stages'][s]:>5}{miss}")
    L += ["", f"ORPHAN  (a signal that lands in no output): {st['orphans']}"]
    if result["orphans"]:
        for rec in result["orphans"]:
            L.append(f"  {rec['source']:<14} {rec['script_type']:<6} {rec['FLD']}  "
                     f"[{rec['address']}]  type={rec['type']}")
    else:
        L.append("  -- none: every staged signal lands in at least one output --")
    L += ["", f"UNPLACED  (an output references a generated-DB member that 520 never created): "
              f"{st['unplaced']}"]
    if result["unplaced"]:
        for u in result["unplaced"]:
            L.append(f'  [{u["db"]}] "{u["member"]}"   <- {u["referenced_by"]}')
    else:
        L.append("  -- none: every emitted DB-member reference resolves to an actual member --")
    L += ["",
          f"Pipeline3-generated data blocks: {', '.join(st['generated_dbs']) or '(none found)'}",
          "",
          "Per-signal detail -> io_project_coverage_report.csv"]
    return "\n".join(L) + "\n"


def generate(rows, out_root) -> dict:
    """910 entry: trace the on-disk outputs and write Reports/io_project_coverage_report.{csv,txt}.
    Returns the trace result + {'csv_path', 'txt_path'}."""
    result = trace(rows, out_root)
    base = config.out_path(out_root, "coverage_report")
    os.makedirs(os.path.dirname(base), exist_ok=True)
    header, table = render_csv(result)
    csv_path = csv_tables.write_table(base + ".csv", header, table)
    txt_path = base + ".txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(render_txt(result))
    result["csv_path"], result["txt_path"] = csv_path, txt_path
    return result
