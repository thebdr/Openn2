"""ColumnResolver: canonical name <-> column letter <-> 1-based index, from column_map.csv
(document="IoList"). NOTHING else in the populator hardcodes a column letter.
"""
from __future__ import annotations
from openpyxl.utils import column_index_from_string

from pipeline3.core import config


class ColumnResolver:
    def __init__(self, document: str = "IoList"):
        self._letter: dict = {}
        self._required: dict = {}
        self._expected: dict = {}
        for r in config.load_column_map(document):
            self._letter[r["canonical"]] = r["column"]
            self._required[r["canonical"]] = r["required"]
            self._expected[r["canonical"]] = r["expected_header"]

    def __contains__(self, canonical: str) -> bool:
        return canonical in self._letter

    def letter(self, canonical: str) -> str:
        return self._letter[canonical]

    def idx(self, canonical: str) -> int:
        return column_index_from_string(self._letter[canonical])

    def expected_header(self, canonical: str) -> str:
        return self._expected.get(canonical, "")

    def required(self, canonical: str) -> bool:
        return bool(self._required.get(canonical, False))

    # --- worksheet access by canonical name --- #
    def get(self, ws, row: int, canonical: str):
        if canonical not in self._letter:
            return None
        return ws.cell(row=row, column=self.idx(canonical)).value

    def cell(self, ws, row: int, canonical: str):
        return ws.cell(row=row, column=self.idx(canonical))

    def set(self, ws, row: int, canonical: str, value) -> None:
        ws.cell(row=row, column=self.idx(canonical)).value = value
