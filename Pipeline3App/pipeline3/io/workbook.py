"""THE one shared workbook reader (plan §4.1). Every phase opens sheets through here.

It owns the openpyxl load-mode policy, regex sheet resolution, by-position cell access, strike
detection, and the header check (which normalizes newline->space, then prefix-matches - real I/O
List headers wrap across lines and carry suffixes like "(1° part)").

Load modes:
- read (default): data_only=True, read_only=False -> cached values AND cell fonts (strike) available.
- write: load_for_write(path) -> data_only=False, read_only=False -> formulas preserved for the
  populator's edit+save path.
"""
from __future__ import annotations
import os
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from pipeline3.core import config


def _norm_header(s) -> str:
    """Newline -> space, then whitespace-collapse + strip. The canonical header form for matching."""
    text = str(s if s is not None else "").replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def _colidx(col) -> int:
    """Accept a 1-based int or a column letter ('A','AB'); return the 1-based index."""
    if isinstance(col, int):
        return col
    return column_index_from_string(str(col).strip())


def open_workbook(path: str, *, data_only: bool = True, read_only: bool = False):
    return load_workbook(path, data_only=data_only, read_only=read_only)


def load_for_write(path: str):
    """Formula-preserving load for editing + saving (the populator). Not data_only, not read_only."""
    return load_workbook(path, data_only=False, read_only=False)


def available_sheets(path: str) -> list:
    wb = load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


class SheetView:
    """A positional view over one worksheet. Cells are addressed by (row, col) where col is a
    1-based int or a column letter. Row/col are 1-based to match Excel + openpyxl."""

    def __init__(self, ws, doc_label: str = "", header_row: int = 1, first_data_row: int | None = None):
        self.ws = ws
        self.name = ws.title
        self.doc = doc_label
        self.header_row = header_row
        self.first_data_row = first_data_row if first_data_row is not None else header_row + 1

    # --- dimensions --- #
    @property
    def max_row(self) -> int:
        return self.ws.max_row

    @property
    def max_col(self) -> int:
        return self.ws.max_column

    # --- values --- #
    def cell(self, row: int, col):
        """The cell value (by position). None when empty."""
        return self.ws.cell(row=row, column=_colidx(col)).value

    def text(self, row: int, col) -> str:
        """The cell value as a stripped string ('' when empty)."""
        v = self.cell(row, col)
        return "" if v is None else str(v).strip()

    # --- headers --- #
    def header(self, col) -> str:
        """The normalized (newline->space, collapsed) header text at `col`."""
        return _norm_header(self.cell(self.header_row, col))

    def header_matches(self, col, expected: str) -> bool:
        """True if the header at `col` starts with `expected` (both normalized, case-insensitive)."""
        return self.header(col).lower().startswith(_norm_header(expected).lower())

    # --- strikethrough --- #
    def cell_struck(self, row: int, col) -> bool:
        f = self.ws.cell(row=row, column=_colidx(col)).font
        return bool(f and f.strike)

    def row_struck(self, row: int, cols=None) -> bool:
        """True if ANY cell in the row (over `cols`, default the used columns) is struck through."""
        cols = cols if cols is not None else range(1, self.max_col + 1)
        return any(self.cell_struck(row, c) for c in cols)

    # --- iteration + links --- #
    def data_rows(self):
        """Row indices from first_data_row through max_row (inclusive)."""
        return range(self.first_data_row, self.max_row + 1)

    def location(self, row: int, col) -> str:
        """A clickable 'doc!sheet!cell' link string, e.g. 'IOList.xlsx!NET SAFETY 50!G7'."""
        return f"{self.doc}!{self.name}!{get_column_letter(_colidx(col))}{row}"

    def close(self):
        try:
            self.ws.parent.close()
        except Exception:  # noqa: BLE001
            pass


def open_sheet(path: str, sheet_pattern, header_row: int = 1, *, doc_label: str | None = None,
               first_data_row: int | None = None, data_only: bool = True,
               read_only: bool = False) -> SheetView:
    """Open the first sheet of `path` matching `sheet_pattern` (regex via config.resolve_sheet).
    Raises ValueError when no sheet matches. Caller may .close() the view when done."""
    wb = open_workbook(path, data_only=data_only, read_only=read_only)
    name = config.resolve_sheet(sheet_pattern, wb.sheetnames)
    if name is None:
        wb.close()
        raise ValueError(f"no sheet matching {sheet_pattern!r} in {os.path.basename(path)} "
                         f"(have: {wb.sheetnames})")
    return SheetView(wb[name], doc_label or os.path.basename(path), header_row, first_data_row)


def open_sheets(path: str, sheet_pattern, header_row: int = 1, *, doc_label: str | None = None,
                first_data_row: int | None = None, data_only: bool = True,
                read_only: bool = False) -> list:
    """Open ALL sheets matching `sheet_pattern` (the I/O List sheet may be a list/regex). Returns a
    list of SheetView (possibly empty). Each view shares the one underlying workbook."""
    wb = open_workbook(path, data_only=data_only, read_only=read_only)
    names = config.resolve_sheets(sheet_pattern, wb.sheetnames)
    label = doc_label or os.path.basename(path)
    return [SheetView(wb[n], label, header_row, first_data_row) for n in names]
