"""xlsxview.py - the read-only Excel preview for the Files tab.

Shows an .xlsx/.xlsm workbook in a tksheet grid - the cached VALUES (openpyxl `data_only`), or the
raw **formulas** when *Show formulas* is ticked. A **sheet selector** switches sheets, and **Tab /
Shift+Tab** step to the next / previous sheet (the read-only grid doesn't need Tab for cell entry -
arrow keys move cells). Columns get the zebra + right-click **Sort/Filter** from `grid.decorate`.
Editing always hands off to LibreOffice/Excel (the workbook's formulas / array-formula / VBA can't
survive an openpyxl save) via the header's **Open folder** + **Edit externally**. Capped to
3000 rows x 80 cols for responsiveness.
"""
from __future__ import annotations
import tkinter as tk
from tkinter import ttk

from openpyxl import load_workbook
from tksheet import Sheet

from pipeline3.gui import grid, widgets

_MAX_ROWS, _MAX_COLS = 3000, 80
# read-only-friendly bindings: select / navigate / copy / the right-click menu (for Sort/Filter) -
# no cell editing, no row/column insert-delete.
_BINDINGS = ("single_select", "drag_select", "row_select", "column_select",
             "column_width_resize", "arrowkeys", "copy", "right_click_popup_menu", "rc_select")


class XlsxViewer(ttk.Frame):
    def __init__(self, parent, path, pal, font_family, dark, on_status, on_external, **kw):
        super().__init__(parent, **kw)
        self.path = path
        self.pal = pal
        self.font = font_family
        self.dark = dark
        self.on_status = on_status
        try:
            wb = load_workbook(path, read_only=True)
            self.sheets = list(wb.sheetnames)
            wb.close()
        except Exception as e:  # noqa: BLE001
            ttk.Label(self, text=f"could not open {path}: {e}", padding=12).pack(anchor="nw")
            self.sheets = []
            return
        self.cur = self.sheets[0] if self.sheets else None
        self.formulas = tk.BooleanVar(value=False)
        self._grid = None

        header, bar = widgets.editor_header(self, path, on_status)
        header.pack(side="top", fill="x")
        ttk.Button(bar, text="Reload", command=self._reload).pack(side="right")
        ttk.Button(bar, text="Edit externally", command=lambda: on_external(path)).pack(side="right", padx=6)

        sel = ttk.Frame(self, padding=(6, 2))
        sel.pack(side="top", fill="x")
        ttk.Label(sel, text="Sheet:").pack(side="left")
        self.box = ttk.Combobox(sel, values=self.sheets, state="readonly", width=26)
        self.box.set(self.cur or "")
        self.box.pack(side="left", padx=(4, 12))
        self.box.bind("<<ComboboxSelected>>", lambda e: self._show(self.box.get()))
        ttk.Checkbutton(sel, text="Show formulas", variable=self.formulas,
                        command=self._reload).pack(side="left")
        ttk.Label(sel, text="Tab / Shift+Tab: next / previous sheet").pack(side="left", padx=12)

        self.holder = ttk.Frame(self)
        self.holder.pack(side="top", fill="both", expand=True)
        self._show(self.cur)

    def _show(self, name, focus_grid: bool = False):
        if not name:
            return
        self.cur = name
        try:
            wb = load_workbook(self.path, data_only=not self.formulas.get(), read_only=True)
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _MAX_ROWS:
                    break
                rows.append(["" if c is None else str(c) for c in row[:_MAX_COLS]])
            wb.close()
        except Exception as e:  # noqa: BLE001
            self.on_status(f"could not read sheet '{name}': {e}")
            return
        for w in self.holder.winfo_children():
            w.destroy()
        sheet = Sheet(self.holder, theme="dark blue" if self.dark else "light blue", data=rows or [[""]])
        sheet.enable_bindings(*_BINDINGS)
        sheet.pack(side="top", fill="both", expand=True)
        self._grid = grid.decorate(sheet, self.dark, on_status=self.on_status)
        # repurpose Tab / Shift+Tab on the table -> step sheets (read-only: cells use the arrow keys)
        for w in (sheet, sheet.MT):
            w.bind("<Tab>", lambda e: self._step(1))
            w.bind("<Shift-Tab>", lambda e: self._step(-1))
            w.bind("<ISO_Left_Tab>", lambda e: self._step(-1))   # Linux Shift+Tab
        self.box.set(name)
        if focus_grid:
            sheet.MT.focus_set()

    def _step(self, delta: int):
        if self.sheets:
            i = self.sheets.index(self.cur)
            self._show(self.sheets[(i + delta) % len(self.sheets)], focus_grid=True)
        return "break"

    def _reload(self):
        self._show(self.cur)
