"""grid.py - decorate a tksheet `Sheet` with zebra striping + right-click column sort/filter.

- **Filter** is VIEW-only and SAVE-safe: `display_rows` hides non-matching rows (the data is untouched,
  so `get_sheet_data()` still returns every row - a Save of an edited CSV keeps all rows). The filter
  text box is a **regex** (case-insensitive `re.search`), or pick from the column's distinct values.
- **Sort** physically REORDERS the data rows (tksheet's `display_rows` can't reorder, only subset).
  Row 0 (our CSV/xlsx header row) is kept pinned; the original order is snapshotted on first sort/filter
  so **Clear** restores it. A Save after a sort writes the sorted order (Clear / Reload to revert).
The active column shows the state on its header (the A/B/C cell): ` ↑`/` ↓` sort, ` ▽` filter, + a
status line. Re-stripes after every change so the zebra follows the visible rows.
"""
from __future__ import annotations
import re
import tkinter as tk
from tkinter import ttk

from openpyxl.utils import get_column_letter

_STRIPE_DARK = "#2a2a2a"
_STRIPE_LIGHT = "#eef2f7"
_HEADERTINT_DARK = "#333333"
_HEADERTINT_LIGHT = "#dbe5f0"
_SORT_ASC, _SORT_DESC, _FILTER = " ↑", " ↓", " ▽"   # ↑ ↓ ▽


def _key(v):
    """Sort key: numeric when the whole cell is a number, else case-folded text."""
    s = str(v).strip()
    try:
        return (0, float(s))
    except ValueError:
        return (1, s.lower())


class Grid:
    """Wraps a tksheet `Sheet`; adds zebra + header-menu sort/filter + active-column indicators."""

    def __init__(self, sheet, dark: bool, header_row: bool = True, on_status=lambda *_a: None):
        self.s = sheet
        self.dark = dark
        self.header_row = header_row
        self.on_status = on_status
        self._orig = None                                   # snapshot of the file's row order (lazy)
        self.sort_col = None
        self.sort_rev = False
        self.filter_col = None
        self._displayed = list(range(self._n()))            # data-row indices currently shown (in order)
        self._add_menu()
        self._restyle()

    def _n(self) -> int:
        return len(self.s.get_sheet_data())

    def _ncols(self) -> int:
        return max((len(r) for r in self.s.get_sheet_data()), default=0)

    def _head(self):
        return [0] if (self.header_row and self._n()) else []

    def _col(self) -> int:
        cols = self.s.get_selected_columns(return_tuple=True)
        return cols[0] if cols else 0

    def _ensure_orig(self):
        if self._orig is None:
            self._orig = [list(r) for r in self.s.get_sheet_data()]

    # ---- actions --------------------------------------------------------- #
    def sort(self, reverse: bool):
        self._ensure_orig()
        c = self._col()
        data = self.s.get_sheet_data()
        head = data[:1] if self.header_row else []
        body = sorted(data[len(head):], key=lambda r: _key(r[c] if c < len(r) else ""), reverse=reverse)
        self.s.set_sheet_data(head + body, redraw=False)    # physically reorder (display_rows can't)
        self._displayed = list(range(self._n()))
        self.sort_col, self.sort_rev, self.filter_col = c, reverse, None
        self._restyle()
        self.on_status(f"sorted by column {get_column_letter(c + 1)} "
                       f"{'descending' if reverse else 'ascending'} (Clear/Reload to revert)")

    def filter(self):
        c = self._col()
        data = self.s.get_sheet_data()
        header = str(data[0][c]) if (self.header_row and data and c < len(data[0])) \
            else f"column {get_column_letter(c + 1)}"
        body = range(1, len(data)) if self.header_row else range(len(data))
        distinct = sorted({str(data[i][c]) for i in body if c < len(data[i])}, key=_key)
        self._dialog(c, header, distinct)

    def _apply(self, c, pred, label):
        self._ensure_orig()
        data = self.s.get_sheet_data()
        keep = self._head() + [i for i in (range(1, len(data)) if self.header_row else range(len(data)))
                               if pred(str(data[i][c] if c < len(data[i]) else ""))]
        self._displayed = keep
        self.filter_col = c
        self._restyle()
        total = len(data) - len(self._head())
        self.on_status(f"filtered column {get_column_letter(c + 1)} ({label}): "
                       f"{len(keep) - len(self._head())} of {total} row(s)")

    def clear(self):
        if self._orig is not None:
            self.s.set_sheet_data([list(r) for r in self._orig], redraw=False)
            self._orig = None
        self._displayed = list(range(self._n()))
        self.sort_col, self.filter_col = None, None
        self._restyle()
        self.on_status("sort/filter cleared")

    # ---- display + indicators -------------------------------------------- #
    def _restyle(self):
        self.s.display_rows(self._displayed, all_rows_displayed=False, reset_row_positions=True, redraw=False)
        self._headers()
        self._zebra()
        self.s.redraw()

    def _headers(self):
        heads = []
        for c in range(self._ncols()):
            h = get_column_letter(c + 1)
            if c == self.sort_col:
                h += _SORT_DESC if self.sort_rev else _SORT_ASC
            if c == self.filter_col:
                h += _FILTER
            heads.append(h)
        if heads:
            self.s.headers(heads)

    def _zebra(self):
        self.s.dehighlight_all(redraw=False)
        head = self._head()
        if head:
            self.s.highlight_rows(head, bg=_HEADERTINT_DARK if self.dark else _HEADERTINT_LIGHT, redraw=False)
        body = [d for d in self._displayed if d not in head]
        stripe = body[1::2]                       # every other VISIBLE body row
        if stripe:
            self.s.highlight_rows(stripe, bg=_STRIPE_DARK if self.dark else _STRIPE_LIGHT, redraw=False)

    def _add_menu(self):
        opts = dict(header_menu=True, table_menu=False, index_menu=False, empty_space_menu=False)
        self.s.popup_menu_add_command("Sort ↑ (this column)", lambda: self.sort(False), **opts)
        self.s.popup_menu_add_command("Sort ↓ (this column)", lambda: self.sort(True), **opts)
        self.s.popup_menu_add_command("Filter… (this column)", self.filter, **opts)
        self.s.popup_menu_add_command("Clear sort/filter", lambda: self.clear(), **opts)

    # ---- the filter dialog (regex OR a pick of distinct values) ---------- #
    def _dialog(self, c, header, distinct):
        top = tk.Toplevel(self.s)
        top.title(f"Filter: {header}")
        top.transient(self.s.winfo_toplevel())
        ttk.Label(top, text=f"Filter '{header}'", padding=(8, 6)).pack(anchor="w")
        row = ttk.Frame(top, padding=(8, 0))
        row.pack(fill="x")
        ttk.Label(row, text="regex:").pack(side="left")
        rx = tk.StringVar()
        e = ttk.Entry(row, textvariable=rx, width=26)
        e.pack(side="left", padx=4)
        ttk.Label(top, text="(case-insensitive regex — e.g.  ^KQ   contactor|relay   \\d{2,})",
                  padding=(8, 0)).pack(anchor="w")
        ttk.Label(top, text="or pick values (none = all):", padding=(8, 4)).pack(anchor="w")
        lb = tk.Listbox(top, selectmode="extended", height=min(12, max(3, len(distinct))), width=40,
                        exportselection=False)
        for v in distinct:
            lb.insert("end", v)
        lb.pack(fill="both", expand=True, padx=8)
        btns = ttk.Frame(top, padding=8)
        btns.pack(fill="x")

        def apply_():
            pat = rx.get().strip()
            picks = {distinct[i] for i in lb.curselection()}
            if pat:
                try:
                    rec = re.compile(pat, re.IGNORECASE)
                except re.error as err:
                    self.on_status(f"invalid regex: {err}")
                    return
                self._apply(c, lambda v, r=rec: bool(r.search(v)), f"/{pat}/")
            elif picks:
                self._apply(c, lambda v, p=picks: v in p, f"{len(picks)} value(s)")
            else:
                self.clear()
            top.destroy()

        ttk.Button(btns, text="Apply", command=apply_).pack(side="right")
        ttk.Button(btns, text="Cancel", command=top.destroy).pack(side="right", padx=6)
        ttk.Button(btns, text="Clear", command=lambda: (top.destroy(), self.clear())).pack(side="left")
        top.bind("<Return>", lambda _e: apply_())
        top.bind("<Escape>", lambda _e: top.destroy())
        e.focus_set()


def decorate(sheet, dark: bool, header_row: bool = True, on_status=lambda *_a: None) -> Grid:
    """Attach zebra + right-click sort/filter (with active-column indicators) to `sheet`."""
    return Grid(sheet, dark, header_row=header_row, on_status=on_status)
