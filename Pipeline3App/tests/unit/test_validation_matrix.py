"""M6 gate (data-independent): 120 Validate C&E Matrix standalone.

The CAUSE&EFFECT MATRIX must carry only input (I) addresses; the AREA sheets only output (Q/O)
addresses; no address repeats within a sheet. C&E absent -> one ce_absent (SKIP), no crash.
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as CI

from pipeline3.context import PipelineContext
from pipeline3.domain.validation import run_ce_matrix


def _set(ws, col, row, val):
    ws.cell(row=row, column=CI(col), value=val)


def _make_ce(path):
    wb = Workbook()
    m = wb.active; m.title = "CAUSE&EFFECT MATRIX"
    _set(m, "F", 5, "I20.0")     # input ok
    _set(m, "F", 6, "Q0.0")      # output on the matrix -> matrix_addr_kind
    _set(m, "F", 7, "I20.0")     # dup of row 5 -> dup_addr
    a = wb.create_sheet("AREA 1")
    _set(a, "C", 4, "Q1.0")      # output ok
    _set(a, "C", 5, "I50.0")     # input on an AREA sheet -> area_addr_kind
    _set(a, "C", 6, "Q1.0")      # dup -> dup_addr
    wb.save(path)


def _ctx(ce_path):
    return PipelineContext(
        params={"io_list": {"path": "X/IOList.xlsx", "sheet": r"NET SAFETY \d+", "header_row": 1},
                "ce": {"path": ce_path, "matrix_sheet": "CAUSE&EFFECT MATRIX", "matrix_header_row": 2,
                       "matrix_data_row": 5, "area_header_row": 3, "area_data_row": 4},
                "areas": r"AREA\s?\d{1,2}"},
        out_root="X", lang="en", profile="main")


def _run():
    with tempfile.TemporaryDirectory() as d:
        ce = os.path.join(d, "ce.xlsx")
        _make_ce(ce)
        return run_ce_matrix(_ctx(ce))


def test_address_kind_and_dups():
    e = _run()
    types = [x.type for x in e]
    ok("matrix_addr_kind" in types, "Q address on the matrix flagged")
    ok("area_addr_kind" in types, "I address on an AREA sheet flagged")
    eq(types.count("dup_addr"), 2, "one dup per sheet")
    ok("ce_summary" in types)


def test_absent_ce_skips():
    e = run_ce_matrix(_ctx(os.path.join("nope", "missing.xlsx")))
    eq(len(e), 1, "exactly one entry")
    eq(e[0].type, "ce_absent")
    eq(e[0].level, "SKIP", "absent C&E -> the sub-phase is skipped")


if __name__ == "__main__":
    raise SystemExit(run("validation_matrix", [
        ("address_kind_and_dups", test_address_kind_and_dups),
        ("absent_ce_skips", test_absent_ce_skips),
    ]))
