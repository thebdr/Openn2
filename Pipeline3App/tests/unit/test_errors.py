"""M6b gate (data-independent): the phase-100 treatment registry.

Covers warn/skip downgrades, the reconcile policy (new untreated rows recorded, untreated-gone
pruned, treated-gone marked stale), treatment round-trip by uid, designer suppression, and the
document-mutating `accept` (idempotent + gated to the address-match/device-mismatch type).
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string as CI

from pipeline3.core import errors
from pipeline3.core.model import LogEntry
from pipeline3.context import PipelineContext


def _fail(phase, type_, location, detail="x"):
    return LogEntry(level="FAIL", phase=phase, type=type_, location=location, detail=detail)


def _ctx(rows=None, io_path="X/IOList.xlsx"):
    return PipelineContext(params={"io_list": {"path": io_path}}, out_root="X",
                           rows=rows or [], lang="en", profile="main")


def _csv(d):
    return os.path.join(d, "error_management.csv")


def test_warn_and_skip_downgrade():
    with tempfile.TemporaryDirectory() as d:
        log = [_fail(110, "ts_missing", "S!U1"), _fail(110, "pp_empty", "S!K2"), _fail(110, "addr_format", "S!G3")]
        errors._write(_csv(d), {log[0].uid: errors.Treatment(uid=log[0].uid, treatment="warn"),
                                log[1].uid: errors.Treatment(uid=log[1].uid, treatment="skip")})
        errors.apply_and_reconcile(_ctx(), log, path=_csv(d))
        eq(log[0].level, "WARN")
        eq(log[1].level, "SKIP")
        eq(log[2].level, "FAIL", "untreated finding stays FAIL")


def test_reconcile_prune_stale_new():
    with tempfile.TemporaryDirectory() as d:
        seed = {"aaaaaaaaaa": errors.Treatment(uid="aaaaaaaaaa", treatment="skip", type="old", location="S!A1"),
                "bbbbbbbbbb": errors.Treatment(uid="bbbbbbbbbb", treatment="", type="old2", location="S!A2")}
        errors._write(_csv(d), seed)
        log = [_fail(130, "cem_fld_none", "CEM!F5")]              # a NEW finding this run
        errors.apply_and_reconcile(_ctx(), log, path=_csv(d))
        back = errors.load(_csv(d))
        ok(log[0].uid in back and back[log[0].uid].treatment == "", "new FAIL recorded, untreated")
        ok(back.get("aaaaaaaaaa") and back["aaaaaaaaaa"].status == "stale", "treated-but-gone -> stale")
        ok("bbbbbbbbbb" not in back, "untreated-but-gone -> pruned")


def test_treatment_roundtrip_by_uid():
    with tempfile.TemporaryDirectory() as d:
        errors.apply_and_reconcile(_ctx(), [_fail(140, "iol_cem_missing_mandatory", "S!O9")], path=_csv(d))
        back = errors.load(_csv(d))
        uid = next(iter(back))
        back[uid].treatment = "skip"
        errors._write(_csv(d), back)                              # the user records a treatment
        log2 = [_fail(140, "iol_cem_missing_mandatory", "S!O9")]  # same finding next run -> same uid
        errors.apply_and_reconcile(_ctx(), log2, path=_csv(d))
        eq(log2[0].level, "SKIP", "treatment persists across runs, keyed by uid")


def test_designer_suppression():
    with tempfile.TemporaryDirectory() as d:
        log = [_fail(110, "ts_missing", "S!U1"),
               LogEntry(level="WARN", phase=140, type="iol_cem_missing_plain", location="S!O2")]
        sup = os.path.join(d, "designer_suppressions.csv")
        errors._write(sup, {log[0].uid: errors.Treatment(uid=log[0].uid),
                            log[1].uid: errors.Treatment(uid=log[1].uid)})
        n = errors.apply_suppressions(log, path=sup)
        eq(n, 2)
        eq(log[0].level, "SKIP")
        eq(log[1].level, "SKIP")


def _make_iolist(path):
    wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
    ws.cell(row=5, column=CI("L"), value="DOOR OPEN")             # desc_l1b cell to annotate
    wb.save(path)


def _accept_row():
    return {"source_cell": "NET SAFETY 50!O5", "_source_sheet": "NET SAFETY 50", "_source_row": 5,
            "ce_functional_unit": "=S9", "ce_location": "+X", "ce_device": "-Z"}


def test_accept_mutates_doc_idempotently():
    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "io.xlsx"); _make_iolist(io)
        log = [_fail(140, "iol_cem_addr_only", "NET SAFETY 50!O5")]
        errors._write(_csv(d), {log[0].uid: errors.Treatment(uid=log[0].uid, treatment="accept")})
        ctx = _ctx(rows=[_accept_row()], io_path=io)
        errors.apply_and_reconcile(ctx, log, path=_csv(d))
        eq(log[0].level, "SKIP", "accept resolves the finding to SKIP")
        eq(load_workbook(io)["NET SAFETY 50"]["L5"].value, "DOOR OPEN [C&E: =S9 +X -Z]")
        log2 = [_fail(140, "iol_cem_addr_only", "NET SAFETY 50!O5")]
        errors.apply_and_reconcile(ctx, log2, path=_csv(d))
        eq(load_workbook(io)["NET SAFETY 50"]["L5"].value.count("[C&E:"), 1, "idempotent: no double-append")


def test_accept_gated_to_addr_only():
    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "io.xlsx"); _make_iolist(io)
        log = [_fail(140, "iol_cem_missing_mandatory", "NET SAFETY 50!O5")]   # NOT the addr_only type
        errors._write(_csv(d), {log[0].uid: errors.Treatment(uid=log[0].uid, treatment="accept")})
        errors.apply_and_reconcile(_ctx(rows=[_accept_row()], io_path=io), log, path=_csv(d))
        eq(log[0].level, "FAIL", "accept only applies to the address-match/device-mismatch finding")
        eq(load_workbook(io)["NET SAFETY 50"]["L5"].value, "DOOR OPEN", "no mutation for a gated-out type")


def test_set_treatment_quick_treat():
    with tempfile.TemporaryDirectory() as d:
        log = [_fail(130, "cem_fld_none", "CEM!F5")]
        uid = log[0].uid
        errors.apply_and_reconcile(_ctx(), log, path=_csv(d))            # seed the (untreated) row
        ok(errors.set_treatment(_csv(d), uid, "warn"), "the finding's row is found + set")
        eq(errors.load(_csv(d))[uid].treatment, "warn", "warn written, other rows preserved")
        ok(errors.set_treatment(_csv(d), uid, ""), "clear returns True")
        eq(errors.load(_csv(d))[uid].treatment, "", "treatment cleared")
        ok(not errors.set_treatment(_csv(d), "deadbeef00", "warn"), "an unknown uid -> False (no write)")
        errors.set_treatment(_csv(d), uid, "warn")
        log2 = [_fail(130, "cem_fld_none", "CEM!F5")]                    # the SAME finding next run
        errors.apply_and_reconcile(_ctx(), log2, path=_csv(d))
        eq(log2[0].level, "WARN", "the recorded warn downgrades the finding on the next validation run")


if __name__ == "__main__":
    raise SystemExit(run("errors", [
        ("set_treatment_quick_treat", test_set_treatment_quick_treat),
        ("warn_and_skip_downgrade", test_warn_and_skip_downgrade),
        ("reconcile_prune_stale_new", test_reconcile_prune_stale_new),
        ("treatment_roundtrip_by_uid", test_treatment_roundtrip_by_uid),
        ("designer_suppression", test_designer_suppression),
        ("accept_mutates_doc_idempotently", test_accept_mutates_doc_idempotently),
        ("accept_gated_to_addr_only", test_accept_gated_to_addr_only),
    ]))
