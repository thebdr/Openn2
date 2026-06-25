"""Clean Addresses & Electrical Names (the manual 165/175 buttons): the pure strip rule + the
surgical in-place clean of a synthetic I/O List + C&E (apostrophes/spaces removed, a formula cell left
untouched, idempotent re-run drops its backup). Data-independent."""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string as CI

from pipeline3.core import config
from pipeline3.domain import clean


def _set(ws, col, row, val, text=False):
    c = ws.cell(row=row, column=CI(col), value=val)
    if text:
        c.data_type = "s"
    return c


def test_clean_value_rule():
    eq(clean._clean_value("-S1' 2"), "-S12", "apostrophe + spaces removed")
    eq(clean._clean_value("  =S1 "), "=S1", "outer whitespace collapsed; leading = kept")
    eq(clean._clean_value("=S1"), "=S1", "clean value unchanged")
    ok(clean._clean_value(7) is None and clean._clean_value(None) is None and clean._clean_value(True) is None,
       "numeric / None / bool -> None (skipped)")
    eq(clean._clean_value("ABC"), "ABC", "case preserved")
    clean.CLEAN_SYMBOLS.append("*")
    try:
        eq(clean._clean_value("A*B"), "AB", "a newly-added symbol is stripped too")
    finally:
        clean.CLEAN_SYMBOLS.remove("*")


def _make_iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    for col, name in (("G", "Bit"), ("O", "Functional unit"), ("P", "Location"), ("Q", "Device")):
        _set(ws, col, 1, name)
    _set(ws, "G", 2, "I 1.0", text=True)         # a space
    _set(ws, "O", 2, "'=S1", text=True)          # a leading apostrophe (on a leading-= value)
    _set(ws, "P", 2, "+MS1'.CC1", text=True)     # a mid apostrophe
    _set(ws, "Q", 2, "-F09001", text=True)       # already clean
    ws.cell(3, CI("Q")).value = "=A1"            # a real FORMULA in a target column -> must be untouched
    wb.save(path)


def test_clean_iolist_end_to_end():
    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "io.xlsx")
        _make_iolist(io)
        params = config.load_params()
        params["io_list"] = {"path": io, "sheet": ["NET SAFETY 50"], "header_row": 1}

        logs = []
        res = clean.clean_iolist(params, emit=logs.append)
        ok(any("cleaning columns" in m and "G:bit" in m and "O:functional_unit" in m for m in logs),
           "the scanned columns are logged per sheet")
        ok(res.changed >= 3, f"the 3 dirty cells were cleaned (got {res.changed})")
        ok(res.backup and os.path.exists(res.backup), "a backup is kept when something changed")
        eq(res.skipped_formula, 1, "the formula cell was skipped, not flattened")

        wb = load_workbook(io)
        ws = wb["NET SAFETY 50"]
        eq(ws.cell(2, CI("G")).value, "I1.0", "the space is collapsed")
        eq(ws.cell(2, CI("O")).value, "=S1", "the leading apostrophe is removed")
        ok(ws.cell(2, CI("O")).data_type != "f", "the cleaned leading-= value stays TEXT, not a formula")
        eq(ws.cell(2, CI("P")).value, "+MS1.CC1", "the mid apostrophe is removed")
        eq(ws.cell(3, CI("Q")).value, "=A1", "the formula cell is untouched")
        eq(ws.cell(3, CI("Q")).data_type, "f", "the formula cell is still a formula")
        wb.close()

        res2 = clean.clean_iolist(params, emit=lambda *_: None)
        eq(res2.changed, 0, "a re-run finds nothing to clean")
        eq(res2.backup, "", "the no-op re-run drops its backup")


def _make_ce(path):
    wb = Workbook()
    m = wb.active
    m.title = "CAUSE&EFFECT MATRIX"
    for col, name in (("F", "BIT"), ("J", "FU"), ("K", "LOC"), ("L", "DEV")):
        _set(m, col, 2, name)
    _set(m, "F", 5, "I 1.0", text=True)
    _set(m, "J", 5, "'=S1", text=True)
    _set(m, "K", 5, "+MS1 .CC1", text=True)
    _set(m, "L", 5, "-F0'9001", text=True)
    a = wb.create_sheet("AREA 1")
    _set(a, "A", 3, "SIGLA")
    _set(a, "C", 3, "OUT")
    _set(a, "A", 4, "+K1'", text=True)
    _set(a, "C", 4, "Q 2.3", text=True)
    wb.save(path)


def test_clean_cematrix_end_to_end():
    with tempfile.TemporaryDirectory() as d:
        ce = os.path.join(d, "ce.xlsx")
        _make_ce(ce)
        params = config.load_params()
        params["ce"] = {"path": ce, "matrix_sheet": "CAUSE&EFFECT MATRIX",
                        "matrix_data_row": 5, "area_data_row": 4}
        params["areas"] = ["AREA 1"]

        res = clean.clean_cematrix(params, emit=lambda *_: None)
        ok(res.changed >= 6, f"matrix + area dirty cells cleaned (got {res.changed})")
        wb = load_workbook(ce)
        eq(wb["CAUSE&EFFECT MATRIX"].cell(5, CI("J")).value, "=S1", "matrix FU cleaned")
        eq(wb["CAUSE&EFFECT MATRIX"].cell(5, CI("L")).value, "-F09001", "matrix Device apostrophe removed")
        eq(wb["AREA 1"].cell(4, CI("A")).value, "+K1", "AREA contactor sigla cleaned")
        eq(wb["AREA 1"].cell(4, CI("C")).value, "Q2.3", "AREA address cleaned")
        wb.close()


def test_clean_cematrix_absent_is_noop():
    params = config.load_params()
    params["ce"] = {"path": ""}
    res = clean.clean_cematrix(params, emit=lambda *_: None)
    eq(res.changed, 0, "no C&E -> nothing changed")
    ok(res.note and not res.backup, "a clear note, no backup, no exception")


if __name__ == "__main__":
    raise SystemExit(run("clean", [
        ("clean_value_rule", test_clean_value_rule),
        ("clean_iolist_end_to_end", test_clean_iolist_end_to_end),
        ("clean_cematrix_end_to_end", test_clean_cematrix_end_to_end),
        ("clean_cematrix_absent_is_noop", test_clean_cematrix_absent_is_noop),
    ]))
