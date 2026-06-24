"""The single surgical Excel-write primitive (io/xlsx_edit). The engine `set_cells` rewrites chosen
cells in a worksheet's XML while preserving every untouched cell/formula, and - when an edit lands in
an array formula's spill range - FREEZES that array to its cached values (keeping ALL spilled values,
dropping the formula) instead of leaving a master+literal overlap (the openpyxl-flatten bug that
corrupted the I/O List). Pure-string, data-independent."""
import xml.etree.ElementTree as ET
from _harness import run, ok, eq
from pipeline3.io import xlsx_edit as xe

# a worksheet with: a header row, an ARRAY master A2:A4 (spill cells A3/A4 carry cached values), a
# SHARED formula B2:B3, and styled cells.
_SHEET = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<sheetData>'
    '<row r="1"><c r="A1" t="inlineStr"><is><t>H</t></is></c><c r="C1" s="5"/></row>'
    '<row r="2"><c r="A2" cm="1"><f t="array" ref="A2:A4">SEQ(3)</f><v>10</v></c>'
    '<c r="B2" s="7"><f t="shared" ref="B2:B3" si="0">A2+1</f><v>11</v></c></row>'
    '<row r="3"><c r="A3" t="str"><v>20</v></c><c r="B3"><f t="shared" si="0"/><v>21</v></c></row>'
    '<row r="4"><c r="A4" t="str"><v>30</v></c></row>'
    '</sheetData></worksheet>'
)


def _wf(xml):
    ET.fromstring(xml)  # raises on malformed -> the test fails


def test_edit_existing_cell_preserves_style_neighbours_and_unrelated_array():
    out, mats = xe.set_cells(_SHEET, {"C1": "PA"})
    _wf(out)
    eq(mats, [], "no array touched -> nothing materialized")
    ok('<c r="C1" s="5" t="inlineStr"><is><t xml:space="preserve">PA</t></is></c>' in out, "C1 set, style kept")
    ok('<c r="A1" t="inlineStr"><is><t>H</t></is></c>' in out, "A1 untouched")
    ok('t="array"' in out, "the A2 array master is preserved (C1 is not in its range)")
    ok('<f t="shared" ref="B2:B3" si="0">A2+1</f>' in out, "shared formula preserved verbatim")


def test_edit_into_array_freezes_to_cached_values_no_overlap():
    out, mats = xe.set_cells(_SHEET, {"A3": "X"})          # A3 is inside the A2:A4 array
    _wf(out)
    eq(mats, [("A2", "A2:A4")], "the frozen array is reported for the WARN log (cell + range)")
    ok('t="array"' not in out, "the array formula is gone (no master+literal overlap)")
    ok('cm="1"' not in out, "the dangling dynamic-array metadata ref is stripped")
    ok('<c r="A2"><v>10</v></c>' in out, "master kept its cached value, formula dropped")
    ok('<v>30</v>' in out, "the OTHER spilled value (A4) is kept - no data lost")
    ok('<is><t xml:space="preserve">X</t></is>' in out, "A3 overwritten with the edit")
    ok('<f t="shared"' in out, "the unrelated shared formula is left alone")


def test_insert_cell_into_row_column_sorted():
    out, _ = xe.set_cells(_SHEET, {"B1": "mid"})           # B1 doesn't exist; lands between A1 and C1
    _wf(out)
    r1 = out.split('<row r="1"')[1].split('</row>')[0]
    eq(__import__("re").findall(r'<c r="([A-Z]+)1"', r1), ["A", "B", "C"], "inserted cell is column-sorted")


def test_insert_new_row_sorted():
    out, _ = xe.set_cells(_SHEET, {"A9": "tail"})
    _wf(out)
    ok('<row r="9">' in out and 'tail' in out, "new row created")
    ok(out.index('<row r="4"') < out.index('<row r="9"'), "row 9 after row 4")


def test_html_escaping():
    out, _ = xe.set_cells(_SHEET, {"C1": "a&b<c>"})
    _wf(out)
    ok("a&amp;b&lt;c&gt;" in out, "value is XML-escaped")


def test_build_sheet_xml_grid():
    x = xe.build_sheet_xml([["H1", "H2"], ["=text", 7], [None, "", "x"]], hyperlinks=[("A2", "S!A1", "d")])
    _wf(x)
    ok('<dimension ref="A3:C3"' not in x and '<dimension ref="' in x, "has a dimension")
    ok("<sheetViews>" in x and "<sheetFormatPr" in x and "<pageMargins" in x, "full Excel worksheet structure")
    ok('<c r="A1" t="inlineStr"><is><t>H1</t></is></c>' in x, "header cell (no xml:space when unneeded)")
    ok('<c r="A2" t="inlineStr"><is><t>=text</t></is></c>' in x, "leading = is TEXT not a formula")
    ok('<c r="B2"><v>7</v></c>' in x, "int -> numeric cell")
    ok('<c r="C3"' in x and '<c r="A3"' not in x, "blank cells skipped, later cell still placed")
    ok('<hyperlink ref="A2" location="S!A1" display="d"/>' in x, "internal hyperlink")


def test_new_sheets_add_replace_delete_preserves_arrays():
    import os, tempfile, zipfile, warnings
    warnings.simplefilter("ignore")
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    path = os.path.join(tempfile.mkdtemp(), "w.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "keep"
    ws["C2"] = ArrayFormula("C2:C4", "=A1")          # an array we must NOT touch -> must survive
    wb.create_sheet("Legacy")["A1"] = "old"
    wb.save(path)
    wb.close()

    xe.edit_workbook(path, new_sheets=[
        {"name": "Gen", "rows": [["ID", "Count"], ["=S1+X", 5]], "hyperlinks": [("A2", "Data!A1", "go")]}],
        delete_sheets=["Legacy"])
    with zipfile.ZipFile(path) as z:
        ok(z.testzip() is None, "valid zip after add/delete sheets")
        dpart = xe._sheet_name_to_part(open(path, "rb").read())["Data"]
        ok('t="array"' in z.read(dpart).decode("utf-8"), "the untouched Data array is byte-preserved")
    wb = openpyxl.load_workbook(path)
    ok("Gen" in wb.sheetnames and "Legacy" not in wb.sheetnames, "Gen added, Legacy deleted")
    ok("Data" in wb.sheetnames, "untouched sheet kept")
    g = wb["Gen"]
    eq(g["A1"].value, "ID", "grid header")
    eq(g["B2"].value, 5, "numeric cell read back as a number")
    eq(g["A2"].value, "=S1+X", "a leading '=' stays TEXT, not a formula")
    wb.close()

    xe.edit_workbook(path, new_sheets=[{"name": "Gen", "rows": [["ONLY"]]}])   # re-run -> replace Gen
    wb = openpyxl.load_workbook(path)
    eq(wb["Gen"]["A1"].value, "ONLY", "Gen replaced in place")
    ok(wb["Gen"]["A2"].value is None, "the old Gen rows are gone")
    wb.close()


if __name__ == "__main__":
    raise SystemExit(run("xlsx_edit", [
        ("edit_existing_cell_preserves_style_neighbours_and_unrelated_array",
         test_edit_existing_cell_preserves_style_neighbours_and_unrelated_array),
        ("edit_into_array_freezes_to_cached_values_no_overlap", test_edit_into_array_freezes_to_cached_values_no_overlap),
        ("insert_cell_into_row_column_sorted", test_insert_cell_into_row_column_sorted),
        ("insert_new_row_sorted", test_insert_new_row_sorted),
        ("html_escaping", test_html_escaping),
        ("build_sheet_xml_grid", test_build_sheet_xml_grid),
        ("new_sheets_add_replace_delete_preserves_arrays", test_new_sheets_add_replace_delete_preserves_arrays),
    ]))
