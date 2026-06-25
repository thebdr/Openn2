"""M4 gate (data-independent): the phase-200 populator on a synthetic I/O List.

Covers: node-meta row skipped, the suggested-type ladder (alarm A, emergency-PB channel pair),
canonical mapping, index assignment (channel pairing -> shared index), diagnosis allocation
(node cabinet for A; Type-2 block for E1/2), idempotency, and DiagnosisBlocks creation.
"""
import os
import shutil
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string as CI

from pipeline3.core import config
from pipeline3.domain.iolist_diag import populate as pop
from pipeline3.domain.iolist_diag import script_type as st
from pipeline3.domain.iolist_diag import diag_alloc
from pipeline3.domain.iolist_diag.blocks import HEADERS as DIAG_HEADERS
from pipeline3.domain.iolist_diag.models import DIAGBLOCKS_SHEET, IoRow, RowResult


def _set(ws, col, row, val, text=False):
    c = ws.cell(row=row, column=CI(col), value=val)
    if text:
        c.data_type = "s"
    return c


def _make_iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    # header row 1 (positions matter, not text)
    for col, name in (("F", "ID"), ("G", "Bit"), ("K", "Desc1"), ("L", "Desc1b"),
                      ("O", "Functional unit"), ("P", "Location"), ("Q", "Device"), ("R", "Type")):
        _set(ws, col, 1, name)
    # row 2: node-meta "50" line -> must be skipped (no FU, no I/Q addr, no script_type)
    _set(ws, "F", 2, 50)
    _set(ws, "K", 2, "Profinet Network")
    # row 3: an alarm input (type A in col R)
    _set(ws, "G", 3, "I1.0"); _set(ws, "K", 3, "CIRCUIT BREAKER TRIPPED")
    _set(ws, "O", 3, "=S1", text=True); _set(ws, "P", 3, "+MS1.CC1", text=True)
    _set(ws, "Q", 3, "-F09001", text=True); _set(ws, "R", 3, "A")
    # rows 4/5: an emergency push-button channel pair (same FLD)
    for r, ch, dev in ((4, "CH1", "-S001"), (5, "CH2", "-S001")):
        _set(ws, "G", r, f"I2.{r-4}")
        _set(ws, "K", r, "EMERGENCY PUSH-BUTTON PRESSED")
        _set(ws, "L", r, ch)
        _set(ws, "O", r, "=S1", text=True); _set(ws, "P", r, "+MS1.CC1", text=True)
        _set(ws, "Q", r, dev, text=True)
    wb.save(path)


def _params(io_path):
    p = config.load_params()
    p["io_list"] = {"path": io_path, "sheet": ["NET SAFETY 50"], "header_row": 1}
    return p


def _by_addr(results, addr):
    return next(r for r in results if r.iorow.addr == addr)


def _run(fn):
    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "iolist.xlsx")
        out = os.path.join(d, "out")
        _make_iolist(io)
        params = _params(io)
        params["output_dir"] = out
        fn(params, out, d)


def test_populates_and_skips_nodemeta():
    def check(params, out, d):
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)
        eq(res.processed, 3, "node-meta row 2 skipped; 3 data rows")
        eq(res.unresolved, 0)
        ok(not res.halt)
    _run(check)


def test_script_type_and_index():
    def check(params, out, d):
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)
        a = _by_addr(res.results, "I1.0")
        e1 = _by_addr(res.results, "I2.0")
        e2 = _by_addr(res.results, "I2.1")
        eq(a.script_type, "A")
        eq(e1.script_type, "E1/2")
        eq(e2.script_type, "E2/2")
        eq(e1.index, "0001")
        eq(e2.index, "0001", "channel-2 inherits the anchor's index by FLD")
        eq(a.index, "0001", "standalone A on its own counter")
    _run(check)


def test_diag_allocation():
    def check(params, out, d):
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)
        a = _by_addr(res.results, "I1.0")
        e1 = _by_addr(res.results, "I2.0")
        e2 = _by_addr(res.results, "I2.1")
        ok(a.diag_cabinet and a.diag_bit, "alarm A diagnosed in a node cabinet")
        eq(a.diag_cabinet, "000")
        eq(a.diag_bit, "00")
        ok(e1.diag_cabinet and e1.diag_cabinet != "000", "E1/2 in a Type-2 block")
        eq(e2.diag_cabinet, "", "E2/2 is not in_diag")
        ok(res.blocks >= 2, f"node + Type-2 blocks, got {res.blocks}")
    _run(check)


def test_writes_columns_and_diagblocks():
    def check(params, out, d):
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        ok(DIAGBLOCKS_SHEET in wb.sheetnames, "DiagnosisBlocks sheet written")
        ws = wb["NET SAFETY 50"]
        eq(ws.cell(3, CI("AB")).value, "A")
        eq(ws.cell(4, CI("AB")).value, "E1/2")
        eq(ws.cell(4, CI("AD")).value, "0001")     # index written
        # AC = the auto-inferred canonical, kept for reference (== AB in Mode-1)
        eq(ws.cell(3, CI("AC")).value, "A")
        eq(ws.cell(4, CI("AC")).value, "E1/2")
        # the output-column headers are written from column_map (template left them blank)
        eq(ws.cell(1, CI("AA")).value, "Skip Reason")
        eq(ws.cell(1, CI("AB")).value, "Script Type")
        eq(ws.cell(1, CI("AC")).value, "Suggested Type")
        eq(ws.cell(1, CI("AD")).value, "Index")
        eq(ws.cell(1, CI("AE")).value, "Diag Cabinet")
        eq(ws.cell(1, CI("AF")).value, "Diag Bit")
        eq(ws.cell(1, CI("AG")).value, "Hardware Parameters")
        wb.close()
    _run(check)


def test_fire_alarm_rung():
    r1 = IoRow(sheet="s", row=1, addr="I20.3", desc_l1="FIRE ALARM", desc_l1b="CH1")
    r2 = IoRow(sheet="s", row=2, addr="I20.7", desc_l1="FIRE ALARM", desc_l1b="CH2")
    eq(st.suggested_type(r1), "F1/2", "fire alarm CH1 -> F1/2 (like E1/2)")
    eq(st.suggested_type(r2), "F2/2", "fire alarm CH2 -> F2/2")


def test_ladder_yields_canonical_directly():
    # the legacy ENC/FA/RES intermediates are gone: the ladder returns N/Z/R directly (== AB).
    enc = IoRow(sheet="s", row=1, addr="I3.0", desc_l1="SAFETY ENCODER PHOTOCELL", desc_l1b="CELL 1")
    eq(st.suggested_type(enc), "N1/2", "encoder -> N, not legacy ENC")
    fa = IoRow(sheet="s", row=2, addr="Q3.0", desc_l1="INDICATOR", desc_l1b="EMERGENCY AREA 3")
    eq(st.suggested_type(fa), "Z3", "emergency-area output -> Z, not legacy FA")
    res = IoRow(sheet="s", row=3, addr="I3.1", desc_l1="EMERGENCY RESET", desc_l1b="")
    eq(st.suggested_type(res), "R*", "reset without an area -> R*, not legacy RES")
    res2 = IoRow(sheet="s", row=4, addr="I3.2", desc_l1="EMERGENCY RESET", desc_l1b="AREA 2")
    eq(st.suggested_type(res2), "R2", "reset with an area -> R<area>")


def test_strike_handling():
    def build(path):
        wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
        for col, name in (("G", "Bit"), ("K", "D1"), ("O", "FU"), ("P", "Loc"), ("Q", "Dev"), ("R", "Type")):
            _set(ws, col, 1, name)
        _set(ws, "G", 2, "I1.0"); _set(ws, "K", 2, "CIRCUIT BREAKER TRIPPED")
        _set(ws, "O", 2, "S1"); _set(ws, "P", 2, "M1"); _set(ws, "Q", 2, "-F1", text=True); _set(ws, "R", 2, "A")
        _set(ws, "G", 3, "I1.1"); _set(ws, "K", 3, "CIRCUIT BREAKER TRIPPED")
        _set(ws, "O", 3, "S1"); _set(ws, "P", 3, "M1"); _set(ws, "Q", 3, "-F2", text=True); _set(ws, "R", 3, "A")
        ws.cell(3, CI("G")).font = Font(strike=True)        # row 3 is struck
        wb.save(path)

    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "io.xlsx"); out = os.path.join(d, "out"); build(io)
        params = _params(io); params["output_dir"] = out

        params["strike_handling"] = "exclude"
        r = pop.populate(params, out_dir=out, emit=lambda *_: None)
        s = _by_addr(r.results, "I1.1")
        ok(s.skipped and s.skip_reason == "struck", "exclude -> struck row skipped")

        shutil.rmtree(out, ignore_errors=True)              # fresh copy for the second run
        params["strike_handling"] = "ignore"
        r2 = pop.populate(params, out_dir=out, emit=lambda *_: None)
        s2 = _by_addr(r2.results, "I1.1")
        ok(not s2.skipped, "ignore -> struck row kept")
        eq(s2.script_type, "A", "kept struck row is populated normally")


def test_multi_sheet_global_and_skip_log():
    def build(path):
        wb = Workbook()
        cover = wb.active; cover.title = "COVER"; cover["A1"] = "cover page"
        for sname in ("NET SAFETY 50", "NET SAFETY 51"):
            ws = wb.create_sheet(sname)
            for col, name in (("G", "Bit"), ("K", "D1"), ("O", "FU"), ("P", "Loc"), ("Q", "Dev"), ("R", "Type")):
                _set(ws, col, 1, name)
            _set(ws, "G", 2, "I1.0"); _set(ws, "K", 2, "CIRCUIT BREAKER TRIPPED")
            _set(ws, "O", 2, "S1"); _set(ws, "P", 2, "M" + sname[-2:]); _set(ws, "Q", 2, "-F1", text=True)
            _set(ws, "R", 2, "A")
        wb.save(path)

    with tempfile.TemporaryDirectory() as d:
        io = os.path.join(d, "io.xlsx"); out = os.path.join(d, "out"); build(io)
        params = config.load_params()
        params["io_list"] = {"path": io, "sheet": r"NET SAFETY \d{1,3}/gmi", "header_row": 1}
        params["output_dir"] = out
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)
        eq(res.processed, 2, "rows from BOTH matching sheets")
        eq(sorted(res.sheets), ["NET SAFETY 50", "NET SAFETY 51"])
        ok("COVER" in res.skipped_sheets, "non-matching COVER logged as skipped")
        eq(sorted(r.index for r in res.results), ["0001", "0002"], "index is GLOBAL across sheets")
        # two distinct node cabinets, global numbering 000/001
        eq(sorted(r.diag_cabinet for r in res.results), ["000", "001"])


def test_idempotent_rerun():
    def check(params, out, d):
        r1 = pop.populate(params, out_dir=out, emit=lambda *_: None)
        r2 = pop.populate(params, out_dir=out, emit=lambda *_: None)
        eq(r1.indexed, r2.indexed, "re-run does not double-index")
        eq(_by_addr(r2.results, "I2.0").index, "0001", "index stable across re-run")
    _run(check)


def test_fills_source_in_place_and_backs_up():
    def check(params, out, d):
        io = params["io_list"]["path"]
        baks = lambda: [f for f in os.listdir(d) if ".bak_" in f]
        r1 = pop.populate(params, out_dir=out, emit=lambda *_: None)
        eq(r1.output_path, io, "the SOURCE is filled in place (it becomes the populated doc)")
        ok(r1.backup and os.path.exists(r1.backup), "a changing fill keeps a timestamped backup")
        eq(len(baks()), 1, "exactly one backup after the changing run")
        r2 = pop.populate(params, out_dir=out, emit=lambda *_: None)
        eq(r2.backup, "", "a no-op re-fill removes its own backup")
        eq(len(baks()), 1, "the no-op run leaves no new backup")
    _run(check)


def _diag_row(sheet, row, fu, loc, ex_cab="", ex_bit=""):
    """A minimal diagnosed node RowResult (type A, no family) for the pure diag_alloc tests."""
    io = IoRow(sheet=sheet, row=row, functional_unit=fu, location=loc,
               ex_diag_cabinet=ex_cab, ex_diag_bit=ex_bit)
    return RowResult(iorow=io, type_def={"in_diagnosis": "1", "type_id": "A"}, family=None, index="0001")


def test_allocate_reuses_and_extends_ids():
    fn = "=S1" "+A1"
    r = _diag_row("s", 2, "=S1", "+A1")
    diag_alloc.allocate([r], {}, existing={fn: 7})
    eq(r.diag_cabinet, "007", "an existing cabinet's ID_Local is REUSED (keyed by full_name)")

    r2 = _diag_row("s", 3, "=S2", "+B2")               # a new cabinet, with an occupied id gap
    diag_alloc.allocate([r2], {}, existing={"X": 0, "Y": 2})
    eq(r2.diag_cabinet, "003", "a new cabinet takes max+1 (above every existing id)")

    r3 = _diag_row("s", 4, "=S3", "+C3")
    diag_alloc.allocate([r3], {})
    eq(r3.diag_cabinet, "000", "no existing sheet -> 0-based numbering (regression)")


def test_allocate_bit_continues_past_used():
    fn = "=S1" "+A1"
    filled = _diag_row("s", 2, "=S1", "+A1", ex_cab="000", ex_bit="00")
    filled2 = _diag_row("s", 3, "=S1", "+A1", ex_cab="000", ex_bit="01")
    blank = _diag_row("s", 4, "=S1", "+A1")
    diag_alloc.allocate([filled, filled2, blank], {}, existing={fn: 0})
    eq(filled.diag_bit, "00", "an already-assigned bit is kept")
    eq(blank.diag_bit, "02", "a new blank row CONTINUES past the bits already used (0,1) -> 2")

    a, b = _diag_row("s", 2, "=S1", "+A1"), _diag_row("s", 3, "=S1", "+A1")
    diag_alloc.allocate([a, b], {})
    eq([a.diag_bit, b.diag_bit], ["00", "01"], "fresh fill -> 00,01 (regression)")


def _seed_diagblocks(path, data_rows, headers=None):
    """Add a DiagnosisBlocks sheet (`headers` or DIAG_HEADERS + `data_rows`) to an existing I/O List;
    string cells are written as TEXT (so a FullName like '=S1+A1' isn't stored as a formula). Short
    `data_rows` just omit trailing cells. The synthetic I/O List has no dynamic arrays, so an openpyxl
    save here is safe."""
    wb = load_workbook(path)
    if DIAGBLOCKS_SHEET in wb.sheetnames:
        del wb[DIAGBLOCKS_SHEET]
    ws = wb.create_sheet(DIAGBLOCKS_SHEET)
    ws.append(list(headers) if headers is not None else list(DIAG_HEADERS))
    for ri, row in enumerate(data_rows, 2):
        for ci, val in enumerate(row, 1):
            c = ws.cell(ri, ci, val)
            if isinstance(val, str):
                c.data_type = "s"
    wb.save(path)
    wb.close()


def _set_iolist_cell(path, sheet, row, canonical, value):
    """Set an I/O-List cell by CANONICAL column name (e.g. 'diag_cabinet' -> col AE), as TEXT."""
    from pipeline3.domain.iolist_diag.columns import ColumnResolver
    col = ColumnResolver("IoList").letter(canonical)
    wb = load_workbook(path)
    c = wb[sheet].cell(row=row, column=CI(col), value=value)
    c.data_type = "s"
    wb.save(path)
    wb.close()


def test_diagblocks_id_reused_not_renumbered():
    def check(params, out, d):
        io = params["io_list"]["path"]
        # the alarm node's full_name = FU+Location = '=S1+MS1.CC1'; pre-seed it at a NON-zero ID_Local
        _seed_diagblocks(io, [[5, 5, "=S1", "+MS1.CC1", "=S1+MS1.CC1", 1, "", 1]])
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        eq(_by_addr(res.results, "I1.0").diag_cabinet, "005",
           "the node cabinet REUSES its existing ID_Local (5), it is not renumbered to 000")
    _run(check)


def test_diagblocks_appends_new_cabinet_preserving_existing():
    def check(params, out, d):
        io = params["io_list"]["path"]
        _seed_diagblocks(io, [[0, 777, "ZZ", "ZZ", "ZZZZ", 1, "manual", 9]])   # unrelated cabinet, hand-edited
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        grid = [[c.value for c in r] for r in wb[DIAGBLOCKS_SHEET].iter_rows()]
        wb.close()
        zz = next(r for r in grid if r[4] == "ZZZZ")
        eq([zz[0], zz[1], zz[6]], [0, 777, "manual"], "the orphan row's identity (ID_Local/ID_SWP/Notes) is preserved")
        eq(zz[7], 0, "its Count is refreshed to 0 (no live signals)")
        eq(zz[8], "0-62", "its unused_bits = the full range")
        ok(zz[9] in (None, ""), "its non_unique_bits is empty")
        ok("=S1+MS1.CC1" in [r[4] for r in grid], "the new node cabinet was APPENDED")
        eq(_by_addr(res.results, "I1.0").diag_cabinet, "001", "the appended cabinet gets a new id above 0")
    _run(check)


def test_diagblocks_count_rewritten_each_run():
    def check(params, out, d):
        io = params["io_list"]["path"]
        # seed the real alarm-node cabinet with a STALE Count (99); populate must rewrite it to the truth
        _seed_diagblocks(io, [[0, 0, "=S1", "+MS1.CC1", "=S1+MS1.CC1", 1, "", 99]])
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        grid = [[c.value for c in r] for r in wb[DIAGBLOCKS_SHEET].iter_rows()]
        wb.close()
        node = next(r for r in grid if r[4] == "=S1+MS1.CC1")
        eq(node[7], 1, "the derived Count is REWRITTEN to the current signal count (1), not the stale 99")
        eq(node[0], 0, "the cabinet's identity columns (ID_Local) are preserved")
    _run(check)


def test_diagblocks_preserved_and_no_duplicate_on_rerun():
    def check(params, out, d):
        io = params["io_list"]["path"]
        pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(io)
        ws = wb[DIAGBLOCKS_SHEET]
        n1 = ws.max_row
        ws.cell(2, 2, 4242)                            # a hand-edit to an existing row's ID_SWP
        wb.save(io)
        wb.close()
        pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(io)
        ws = wb[DIAGBLOCKS_SHEET]
        eq(ws.max_row, n1, "a re-run appends NO duplicate rows")
        eq(ws.cell(2, 2).value, 4242, "the hand-edited ID_SWP survives the re-run (the row is untouched)")
        wb.close()
    _run(check)


def test_diagblocks_append_after_header_only_sheet():
    def check(params, out, d):
        io = params["io_list"]["path"]
        _seed_diagblocks(io, [])                       # header only, no data rows
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        ws = wb[DIAGBLOCKS_SHEET]
        eq(ws.cell(1, 1).value, "ID_Local", "the header row is preserved")
        ok(ws.max_row >= 2, "blocks appended after the header")
        wb.close()
        eq(_by_addr(res.results, "I1.0").diag_cabinet, "000", "an empty sheet -> ids start at 000")
    _run(check)


def test_fmt_ranges():
    from pipeline3.domain.iolist_diag import blocks as b
    eq(b._fmt_ranges([]), "")
    eq(b._fmt_ranges([5]), "5")
    eq(b._fmt_ranges([0, 1, 2, 3, 4, 5, 6, 7, 9, 13] + list(range(18, 63))), "0-7, 9, 13, 18-62")
    eq(b._fmt_ranges([3, 3, 1]), "1, 3", "deduped + sorted")


def test_analyze_and_norm():
    from pipeline3.domain.iolist_diag import blocks as b
    eq(b._norm_cab("0"), "000")
    eq(b._norm_cab("7"), "007")
    eq(b._norm_cab("xx"), "xx")
    eq(b._bit_int("05"), 5)
    ok(b._bit_int("") is None and b._bit_int("<input required>") is None)
    d = b.analyze([("000", 0), ("000", 1), ("000", 1), ("001", None)], 0, 62)
    eq(d["000"], (3, "2-62", "1"), "3 rows; bits {0,1} used (1 twice) -> unused 2-62, non_unique 1")
    eq(d["001"], (1, "0-62", ""), "a row with no bit counts once; full unused; no collision")
    eq(b.analyze([("002", 70), ("002", 70)], 0, 62)["002"], (2, "0-62", "70"),
       "an out-of-range duplicated bit shows in non_unique, doesn't shrink unused")


def test_real_count_from_effective_diag_cabinet():
    def check(params, out, d):
        io = params["io_list"]["path"]
        _set_iolist_cell(io, "NET SAFETY 50", 3, "diag_cabinet", "005")    # alarm A manually homed to cab 005
        _seed_diagblocks(io, [[5, 5, "M", "M", "MANUAL5", 1, "", 0]])
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        grid = [[c.value for c in r] for r in wb[DIAGBLOCKS_SHEET].iter_rows()]
        wb.close()
        m5 = next(r for r in grid if r[4] == "MANUAL5")
        eq(m5[7], 1, "Count follows the REAL Diag_Cabinet column (the manual 005), not the computed cabinet")
    _run(check)


def test_unused_nonunique_columns_fresh():
    def check(params, out, d):
        res = pop.populate(params, out_dir=out, emit=lambda *_: None)   # full fill, no prior sheet -> fresh
        wb = load_workbook(res.output_path)
        ws = wb[DIAGBLOCKS_SHEET]
        eq([c.value for c in ws[1]], list(DIAG_HEADERS), "fresh sheet carries the 10-column header")
        grid = [[c.value for c in r] for r in ws.iter_rows()]
        wb.close()
        node = next(r for r in grid if r[4] == "=S1+MS1.CC1")           # alarm A at bit 0
        eq(node[8], "1-62", "unused_bits = full range minus the used bit 0")
        ok(node[9] in (None, ""), "non_unique_bits empty (the allocator never collides)")
    _run(check)


def test_add_columns_to_existing_8col_sheet():
    def check(params, out, d):
        io = params["io_list"]["path"]
        _seed_diagblocks(io, [[0, 0, "=S1", "+MS1.CC1", "=S1+MS1.CC1", 1, "note", 99]],
                         headers=DIAG_HEADERS[:8])                      # OLD 8-column sheet (no I/J)
        res = pop.populate(params, write={"diag_cabinet"}, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(res.output_path)
        ws = wb[DIAGBLOCKS_SHEET]
        eq(ws.cell(1, 9).value, "unused_bits", "I1 header added")
        eq(ws.cell(1, 10).value, "non_unique_bits", "J1 header added")
        row = [c.value for c in ws[2]]
        wb.close()
        eq(row[:7], [0, 0, "=S1", "+MS1.CC1", "=S1+MS1.CC1", 1, "note"], "identity columns A..G preserved")
        eq(row[7], 1, "Count refreshed at H (was 99)")
        ok(row[8], "unused_bits populated at I")
    _run(check)


def test_rerun_no_duplicate_columns():
    def check(params, out, d):
        io = params["io_list"]["path"]
        pop.populate(params, out_dir=out, emit=lambda *_: None)
        pop.populate(params, out_dir=out, emit=lambda *_: None)
        wb = load_workbook(io)
        ws = wb[DIAGBLOCKS_SHEET]
        hdr = [c.value for c in ws[1]]
        wb.close()
        eq(ws.max_column, 10, "still exactly 10 columns after a re-run")
        eq(hdr.count("unused_bits"), 1, "no duplicate unused_bits column")
        eq(hdr.count("non_unique_bits"), 1, "no duplicate non_unique_bits column")
    _run(check)


if __name__ == "__main__":
    raise SystemExit(run("iolist_diag", [
        ("fmt_ranges", test_fmt_ranges),
        ("analyze_and_norm", test_analyze_and_norm),
        ("real_count_from_effective_diag_cabinet", test_real_count_from_effective_diag_cabinet),
        ("unused_nonunique_columns_fresh", test_unused_nonunique_columns_fresh),
        ("add_columns_to_existing_8col_sheet", test_add_columns_to_existing_8col_sheet),
        ("rerun_no_duplicate_columns", test_rerun_no_duplicate_columns),
        ("allocate_reuses_and_extends_ids", test_allocate_reuses_and_extends_ids),
        ("allocate_bit_continues_past_used", test_allocate_bit_continues_past_used),
        ("diagblocks_id_reused_not_renumbered", test_diagblocks_id_reused_not_renumbered),
        ("diagblocks_appends_new_cabinet_preserving_existing", test_diagblocks_appends_new_cabinet_preserving_existing),
        ("diagblocks_count_rewritten_each_run", test_diagblocks_count_rewritten_each_run),
        ("diagblocks_preserved_and_no_duplicate_on_rerun", test_diagblocks_preserved_and_no_duplicate_on_rerun),
        ("diagblocks_append_after_header_only_sheet", test_diagblocks_append_after_header_only_sheet),
        ("populates_and_skips_nodemeta", test_populates_and_skips_nodemeta),
        ("script_type_and_index", test_script_type_and_index),
        ("diag_allocation", test_diag_allocation),
        ("writes_columns_and_diagblocks", test_writes_columns_and_diagblocks),
        ("fire_alarm_rung", test_fire_alarm_rung),
        ("ladder_yields_canonical_directly", test_ladder_yields_canonical_directly),
        ("strike_handling", test_strike_handling),
        ("multi_sheet_global_and_skip_log", test_multi_sheet_global_and_skip_log),
        ("idempotent_rerun", test_idempotent_rerun),
        ("fills_source_in_place_and_backs_up", test_fills_source_in_place_and_backs_up),
    ]))
