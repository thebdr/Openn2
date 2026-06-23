"""M6 gate (data-independent): 130 CEM->IOL forward + 140 IOL->CEM reverse cross-checks.

Synthetic staged rows + a synthetic Cause&Effect matrix exercise every match state and assert the
second-workbook link (location2/doc2) is the matched cell in the OTHER workbook (sheet!cell only).
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as CI

from pipeline3.context import PipelineContext
from pipeline3.domain.validation import run_xcheck_cem_iol, run_xcheck_iol_cem

IO_PATH = os.path.join("somewhere", "IOList.xlsx")     # only the basename is used


def _set(ws, col, row, val, text=False):
    c = ws.cell(row=row, column=CI(col), value=val)
    if text:
        c.data_type = "s"
    return c


def _mrow(m, row, addr, fu, lo, de):
    _set(m, "F", row, addr)
    _set(m, "J", row, fu, text=True)
    _set(m, "K", row, lo, text=True)
    _set(m, "L", row, de, text=True)


def _row(fu, lo, de, bit, srow, ttype=None, **extra):
    r = {"functional_unit": fu, "location": lo, "device": de, "bit": bit,
         "source_cell": f"NET SAFETY 50!O{srow}", "_source_sheet": "NET SAFETY 50", "_source_row": srow,
         "_type": ttype, "desc_l1": "", "desc_l1b": "", "desc_l2": "", "desc_l2b": "", "mnemonic": ""}
    r.update(extra)
    return r


def _ctx(rows, ce_path, **over):
    params = {"io_list": {"path": IO_PATH, "sheet": r"NET SAFETY \d+", "header_row": 1},
              "ce": {"path": ce_path, "matrix_sheet": "CAUSE&EFFECT MATRIX", "matrix_header_row": 2,
                     "matrix_data_row": 5, "area_header_row": 3, "area_data_row": 4},
              "areas": r"AREA\s?\d{1,2}"}
    params.update(over)
    return PipelineContext(params=params, out_root="X", rows=rows, lang="en", profile="main")


# --- 130 forward ------------------------------------------------------------ #
def _make_ce_fwd(path):
    wb = Workbook(); m = wb.active; m.title = "CAUSE&EFFECT MATRIX"
    _mrow(m, 5, "I20.0", "=S1", "+M1", "-D1")    # full match
    _mrow(m, 6, "I20.1", "=S1", "+M1", "-D1")    # device matches, address differs
    _mrow(m, 7, "I20.2", "=S9", "+X", "-Z")      # device not in IOL
    _mrow(m, 8, "I20.3", "", "", "")             # no device
    wb.save(path)


def test_130_forward():
    with tempfile.TemporaryDirectory() as d:
        ce = os.path.join(d, "ce.xlsx"); _make_ce_fwd(ce)
        rows = [_row("=S1", "+M1", "-D1", "I20.0", 5)]
        out = run_xcheck_cem_iol(_ctx(rows, ce))
    by = {x.type: x for x in out}
    eq(by["cem_match"].level, "PASS")
    eq(by["cem_dev_missing"].level, "FAIL")
    eq(by["cem_ref_empty"].level, "FAIL")
    mm = by["cem_fld_addr_mismatch"]
    eq(mm.level, "FAIL")
    eq(mm.location2, "NET SAFETY 50!O5", "links the matched I/O List cell")
    eq(mm.doc2, "IOList.xlsx", "doc2 is the I/O List workbook")
    ok(mm.cmp is not None and not mm.cmp.addr_eq and mm.cmp.fld_eq, "structured cmp: addr differs, fld matches")
    ok(".xlsx" not in mm.location, "C&E link is sheet!cell only")
    # full match links BOTH workbooks: the C&E ref cell + the matched I/O List cell
    eq((by["cem_match"].location2, by["cem_match"].doc2), ("NET SAFETY 50!O5", "IOList.xlsx"))


def test_130_absent_skips():
    out = run_xcheck_cem_iol(_ctx([], os.path.join("no", "ce.xlsx")))
    eq(len(out), 1)
    eq((out[0].type, out[0].level), ("ce_absent", "SKIP"))


# --- 140 reverse ------------------------------------------------------------ #
def _make_ce_rev(path):
    wb = Workbook(); m = wb.active; m.title = "CAUSE&EFFECT MATRIX"
    _mrow(m, 5, "I10.0", "=A", "+B", "-C")       # the only device/address present in the C&E
    wb.save(path)


def test_140_reverse():
    rows = [
        _row("=A", "+B", "-C", "I10.0", 5, {"ce_mandatory": "yes"}),                  # full -> PASS
        _row("=M", "+M", "-M", "I99.0", 6, {"ce_mandatory": "yes"}),                  # missing_mandatory
        _row("=W", "+W", "-W", "I98.0", 7, {"ce_mandatory": "warn"}),                 # missing_plain (WARN)
        _row("=S", "+S", "-S", "I97.0", 8, None, desc_l1="EMERGENCY STOP"),           # missing_safety
        _row("=X", "+X", "-X", "I96.0", 9, None, desc_l1="spare channel CH2"),        # skipped (always-excl)
        _row("=A", "+B", "-C", "I10.9", 10, {"ce_mandatory": "yes"}),                 # fld_only
        _row("=O", "+O", "-O", "I10.0", 11, {"ce_mandatory": "yes"}),                 # addr_only
    ]
    with tempfile.TemporaryDirectory() as d:
        ce = os.path.join(d, "ce.xlsx"); _make_ce_rev(ce)
        out = run_xcheck_iol_cem(_ctx(rows, ce, ce_mandatory_words=["emergency", "safety"],
                                      ce_always_excluded_words=["CH2"], ce_fuzzy_chars=0,
                                      ce_full_check=True))
    lvl = {x.type: x.level for x in out}
    eq(lvl["iol_cem_match"], "PASS")
    eq(lvl["iol_cem_missing_mandatory"], "FAIL")
    eq(lvl["iol_cem_missing_plain"], "WARN")
    eq(lvl["iol_cem_missing_safety"], "FAIL")
    eq(lvl["iol_cem_skipped"], "SKIP")
    eq(lvl["iol_cem_fld_only"], "FAIL")
    eq(lvl["iol_cem_addr_only"], "FAIL")
    fld = next(x for x in out if x.type == "iol_cem_fld_only")
    eq(fld.location2, "CAUSE&EFFECT MATRIX!F5", "links the C&E cell")
    eq(fld.doc2, "ce.xlsx", "doc2 is the C&E workbook")
    eq(fld.location, "NET SAFETY 50!O10", "primary link is the I/O List row, sheet!cell only")
    miss = next(x for x in out if x.type == "iol_cem_missing_mandatory")
    eq((miss.location2, miss.doc2), ("C&E Matrix", "ce.xlsx"), "a miss links the C&E FILE (label, no cell)")


def test_140_full_check_false_gates_untyped():
    # ce_full_check=false: typed -> CHECK; untyped only if safety-word; untyped non-safety -> NO_CHECK
    rows = [
        _row("=A", "+B", "-C", "I10.0", 5, {"ce_mandatory": "yes"}),                   # typed yes, in C&E -> PASS
        _row("=N", "+N", "-N", "I52.0", 8, {"ce_mandatory": "no"}),                    # typed no -> NO_CHECK (skip)
        _row("=P", "+P", "-P", "I50.0", 6, None, desc_l1="spare channel"),             # untyped non-safety -> SKIP
        _row("=S", "+S", "-S", "I51.0", 7, None, desc_l1="EMERGENCY STOP"),            # untyped safety -> checked
    ]
    with tempfile.TemporaryDirectory() as d:
        ce = os.path.join(d, "ce.xlsx"); _make_ce_rev(ce)
        out = run_xcheck_iol_cem(_ctx(rows, ce, ce_mandatory_words=["emergency", "safety"],
                                      ce_always_excluded_words=["ch2"], ce_fuzzy_chars=0, ce_full_check=False))
    by = {x.type: x.level for x in out}
    eq(by["iol_cem_match"], "PASS", "typed yes/warn row is cross-checked")
    eq(by.get("iol_cem_not_required"), "SKIP", "typed ce_mandatory=no -> not cross-checked")
    eq(by.get("iol_cem_unclassified"), "SKIP", "untyped non-safety -> not cross-checked")
    eq(by["iol_cem_missing_safety"], "FAIL", "untyped safety is still cross-checked")
    ok("iol_cem_missing_plain" not in by, "neither the typed-no nor the undescribed row is WARN'd")


def test_140_absent_skips():
    out = run_xcheck_iol_cem(_ctx([], os.path.join("no", "ce.xlsx")))
    eq(len(out), 1)
    eq((out[0].type, out[0].level), ("ce_absent", "SKIP"))


if __name__ == "__main__":
    raise SystemExit(run("validation_xcheck", [
        ("130_forward", test_130_forward),
        ("130_absent_skips", test_130_absent_skips),
        ("140_reverse", test_140_reverse),
        ("140_full_check_false_gates_untyped", test_140_full_check_false_gates_untyped),
        ("140_absent_skips", test_140_absent_skips),
    ]))
