"""M6 gate (data-independent): 110 Validate I/O List standalone.

Builds a synthetic I/O List with one of each defect and asserts the exact log-type codes/levels.
Required-column headers are set correctly; column C's header is deliberately wrong; AB (a
preliminary_check_exclude column) gets a wrong header that must NOT be flagged.
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string as CI

from pipeline3.core import config
from pipeline3.context import PipelineContext
from pipeline3.domain.validation import run_iolist


def _set(ws, col, row, val, text=False):
    c = ws.cell(row=row, column=CI(col), value=val)
    if text:
        c.data_type = "s"
    return c


def _fld(ws, row, fu, lo, de):
    _set(ws, "O", row, fu, text=True)
    _set(ws, "P", row, lo, text=True)
    _set(ws, "Q", row, de, text=True)


def _make(path):
    perm0 = config.load_permanent_parts()[0]
    wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
    # required headers correct; column C deliberately wrong; AB (excluded) wrong but must be ignored
    _set(ws, "C", 1, "WRONGPART")              # col 3 -> col_wrong
    for col, hdr in (("F", "ID"), ("O", "Functional unit"), ("P", "Location"), ("Q", "Device")):
        _set(ws, col, 1, hdr)
    _set(ws, "AB", 1, "Definitely Not Script Type")   # excluded col -> must NOT be flagged
    _set(ws, "AH", 1, "Extra Thing")           # outside map -> col_unexpected
    _set(ws, "AI", 1, "Device")                # duplicates a standard name -> col_duplicated

    _set(ws, "V", 2, "#REF!", text=True); _fld(ws, 2, "=S0", "+M0", "-D0")   # ip_error
    ws["G2"].font = Font(strike=True)          # struck (must be IGNORED under strike_handling=ignore)
    _set(ws, "V", 3, "192.168.50.6"); _fld(ws, 3, "=S1", "+M1", "-D1")
    _set(ws, "V", 4, "192.168.50.6"); _fld(ws, 4, "=S2", "+M2", "-D2")        # ip_duplicated
    _set(ws, "V", 5, "192.168.50.7"); _fld(ws, 5, "=S3", "+M3", "-D3")
    _set(ws, "V", 6, "192.168.50.8"); _fld(ws, 6, "=S3", "+M3", "-D3")        # dup_fld
    _set(ws, "F", 7, "ID123"); _set(ws, "G", 7, "I1.0")                       # fg_exclusion
    _set(ws, "G", 8, "I0.0.0")                                                # addr_format (3 parts)
    _set(ws, "R", 9, "A"); _set(ws, "U", 9, "TS_X")                           # pp_empty (K blank)
    _set(ws, "R", 10, "A"); _set(ws, "K", 10, "NOT A PERMANENT PART"); _set(ws, "U", 10, "TS_X")  # pp_unknown
    _set(ws, "R", 11, "W"); _set(ws, "K", 11, perm0)                          # ts_missing (U blank, pp valid)
    _set(ws, "R", 12, "PA"); _set(ws, "V", 12, "192.168.50.9"); _set(ws, "W", 12, "badname"); _set(ws, "U", 12, "TS_X")  # pname_wrong
    # 61 extra valid nodes -> total node count > 64 (nodes_over_64 WARN, not too_many_nodes)
    for i in range(20, 81):
        _set(ws, "V", i, f"192.168.51.{i}"); _fld(ws, i, f"=N{i}", f"+N{i}", f"-N{i}")
    wb.save(path)


def _types(io):
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "io.xlsx")
        _make(path)
        ctx = PipelineContext(params={"io_list": {"path": path, "sheet": "NET SAFETY 50", "header_row": 1},
                                      "strike_handling": "ignore"}, out_root=d, lang="en", profile="main")
        return run_iolist(ctx)


def _present(entries):
    return {e.type for e in entries}


def test_each_defect_flagged():
    e = _types(None)
    types = _present(e)
    for t in ("col_wrong", "col_unexpected", "col_duplicated", "ip_error", "ip_duplicated",
              "dup_fld", "fg_exclusion", "addr_format", "pp_empty", "pp_unknown", "ts_missing",
              "pname_wrong", "nodes_over_64"):
        ok(t in types, f"expected a {t} finding")


def test_strike_ignored_and_no_too_many_nodes():
    types = _present(_types(None))
    ok("struck_row" not in types, "strike_handling=ignore -> no struck_row")
    ok("too_many_nodes" not in types, "67 nodes is over 64 but under 126")


def test_excluded_column_not_flagged_and_c_is():
    e = _types(None)
    cols = [x for x in e if x.type == "col_wrong"]
    ok(any(x.detail and "3" in x.detail for x in cols), "column C (3) wrong header flagged")
    # AB is column 28 and preliminary_check_exclude -> never a col_wrong/unexpected/duplicated
    ok(not any(("28" in (x.detail or "")) for x in e if x.type in ("col_wrong", "col_unexpected", "col_duplicated")),
       "the excluded AB column must not be flagged")


def test_clean_rows_emit_all_checks_passed():
    e = _types(None)
    passes = [x for x in e if x.type == "row_ok" and x.level == "PASS"]
    ok(len(passes) >= 1, "a clean row emits a PASS 'all checks passed'")
    ok(all("passed" in x.detail.lower() for x in passes), "PASS detail is localized, not the bare key")
    # a row that raised a finding must NOT also be marked all-checks-passed (same location)
    fail_locs = {x.location for x in e if x.level == "FAIL"}
    ok(not (fail_locs & {x.location for x in passes}), "a flagged row gets no PASS")


def test_locations_have_no_workbook_name():
    for x in _types(None):
        if x.location:
            ok(".xlsx" not in x.location, f"link must be sheet!cell: {x.location}")


if __name__ == "__main__":
    raise SystemExit(run("validation_iolist", [
        ("each_defect_flagged", test_each_defect_flagged),
        ("strike_ignored_and_no_too_many_nodes", test_strike_ignored_and_no_too_many_nodes),
        ("excluded_column_not_flagged_and_c_is", test_excluded_column_not_flagged_and_c_is),
        ("clean_rows_emit_all_checks_passed", test_clean_rows_emit_all_checks_passed),
        ("locations_have_no_workbook_name", test_locations_have_no_workbook_name),
    ]))
