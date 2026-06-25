"""M9 gate (data-independent): the phase-800 software-block skeleton.

Covers the Database, the Table, the @builds registry, the engine ($/#/%/@ CreationInfo CSV with an
ABSOLUTE $ template= ref + InstanceDBs + the horizontal ITERATOR), and the 810 shells (the key
inventory: scans the template *.xml for !!key$$ - block_templates.json is deprecated). Shell-
dependent cases point BLOCK_TEMPLATES_DIR at a synthetic template so they stay data-independent.
NOT the per-template builder logic (the user writes that in domain/blocks/builders.py).
"""
import csv
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import load_workbook

import re
import xml.etree.ElementTree as ET

from pipeline3.core import config
from pipeline3.domain.blocks.database import Database
from pipeline3.domain.blocks.table import Table
from pipeline3.domain.blocks import registry, engine, shells, xml_emit

_TEMPLATE_XML = (
    "<Document>!!nameOfDB$$ !!02_COM.{db_element}$$ !!NetworkComment$$ "
    "!!instanceOf-FDBACK$$ !!ITERATOR_STRINGS$$</Document>"
)


def _rows():
    return [
        {"script_type": "KQ", "datablocks": "03_FDBACK", "matrix_areas": "AREA 1",
         "name_in_db": "Contactor Feedback Error", "tagtable": "SAFETY_Contactors"},
        {"script_type": "E1/2", "datablocks": "01_Pushbutton", "matrix_areas": "AREA 1|AREA 2",
         "name_in_db": "Emergency PB"},
        {"script_type": "A", "datablocks": "", "matrix_areas": "AREA 2", "name_in_db": ""},
    ]


def _read(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def _with_synthetic_templates(fn):
    """Run fn(out_root) with BLOCK_TEMPLATES_DIR pointed at a temp dir holding one synthetic
    template (TEMPLATE--v1.0--T1.xml), and out_root a temp dir. Restores config + registry."""
    saved = config.BLOCK_TEMPLATES_DIR
    registry.clear()
    with tempfile.TemporaryDirectory() as tdir, tempfile.TemporaryDirectory() as out:
        with open(os.path.join(tdir, "TEMPLATE--v1.0--T1.xml"), "w", encoding="utf-8") as f:
            f.write(_TEMPLATE_XML)
        config.BLOCK_TEMPLATES_DIR = tdir
        try:
            fn(out)
        finally:
            config.BLOCK_TEMPLATES_DIR = saved
            registry.clear()


def test_database_queries():
    db = Database(_rows())
    eq(len(db.by_type("kq", "e1/2")), 2)
    eq(len(db.by_db("03_FDBACK")), 1)
    eq(len(db.by_area("AREA 1")), 2)
    eq(db.areas(), ["AREA 1", "AREA 2"])


def test_table_add():
    t = Table("X")
    t.add(template_type="01", NetworkComment="hi", **{"02_COM.{db_element}": "v"})
    eq(t.columns, ["TemplateType", "NetworkComment", "02_COM.{db_element}"])
    eq(t.rows[0]["NetworkComment"], "hi")


def test_engine_absolute_ref_and_instances():
    registry.clear()

    @registry.builds("TestBlock")
    def _b(db):
        t = Table("TestBlock")
        for area in db.areas():
            t.add(template_type="02", NetworkComment=area,
                  **{"instanceOf-FDBACK": f"INST_{area}"},
                  ITERATOR_STRINGS=[r["name_in_db"] for r in db.by_area(area) if r.get("name_in_db")])
        return t

    try:
        with tempfile.TemporaryDirectory() as d:
            res = engine.generate_blocks(_rows(), d, emit=lambda *_: None)
            grid = _read(os.path.join(res["dir"], "TestBlock.csv"))
            eq(grid[0][0], "$")
            ref = grid[0][1].split("=", 1)[1]
            ok(os.path.isabs(ref) and ref.endswith("TestBlock.xml"), f"absolute $ ref: {ref}")
            eq(grid[2][0], "%")
            eq(grid[2][-1], "!!ITERATOR_STRINGS$$", "iterator column last")
            eq(len([r for r in grid if r[0] == "@"]), 2)
            inst = _read(os.path.join(res["dir"], "InstanceDBs.csv"))
            ok(any(r[0] == "@" and r[2] == "FDBACK" for r in inst))
    finally:
        registry.clear()


def test_engine_writes_com_db_with_constants():
    registry.clear()

    @registry.builds("ZC")
    def _b(db):
        t = Table("ZC")
        t.add(template_type="01", **{"02_COM.{db_element}": "AREA 1 PB"})
        t.add(template_type="01", **{"02_COM.{db_element}": "AREA 1 FDB"})
        return t

    try:
        from pipeline3.domain import signals
        with tempfile.TemporaryDirectory() as d:
            engine.generate_blocks(_rows(), d, emit=lambda *_: None)
            com = os.path.join(config.out_path(d, "blocks_import_dir"), "02_COM.xml")
            ok(os.path.exists(com), "the block engine wrote the 02_COM custom DB")
            xml = open(com, encoding="utf-8-sig").read()
            for m in signals.DB_CONSTANTS:
                ok(f'<Member Name="{m}"' in xml, f"standard seed member {m} present")
            ok('<Member Name="AREA 1 PB"' in xml and '<Member Name="AREA 1 FDB"' in xml, "cumulatives present")
            ok(xml.index("No Operation") < xml.index("AREA 1 PB"), "standard seeds precede the cumulatives")
    finally:
        registry.clear()


def test_xml_emit_flgnet_exactly_sized():
    # the FlgNet core: an AND of N inputs -> a coil, sized to N (no tiers/padding), wiring resolves
    flg = "\n".join(xml_emit._flgnet_lines(
        [("01_Pushbutton", "PB a"), ("01_Pushbutton", "PB b")], "02_COM", "AREA 1 PB", ""))
    eq(re.search(r"Card[^>]*>(\d+)<", flg).group(1), "2", "A-box Card = the real input count")
    ok('Name="in1"' in flg and 'Name="in2"' in flg and 'Name="in3"' not in flg, "exactly N AND pins")
    ok('<Component Name="02_COM" />' in flg and '<Component Name="AREA 1 PB" />' in flg,
       "coil writes 02_COM.<output>")
    parts, wrs = flg.split("<Wires>")
    uids = set(re.findall(r'UId="(\d+)"', parts))
    refs = set(re.findall(r'(?:IdentCon|NameCon) UId="(\d+)"', wrs))
    ok(refs and refs <= uids, "every wire UId resolves to a Part/Access")


def test_xml_emit_flgnet_chunks_large_and():
    # TIA caps an instruction at 100 inputs: a large AND splits into leaf ANDs of <= AND_CHUNK + one
    # combiner AND of their outputs (the template's two-level topology). No AND may exceed the cap.
    C = xml_emit.AND_CHUNK
    inputs = [("01_Pushbutton", f"PB{i}") for i in range(C + 30)]     # 2 chunks: [C, 30]
    flg = "\n".join(xml_emit._flgnet_lines(inputs, "02_COM", "AREA 1 PB", ""))
    cards = [int(c) for c in re.findall(r"Card[^>]*>(\d+)<", flg)]
    eq(cards, [C, 30, 2], "two leaf ANDs (<=AND_CHUNK) + a combiner AND of the 2 leaf outputs")
    ok(all(c <= 100 for c in cards), "no AND exceeds TIA's 100-input cap")
    ok(f'Name="in{C}"' in flg, "a full leaf carries AND_CHUNK pins")
    eq(flg.count('Name="out"'), 3, "2 leaf outs + 1 combiner out drive the next stage / coil")
    parts, wrs = flg.split("<Wires>")
    uids = set(re.findall(r'UId="(\d+)"', parts))
    refs = set(re.findall(r'(?:IdentCon|NameCon) UId="(\d+)"', wrs))
    ok(refs <= uids, "every wire UId resolves to a Part/Access")


def test_xml_emit_and_coil_fc_from_table():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "T.xml")
        with open(tpl, "w", encoding="utf-8") as f:                 # a minimal FC header (Name + AttributeList)
            f.write('<?xml version="1.0" encoding="utf-8"?><Document><SW.Blocks.FC ID="0"><AttributeList>'
                    '<Name>TEMPLATE--v1.0--T</Name><ProgrammingLanguage>F_FBD</ProgrammingLanguage>'
                    '</AttributeList><ObjectList /></SW.Blocks.FC></Document>')
        t = Table("T")
        t.add(template_type="01", nameOfDB="01_Pushbutton", NetworkComment="zone one",
              **{"02_COM.{db_element}": "AREA 1 PB"}, ITERATOR_STRINGS=["PB a", "PB b"])
        t.add(template_type="01", nameOfDB="03_FDBACK", NetworkComment="zone two",
              **{"02_COM.{db_element}": "AREA 1 FDB"}, ITERATOR_STRINGS=["FB a"])
        xml = xml_emit.and_coil_fc(t, tpl, "T")
        ET.fromstring(xml)                                          # well-formed
        ok("\r\n" in xml and not xml.startswith("﻿"), "CRLF body, no BOM char (BOM added on write)")
        eq(xml.count("<SW.Blocks.CompileUnit"), 2, "one network per @ row")
        eq(re.search(r"<Name>([^<]*)</Name>", xml).group(1), "T", "block name swapped (prefix dropped)")
        ok("AREA 1 PB" in xml and "AREA 1 FDB" in xml, "both coil outputs present")
        ok("zone one" in xml and "zone two" in xml, "per-network titles")
        # global object IDs are unique (TIA reassigns, but they must not collide)
        gids = re.findall(r'<(?:SW\.Blocks\.CompileUnit|MultilingualText|MultilingualTextItem) ID="(\d+)"', xml)
        eq(len(gids), len(set(gids)), "global object IDs unique")


def test_xml_emitted_block_skips_creation_csv():
    registry.clear()

    @registry.builds("TestBlock")
    def _b(db):
        t = Table("TestBlock")
        t.add(template_type="01", nameOfDB="X")
        return t

    orig = xml_emit.write_fc_xml
    xml_emit.write_fc_xml = (lambda name, table, ref, out_dir:
                             os.path.join(out_dir, f"{name}.xml") if name == "TestBlock" else "")
    try:
        with tempfile.TemporaryDirectory() as d:
            res = engine.generate_blocks(_rows(), d, emit=lambda *_: None)
            ok(not os.path.exists(os.path.join(res["dir"], "TestBlock.csv")),
               "a direct-XML block writes NO CreationInfo CSV")
            ok("TestBlock" not in [os.path.basename(f) for f in res["files"]], "CSV not in the result list")
            ok(any("TestBlock.xml" in x for x in res["xml_files"]), "the FC XML is recorded instead")
    finally:
        xml_emit.write_fc_xml = orig
        registry.clear()


def test_standard_sheet_mirrors_csv_and_reads_back():
    from pipeline3.domain.blocks.engine import _ordered_columns, _wrap, _at_row_cells
    with tempfile.TemporaryDirectory() as out:
        t = Table("00_Only for Commissioning")
        t.add(template_type="01", NetworkComment="n1 1.2.3.4", **{"00_Commissioning.{db_element}": "n1 1.2.3.4"})
        t.add(template_type="01", NetworkComment="n2 1.2.3.5", **{"00_Commissioning.{db_element}": "n2 1.2.3.5"})
        keys = ["TemplateType", "00_Commissioning.{db_element}", "NetworkComment"]
        ordered = _ordered_columns(keys, t)
        p = shells.write_standard_sheet(out, "00_Only for Commissioning", [_wrap(c) for c in ordered],
                                        [_at_row_cells(r, ordered) for r in t.rows], r"C:\t\00.xml")
        ws = load_workbook(p)["00_Only for Commissioning"]
        eq([ws["A1"].value, ws["A2"].value, ws["A3"].value], ["$", "$", "%"], "$ template / $ mode / % markers")
        eq(ws["B2"].value, "fill", "mode preserved (default fill)")
        eq([ws["B3"].value, ws["A4"].value], ["TemplateType", "@"], "% header + @ data rows")
        ws.parent.close()
        back = shells.read_override_table(out, "00_Only for Commissioning")
        eq(len(back), 2, "the mirrored @ rows read back via override")
        eq(back.rows[0]["00_Commissioning.{db_element}"], "n1 1.2.3.4")
        eq(back.rows[1]["NetworkComment"], "n2 1.2.3.5")


def test_coil_columns_round_trip():
    with tempfile.TemporaryDirectory() as out:
        t = Table("03_Zone Cumulative")
        t.add(template_type="01", nameOfDB="01_Pushbutton", NetworkComment="AREA 1 PB - Sorter",
              **{"02_COM.{db_element}": "AREA 1 PB"}, ITERATOR_STRINGS=["PB a", "PB b", "PB c"])
        t.add(template_type="01", nameOfDB="03_FDBACK", NetworkComment="AREA 1 FDB - Sorter",
              **{"02_COM.{db_element}": "AREA 1 FDB"}, ITERATOR_STRINGS=["FB a"])
        p = shells.write_coil_columns(out, "03_Zone Cumulative", t, r"C:\t\03.xml")
        ok(os.path.exists(p), "coil-columns shell written")
        ws = load_workbook(p)["03_Zone Cumulative"]
        eq([ws["A3"].value, ws["A4"].value, ws["A5"].value], ["%coil", "nameOfDB", "comment"], "row markers")
        eq([ws["B3"].value, ws["C3"].value], ["AREA 1 PB", "AREA 1 FDB"], "one column per coil")
        eq([ws["B6"].value, ws["B7"].value, ws["B8"].value], ["PB a", "PB b", "PB c"], "inputs down the column")
        eq(ws["C6"].value, "FB a")
        ws.parent.close()
        # round-trip: read it back to an equivalent table
        back = shells.read_coil_columns(out, "03_Zone Cumulative")
        m = {r["02_COM.{db_element}"]: r for r in back.rows}
        eq(sorted(m), ["AREA 1 FDB", "AREA 1 PB"])
        eq(m["AREA 1 PB"]["nameOfDB"], "01_Pushbutton")
        eq(m["AREA 1 PB"]["NetworkComment"], "AREA 1 PB - Sorter")
        eq(m["AREA 1 PB"]["ITERATOR_STRINGS"], ["PB a", "PB b", "PB c"], "AND inputs recovered, in order")
        eq(m["AREA 1 FDB"]["ITERATOR_STRINGS"], ["FB a"])


def test_override_writes_csv_verbatim_with_iterator():
    def check(out):
        shells.generate_shells(out, emit=lambda *_: None)        # T1 standard sheet (mode fill)
        path = shells._shell_path(out)
        wb = load_workbook(path)
        ws = wb["T1"]
        ws["B2"] = "override"
        ws["A4"] = "@"; ws["B4"] = "01"; ws["C4"] = "DB1"
        for i, m in enumerate(["m1", "m2", "m3"]):                # a spread ITERATOR (3 members)
            ws.cell(row=4, column=4 + i, value=m)                 # D4, E4, F4
        wb.save(path)
        res = engine.generate_blocks([], out, emit=lambda *_: None)
        at = [r for r in _read(os.path.join(res["dir"], "T1.csv")) if r and r[0] == "@"]
        eq(len(at), 1)
        eq(at[0], ["@", "01", "DB1", "m1", "m2", "m3"], "override @ row copied verbatim, iterator preserved")
    _with_synthetic_templates(check)


def test_shells_scan_and_inventory():
    def check(out):
        res = shells.generate_shells(out, emit=lambda *_: None)
        eq(res["created"], ["T1"], "one shell per scanned template")
        info = shells.read_shells(out)["T1"]
        eq(info["mode"], "fill")
        ok(os.path.isabs(info["ref"]) and info["ref"].endswith("T1.xml"), info["ref"])
        # the % header is the FULL key set scanned from the .xml (block_templates.json NOT used)
        eq(info["keys"], ["TemplateType", "nameOfDB", "02_COM.{db_element}",
                          "NetworkComment", "instanceOf-FDBACK", "ITERATOR_STRINGS"])
        # additive: a 2nd run keeps it
        eq(shells.generate_shells(out, emit=lambda *_: None)["created"], [])
    _with_synthetic_templates(check)


def test_output_header_is_full_shell_keyset():
    def check(out):
        shells.generate_shells(out, emit=lambda *_: None)

        @registry.builds("T1")
        def _b(db):
            t = Table("T1")
            t.add(template_type="01", NetworkComment="only this one set")   # builder fills 1 of 5 keys
            return t

        res = engine.generate_blocks(_rows(), out, emit=lambda *_: None)
        grid = _read(os.path.join(res["dir"], "T1.csv"))
        header = grid[2]
        for k in ("TemplateType", "!!nameOfDB$$", "!!02_COM.{db_element}$$",
                  "!!NetworkComment$$", "!!instanceOf-FDBACK$$", "!!ITERATOR_STRINGS$$"):
            ok(k in header, f"% header carries the full shell key set: missing {k}")
    _with_synthetic_templates(check)


def test_override_uses_shell_rows():
    def check(out):
        shells.generate_shells(out, emit=lambda *_: None)
        path = shells._shell_path(out)
        wb = load_workbook(path)
        ws = wb["T1"]
        ws["B2"] = "override"
        ws["A4"] = "@"; ws["B4"] = "01"; ws["C4"] = "MY_DB"     # B4=TemplateType, C4=nameOfDB
        wb.save(path)
        res = engine.generate_blocks([], out, emit=lambda *_: None)   # no builder -> override drives it
        grid = _read(os.path.join(res["dir"], "T1.csv"))
        ok(any(r[0] == "@" and "MY_DB" in r for r in grid), "override @ row carried through")
    _with_synthetic_templates(check)


if __name__ == "__main__":
    raise SystemExit(run("blocks", [
        ("database_queries", test_database_queries),
        ("table_add", test_table_add),
        ("engine_absolute_ref_and_instances", test_engine_absolute_ref_and_instances),
        ("engine_writes_com_db_with_constants", test_engine_writes_com_db_with_constants),
        ("xml_emit_flgnet_exactly_sized", test_xml_emit_flgnet_exactly_sized),
        ("xml_emit_flgnet_chunks_large_and", test_xml_emit_flgnet_chunks_large_and),
        ("xml_emit_and_coil_fc_from_table", test_xml_emit_and_coil_fc_from_table),
        ("xml_emitted_block_skips_creation_csv", test_xml_emitted_block_skips_creation_csv),
        ("standard_sheet_mirrors_csv_and_reads_back", test_standard_sheet_mirrors_csv_and_reads_back),
        ("coil_columns_round_trip", test_coil_columns_round_trip),
        ("override_writes_csv_verbatim_with_iterator", test_override_writes_csv_verbatim_with_iterator),
        ("shells_scan_and_inventory", test_shells_scan_and_inventory),
        ("output_header_is_full_shell_keyset", test_output_header_is_full_shell_keyset),
        ("override_uses_shell_rows", test_override_uses_shell_rows),
    ]))
