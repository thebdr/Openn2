"""M3 gate: the one shared workbook reader - newline->space headers, by-position, strike."""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook
from openpyxl.styles import Font
from pipeline3.io import workbook as wbk


def _make(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    # header row 1 (col G = "Bit ...", col O = newline-wrapped "Functional\nUnit")
    ws["G1"] = "Bit\n(full address)"
    ws["O1"] = "Functional\nUnit"
    # data rows 2,3
    ws["G2"] = "I0.0"; ws["O2"] = "MS1.CC1"
    ws["G3"] = "I0.1"; ws["O3"] = "MS2"
    ws["G3"].font = Font(strike=True)       # struck cell on row 3
    wb.save(path)


def _with_wb(fn):
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "io.xlsx")
        _make(path)
        fn(path)


def test_header_newline_to_space():
    def check(path):
        v = wbk.open_sheet(path, "NET SAFETY 50", header_row=1)
        try:
            eq(v.header("O"), "Functional Unit", "newline normalized to space")
            eq(v.header("G"), "Bit (full address)")
        finally:
            v.close()
    _with_wb(check)


def test_header_prefix_match():
    def check(path):
        v = wbk.open_sheet(path, "NET SAFETY 50")
        try:
            ok(v.header_matches("O", "Functional"))      # prefix
            ok(v.header_matches("G", "Bit"))
            ok(not v.header_matches("O", "Location"))
        finally:
            v.close()
    _with_wb(check)


def test_by_position_letter_and_index():
    def check(path):
        v = wbk.open_sheet(path, "NET SAFETY 50")
        try:
            eq(v.cell(2, "G"), "I0.0")
            eq(v.cell(2, 7), "I0.0")          # column G == index 7
            eq(v.text(2, "O"), "MS1.CC1")
        finally:
            v.close()
    _with_wb(check)


def test_strike():
    def check(path):
        v = wbk.open_sheet(path, "NET SAFETY 50")
        try:
            ok(v.cell_struck(3, "G"), "row 3 col G is struck")
            ok(not v.cell_struck(2, "G"))
            ok(v.row_struck(3))
            ok(not v.row_struck(2))
        finally:
            v.close()
    _with_wb(check)


def test_data_rows_and_location():
    def check(path):
        v = wbk.open_sheet(path, "NET SAFETY 50")
        try:
            eq(list(v.data_rows()), [2, 3])
            loc = v.location(2, "G")
            ok(loc.endswith("G2"), loc)
            ok("NET SAFETY 50" in loc, loc)
            ok(".xlsx" not in loc, "the workbook name is NOT in the link")
            eq(loc.count("!"), 1, "link is sheet!cell, not doc!sheet!cell")
        finally:
            v.close()
    _with_wb(check)


def test_regex_sheet_resolve():
    def check(path):
        v = wbk.open_sheet(path, r"NET SAFETY \d+")     # regex pattern
        try:
            eq(v.name, "NET SAFETY 50")
        finally:
            v.close()
    _with_wb(check)


if __name__ == "__main__":
    raise SystemExit(run("workbook", [
        ("header_newline_to_space", test_header_newline_to_space),
        ("header_prefix_match", test_header_prefix_match),
        ("by_position_letter_and_index", test_by_position_letter_and_index),
        ("strike", test_strike),
        ("data_rows_and_location", test_data_rows_and_location),
        ("regex_sheet_resolve", test_regex_sheet_resolve),
    ]))
