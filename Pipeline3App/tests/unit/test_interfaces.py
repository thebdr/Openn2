"""M7 gate (data-independent): Phase 400 interface generation.

Builds a synthetic MachineInterfaces template (<GENERIC> + SORTER, each with Side / Index /
Base Node / Base Address headers and <index> tokens) and synthetic staged IOC rows, then checks:
parse_instance / choose_sheet / find_interfaces, the 410 auto path (named-sheet hit + GENERIC
fallback + blank-Index warning + preserve-on-rerun), the 430 generate_one primitive (Side-1 +
Side-2 nodes), and the phase wiring (registration + run -> artifacts/files).
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string as CI
from openpyxl.worksheet.table import Table

from pipeline3.core import config
from pipeline3.context import PipelineContext
from pipeline3.domain import interfaces
import pipeline3.phases  # registers all phases
from pipeline3.phases import p400_interfaces
from pipeline3.registry import registry


def _set(ws, col, row, val):
    ws.cell(row=row, column=CI(col), value=val)


def _make_template(path):
    """Two sheets; headers Side(A)/Index(B)/Base Node(C)/Base Address(D), a signal-name col E with
    <index> tokens. Row 2 = Side 1, row 3 = Side 2 (matching the real template's shape)."""
    wb = Workbook()
    for i, sh in enumerate(("<GENERIC>", "SORTER")):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = sh
        for col, hdr in (("A", "Side"), ("B", "Index"), ("C", "Base Node"), ("D", "Base Address"),
                         ("E", "Signal Name")):
            _set(ws, col, 1, hdr)
        tag = "Sorter" if sh == "SORTER" else "Generic"
        _set(ws, "A", 2, 1); _set(ws, "B", 2, "01"); _set(ws, "C", 2, 10); _set(ws, "D", 2, 10000)
        _set(ws, "E", 2, f"PNC_Q_{tag}<index> HEARTBEAT")
        _set(ws, "A", 3, 2); _set(ws, "B", 3, "01"); _set(ws, "C", 3, 10); _set(ws, "D", 3, 0)
        _set(ws, "E", 3, f"PNC_Q_{tag}<index> POWER CUT")
    wb.save(path)


def _rows():
    return [
        {"script_type": "IOC", "index": "SORTER-01", "bit": "200", "id_node": "7",
         "device": "-K1", "profinet_ip": "192.168.5.7", "_source_sheet": "NET SAFETY 50", "_source_row": 5},
        {"script_type": "IOC", "index": "FOO-02", "bit": "300", "id_node": "8",
         "device": "-K2", "profinet_ip": "192.168.5.8", "_source_sheet": "NET SAFETY 50", "_source_row": 6},
        {"script_type": "A", "index": "0001", "bit": "I0.0", "_source_sheet": "NET SAFETY 50", "_source_row": 7},
        {"script_type": "IOC", "index": "", "bit": "400", "id_node": "9",
         "_source_sheet": "NET SAFETY 50", "_source_row": 8},
    ]


# --- pure helpers -------------------------------------------------------------------------- #

def test_parse_instance():
    eq(interfaces.parse_instance("SORTER-01"), ("SORTER", "01"))
    eq(interfaces.parse_instance("GEN 7"), ("GEN", "7"))
    eq(interfaces.parse_instance("SORTER-012"), ("SORTER", "012"))
    eq(interfaces.parse_instance(""), ("", ""))
    eq(interfaces.parse_instance(None), ("", ""))


def test_choose_sheet():
    sheets = ["<GENERIC>", "SORTER"]
    eq(interfaces.choose_sheet(sheets, "SORTER"), "SORTER")
    eq(interfaces.choose_sheet(sheets, "sorter"), "SORTER", "case-insensitive")
    eq(interfaces.choose_sheet(sheets, "FOO"), "<GENERIC>", "no match -> generic")
    eq(interfaces.choose_sheet(sheets, ""), "<GENERIC>")


def test_find_interfaces():
    recs, warnings = interfaces.find_interfaces(_rows())
    eq(len(recs), 2, "only the two IOC rows with an Index")
    eq([r["instance"] for r in recs], ["SORTER-01", "FOO-02"])
    eq(recs[0]["base"], "200"); eq(recs[0]["base_node"], "7"); eq(recs[0]["device"], "-K1")
    eq(len(warnings), 1, "blank-Index IOC row warned")
    ok("no Index" in warnings[0])


# --- 410 auto path ------------------------------------------------------------------------- #

def test_generate_auto_path():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        res = interfaces.generate(_rows(), out, tpl)
        eq(sorted(res["created"]), ["IF_FOO-02.xlsx", "IF_SORTER-01.xlsx"])
        eq(res["fallback"], ["FOO-02"], "FOO has no sheet -> GENERIC")
        eq(len(res["warnings"]), 1, "blank-Index warning carried through")

        # SORTER-01: the SORTER sheet, retitled, base/node/index plugged, <index> replaced
        wb = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))
        eq(wb.sheetnames, ["SORTER-01"], "only the chosen sheet, retitled to the instance")
        ws = wb["SORTER-01"]
        eq(ws["D2"].value, 200, "Side-1 Base Address = the IOC row's Bit")
        eq(ws["C2"].value, 7, "Side-1 Base Node = the IOC row's ID")
        eq(ws["B2"].value, "01"); eq(ws["B3"].value, "01", "Index set on the Side-1 and Side-2 rows")
        eq(ws["E2"].value, "PNC_Q_Sorter01 HEARTBEAT", "<index> token replaced")
        ok("<index>" not in (ws["E3"].value or ""))
        wb.close()

        # FOO-02: GENERIC fallback
        wb = load_workbook(os.path.join(out, "IF_FOO-02.xlsx"))
        eq(wb.sheetnames, ["FOO-02"])
        eq(wb["FOO-02"]["E2"].value, "PNC_Q_Generic02 HEARTBEAT", "GENERIC sheet, index 02")
        wb.close()


def test_index_written_to_side_rows_only():
    # The Index value goes to the Side-1 (row 2) and Side-2 (row 3) rows ONLY - a populated Index
    # cell on a non-Side row (here row 4) must be left untouched (regression: was "every populated
    # Index cell").
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx")
        wb = Workbook(); ws = wb.active; ws.title = "SORTER"
        for col, hdr in (("A", "Side"), ("B", "Index"), ("C", "Base Node"), ("D", "Base Address")):
            _set(ws, col, 1, hdr)
        _set(ws, "A", 2, 1); _set(ws, "B", 2, "TPL")     # Side 1, placeholder index
        _set(ws, "A", 3, 2); _set(ws, "B", 3, "TPL")     # Side 2, placeholder index
        _set(ws, "B", 4, "TPL")                          # a populated Index cell on a NON-Side row
        wb.save(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-07", "bit": "5", "id_node": "9",
                 "_source_sheet": "S", "_source_row": 2}]
        interfaces.generate(rows, out, tpl)
        ws = load_workbook(os.path.join(out, "IF_SORTER-07.xlsx"))["SORTER-07"]
        eq(ws["B2"].value, "07", "Side-1 Index written")
        eq(ws["B3"].value, "07", "Side-2 Index written")
        eq(ws["B4"].value, "TPL", "non-Side Index cell left untouched")


def test_preserve_on_rerun():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        interfaces.generate(_rows(), out, tpl)
        target = os.path.join(out, "IF_SORTER-01.xlsx")
        before = open(target, "rb").read()
        res2 = interfaces.generate(_rows(), out, tpl)
        eq(sorted(res2["preserved"]), ["IF_FOO-02.xlsx", "IF_SORTER-01.xlsx"])
        eq(res2["created"], [], "nothing re-created")
        eq(open(target, "rb").read(), before, "existing interface left byte-for-byte intact")


def test_float_string_base_and_node():
    # openpyxl widens cached/exported numeric cells to float -> staging yields '10000.0' / '7.0';
    # those must still plug as ints (regression: _to_int float tolerance).
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "10000.0", "id_node": "7.0",
                 "_source_sheet": "S", "_source_row": 2}]
        res = interfaces.generate(rows, out, tpl)
        eq(res["created"], ["IF_SORTER-01.xlsx"])
        ws = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))["SORTER-01"]
        eq(ws["D2"].value, 10000, "float-string Base Address plugged as int")
        eq(ws["C2"].value, 7, "float-string Base Node plugged as int")


def test_zero_base_address():
    # a legitimate 0 must plug (not be treated as blank/None).
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "0", "id_node": "0",
                 "_source_sheet": "S", "_source_row": 2}]
        interfaces.generate(rows, out, tpl)
        ws = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))["SORTER-01"]
        eq(ws["D2"].value, 0, "Base Address 0 plugged"); eq(ws["C2"].value, 0)


def test_no_generic_fallback_degrades_gracefully():
    # a template with no <GENERIC>: an unknown machine type must skip that row with a warning, not
    # crash the phase, and must not stop the valid rows around it (regression: per-row degradation).
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx")
        wb = Workbook(); ws = wb.active; ws.title = "SORTER"
        for col, hdr in (("A", "Side"), ("B", "Index"), ("C", "Base Node"), ("D", "Base Address")):
            _set(ws, col, 1, hdr)
        _set(ws, "A", 2, 1); _set(ws, "B", 2, "01")
        wb.save(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [
            {"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "S", "_source_row": 2},
            {"script_type": "IOC", "index": "FOO-02", "bit": "2", "_source_sheet": "S", "_source_row": 3},
            {"script_type": "IOC", "index": "SORTER-03", "bit": "3", "_source_sheet": "S", "_source_row": 4},
        ]
        res = interfaces.generate(rows, out, tpl)  # must NOT raise
        eq(sorted(res["created"]), ["IF_SORTER-01.xlsx", "IF_SORTER-03.xlsx"], "valid rows still generated")
        eq(len(res["warnings"]), 1, "the FOO row warned, not fatal")
        ok("FOO" in res["warnings"][0] and "skipped" in res["warnings"][0])
        ok(not os.path.exists(os.path.join(out, "IF_FOO-02.xlsx")))


def test_duplicate_instance_warns():
    # two IOC rows with the same Index: first wins, second is reported (data dropped), not silent.
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [
            {"script_type": "IOC", "index": "SORTER-01", "bit": "200", "_source_sheet": "NS50", "_source_row": 5},
            {"script_type": "IOC", "index": "SORTER-01", "bit": "999", "_source_sheet": "NS51", "_source_row": 9},
        ]
        res = interfaces.generate(rows, out, tpl)
        eq(res["created"], ["IF_SORTER-01.xlsx"], "only one file created")
        eq(res["preserved"], [], "the collision is not counted as a normal preserve")
        eq(len(res["warnings"]), 1); ok("duplicate IOC instance SORTER-01" in res["warnings"][0])
        ws = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))["SORTER-01"]
        eq(ws["D2"].value, 200, "first row wins")


def test_generate_no_ioc_rows():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        res = interfaces.generate([{"script_type": "A", "index": "0001"}], out, tpl)
        eq(res["created"], []); eq(res["warnings"], [])
        ok(os.path.isdir(out), "the (empty) interfaces folder is still created")


# --- 430 primitive ------------------------------------------------------------------------- #

def test_generate_one_both_sides():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        path, status = interfaces.generate_one(
            tpl, out, machine_type="SORTER", base_address="100",
            node_side1="3", node_side2="4", index="09", instance="SORTER-09")
        eq(status, "created"); eq(os.path.basename(path), "IF_SORTER-09.xlsx")
        ws = load_workbook(path)["SORTER-09"]
        eq(ws["D2"].value, 100, "Side-1 Base Address")
        eq(ws["C2"].value, 3, "Side-1 Base Node")
        eq(ws["C3"].value, 4, "Side-2 Base Node (explicit, custom path)")
        eq(ws["E2"].value, "PNC_Q_Sorter09 HEARTBEAT")


def test_generate_one_default_instance():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        out = os.path.join(d, "Interfaces")
        # no instance, no index -> instance defaults to the machine type
        path, _ = interfaces.generate_one(tpl, out, machine_type="SORTER", base_address="5")
        eq(os.path.basename(path), "IF_SORTER.xlsx")


# --- phase wiring -------------------------------------------------------------------------- #

def test_phase_registered():
    p = registry().get(400)
    eq(p.number, 400); eq(p.key, "interfaces"); eq(p.requires, (300,))
    eq([b.number for b in p.buttons], [410, 420, 430])
    eq([s.number for s in p.sub_phases], [410])


def test_phase_run():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_template(tpl)
        ctx = PipelineContext(params={"interface_template": tpl}, out_root=os.path.join(d, "out"))
        ctx.rows = _rows()
        res = p400_interfaces.run(ctx)
        ok(res.ok)
        out_dir = config.out_path(ctx.out_root, "interfaces_dir")
        eq(res.artifacts["interfaces_dir"], out_dir)
        ok(os.path.exists(os.path.join(out_dir, "IF_SORTER-01.xlsx")))
        ok(os.path.exists(os.path.join(out_dir, "IF_FOO-02.xlsx")))


# ============================================================================================== #
# Signal mirroring (the "Interface Mapping" feature)
# ============================================================================================== #

def _srow(script_type, mapping="", *, in_diag=False, db_kind="", db_names=(), name_in_db="",
          tag="", datablocks="", fu="=S1", loc="+L1", dev="-D1", diag_cabinet="", swp_cabinet="",
          diag_bit="", bit="I0.0", sheet="NS50", row=2, interface_tagname="", type_tag_name=""):
    """A synthetic staged row carrying the fields the mirroring engine reads. `interface_tagname` is
    the staging-computed value (may still carry {interface_name}/{interface_id} tokens)."""
    return {
        "script_type": script_type, "interface_mapping": mapping,
        "_type": {"in_diagnosis": in_diag, "db_kind": db_kind, "db_names": list(db_names),
                  "tag_name": type_tag_name},
        "name_in_db": name_in_db, "name_in_tagtable": tag, "datablocks": datablocks,
        "functional_unit": fu, "location": loc, "device": dev, "interface_tagname": interface_tagname,
        "diag_cabinet": diag_cabinet, "swp_cabinet": swp_cabinet, "diag_bit": diag_bit,
        "bit": bit, "_source_sheet": sheet, "_source_row": row,
    }


def _make_mirror_template(path):
    """A template whose SORTER sheet carries a real Excel data table (SORTER_SIGNALS) WITH the 3 diag
    columns + a Side/Index/Base Node/Base Address config block + an address formula on the data rows;
    GENERIC (Table136) has NO diag columns. Standard rows give last-used byte: SORTER Q=2/I=1."""
    wb = Workbook()
    ws = wb.active; ws.title = "SORTER"
    cols = ["Category", "Description", "Functional Unit", "Location", "Device",
            "diag_cabinet", "swp_cabinet", "diag_bit", "Data Type", "Direction </>",
            "I/O Offset Byte", "I/O Bit", "I/O Address Side 1", "I/O Address Side 2",
            "Signal Name Side 1", "Expression Side 1"]
    for i, h in enumerate(cols, start=1):
        ws.cell(1, i, h)
    addr = "=SORTER_SIGNALS[[#This Row],[I/O Offset Byte]]"
    for r, (d, off) in enumerate([(">", 0), (">", 2), ("<", 0), ("<", 1)], start=2):
        ws.cell(r, 9, "BOOL"); ws.cell(r, 10, d); ws.cell(r, 11, off); ws.cell(r, 12, 0)
        ws.cell(r, 13, addr); ws.cell(r, 14, addr); ws.cell(r, 15, "PNC<index> STD")
    ws.add_table(Table(displayName="SORTER_SIGNALS", ref="A1:P5"))
    for i, h in enumerate(["Side", "Index", "Base Node", "Base Address"], start=31):  # AE..AH
        ws.cell(1, i, h)
    ws.cell(2, 31, 1); ws.cell(2, 32, "01"); ws.cell(2, 33, 10); ws.cell(2, 34, 10000)
    ws.cell(3, 31, 2); ws.cell(3, 32, "01"); ws.cell(3, 33, 10); ws.cell(3, 34, 0)

    g = wb.create_sheet("<GENERIC>")
    gcols = ["Category", "Description", "Functional Unit", "Location", "Device", "Data Type",
             "Direction </>", "I/O Offset Byte", "I/O Bit", "I/O Address Side 1",
             "I/O Address Side 2", "Signal Name Side 1"]
    for i, h in enumerate(gcols, start=1):
        g.cell(1, i, h)
    gaddr = "=Table136[[#This Row],[I/O Offset Byte]]"
    for r, (d, off) in enumerate([(">", 0), ("<", 0)], start=2):
        g.cell(r, 6, "BOOL"); g.cell(r, 7, d); g.cell(r, 8, off); g.cell(r, 9, 0)
        g.cell(r, 10, gaddr); g.cell(r, 11, gaddr); g.cell(r, 12, "PNC<index> GEN")
    g.add_table(Table(displayName="Table136", ref="A1:L3"))
    for i, h in enumerate(["Side", "Index", "Base Node", "Base Address"], start=28):  # AB..AE
        g.cell(1, i, h)
    g.cell(2, 28, 1); g.cell(2, 29, "01"); g.cell(2, 30, 10); g.cell(2, 31, 10000)
    g.cell(3, 28, 2); g.cell(3, 29, "01"); g.cell(3, 30, 10); g.cell(3, 31, 0)
    wb.save(path)


def test_detect_diag_and_parse():
    ok(interfaces._detect_diag("SORTER+DIAG-02")); ok(not interfaces._detect_diag("SORTER-01"))
    eq(interfaces.parse_instance("SORTER+DIAG-02"), ("SORTER", "02"))
    eq(interfaces.parse_instance("SORTER-01"), ("SORTER", "01"), "legacy unchanged")


def test_mirror_membership_mapping():
    rows = [_srow("B1/2", "01", tag="BreakerTag"), _srow("F1/2", "02", tag="FireTag"),
            _srow("Z1", "01|02", tag="ZTag"), _srow("KQ", "", tag="KqTag"),
            _srow("IOC", "01")]
    e1, _ = interfaces.collect_mirror_set(rows, index="01", is_diag=False,
                                          db_rules=[], diag_rules=[], if_rules=[])
    eq(sorted(x.mirror_name for x in e1), ['"BreakerTag"', '"ZTag"'], "only 01-mapped non-IOC")
    ok(all(x.direction == "Q" for x in e1), "I/O-list mirrors are Q")
    e2, _ = interfaces.collect_mirror_set(rows, index="2", is_diag=False,
                                          db_rules=[], diag_rules=[], if_rules=[])
    eq(sorted(x.mirror_name for x in e2), ['"FireTag"', '"ZTag"'], "'2' == '02' (int compare)")


def test_mirror_membership_diag():
    rows = [_srow("A", "", in_diag=True, tag="AlarmTag"), _srow("W", "", in_diag=True, tag="WarnTag"),
            _srow("KQ", "", in_diag=False, tag="KqTag"), _srow("B1/2", "07", in_diag=True, tag="Brk")]
    e, _ = interfaces.collect_mirror_set(rows, index="99", is_diag=True,
                                         db_rules=[], diag_rules=[], if_rules=[])
    eq(sorted(x.script_type for x in e), ["A", "B1/2", "W"], "all in_diag, KQ excluded")
    e0, _ = interfaces.collect_mirror_set(rows, index="99", is_diag=False,
                                          db_rules=[], diag_rules=[], if_rules=[])
    eq(e0, [], "without is_diag, index 99 matches nothing")


def test_byte_packing_by_script_type():
    elems = ([interfaces._Elem(script_type="A", mirror_name=f"a{i}") for i in range(9)] +
             [interfaces._Elem(script_type="W", mirror_name="w")] +
             [interfaces._Elem(script_type="SPD", mirror_name="spd", data_type="WORD")])
    interfaces.allocate_bytes(elems, {"Q": 2, "I": -1}, 8)
    a = [x for x in elems if x.script_type == "A"]
    # custom area 2-byte-aligned start = last_Q(2)+1+gap(8)=11 -> 12; A's 9 bits fill its 2-byte block
    eq([x.offset_byte for x in a], [12] * 8 + [13], "9 A bits across the 2-byte block 12-13")
    eq([x.bit for x in a], [0, 1, 2, 3, 4, 5, 6, 7, 0])
    w = next(x for x in elems if x.script_type == "W")
    eq((w.offset_byte, w.bit), (14, 0), "W on its OWN 2-byte block (A reserved a full 2 bytes)")
    spd = next(x for x in elems if x.script_type == "SPD")
    eq((spd.offset_byte, spd.bit), (16, None), "WORD even-aligned, its own slot, no bit")
    by_byte = {}
    for x in elems:
        for bb in ([x.offset_byte] if x.bit is not None else [x.offset_byte, x.offset_byte + 1]):
            by_byte.setdefault(bb, set()).add(x.script_type)
    ok(all(len(v) == 1 for v in by_byte.values()), "never two script_types in one byte")


def test_follower_and_if_rules():
    rows = [_srow("DI1/2", "01", db_kind="db", db_names=["07_DOOR"], name_in_db="Door Closed",
                  datablocks="07_DOOR", fu="=S1", loc="+SG1", dev="-B1")]
    db_rules = [{"name": "Door Alarm", "required_types": ["DI1/2"], "dev_type": "",
                 "db_name": "07_DOOR", "member": "Door Alarm [ {functional_unit}{location}{device} ]"}]
    if_rules = [{"name": "Door Cmd", "required_types": ["DI1/2"], "dev_type": "", "direction": "I",
                 "data_type": "BOOL", "script_type": "IF_DOOR_CMD",
                 "member": "Open [ {functional_unit}{location}{device} ]"}]
    e, _ = interfaces.collect_mirror_set(rows, index="01", is_diag=False,
                                         db_rules=db_rules, diag_rules=[], if_rules=if_rules)
    eq(next(x for x in e if x.source == "mirror").mirror_name, '"07_DOOR"."Door Closed"', "db-qualified")
    fol = next(x for x in e if x.source == "follow")
    eq(fol.mirror_name, '"07_DOOR"."Door Alarm [ =S1+SG1-B1 ]"', "follower db-qualified"); eq(fol.direction, "Q")
    ifr = next(x for x in e if x.source == "if_rule")
    eq((ifr.direction, ifr.script_type), ("I", "IF_DOOR_CMD"), "if-rule may add an input")
    ok("Open [ =S1+SG1-B1 ]" in ifr.mirror_name)


def test_mirror_name_precedence():
    db = _srow("X", db_kind="db", db_names=["DB1"], name_in_db="Member", datablocks="DB1", tag="TagX")
    eq(interfaces._mirror_name(db), '"DB1"."Member"', "db_element wins over tag")
    eq(interfaces._mirror_name(_srow("Y", name_in_db="", tag="TagY")), '"TagY"', "tag fallback")
    eq(interfaces._mirror_name(_srow("Z", name_in_db="", tag="")), "", "neither -> empty")


def test_generate_with_mirror():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "100", "id_node": "5",
                 "_source_sheet": "NS", "_source_row": 2},
                _srow("B1/2", "01", db_kind="db", db_names=["DB1"], name_in_db="Breaker",
                      datablocks="DB1", diag_cabinet="012", swp_cabinet="12", diag_bit="03", row=10,
                      interface_tagname="PNC_Q_{interface_name}-{interface_id}_Breaker")]
        res = interfaces.generate(rows, out, tpl)
        eq(res["errors"], []); eq(res["created"], ["IF_SORTER-01.xlsx"])
        ws = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))["SORTER-01"]
        name, c1, r1, c2, r2, hdr = interfaces._find_data_table(ws)
        eq(ws.tables[name].ref, "A1:P21", "B1/2 fills a full 2-byte block (16 rows) + separator")
        crow = 6   # first custom row = original last data row (5) + 1; the B1/2 signal sits there
        eq(ws.cell(crow, hdr["Direction </>"]).value, ">", "mirror is Q")
        eq(ws.cell(crow, hdr["I/O Offset Byte"]).value, 12, "2-byte-aligned block: 2+1+8=11 -> 12")
        eq(ws.cell(crow, hdr["diag_cabinet"]).value, "012")
        eq(ws.cell(crow, hdr["swp_cabinet"]).value, "12")
        eq(ws.cell(crow, hdr["diag_bit"]).value, "03")
        eq(ws.cell(crow, hdr["Signal Name Side 1"]).value, "PNC_Q_SORTER-01_Breaker",
           "Signal Name = interface_tagname with {interface_name}/{interface_id} filled")
        eq(ws.cell(crow, hdr["Expression Side 1"]).value, '"DB1"."Breaker"', "Expression = the binding")
        ok(str(ws.cell(crow, hdr["I/O Address Side 1"]).value or "").startswith("="), "addressed")
        ok(ws.cell(7, hdr["Signal Name Side 1"]).value in (None, ""), "padding rows blank for input")
        ok(str(ws.cell(7, hdr["I/O Address Side 1"]).value or "").startswith("="), "padding rows addressed")


def test_generic_mirror_omits_diag_columns():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "FOO-03", "bit": "50", "_source_sheet": "NS", "_source_row": 2},
                _srow("B1/2", "03", db_kind="db", db_names=["DB1"], name_in_db="Breaker",
                      datablocks="DB1", diag_cabinet="012", swp_cabinet="12", diag_bit="03", row=10,
                      interface_tagname="PNC_Q_{interface_name}-{interface_id}_Breaker")]
        res = interfaces.generate(rows, out, tpl)
        eq(res["fallback"], ["FOO-03"], "FOO -> GENERIC")
        ws = load_workbook(os.path.join(out, "IF_FOO-03.xlsx"))["FOO-03"]
        name, c1, r1, c2, r2, hdr = interfaces._find_data_table(ws)
        ok("diag_cabinet" not in hdr, "GENERIC has no diag columns")
        # GENERIC table was A1:L3 -> the mirrored signal is the first custom row (4)
        eq(ws.cell(4, hdr["Signal Name Side 1"]).value, "PNC_Q_FOO-03_Breaker", "still mirrored on GENERIC")


def test_duplicate_index_error():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-02", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                {"script_type": "IOC", "index": "SORTER+DIAG-02", "bit": "2", "_source_sheet": "NS", "_source_row": 3}]
        res = interfaces.generate(rows, out, tpl)
        eq(len(res["errors"]), 1)
        ok("SORTER-02" in res["errors"][0] and "SORTER+DIAG-02" in res["errors"][0],
           "the error names both colliding instances")
        ok(os.path.exists(os.path.join(out, "IF_SORTER-02.xlsx")), "full names still produced")
        ok(os.path.exists(os.path.join(out, "IF_SORTER+DIAG-02.xlsx")))


def test_idx_key():
    eq(interfaces._idx_key("01"), 1); eq(interfaces._idx_key("1"), 1); eq(interfaces._idx_key("02"), 2)
    eq(interfaces._idx_key(""), None); eq(interfaces._idx_key("FOO"), "FOO")


def test_duplicate_index_zero_padding():
    # '02' and '2' normalize to one interface, so a mapping '2' would resolve to both -> ERROR.
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-02", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                {"script_type": "IOC", "index": "SORTER-2", "bit": "2", "_source_sheet": "NS", "_source_row": 3}]
        res = interfaces.generate(rows, out, tpl)
        eq(len(res["errors"]), 1, "'02' vs '2' is the same index -> ERROR")
        ok("SORTER-02" in res["errors"][0] and "SORTER-2" in res["errors"][0])


def test_mirror_warning_surfaced():
    # a mapped signal with neither db_element nor tag is dropped WITH a warning that reaches result.
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                _srow("XX", "01", name_in_db="", tag="", row=3)]
        res = interfaces.generate(rows, out, tpl)
        eq(res["created"], ["IF_SORTER-01.xlsx"])
        ok(any("not mirrored" in w for w in res["warnings"]), "skip warning reaches result['warnings']")
        ok(any("SORTER-01:" in w for w in res["warnings"]), "warning is instance-prefixed")


def test_dangling_nonnumeric_mapping_warns():
    # a blank-index IOC row must not poison the dangling-mapping check (None excluded from ioc_idx).
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                {"script_type": "IOC", "index": "SORTER-01", "bit": "2", "_source_sheet": "NS", "_source_row": 3},
                _srow("B1/2", "FOO", tag="X", row=4)]
        res = interfaces.generate(rows, out, tpl)
        ok(any("references no IOC interface" in w and "FOO" in w for w in res["warnings"]),
           "dangling 'FOO' surfaced despite the blank-index IOC row")


def test_no_match_index_warns():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                _srow("B1/2", "99", tag="X", row=3)]
        res = interfaces.generate(rows, out, tpl)
        ok(any("references no IOC interface" in w for w in res["warnings"])); eq(res["errors"], [])


_DB_RULE = {"name": "DoorAlarm", "required_types": ["DI1/2"], "dev_type": "", "db_name": "05",
            "member": "Alarm [ {functional_unit}{location}{device} ]", "interface_tagname": "PNC_Q_{member}"}
_IF_RULE = {"name": "Cmd", "required_types": ["DI1/2"], "dev_type": "", "direction": "I",
            "data_type": "BOOL", "script_type": "IF_CMD",
            "member": "Cmd [ {functional_unit}{location}{device} ]", "interface_tagname": "PNC_{direction}_{member}"}
_WORD_RULE = {"name": "Spd", "required_types": ["DI1/2"], "dev_type": "", "direction": "Q",
              "data_type": "WORD", "script_type": "IF_SPD",
              "member": "Speed [ {functional_unit}{location}{device} ]", "interface_tagname": "PNC_{direction}_{member}"}


def test_signal_name_sources():
    rows = [_srow("DI1/2", "01", db_kind="db", db_names=["DB1"], name_in_db="DC", datablocks="DB1",
                  fu="=S1", loc="+SG1", dev="-B1",
                  interface_tagname="PNC_Q_{interface_name}-{interface_id}_DoorIn")]
    elems, _ = interfaces.collect_mirror_set(rows, index="01", is_diag=False,
                                             db_rules=[_DB_RULE], diag_rules=[], if_rules=[_IF_RULE])
    direct = next(e for e in elems if e.source == "mirror")
    ok("{interface_name}" in direct.signal_name and "{interface_id}" in direct.signal_name,
       "direct keeps the per-interface tokens for the generator to fill")
    eq(direct.mirror_name, '"DB1"."DC"', "Expression = the binding")
    fol = next(e for e in elems if e.source == "follow")
    eq(fol.signal_name, "PNC_Q_Alarm [ =S1+SG1-B1 ]", "follower Signal Name from the rule interface_tagname")
    eq(fol.mirror_name, '"05"."Alarm [ =S1+SG1-B1 ]"', "follower Expression = db-qualified")
    ifr = next(e for e in elems if e.source == "if_rule")
    eq(ifr.signal_name, "PNC_I_Cmd [ =S1+SG1-B1 ]", "if-rule uses {direction}=I")


def _gen_one(d, cfg, rows):
    tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
    out = os.path.join(d, "Interfaces")
    interfaces.generate_one(tpl, out, machine_type="SORTER", index="01", instance="SORTER-01",
                            template_sheets=interfaces.template_sheet_names(tpl),
                            mirror_rows=rows, mirror_cfg=cfg)
    return load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))["SORTER-01"]


def test_word_single_row():
    # a WORD is ONE row (no bit expansion) + a blank separator.
    with tempfile.TemporaryDirectory() as d:
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                _srow("DI1/2", "01", db_kind="db", db_names=["DB1"], name_in_db="DC", datablocks="DB1",
                      fu="=S1", loc="+SG1", dev="-B1")]
        ws = _gen_one(d, {"db_rules": [], "diag_rules": [], "if_rules": [_WORD_RULE], "gap": 8}, rows)
        name, c1, r1, c2, r2, hdr = interfaces._find_data_table(ws)
        wrows = [r for r in range(r1 + 1, r2 + 2) if str(ws.cell(r, hdr["Category"]).value or "") == "IF_SPD"]
        eq(len(wrows), 1, "WORD = exactly one row, NOT bit-expanded")
        eq(ws.cell(wrows[0], hdr["Data Type"]).value, "WORD")
        ok(ws.cell(wrows[0], hdr["I/O Bit"]).value in (None, ""), "WORD has no bit")
        eq(ws.cell(wrows[0], hdr["I/O Offset Byte"]).value % 2, 0, "WORD even-aligned")
        ok(ws.cell(wrows[0], hdr["Signal Name Side 1"]).value, "WORD carries its name")
        ok(str(ws.cell(wrows[0] + 1, hdr["Category"]).value or "") == "", "blank separator after the word")


def test_bool_block_full_2bytes():
    # a BOOL group fills a FULL 2-byte block: used bit(s) + blank addressed rows + a separator.
    with tempfile.TemporaryDirectory() as d:
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                _srow("DI1/2", "01", db_kind="db", db_names=["DB1"], name_in_db="DC", datablocks="DB1",
                      fu="=S1", loc="+SG1", dev="-B1", interface_tagname="PNC_Q_{interface_id}_DoorIn")]
        ws = _gen_one(d, {"db_rules": [], "diag_rules": [], "if_rules": [], "gap": 8}, rows)
        name, c1, r1, c2, r2, hdr = interfaces._find_data_table(ws)
        drows = [r for r in range(r1 + 1, r2 + 1) if str(ws.cell(r, hdr["Category"]).value or "") == "DI1/2"]
        eq(len(drows), 16, "one BOOL signal still fills a full 2-byte block (16 rows)")
        eq(ws.cell(drows[0], hdr["Data Type"]).value, "BOOL")
        ok(ws.cell(drows[0], hdr["Signal Name Side 1"]).value, "the used signal is the first row")
        ok(ws.cell(drows[5], hdr["Signal Name Side 1"]).value in (None, ""), "the rest blank for input")
        ok(str(ws.cell(drows[5], hdr["I/O Address Side 1"]).value or "").startswith("="), "blank rows addressed")
        start = ws.cell(drows[0], hdr["I/O Offset Byte"]).value
        eq(start % 2, 0, "block 2-byte aligned")
        eq([(ws.cell(r, hdr["I/O Offset Byte"]).value, ws.cell(r, hdr["I/O Bit"]).value) for r in drows],
           [(start + b // 8, b % 8) for b in range(16)], "16 bits across the 2 bytes")
        ok(str(ws.cell(drows[-1] + 1, hdr["Category"]).value or "") == "", "blank separator after the block")


def test_insert_interface_sheets():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        iol = os.path.join(d, "iolist.xlsx")
        wb = Workbook(); wb.active.title = "NET SAFETY 50"; wb.active["A1"] = "src"; wb.save(iol)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2},
                _srow("B1/2", "01", db_kind="db", db_names=["DB1"], name_in_db="Breaker", datablocks="DB1",
                      interface_tagname="PNC_Q_{interface_id}_Breaker", row=3)]
        res = interfaces.generate(rows, out, tpl, overwrite=True, iolist_path=iol)
        ok(any("inserted" in a and "IF_SORTER-01" in a for a in res["iolist"]), "sheet inserted")
        ok(any(".bak_" in a for a in res["iolist"]) and
           any(f.startswith("iolist.xlsx.bak_") for f in os.listdir(d)), "I/O List backed up first")
        wb2 = load_workbook(iol)
        ok("IF_SORTER-01" in wb2.sheetnames and "NET SAFETY 50" in wb2.sheetnames,
           "inserted sheet is IF_-prefixed + original kept")
        ok("SORTER-01" not in wb2.sheetnames, "the un-prefixed name is NOT used")
        ok("SORTER_SIGNALS_IF_SORTER_01" in list(wb2["IF_SORTER-01"].tables), "table renamed per-instance")
        # idempotent: a second run leaves the sheet alone
        res2 = interfaces.generate(rows, out, tpl, overwrite=True, iolist_path=iol)
        ok(any("already present" in a and "IF_SORTER-01" in a for a in res2["iolist"]), "second run skips")
        eq(load_workbook(iol).sheetnames.count("IF_SORTER-01"), 1, "no duplicate sheet")


def test_insert_preserves_formula_values():
    # data-dependent (skips when the real docs are absent): inserting the interface sheets must be
    # LOSSLESS - the I/O List's existing formulas AND their cached values both survive.
    from pipeline3.core import config as c
    iol_src = (c.load_params().get("io_list") or {}).get("path")
    if not (iol_src and os.path.exists(iol_src)):
        print("    (skipped: real I/O List absent)"); return
    import shutil
    from pipeline3 import app
    from pipeline3.context import PipelineContext
    from pipeline3.domain import staging
    ctx = PipelineContext.create(profile="main", emit=lambda *a, **k: None)
    ctx.completed.add(200); app.run_phase(ctx, 300)
    ips_before = sorted(r["profinet_ip"] for r in ctx.rows if r.get("profinet_ip"))
    if not ips_before:
        print("    (skipped: no profinet_ip values)"); return
    # find a sheet+cell that holds a formula in the source, to assert the formula itself survives
    src_wb = load_workbook(iol_src)
    fcell = next(((sh, cell.coordinate) for sh in src_wb.sheetnames for row in src_wb[sh].iter_rows()
                  for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")), None)
    src_wb.close()
    with tempfile.TemporaryDirectory() as d:
        iol = os.path.join(d, "iol.xlsx"); shutil.copyfile(iol_src, iol)
        out = os.path.join(d, "Interfaces")
        interfaces.generate(ctx.rows, out, ctx.params["interface_template"], overwrite=True, iolist_path=iol)
        p = dict(ctx.params); p["io_list"] = dict(ctx.params["io_list"], path=iol)
        rows2, *_ = staging.load_io_list(p, c.load_signal_types(), iol)
        ips_after = sorted(r["profinet_ip"] for r in rows2 if r.get("profinet_ip"))
        eq(ips_after, ips_before, "cached Profinet IP values survive (data_only readers)")
        if fcell:
            sh, coord = fcell
            v = load_workbook(iol)[sh][coord].value
            ok(isinstance(v, str) and v.startswith("="), "the formula itself survives (lossless)")


def test_insert_strips_table_dxf_and_calc():
    # Regression: a copied table must NOT carry the source's dataDxfId (an index into the SOURCE
    # workbook's diff-format records - out of range in the I/O List) or a calculatedColumnFormula
    # naming the OLD table; either makes Excel drop the whole table ("Removed Records: Table").
    import zipfile
    from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo, TableFormula
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx")
        wb = Workbook(); ws = wb.active; ws.title = "SORTER"
        hdr = ["Category", "Description", "Data Type", "Direction </>", "I/O Offset Byte",
               "I/O Bit", "I/O Address Side 1", "Signal Name Side 1"]
        for i, h in enumerate(hdr, start=1):
            ws.cell(1, i, h)
        ws.cell(2, 3, "BOOL"); ws.cell(2, 4, ">"); ws.cell(2, 5, 0); ws.cell(2, 6, 0)
        ws.cell(2, 7, "=DT[[#This Row],[I/O Offset Byte]]")
        cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(hdr)]
        cols[6].dataDxfId = 99                                                    # out-of-range dxf ref
        cols[6].calculatedColumnFormula = TableFormula(attr_text="DT[[#This Row],[I/O Offset Byte]]")
        t = Table(displayName="DT", ref="A1:H2"); t.tableColumns = cols
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
        for i, h in enumerate(["Side", "Index", "Base Node", "Base Address"], start=28):
            ws.cell(1, i, h)
        ws.cell(2, 28, 1); ws.cell(2, 29, "01"); ws.cell(2, 30, 10); ws.cell(2, 31, 0)
        wb.save(tpl)
        iol = os.path.join(d, "iol.xlsx"); w = Workbook(); w.active.title = "NET SAFETY 50"; w.active["A1"] = "x"; w.save(iol)
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2}]
        interfaces.generate(rows, os.path.join(d, "O"), tpl, overwrite=True, iolist_path=iol)
        with zipfile.ZipFile(iol) as z:
            tx = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/tables/")]
        ok(tx, "a table part was inserted")
        ok(not any("DxfId" in x for x in tx), "source dataDxfId is stripped (else Excel drops the table)")
        ok(not any("calculatedColumnFormula" in x for x in tx), "stale calc-column formula is stripped")
        load_workbook(iol)   # must reload cleanly


def test_overwrite_vs_preserve():
    with tempfile.TemporaryDirectory() as d:
        tpl = os.path.join(d, "tpl.xlsx"); _make_mirror_template(tpl)
        out = os.path.join(d, "Interfaces")
        rows = [{"script_type": "IOC", "index": "SORTER-01", "bit": "1", "_source_sheet": "NS", "_source_row": 2}]
        eq(interfaces.generate(rows, out, tpl)["created"], ["IF_SORTER-01.xlsx"])
        target = os.path.join(out, "IF_SORTER-01.xlsx")
        os.utime(target, (0, 0))
        interfaces.generate(rows, out, tpl)                       # default -> preserved
        eq(os.path.getmtime(target), 0, "default preserves the existing file (no rewrite)")
        r = interfaces.generate(rows, out, tpl, overwrite=True)   # overwrite -> regenerated
        ok(os.path.getmtime(target) > 0, "overwrite regenerates the file")
        eq(r["preserved"], [], "overwrite never reports preserved")


def test_io_address_mirror():
    f = interfaces._io_address
    isy = "I<base+offset>/.<bit>"
    eq(f(isy, 10000, ">", "BOOL", 0, 0), "Q10000.0", "BOOL output -> Q with bit")
    eq(f(isy, 10000, "<", "BOOL", 5, 3), "I10005.3", "BOOL input -> I, base+offset.bit")
    eq(f(isy, 10000, ">", "WORD", 26, None), "Q10026", "WORD -> byte only (TEXTBEFORE '/'), no bit")
    eq(f(isy, 10000, "", "BOOL", 0, 0), "", "no direction -> no address")
    eq(f(isy, None, ">", "BOOL", 0, 0), "", "no base -> no address")
    eq(f(isy, 10000, ">", "BOOL", None, 0), "", "no offset -> no address")


def test_resolve_num_chain():
    wb = Workbook(); ws = wb.active
    ws["C2"], ws["D2"] = 1, 0
    ws["C3"], ws["D3"] = "=C2", "=D2+1"        # the interface table's two offset/bit chain shapes
    ws["C4"], ws["D4"] = "=C3", "=D3+1"
    ws["C5"], ws["C6"] = "12.0", "x"
    memo = {}
    eq(interfaces._resolve_num(ws, "C4", memo), 1, "=C3 -> =C2 -> 1")
    eq(interfaces._resolve_num(ws, "D4", memo), 2, "=D3+1 -> (=D2+1)+1 -> 2")
    eq(interfaces._resolve_num(ws, "C5", memo), 12, "float-string literal -> int")
    ok(interfaces._resolve_num(ws, "C6", memo) is None, "non-numeric -> None")


def _make_if_addr_sheet(path, base=10000):
    """A minimal inserted-IF_ sheet: the address-relevant headers + an Input Format / Base Address on
    row 2, a (placeholder) LET-formula address cell per data row, and an offset/bit chain (=C3 / =D3+1)."""
    wb = Workbook(); ws = wb.active; ws.title = "IF_T"
    cols = ["Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit", "I/O Address Side 1",
            "Base Address", "Input Format", "Signal Name Side 1"]     # A..H ; offset=C bit=D addr=E
    for i, h in enumerate(cols, start=1):
        ws.cell(1, i, h)
    LET = "=$A$1"                                                     # any formula -> data_type 'f'
    def setrow(r, dt, d, off, bit, sig):
        ws.cell(r, 1, dt); ws.cell(r, 2, d); ws.cell(r, 3, off); ws.cell(r, 4, bit)
        ws.cell(r, 5, LET); ws.cell(r, 8, sig)
    setrow(2, "BOOL", ">", 0, 0, "SIG1")          # Q<base>.0
    setrow(3, "BOOL", ">", 1, 0, "SIG2")          # Q<base+1>.0
    setrow(4, "BOOL", ">", "=C3", "=D3+1", "SIG3")  # CHAIN -> off 1, bit 1 -> Q<base+1>.1
    setrow(5, "WORD", ">", 4, None, "SIG4")        # WORD -> Q<base+4>, no bit
    setrow(6, "BOOL", "<", 0, 0, "SIG5")           # I<base>.0
    ws.cell(7, 5, LET)                             # blank separator: LET present, no direction -> ''
    ws.cell(2, 6, base); ws.cell(2, 7, "I<base+offset>/.<bit>")
    wb.save(path)


def test_interface_address_caches():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(p)
        cache = interfaces._interface_address_caches(p)
        eq(cache.get("E2"), ("str", "Q10000.0"), "BOOL output bit 0")
        eq(cache.get("E3"), ("str", "Q10001.0"), "BOOL output byte+1")
        eq(cache.get("E4"), ("str", "Q10001.1"), "offset/bit CHAIN resolved (=C3 / =D3+1)")
        eq(cache.get("E5"), ("str", "Q10004"), "WORD: byte only, no bit")
        eq(cache.get("E6"), ("str", "I10000.0"), "BOOL input -> I")
        ok("E7" not in cache, "blank separator row (no direction) is not seeded")


def test_insert_seeds_interface_address_cache():
    # the reported phase-500 bug: the inserted IF_ address LET has no Excel cache, so phase 510 reads
    # None and skips. The insert now SEEDS a computed cache that a data_only reader sees (no Excel).
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(ifp)
        iol = os.path.join(d, "iol.xlsx")
        wb = Workbook(); wb.active.title = "NET SAFETY 50"; wb.active["A1"] = "src"; wb.save(iol)
        acts = interfaces.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ok(any("seeded" in a and "IF_SORTER-01" in a for a in acts), "seeding is reported")
        ws = load_workbook(iol, data_only=True)["IF_SORTER-01"]
        eq(ws["E2"].value, "Q10000.0", "data_only reader sees the seeded address (no Excel needed)")
        eq(ws["E4"].value, "Q10001.1", "the chained-offset address is seeded too")
        live = load_workbook(iol)["IF_SORTER-01"]["E2"].value                # formula kept for Excel
        ok(isinstance(live, str) and live.startswith("="), "the live LET formula is preserved")


if __name__ == "__main__":
    raise SystemExit(run("interfaces", [
        ("io_address_mirror", test_io_address_mirror),
        ("resolve_num_chain", test_resolve_num_chain),
        ("interface_address_caches", test_interface_address_caches),
        ("insert_seeds_interface_address_cache", test_insert_seeds_interface_address_cache),
        ("parse_instance", test_parse_instance),
        ("choose_sheet", test_choose_sheet),
        ("find_interfaces", test_find_interfaces),
        ("generate_auto_path", test_generate_auto_path),
        ("index_written_to_side_rows_only", test_index_written_to_side_rows_only),
        ("preserve_on_rerun", test_preserve_on_rerun),
        ("float_string_base_and_node", test_float_string_base_and_node),
        ("zero_base_address", test_zero_base_address),
        ("no_generic_fallback_degrades_gracefully", test_no_generic_fallback_degrades_gracefully),
        ("duplicate_instance_warns", test_duplicate_instance_warns),
        ("generate_no_ioc_rows", test_generate_no_ioc_rows),
        ("generate_one_both_sides", test_generate_one_both_sides),
        ("generate_one_default_instance", test_generate_one_default_instance),
        ("phase_registered", test_phase_registered),
        ("phase_run", test_phase_run),
        ("detect_diag_and_parse", test_detect_diag_and_parse),
        ("mirror_membership_mapping", test_mirror_membership_mapping),
        ("mirror_membership_diag", test_mirror_membership_diag),
        ("byte_packing_by_script_type", test_byte_packing_by_script_type),
        ("follower_and_if_rules", test_follower_and_if_rules),
        ("mirror_name_precedence", test_mirror_name_precedence),
        ("generate_with_mirror", test_generate_with_mirror),
        ("generic_mirror_omits_diag_columns", test_generic_mirror_omits_diag_columns),
        ("duplicate_index_error", test_duplicate_index_error),
        ("idx_key", test_idx_key),
        ("duplicate_index_zero_padding", test_duplicate_index_zero_padding),
        ("mirror_warning_surfaced", test_mirror_warning_surfaced),
        ("dangling_nonnumeric_mapping_warns", test_dangling_nonnumeric_mapping_warns),
        ("no_match_index_warns", test_no_match_index_warns),
        ("signal_name_sources", test_signal_name_sources),
        ("word_single_row", test_word_single_row),
        ("bool_block_full_2bytes", test_bool_block_full_2bytes),
        ("insert_interface_sheets", test_insert_interface_sheets),
        ("insert_strips_table_dxf_and_calc", test_insert_strips_table_dxf_and_calc),
        ("insert_preserves_formula_values", test_insert_preserves_formula_values),
        ("overwrite_vs_preserve", test_overwrite_vs_preserve),
    ]))
