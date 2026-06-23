"""Data-independent unit cases for io/xlsx_cache (formula-cache preservation across an openpyxl save).

openpyxl drops the cached `<v>` of formula cells on save; xlsx_cache.restore patches them back so a
data_only reader still sees them (the fix that stops phase 200 Fill from degrading the I/O List).
Builds a synthetic workbook whose formula cell carries a cached value (injected into the sheet XML,
like a real export) - no real documents needed.
"""
import io
import os
import re
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook

from _harness import run, eq, ok
from pipeline3.io import xlsx_cache


def _make_cached_xlsx(path):
    """A workbook with one sheet 'S1', cell A1 = '=1+2' carrying a cached value <v>3</v>."""
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"] = "=1+2"
    wb.save(path)
    wb.close()
    with open(path, "rb") as f:
        data = f.read()
    part = xlsx_cache._sheet_name_to_part(data)["S1"]
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        parts = {n: z.read(n) for n in names}
    xml = parts[part].decode("utf-8")
    xml = re.sub(r'(<c r="A1"[^>]*>)(.*?)(</c>)',
                 lambda m: m.group(1) + m.group(2).replace("</f>", "</f><v>3</v>") + m.group(3),
                 xml, count=1, flags=re.DOTALL)
    parts[part] = xml.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, parts[n])
    with open(path, "wb") as f:
        f.write(out.getvalue())


def test_capture_reads_formula_cache():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.xlsx")
        _make_cached_xlsx(p)
        with open(p, "rb") as f:
            caches = xlsx_cache.capture(f.read())
        ok("S1" in caches and "A1" in caches["S1"], "captured the formula cache")
        eq(caches["S1"]["A1"][1], "3")


def _a1(path):
    wb = load_workbook(path, data_only=True)
    v = wb.active["A1"].value
    wb.close()
    return v


def test_restore_round_trip():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.xlsx")
        _make_cached_xlsx(p)
        ok(_a1(p) == 3, "data_only sees the injected cache")
        with open(p, "rb") as f:
            orig = f.read()
        wb = load_workbook(p)                       # an openpyxl re-save drops the cache
        wb.save(p)
        wb.close()
        eq(_a1(p), None, "cache dropped on resave")
        ok(xlsx_cache.restore(orig, p), "restore reports it patched")
        ok(_a1(p) == 3, "cache restored after the resave")


def test_restore_noop_without_caches():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "t.xlsx")
        wb = Workbook()
        wb.active["A1"] = "plain text"
        wb.save(p)
        wb.close()
        with open(p, "rb") as f:
            eq(xlsx_cache.restore(f.read(), p), False, "no formula caches -> no-op")


if __name__ == "__main__":
    raise SystemExit(run("xlsx_cache", [
        ("capture_reads_formula_cache", test_capture_reads_formula_cache),
        ("restore_round_trip", test_restore_round_trip),
        ("restore_noop_without_caches", test_restore_noop_without_caches),
    ]))
