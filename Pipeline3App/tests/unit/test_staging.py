"""M5 gate (data-independent): staging enrichment + the single IODatabase.csv.

Builds a synthetic I/O List (one input + one output) and a synthetic C&E (matrix row marked X
with its FLD; an AREA sheet listing the output with a numerazione_linea), then checks matrix_areas,
the ce_* FLD columns, numerazione_linea, the identity fields, and the IODatabase.csv contents.
"""
import os
import tempfile
from _harness import run, eq, ok
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as CI

from pipeline3.core import config
from pipeline3.domain import staging, identity


def _set(ws, col, row, val, text=False):
    c = ws.cell(row=row, column=CI(col), value=val)
    if text:
        c.data_type = "s"
    return c


def _make_iolist(path):
    wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
    # required headers must prefix-match column_map (else SystemExit); others optional
    for col, hdr in (("C", "Part No."), ("F", "ID"), ("G", "Bit"), ("O", "Functional unit"),
                     ("P", "Location"), ("Q", "Device"), ("AB", "Script Type"), ("AD", "Index")):
        _set(ws, col, 1, hdr)
    # an emergency-PB input (prefilled type so staging resolves it without the populator)
    _set(ws, "G", 2, "I20.0"); _set(ws, "K", 2, "EMERGENCY PUSH-BUTTON PRESSED"); _set(ws, "L", 2, "CH1")
    _set(ws, "O", 2, "=S1", text=True); _set(ws, "P", 2, "+MS1.CC1", text=True)
    _set(ws, "Q", 2, "-S67001", text=True); _set(ws, "AB", 2, "E1/2"); _set(ws, "AD", 2, "0001")
    # a contactor output (KQ -> DB-backed); diag_cabinet (AE) points at a DiagnosisBlocks row
    _set(ws, "G", 3, "Q0.0")
    _set(ws, "O", 3, "=S1", text=True); _set(ws, "P", 3, "+MS1.CC1", text=True)
    _set(ws, "Q", 3, "-Q67001", text=True); _set(ws, "AB", 3, "KQ"); _set(ws, "AD", 3, "0001")
    _set(ws, "AE", 3, "5")
    # DiagnosisBlocks: ID_Local 5 maps to a distinct ID_SWP 15 (swp_cabinet import)
    db = wb.create_sheet("DiagnosisBlocks")
    for col, hdr in (("A", "ID_Local"), ("B", "ID_SWP"), ("C", "Functional Unit"),
                     ("D", "Location"), ("E", "FullName"), ("F", "TemplateType")):
        _set(db, col, 1, hdr)
    _set(db, "A", 2, 5); _set(db, "B", 2, 15)
    _set(db, "C", 2, "=S1", text=True); _set(db, "D", 2, "+MS1.CC1", text=True)
    _set(db, "E", 2, "=S1+MS1.CC1", text=True); _set(db, "F", 2, 1)
    wb.save(path)


def _make_ce(path):
    wb = Workbook()
    m = wb.active; m.title = "CAUSE&EFFECT MATRIX"
    _set(m, "S", 2, "AREA 1")                              # area column header (row 2)
    _set(m, "F", 5, "I20.0")                               # input address (data row 5)
    _set(m, "J", 5, "=S1", text=True); _set(m, "K", 5, "+MS1.CC1", text=True)
    _set(m, "L", 5, "-S67001", text=True); _set(m, "S", 5, "X")
    a = wb.create_sheet("AREA 1")
    _set(a, "A", 3, "SIGLA"); _set(a, "B", 3, "NUMERAZIONE LINEA"); _set(a, "C", 3, "DIGITAL OUTPUT")
    _set(a, "A", 4, "concat-x", text=True); _set(a, "B", 4, "LINE-7", text=True); _set(a, "C", 4, "Q0.0")
    # CONCEPT sheet: one description per area on row 2 (positional); newline must collapse to space
    cc = wb.create_sheet("CONCEPT")
    _set(cc, "C", 2, "AREA 1\n(Sorter)")
    wb.save(path)


def _params(d):
    io = os.path.join(d, "io.xlsx"); ce = os.path.join(d, "ce.xlsx")
    _make_iolist(io); _make_ce(ce)
    p = config.load_params()
    p["io_list"] = {"path": io, "sheet": "NET SAFETY 50", "header_row": 1}
    p["ce"] = {"path": ce, "matrix_sheet": "CAUSE&EFFECT MATRIX", "matrix_header_row": 2,
               "matrix_data_row": 5, "area_header_row": 3, "area_data_row": 4,
               "concept_sheet": "CONCEPT", "concept_row": 2}
    p["areas"] = r"AREA\s?\d{1,2}"
    p["sorter_areas"] = ["AREA 1"]
    p["output_dir"] = os.path.join(d, "out")
    return p, io


def _by_addr(rows, addr):
    return next(r for r in rows if r.get("bit") == addr)


def _run(fn):
    with tempfile.TemporaryDirectory() as d:
        p, io = _params(d)
        rows, warnings, matched, skipped = staging.load_io_list(p, config.load_signal_types(), io)
        fn(p, rows)


def test_input_matrix_areas_and_ce_fld():
    def check(p, rows):
        e1 = _by_addr(rows, "I20.0")
        eq(e1["matrix_areas"], "AREA 1", "input area from the matrix X")
        eq(e1["ce_functional_unit"], "=S1")
        eq(e1["ce_location"], "+MS1.CC1")
        eq(e1["ce_device"], "-S67001")
        eq(e1["IsSorterArea"], "yes", "AREA 1 is a sorter area")
        eq(e1["areas_description"], "AREA 1 (Sorter)", "CONCEPT desc, newline->space")
    _run(check)


def test_output_areas_and_numerazione():
    def check(p, rows):
        kq = _by_addr(rows, "Q0.0")
        eq(kq["matrix_areas"], "AREA 1", "output area from the AREA sheet listing")
        eq(kq["numerazione_linea"], "LINE-7", "numerazione_linea from AREA col B")
    _run(check)


def test_identity_fields():
    def check(p, rows):
        e1 = _by_addr(rows, "I20.0")
        ok(e1["name_in_tagtable"], "E1/2 is a taggable I/O point")
        ok(e1["tagtable"], "has a tag-table")
        kq = _by_addr(rows, "Q0.0")
        ok(kq["name_in_db"], "KQ is DB-backed -> has a DB member name")
        ok(kq["datablocks"], "KQ lists its DB(s)")
    _run(check)


def test_combined_fld_dedup():
    # the helper: equal iol/ce -> collapse; differ -> keep both (space-joined); no ce -> just iol
    base = {"functional_unit": "=S1", "location": "+A", "device": "-B"}
    eq(identity.combined_fld({**base, "ce_functional_unit": "=S1", "ce_location": "+A",
                              "ce_device": "-B"}), "=S1+A-B", "iol==ce -> collapsed")
    eq(identity.combined_fld({**base, "ce_functional_unit": "=S2", "ce_location": "+C",
                              "ce_device": "-D"}), "=S1+A-B =S2+C-D", "iol!=ce -> both")
    eq(identity.combined_fld(base), "=S1+A-B", "no ce side -> just iol_FLD")

    # staged: the synthetic E1/2 has iol==ce, so name_in_db collapses to a single FLD (no doubling)
    def check(p, rows):
        e1 = _by_addr(rows, "I20.0")
        eq(e1["iol_FLD"], "=S1+MS1.CC1-S67001")
        eq(e1["ce_FLD"], "=S1+MS1.CC1-S67001")
        eq(e1["combined_FLD"], "=S1+MS1.CC1-S67001", "collapsed")
        eq(e1["name_in_db"], "Emergency Push Button [ =S1+MS1.CC1-S67001 ]", "deduped single FLD")
    _run(check)


def test_interface_tagname_staged():
    def check(p, rows):
        e1 = _by_addr(rows, "I20.0")
        ok(e1["interface_tagname"].startswith("PNC_Q_"), "interface_tagname computed at staging")
        ok("Emergency Push Button" in e1["interface_tagname"], "nested {tag_name} resolved")
    _run(check)


def test_interface_tagname_resolves_db_element():
    """{db_element} (a signal-type field, itself a {combined_FLD} template) must be RESOLVED in the
    interface tag name - while {interface_name}/{interface_id} are KEPT for the generator to fill."""
    row = {"combined_FLD": "=S1+MA1-S8",
           "_type": {"interface_tagname": "PNC_Q_{interface_name}-{interface_id}_{db_element}",
                     "db_element": "Door SAFE_STATE [ {combined_FLD} ]", "tag_name": ""}}
    eq(identity.interface_tagname(row),
       "PNC_Q_{interface_name}-{interface_id}_Door SAFE_STATE [ =S1+MA1-S8 ]",
       "{db_element} is resolved; {interface_name}/{interface_id} are kept")


def test_swp_cabinet_from_diagnosis_blocks():
    def check(p, rows):
        kq = _by_addr(rows, "Q0.0")
        eq(kq["diag_cabinet"], "5", "local cabinet (ID_Local) on the row")
        eq(kq["swp_cabinet"], "15", "swp_cabinet = the cabinet's ID_SWP (imported, may differ)")
        eq(kq["diag_block_name"], "=S1+MS1.CC1", "diag_block_name still resolves (via ID_Local)")
        eq(kq["diag_block_template"], "1")
        e1 = _by_addr(rows, "I20.0")
        eq(e1["swp_cabinet"], "", "a row with no diag cabinet -> blank swp_cabinet")
    _run(check)


def test_iodatabase_csv():
    with tempfile.TemporaryDirectory() as d:
        p, io = _params(d)
        rows, *_ = staging.load_io_list(p, config.load_signal_types(), io)
        path = staging.write_io_database(rows, config.output_root(p))
        ok(os.path.exists(path), "IODatabase.csv written")
        from pipeline3.io.csv_tables import read_rows
        recs = read_rows(path)
        eq(len(recs), 2)
        hdr = recs[0].keys()
        for col in ("ce_functional_unit", "ce_location", "ce_device", "numerazione_linea",
                    "iol_FLD", "ce_FLD", "combined_FLD",
                    "matrix_areas", "areas_description", "name_in_db", "type_id_resolved"):
            ok(col in hdr, f"IODatabase has column {col}")
        e1 = next(r for r in recs if r["bit"] == "I20.0")
        eq(e1["ce_device"], "-S67001")
        eq(e1["type_id_resolved"], "E1/2")


def test_node_address_ranges_positional():
    # a node's I/Q byte range = the min/max byte address of the rows BELOW it (by position) until the
    # next node or sheet end - NOT keyed by FLD/location (a safety module's stops carry their own +ES..).
    rows = [
        {"_source_sheet": "S1", "profinet_name": "nodeA", "location": "+A", "bit": ""},
        {"_source_sheet": "S1", "location": "+ES1", "bit": "I900.0"},     # a DIFFERENT location, still nodeA's
        {"_source_sheet": "S1", "location": "+ES2", "bit": "I901.3"},
        {"_source_sheet": "S1", "location": "+A", "bit": "Q10.0"},
        {"_source_sheet": "S1", "profinet_name": "nodeB", "location": "+B", "bit": ""},
        {"_source_sheet": "S1", "location": "+C", "bit": "I50.0"},
        {"_source_sheet": "S2", "location": "+D", "bit": "I999.0"},       # new sheet -> NOT nodeB's
    ]
    staging._add_node_address_ranges(rows)
    a, b = rows[0], rows[4]
    eq((a["I_startByte"], a["I_endByte"]), (900, 901), "nodeA I-range from its positional rows (diff location)")
    eq((a["Q_startByte"], a["Q_endByte"]), (10, 10), "nodeA Q-range")
    eq((b["I_startByte"], b["I_endByte"]), (50, 50), "nodeB span ends at the sheet boundary (I999 excluded)")
    eq(b["Q_startByte"], "", "nodeB has no Q signals")
    eq(rows[1]["I_startByte"], "", "a signal row carries no range of its own")
    eq(rows[6]["I_startByte"], "", "a row under no node (after a sheet break) is unranged")


if __name__ == "__main__":
    raise SystemExit(run("staging", [
        ("node_address_ranges_positional", test_node_address_ranges_positional),
        ("input_matrix_areas_and_ce_fld", test_input_matrix_areas_and_ce_fld),
        ("output_areas_and_numerazione", test_output_areas_and_numerazione),
        ("identity_fields", test_identity_fields),
        ("combined_fld_dedup", test_combined_fld_dedup),
        ("interface_tagname_staged", test_interface_tagname_staged),
        ("interface_tagname_resolves_db_element", test_interface_tagname_resolves_db_element),
        ("swp_cabinet_from_diagnosis_blocks", test_swp_cabinet_from_diagnosis_blocks),
        ("iodatabase_csv", test_iodatabase_csv),
    ]))
