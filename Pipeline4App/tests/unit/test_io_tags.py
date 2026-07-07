"""Phase 510 - the PLCTags.xlsx projection (domain.io_tags): the two SSOT sources (resolved I/O signals
+ the interface_elements), the data-type map, the %-address, the text-forcing, the two-sheet structure.
Hermetic: a synthetic Database in-memory, projected to a temp dir, re-opened + asserted."""
import os
import tempfile

from openpyxl import load_workbook

from _harness import run, eq, ok
from pipeline4.core.database import Database
from pipeline4.domain import io_tags, interfaces
from pipeline4.domain.signals import signals_table

_IOLIST = ["functional_unit", "location", "device", "script_type", "bit"]


def _db(signals=(), elements=()):
    sig = signals_table(_IOLIST)
    for s in signals:
        sig.add(**s)
    el = interfaces.interface_elements_table()
    for e in elements:
        el.add(**e)
    return Database([sig, el])


def _project(signals=(), elements=()):
    db = _db(signals, elements)
    with tempfile.TemporaryDirectory() as d:
        res = io_tags.project(db, out_dir=d)
        wb = load_workbook(res["path"])
        tags = wb["PLC Tags"]
        rows = [[c.value for c in tags[r]] for r in range(1, tags.max_row + 1)]
        props = wb["TagTable Properties"]
        prows = [[c.value for c in props[r]] for r in range(1, props.max_row + 1)]
        wb.close()
    return res, rows, prows


def test_tia_dtype_and_logical_address():
    eq(io_tags._tia_dtype("bool"), "Bool")
    eq(io_tags._tia_dtype("WORD"), "Word")
    eq(io_tags._tia_dtype(""), "Bool", "blank -> Bool default")
    eq(io_tags._tia_dtype("foo"), "Foo", "unknown type capitalizes")
    eq(io_tags._logical_address("I1.0"), "%I1.0")
    eq(io_tags._logical_address("%Q10000.0"), "%Q10000.0", "already %-prefixed -> unchanged")
    eq(io_tags._logical_address(""), "")
    eq(io_tags._logical_address(None), "")


def test_two_sheet_structure_and_headers():
    _res, rows, prows = _project()
    eq(rows[0], io_tags.TAG_COLUMNS, "PLC Tags header = the 10 columns in order")
    eq(prows[0], io_tags.PROP_COLUMNS, "TagTable Properties header = the 3 columns")


def test_source_a_direct_io_signal():
    sig = {"script_type": "DI1/2", "bit": "I1.0", "name_in_tagtable": "PB START",
           "tagtable": "EMERGENCY_PushButtons",
           "type": {"type_id": "DI1/2", "category": "Safety", "io_comment": "[ {$script_type} ]"}}
    _res, rows, _p = _project(signals=[sig])
    eq(rows[1], ["PB START", "EMERGENCY_PushButtons", "Bool", "%I1.0", "[ DI1/2 ]",
                 "True", "True", "True", None, None],
       "the direct-I/O tag row (Bool, %-addr, resolved comment; empty Typeobject/Version read as None)")


def test_source_a_skips_blank_and_non_io():
    sigs = [
        {"script_type": "DI1/2", "bit": "I1.0", "name_in_tagtable": "",            # blank name -> skip
         "type": {"category": "Safety"}},
        {"script_type": "IOC", "bit": "I2.0", "name_in_tagtable": "X",             # Interface category -> not is_io
         "type": {"category": "Interface"}},
        {"script_type": "STD", "bit": "", "name_in_tagtable": "Y",                 # no I/Q bit -> not is_io
         "type": {"category": "Std"}},
    ]
    res, rows, _p = _project(signals=sigs)
    eq(res["io_count"], 0, "none of the three is a taggable I/O signal")
    eq(len(rows), 1, "only the header row")


def test_source_b_interface_tag_from_ssot():
    el = {"interface": "SORTER-01", "signal_name": "PNC_Q_Breaker", "data_type": "BOOL",
          "direction": "Q", "io_address_side1": "Q10000.0", "description": "breaker"}
    res, rows, prows = _project(elements=[el])
    eq(res["iface_count"], 1)
    eq(rows[1], ["PNC_Q_Breaker", "IF_SORTER-01", "Bool", "%Q10000.0", "breaker",
                 "True", "True", "True", None, None], "interface tag: Path=IF_<inst>, stored addr %-prefixed")
    eq(prows[1], ["IF_SORTER-01", None, None], "the IF_ table appears in TagTable Properties")


def test_source_b_word_and_unresolved():
    els = [
        {"interface": "S-01", "signal_name": "SPD", "data_type": "WORD", "direction": "Q",
         "io_address_side1": "Q10026", "description": "speed"},
        {"interface": "S-01", "signal_name": "NOADDR", "data_type": "BOOL", "direction": "Q",
         "io_address_side1": "", "description": "missing"},            # no resolved address -> skip + warn
    ]
    res, rows, _p = _project(elements=els)
    eq(res["iface_count"], 1, "the WORD tag is emitted; the address-less one is skipped")
    eq(rows[1][2], "Word", "WORD -> Word")
    eq(rows[1][3], "%Q10026", "WORD address has no bit")
    ok(any(f.type == "iotag_no_address" and "NOADDR" in f.location and f.severity == "WARN"
           for f in res["findings"]),
       "the unresolved-address element is a iotag_no_address WARN finding")


def test_text_forcing_and_sort_and_props():
    # two tables (sorted), a Name starting with '=' must stay TEXT (data_type 's'), distinct props
    sigs = [
        {"script_type": "ZZ", "bit": "Q5.0", "name_in_tagtable": "=S1+MS1-F1", "tagtable": "Zlast",
         "type": {"category": "Std", "io_comment": ""}},
        {"script_type": "AA", "bit": "I0.0", "name_in_tagtable": "alpha", "tagtable": "Afirst",
         "type": {"category": "Std", "io_comment": ""}},
        {"script_type": "AA", "bit": "I0.1", "name_in_tagtable": "beta", "tagtable": "Afirst",
         "type": {"category": "Std", "io_comment": ""}},
    ]
    res, rows, prows = _project(signals=sigs)
    eq([r[1] for r in rows[1:]], ["Afirst", "Afirst", "Zlast"], "rows grouped + sorted by Path")
    eq([r[0] for r in rows[1:]], ["alpha", "beta", "=S1+MS1-F1"], "insertion order within a table")
    eq([p[0] for p in prows[1:]], ["Afirst", "Zlast"], "distinct tables, sorted")
    with tempfile.TemporaryDirectory() as d:
        path = io_tags.project(_db(signals=sigs), out_dir=d)["path"]
        ws = load_workbook(path)["PLC Tags"]
        # the '=' Name is on the Zlast row (row 4); it must be stored as a string, not a formula
        cell = next(c for row in ws.iter_rows() for c in row if c.value == "=S1+MS1-F1")
        eq(cell.data_type, "s", "a leading '=' Name is forced to text, not a formula")


def test_return_contract_mixes_sources():
    sigs = [{"script_type": "DI1/2", "bit": "I1.0", "name_in_tagtable": "PB",
             "type": {"category": "Safety", "io_comment": ""}}]
    els = [{"interface": "S-01", "signal_name": "PNC", "data_type": "BOOL", "direction": "Q",
            "io_address_side1": "Q10000.0", "description": ""}]
    res, _rows, _p = _project(signals=sigs, elements=els)
    eq((res["io_count"], res["iface_count"], res["total"]), (1, 1, 2), "io + iface == total")
    ok(res["path"].endswith("PLCTags.xlsx"), "writes PLCTags.xlsx")


def test_config_tags_from_tagtable_elements():
    # source (c): the tagtable_elements rules over the signals - the datablock_elements logic aimed
    # at tag tables, io_address as a FULL expression (regex_replace - the user's flagship case)
    from pipeline4.core import config
    sigs = [
        {"script_type": "PA", "bit": "I13.5", "profinet_name": "n0005", "name_in_tagtable": "",
         "source_cell": "NS!O5", "type": {"category": "Std"}},
        {"script_type": "A", "bit": "I2.0", "profinet_name": "", "name_in_tagtable": "",
         "source_cell": "NS!O9", "type": {"category": "Std"}},
    ]
    db = _db(signals=sigs)
    rules = [
        {"tag_table": "PN_Mirror", "name": "MIRROR {$profinet_name}",
         "for_each": "row where $script_type = 'PA'", "datatype": "bool",
         "io_address": "{regex_replace($bit, /^I/, 'Q')}", "comment": "node {$profinet_name}"},
        {"tag_table": "Bad", "name": "{$x", "for_each": "nonsense(", "datatype": "", "io_address": "",
         "comment": ""},
    ]
    orig = config.load_tagtable_elements
    config.load_tagtable_elements = lambda: rules
    try:
        tags, findings = io_tags.config_tags(db)
    finally:
        config.load_tagtable_elements = orig
    eq(len(tags), 1, "one tag per for_each match (the PA row)")
    t = tags[0]
    eq((t["name"], t["path"], t["data_type"]), ("MIRROR n0005", "PN_Mirror", "Bool"))
    eq(t["address"], "%Q13.5", "io_address through the engine (regex_replace) + the % prefix")
    eq(t["comment"], "node n0005")
    eq((t["location"], t["source_uid"]), ("NS!O5", db["signals"].rows[0]["uid"]),
       "the producing I/O-List row rides along (the duplicate gate links it)")
    eq([f.type for f in findings], ["iotag_cfg_for_each"], "a broken for_each is a FAIL finding")
    eq(findings[0].severity, "FAIL")


def test_config_tags_render_error_fails_and_blocks_write():
    from pipeline4.core import config
    sigs = [{"script_type": "PA", "bit": "I1.0", "name_in_tagtable": "", "type": {"category": "Std"}}]
    db = _db(signals=sigs)
    rules = [{"tag_table": "T", "name": "{$missing_column}", "for_each": "row",
              "datatype": "Bool", "io_address": "", "comment": ""}]
    orig_rules, orig_params, orig_dbdir = (config.load_tagtable_elements, config.load_params,
                                           config.database_dir)
    with tempfile.TemporaryDirectory() as d:
        config.load_tagtable_elements = lambda: rules
        config.load_params = lambda *a, **k: {"iolist_path": os.path.join(d, "IO.xlsx")}
        config.database_dir = lambda: d
        try:
            res = io_tags.project(db, out_dir=d)
        finally:
            config.load_tagtable_elements = orig_rules
            config.load_params = orig_params
            config.database_dir = orig_dbdir
        eq(res["path"], "", "a config render FAIL blocks the write (raw-FAIL guard)")
        ok(any(f.type == "iotag_cfg_render" and f.severity == "FAIL" for f in res["findings"]),
           "strict render: the missing column is a located FAIL")
        ok(not os.path.exists(os.path.join(d, io_tags.TAG_TABLE_FILE)), "PLCTags.xlsx NOT on disk")


def test_duplicate_same_table_fails_links_rows_and_blocks_write():
    # the FVX_PL4_Pilot defect: I/O-List rows duplicated verbatim -> the same (tag table, name) twice.
    # Case-insensitive; each 2nd+ occurrence FAILs, location = ITS row, location2 = the FIRST row's.
    from pipeline4.core import config
    sigs = [
        {"script_type": "A", "bit": "I13.0", "name_in_tagtable": "Fire Alarm", "tagtable": "Alarms",
         "source_cell": "NET SAFETY 50!O113", "type": {"category": "Std", "io_comment": ""}},
        {"script_type": "A", "bit": "I14.0", "name_in_tagtable": "FIRE ALARM", "tagtable": "Alarms",
         "source_cell": "NET SAFETY 50!O121", "type": {"category": "Std", "io_comment": ""}},
        {"script_type": "A", "bit": "I15.0", "name_in_tagtable": "fire alarm", "tagtable": "Alarms",
         "source_cell": "NET SAFETY 50!O122", "type": {"category": "Std", "io_comment": ""}},
    ]
    db = _db(signals=sigs)
    orig_params, orig_dbdir = config.load_params, config.database_dir
    with tempfile.TemporaryDirectory() as d:
        config.load_params = lambda *a, **k: {"iolist_path": os.path.join(d, "IOList.xlsx")}
        config.database_dir = lambda: d
        try:
            res = io_tags.project(db, out_dir=d)
        finally:
            config.load_params, config.database_dir = orig_params, orig_dbdir
        eq(res["path"], "", "a duplicate tag blocks the write (raw-FAIL guard)")
        ok(not os.path.exists(os.path.join(d, io_tags.TAG_TABLE_FILE)), "PLCTags.xlsx is NOT on disk")
        dups = [f for f in res["findings"] if f.type == "iotag_duplicate"]
        eq(len(dups), 2, "one FAIL per 2nd+ occurrence (the first row is not flagged)")
        eq({f.severity for f in dups}, {"FAIL"})
        eq(dups[0].location, "NET SAFETY 50!O121", "location = the duplicate's own I/O-List row")
        eq(dups[1].location, "NET SAFETY 50!O122")
        eq({f.location2 for f in dups}, {"NET SAFETY 50!O113"}, "location2 = the FIRST occurrence's row")
        ok(dups[0].detail.endswith("- Count: 3"), "the detail carries the TOTAL occurrence count")
        ok(dups[1].detail.endswith("- Count: 3"), "…on every 2nd+ finding of the same tag")
        ok("in table 'Alarms'" in dups[0].detail)
        eq((dups[0].doc, dups[0].doc2), ("IOList.xlsx", "IOList.xlsx"),
           "both links carry the I/O List basename (the GUI's clickable-cell resolution)")
        eq(dups[0].source_uid, db["signals"].rows[1]["uid"], "FK to the duplicate's producing signal")
        with open(os.path.join(d, "validation_issues.csv"), encoding="utf-8-sig") as fh:
            recorded = [line for line in fh if "iotag_duplicate" in line]
        eq(len(recorded), 2, "the FAILs are recorded (the [FAIL]->Findings jump)")


def test_duplicate_name_across_tables_is_by_design():
    # one signal mirrors into SEVERAL IF_ tables (SORTER-01 + SORTER+DIAG-02) - never a duplicate.
    els = [
        {"interface": "SORTER-01", "signal_name": "PNC_I_Open Door", "data_type": "BOOL",
         "direction": "I", "io_address_side1": "I10010.0", "description": ""},
        {"interface": "SORTER+DIAG-02", "signal_name": "PNC_I_Open Door", "data_type": "BOOL",
         "direction": "I", "io_address_side1": "I20010.0", "description": ""},
    ]
    res, rows, _p = _project(elements=els)
    eq([f for f in res["findings"] if f.type == "iotag_duplicate"], [],
       "the same name on two tables is not flagged")
    eq(len(rows), 3, "both tags written")


def test_duplicate_interface_tag_links_the_source_signal_row():
    # a duplicate WITHIN one interface: the mirror element's link follows source_signal back to the
    # producing I/O-List row; an element with no source signal falls back to <interface>/<name>.
    from pipeline4.core import config
    sig = {"script_type": "IOC", "bit": "I2.0", "name_in_tagtable": "", "source_cell": "IO!O44",
           "type": {"category": "Interface"}}
    db = _db(signals=[sig])
    sig_uid = db["signals"].rows[0]["uid"]
    els = [
        {"interface": "S-01", "signal_name": "PNC_Q_Alarm", "data_type": "BOOL", "direction": "Q",
         "io_address_side1": "Q10000.0", "description": "", "source": "template"},
        {"interface": "S-01", "signal_name": "pnc_q_alarm", "data_type": "BOOL", "direction": "Q",
         "io_address_side1": "Q10000.1", "description": "", "source_signal": sig_uid},
    ]
    for e in els:
        db["interface_elements"].add(**e)
    orig_params, orig_dbdir = config.load_params, config.database_dir
    with tempfile.TemporaryDirectory() as d:
        config.load_params = lambda *a, **k: {"iolist_path": os.path.join(d, "IOList.xlsx")}
        config.database_dir = lambda: d
        try:
            res = io_tags.project(db, out_dir=d)
        finally:
            config.load_params, config.database_dir = orig_params, orig_dbdir
    dups = [f for f in res["findings"] if f.type == "iotag_duplicate"]
    eq(len(dups), 1)
    eq(dups[0].location, "IO!O44", "the mirror duplicate links its source signal's I/O-List row")
    eq(dups[0].location2, "S-01/PNC_Q_Alarm", "the template-native first occurrence: logical locator")
    ok(dups[0].detail.endswith("- Count: 2"), "a pair counts 2")
    eq((dups[0].doc, dups[0].doc2), ("IOList.xlsx", ""), "doc only on a Sheet!Cell location")


if __name__ == "__main__":
    import sys
    sys.exit(run("io_tags", [
        ("tia_dtype_and_logical_address", test_tia_dtype_and_logical_address),
        ("two_sheet_structure_and_headers", test_two_sheet_structure_and_headers),
        ("source_a_direct_io_signal", test_source_a_direct_io_signal),
        ("source_a_skips_blank_and_non_io", test_source_a_skips_blank_and_non_io),
        ("source_b_interface_tag_from_ssot", test_source_b_interface_tag_from_ssot),
        ("source_b_word_and_unresolved", test_source_b_word_and_unresolved),
        ("text_forcing_and_sort_and_props", test_text_forcing_and_sort_and_props),
        ("return_contract_mixes_sources", test_return_contract_mixes_sources),
        ("config_tags_from_tagtable_elements", test_config_tags_from_tagtable_elements),
        ("config_tags_render_error_fails_and_blocks_write",
         test_config_tags_render_error_fails_and_blocks_write),
        ("duplicate_same_table_fails_links_rows_and_blocks_write",
         test_duplicate_same_table_fails_links_rows_and_blocks_write),
        ("duplicate_name_across_tables_is_by_design", test_duplicate_name_across_tables_is_by_design),
        ("duplicate_interface_tag_links_the_source_signal_row",
         test_duplicate_interface_tag_links_the_source_signal_row),
    ]))
