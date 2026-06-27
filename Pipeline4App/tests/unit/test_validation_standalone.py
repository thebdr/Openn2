"""Phase 100b - the two standalone validators: 110 (I/O List) + 120 (C&E Matrix). Hermetic via a StubView
+ monkeypatched workbook openers (the validator logic is the unit under test, not the openpyxl read which
io/workbook already covers). The findings are checked by (type, severity)."""
import re

from openpyxl.utils import column_index_from_string as _ci, get_column_letter

from _harness import run, eq, ok
from pipeline4.core import config
from pipeline4.domain.validation import iolist, matrix


# --- a positional stub matching the SheetView surface the validators use ------------------------ #
class StubView:
    def __init__(self, name, colmap, rows, header_overrides=None, struck=None, extra_cols=0):
        self.name = name
        self.doc = name + ".xlsx"
        self._headers = {m["column"]: m["expected_header"] for m in colmap if m["expected_header"]}
        self._headers.update(header_overrides or {})
        self._rows = rows                                   # {row_int: {col_letter: value}}
        self._struck = struck or set()
        self.max_col = max((_ci(c) for c in self._headers), default=1) + extra_cols
        self.first_data_row = 2

    def _letter(self, c):
        return c if isinstance(c, str) else get_column_letter(c)

    def cell(self, r, c):
        return self._rows.get(r, {}).get(self._letter(c))

    def text(self, r, c):
        v = self.cell(r, c)
        return "" if v is None else str(v).strip()

    def header(self, c):
        return re.sub(r"\s+", " ", str(self._headers.get(self._letter(c), "")).replace("\n", " ")).strip()

    def header_matches(self, c, expected):
        exp = re.sub(r"\s+", " ", str(expected).replace("\n", " ")).strip().lower()
        return self.header(c).lower().startswith(exp)

    def cell_struck(self, r, c):
        return r in self._struck

    def row_struck(self, r, columns=None):
        return r in self._struck

    def data_rows(self):
        return sorted(self._rows)

    def location(self, r, c):
        return f"{self.name}!{self._letter(c)}{r}"

    def close(self):
        pass


_IOCOL = {m["canonical"]: m["column"] for m in config.load_column_map("IoList")}


def _iorow(**canon):
    """A row dict keyed by COLUMN LETTER from canonical kwargs (so tests read naturally)."""
    return {_IOCOL[k]: v for k, v in canon.items()}


def _run_iolist(views, params=None, available=None):
    """Run run_iolist with the workbook openers monkeypatched to return `views`."""
    p = {"iolist_path": "X.xlsx", "iolist_params": {"sheets": "IO", "header_row": 1},
         "validation_params": {"global": {"strike_handling": "ignore"},
                               "iolist": {"permanent_parts": ["VALIDPART"]}}}
    if params:
        p.update(params)
    orig_os, orig_ss, orig_as = iolist.os.path.exists, iolist.wbk.open_sheets, iolist.wbk.available_sheets
    try:
        iolist.os.path.exists = lambda _p: True
        iolist.wbk.open_sheets = lambda *a, **k: list(views)
        iolist.wbk.available_sheets = lambda _p: available if available is not None else [v.name for v in views]
        return iolist.run_iolist(p)
    finally:
        iolist.os.path.exists, iolist.wbk.open_sheets, iolist.wbk.available_sheets = orig_os, orig_ss, orig_as


def _types(findings):
    return [(f.type, f.severity) for f in findings]


# --- 110 ---------------------------------------------------------------------------------------- #
def test_iolist_missing():
    eq(_types(iolist.run_iolist({"iolist_path": ""})), [("iolist_missing", "FAIL")])


def test_iolist_clean_row_passes():
    v = StubView("IO", config.load_column_map("IoList"),
                 {2: _iorow(functional_unit="=S1", bit="I0.0", type_hw="KI")})
    found = _types(_run_iolist([v]))
    ok(("row_ok", "PASS") in found, "a clean row -> row_ok PASS")
    ok(not any(sev == "FAIL" for _t, sev in found), "no FAILs on a clean row")
    ok(("proc_sheet", "INFO") in found and ("iolist_summary", "INFO") in found)


def test_iolist_addr_format_and_fg():
    v = StubView("IO", config.load_column_map("IoList"), {
        2: _iorow(functional_unit="=S1", bit="I:0.0"),                 # Rockwell colon -> addr_format
        3: _iorow(functional_unit="=S2", bit="I0.0", id_node="5"),     # G and F both present -> fg_exclusion
    })
    found = _types(_run_iolist([v]))
    ok(("addr_format", "FAIL") in found and ("fg_exclusion", "FAIL") in found)


def test_iolist_dup_ip_and_fld():
    cm = config.load_column_map("IoList")
    v = StubView("IO", cm, {
        2: _iorow(functional_unit="=S1", location="+L1", device="-D1", profinet_ip="10.0.0.2", profinet_name="n1"),
        3: _iorow(functional_unit="=S1", location="+L1", device="-D1", profinet_ip="10.0.0.2", profinet_name="n2"),
    })
    found = _types(_run_iolist([v]))
    ok(("ip_duplicated", "FAIL") in found, "same non-.1 IP twice -> ip_duplicated")
    ok(("dup_fld", "FAIL") in found, "same FU+LOC+DEV twice -> dup_fld")


def test_iolist_permanent_part_and_ts():
    cm = config.load_column_map("IoList")
    v = StubView("IO", cm, {
        2: _iorow(functional_unit="=S1", type_hw="A", desc_l1="NOTKNOWN", ts_ref="TS_1"),   # pp_unknown
        3: _iorow(functional_unit="=S2", type_hw="A", desc_l1="VALIDPART", ts_ref="x"),       # ts_missing
    })
    found = _types(_run_iolist([v]))
    ok(("pp_unknown", "FAIL") in found, "desc_l1 not in permanent_parts -> pp_unknown")
    ok(("ts_missing", "FAIL") in found, "A/W with no TS_ ref -> ts_missing")


def test_iolist_col_wrong_header():
    cm = config.load_column_map("IoList")
    bit_col = _IOCOL["bit"]
    v = StubView("IO", cm, {2: _iorow(functional_unit="=S1", bit="I0.0")},
                 header_overrides={bit_col: "WRONG HEADER"})
    ok(("col_wrong", "FAIL") in _types(_run_iolist([v])), "a wrong mapped header -> col_wrong")


# --- 120 ---------------------------------------------------------------------------------------- #
def _run_ce(matrix_view, area_views, params=None):
    p = {"matrix_path": "CE.xlsx", "matrix_params": {"ce_sheet": {"name": "MATRIX", "header_row": 2, "data_row": 5},
                                                     "area_sheets": {"name": "AREA", "header_row": 3, "data_row": 4}}}
    if params:
        p.update(params)
    orig_os, orig_o1, orig_os2 = matrix.os.path.exists, matrix.wbk.open_sheet, matrix.wbk.open_sheets
    try:
        matrix.os.path.exists = lambda _p: True
        matrix.wbk.open_sheet = lambda *a, **k: matrix_view
        matrix.wbk.open_sheets = lambda *a, **k: list(area_views)
        return matrix.run_ce_matrix(p)
    finally:
        matrix.os.path.exists, matrix.wbk.open_sheet, matrix.wbk.open_sheets = orig_os, orig_o1, orig_os2


def test_ce_absent_skips():
    eq(_types(matrix.run_ce_matrix({"matrix_path": ""})), [("ce_absent", "SKIP")])


def test_ce_matrix_inputs_only():
    ce = config.load_column_map("CE")
    addr_c = {m["canonical"]: m["column"] for m in ce}["address"]          # F
    mv = StubView("MATRIX", ce, {6: {addr_c: "Q0.0"}, 7: {addr_c: "I1.0"}})  # Q on the matrix -> matrix_addr_kind
    found = _types(_run_ce(mv, []))
    ok(("matrix_addr_kind", "FAIL") in found, "an output address on the CAUSE&EFFECT MATRIX is wrong")


def test_ce_area_outputs_only_and_dup():
    area = config.load_column_map("AREA")
    addr_c = {m["canonical"]: m["column"] for m in area}["address"]        # C
    av = StubView("AREA 1", area, {5: {addr_c: "I0.0"}, 6: {addr_c: "Q1.0"}, 7: {addr_c: "Q1.0"}})
    found = _types(_run_ce(None, [av]))
    ok(("area_addr_kind", "FAIL") in found, "an input address on an AREA sheet is wrong")
    ok(("dup_addr", "FAIL") in found, "a repeated address within a sheet -> dup_addr")


if __name__ == "__main__":
    import sys
    sys.exit(run("validation_standalone", [
        ("iolist_missing", test_iolist_missing),
        ("iolist_clean_row_passes", test_iolist_clean_row_passes),
        ("iolist_addr_format_and_fg", test_iolist_addr_format_and_fg),
        ("iolist_dup_ip_and_fld", test_iolist_dup_ip_and_fld),
        ("iolist_permanent_part_and_ts", test_iolist_permanent_part_and_ts),
        ("iolist_col_wrong_header", test_iolist_col_wrong_header),
        ("ce_absent_skips", test_ce_absent_skips),
        ("ce_matrix_inputs_only", test_ce_matrix_inputs_only),
        ("ce_area_outputs_only_and_dup", test_ce_area_outputs_only_and_dup),
    ]))
