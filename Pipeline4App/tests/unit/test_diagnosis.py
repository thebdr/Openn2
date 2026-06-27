"""Phase 600a - the diagnosis config relocation + staging prereqs:
- config.load_signal_diagnosis (the relocated per-type diag attrs + the new `tristate` flag),
- config.load_diagnosis_columns (the DiagList column model, with the :03d/:02d format specs restored),
- identity.interp honoring a {token:spec} format spec (the diagnosis columns rely on it),
- the diagnosis_entries / diagnosis_cabinets table defs,
- staging.load_diagnosis_blocks (the DiagnosisBlocks-sheet -> diagnosis_cabinets port, over a synthetic xlsx).
"""
import os
import tempfile

from openpyxl import Workbook

from _harness import run, eq, ok
from pipeline4.core import config
from pipeline4.domain import identity, staging
from pipeline4.domain.diagnosis_entries import diagnosis_entries_table, diagnosis_cabinets_table


def test_interp_format_spec():
    eq(identity.interp("{diag_cabinet:03d}", {"diag_cabinet": "0"}), "000", "0 -> 000")
    eq(identity.interp("{diag_bit:02d}", {"diag_bit": "62"}), "62", "62 -> 62 (already 2 digits)")
    eq(identity.interp("{diag_cabinet:03d}", {"diag_cabinet": "5"}), "005")
    eq(identity.interp("{diag_cabinet:03d}", {"diag_cabinet": "1.0"}), "001", "a float-read int coerces")
    eq(identity.interp("{diag_cabinet:03d}", {"diag_cabinet": ""}), "", "a blank value is NOT padded")
    eq(identity.interp("[ {x} {y} ]", {"x": "A", "y": "B"}), "[ A B ]", "bare tokens unchanged")
    eq(identity.interp("{name}", {"name": "x"}), "x", "no spec -> verbatim")


def test_load_signal_diagnosis():
    d = config.load_signal_diagnosis()
    ok("E1/2".upper() in d and "PA" in d, "the in_diag types are present")
    eq(d["E1/2"]["in_diag"], True)
    eq(d["E1/2"]["diag_logic"], "invert")
    eq(d["E1/2"]["tristate"], True, "E1/2 is the tristate type")
    eq(d["E1/2"]["tristate_desc"], "EMERGENCY PUSH-BUTTON RELEASED {combined_FLD}")
    eq(d["PA"]["diag_logic"], "mirror")
    eq(d["PA"]["diag_desc"], "FIELDBUS NODE FAILURE {profinet_name} {profinet_ip}")
    eq(d["B1/2"]["tristate"], False, "a non-tristate in_diag type")
    ok("Z#" in d and d["Z#"]["in_diag"] is True, "a pattern type carries its diag attrs (keyed by type_id)")


def test_load_diagnosis_columns():
    cols = config.load_diagnosis_columns()
    headers = [c["header"] for c in cols]
    eq(headers[0], "Diag Cabinet")
    eq(cols[0]["expression"], "{diag_cabinet:03d}", "the :03d spec is restored")
    eq(cols[1]["expression"], "{diag_bit:02d}", "the :02d spec is restored")
    ok(any(c["expression"].strip() == "$PLC_Binding$" for c in cols), "the PLC_Binding sentinel column")
    eq(len(cols), 19, "the 19 DiagList columns")


def test_table_defs():
    e = diagnosis_entries_table()
    eq(e.name, "diagnosis_entries")
    for c in ("source", "cabinet", "bit", "is_warning", "in_binding", "ml_value", "fl_value", "diag_columns"):
        ok(c in e.columns, f"{c} column present")
    ok("is_warning" in e.json_columns and "diag_columns" in e.json_columns, "the JSON cells")
    c = diagnosis_cabinets_table()
    eq(c.name, "diagnosis_cabinets")
    for col in ("cabinet_id", "index", "fld", "template_type", "swp"):
        ok(col in c.columns, f"{col} column present")
    eq(c.key_columns, ["cabinet_id"])


def _diagblocks_xlsx(path, rows, sheet="DiagnosisBlocks"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["ID_Local", "ID_SWP", "Functional Unit", "Location", "FullName", "TemplateType"])
    for r in rows:
        ws.append(r)
        rr = ws.max_row
        for c, v in enumerate(r, 1):                      # the real sheet stores FU/FullName as TEXT;
            if isinstance(v, str) and v[:1] in ("=", "+", "-"):   # force it (else openpyxl = a formula,
                ws.cell(rr, c).data_type = "s"                    # which a data_only read returns as None)
    wb.save(path)


def test_load_diagnosis_blocks():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "iol.xlsx")
        _diagblocks_xlsx(p, [
            [0, 0, "=S1", "+MC1", "=S1+MC1.CC1", "1"],
            [2, 5, "=S1", "+IN1", "=S1+IN1.CC1", "2"],     # ID_SWP differs from ID_Local
            ["", "", "", "", "orphan", ""],                # non-numeric key -> skipped
        ])
        cabs = staging.load_diagnosis_blocks(p)
        eq(set(cabs), {0, 2}, "keyed by ID_Local (int), the non-numeric row skipped")
        eq(cabs[0], {"index": "000", "fld": "=S1+MC1.CC1", "template_type": "1", "swp": "0"})
        eq(cabs[2]["template_type"], "2", "TemplateType -> the SCL 01-04 variant")
        eq(cabs[2]["swp"], "5", "ID_SWP kept when it differs from ID_Local")


def test_load_diagnosis_blocks_absent():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "no_sheet.xlsx")
        wb = Workbook(); wb.active.title = "NET SAFETY 50"; wb.active["A1"] = "x"; wb.save(p)
        eq(staging.load_diagnosis_blocks(p), {}, "no DiagnosisBlocks sheet -> {}")
        eq(staging.load_diagnosis_blocks(""), {}, "no path -> {}")


if __name__ == "__main__":
    import sys
    sys.exit(run("diagnosis", [
        ("interp_format_spec", test_interp_format_spec),
        ("load_signal_diagnosis", test_load_signal_diagnosis),
        ("load_diagnosis_columns", test_load_diagnosis_columns),
        ("table_defs", test_table_defs),
        ("load_diagnosis_blocks", test_load_diagnosis_blocks),
        ("load_diagnosis_blocks_absent", test_load_diagnosis_blocks_absent),
    ]))
