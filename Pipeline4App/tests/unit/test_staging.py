"""Staging (domain.identity + domain.staging): the FLD/tag identity + the IoList read into rows."""
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font

from _harness import run, eq, ok
from pipeline4.domain import identity, staging
from pipeline4.domain.signals import signals_table
from pipeline4.io import workbook


def test_identity_flds_and_tag():
    row = {"functional_unit": "S1", "location": "+SG1", "device": "-B1", "script_type": "DI1/2", "bit": "I0.0",
           "type": {"category": "Safety", "tag_name": "Door [ {combined_FLD} ]", "tagtable_name": "SAFETY_Doors"}}
    row["combined_FLD"] = identity.combined_fld(row)
    eq(identity.fld(row), "S1+SG1-B1", "iol_FLD = FU+LOC+DEV")
    eq(identity.combined_fld(row), "S1+SG1-B1", "combined == iol when there is no C&E FLD")
    ok(identity.is_io_signal(row), "an I/Q bit on a non-interface type is an I/O signal")
    eq(identity.tag_name(row), "Door [ S1+SG1-B1 ]", "tag_name template interpolated")
    eq(identity.tagtable(row), "SAFETY_Doors", "tagtable from the type")


def test_combined_fld_appends_ce_when_different():
    row = {"functional_unit": "S1", "location": "+SG1", "device": "-B1",
           "ce_functional_unit": "S1", "ce_location": "+CE", "ce_device": "-B1"}
    eq(identity.combined_fld(row), "S1+SG1-B1 S1+CE-B1", "the C&E-side FLD is appended only when it differs")


_COLMAP = [{"column": c, "canonical": n} for c, n in
           [("A", "functional_unit"), ("B", "location"), ("C", "device"),
            ("D", "script_type"), ("E", "bit"), ("F", "skip_reason")]]
_TYPES = {"DI1/2": {"type_id": "DI1/2", "category": "Safety",
                    "tag_name": "Door [ {combined_FLD} ]", "tagtable_name": "SAFETY_Doors"}}


def _iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    for col, head in zip("ABCDEF", ["FU", "Loc", "Dev", "Type", "Bit", "Skip"]):
        ws[f"{col}1"] = head
    ws["A2"] = "S1"; ws["B2"] = "+SG1"; ws["C2"] = "-B1"; ws["D2"] = "DI1/2"; ws["E2"] = "I0.0"
    ws["A3"] = "S2"; ws["B3"] = "+SG2"; ws["C3"] = "-B2"; ws["D3"] = "DI1/2"; ws["E3"] = "I0.1"; ws["F3"] = "skip me"
    ws["A4"] = "S3"; ws["B4"] = "+SG3"; ws["C4"] = "-B3"; ws["D4"] = "DI1/2"; ws["E4"] = "I0.2"
    ws["A4"].font = Font(strike=True)
    wb.save(path)


def test_read_view_drops_skip_and_struck():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx"); _iolist(p)
        v = workbook.open_sheet(p, "NET SAFETY 50", 1)
        rows = staging._read_view(v, _COLMAP, True, _TYPES)        # strike_exclude=True
        eq(len(rows), 1, "the Skip-Reason row + the struck row are dropped -> 1 kept")
        r = rows[0]
        eq(r["functional_unit"], "S1"); eq(r["script_type"], "DI1/2")
        eq(r["type"]["type_id"], "DI1/2", "the resolved type is attached")
        eq(r["source_sheet"], "NET SAFETY 50"); eq(r["source_row"], 2)
        v.close()


def test_read_view_keeps_struck_when_not_excluding():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx"); _iolist(p)
        v = workbook.open_sheet(p, "NET SAFETY 50", 1)
        rows = staging._read_view(v, _COLMAP, False, _TYPES)       # strike_exclude=False
        eq(len(rows), 2, "the struck row is KEPT when not excluding; the Skip-Reason row is still dropped")
        v.close()


def test_is_sorter_area_number_to_name():
    names = staging._sorter_area_names({"matrix_params": {"sorter_areas": [1]}})
    eq(names, {"AREA 1"}, "sorter number 1 -> area name 'AREA 1'")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 1"]}, names), "yes", "matches its sorter area")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 2", "AREA 1"]}, names), "yes", "yes when ANY area is a sorter")
    eq(staging._is_sorter_area({"matrix_areas": ["AREA 3"]}, names), "", "a non-sorter area -> ''")
    eq(staging._is_sorter_area({"matrix_areas": []}, names), "", "no areas -> ''")
    eq(staging._is_sorter_area({}, names), "", "missing matrix_areas -> ''")
    eq(staging._sorter_area_names({}), set(), "no sorter_areas configured -> no names")


def test_node_address_ranges_positional():
    rows = [
        {"source_sheet": "S1", "profinet_name": "n1", "bit": ""},      # node head (owns no address of its own)
        {"source_sheet": "S1", "bit": "I0.0"},
        {"source_sheet": "S1", "bit": "I20.7"},
        {"source_sheet": "S1", "bit": "Q0.0"},
        {"source_sheet": "S1", "profinet_name": "n2", "bit": "Q5.0"},  # next head opens a new span + owns its own Q5
        {"source_sheet": "S1", "bit": "I8.1"},
        {"source_sheet": "S2", "bit": "I3.0"},                          # past the sheet boundary -> owned by no node
    ]
    staging._add_node_address_ranges(rows)
    eq((rows[0]["I_startByte"], rows[0]["I_endByte"]), (0, 20), "n1 I range = min/max byte beneath it")
    eq((rows[0]["Q_startByte"], rows[0]["Q_endByte"]), (0, 0), "n1 owns the Q0 row in its span too")
    eq((rows[4]["I_startByte"], rows[4]["I_endByte"]), (8, 8), "n2 owns the I8 row beneath it on the same sheet")
    eq((rows[4]["Q_startByte"], rows[4]["Q_endByte"]), (5, 5), "n2 owns its own Q5 bit")
    eq(rows[1]["I_startByte"], "", "a non-node row NEVER carries a range")
    eq(rows[6]["I_startByte"], "", "the sheet boundary ends n2's span -> the S2 row is unowned")


def test_node_address_range_empty_when_no_addressed_rows():
    rows = [{"source_sheet": "S", "profinet_name": "n", "bit": ""}]
    staging._add_node_address_ranges(rows)
    eq(rows[0]["I_startByte"], "", "a node owning no addressed row -> empty range (like PL3's 7 empty nodes)")
    eq(rows[0]["Q_endByte"], "")


def test_dup_signal_uid_findings():
    """The post-stage duplicate-uid check (moved from the GUI into stage) is a stg_dup_signal_uid WARN,
    one per shared content-hash uid; the clean case emits nothing."""
    cols = ["combined_FLD", "script_type", "source_cell"]      # the stable key -> the content-hash uid
    t = signals_table(cols)
    for _ in range(2):                                          # two rows share the key -> the SAME uid
        t.add_row({"combined_FLD": "S1+SG1-B1", "script_type": "DI1/2", "source_cell": "NET!O2"})
    t.add_row({"combined_FLD": "S2+SG2-B2", "script_type": "DI1/2", "source_cell": "NET!O3"})   # unique
    fs = staging._dup_findings(t)
    eq(len(fs), 1, "one WARN per shared uid (the unique row produces none)")
    eq((fs[0].phase, fs[0].type, fs[0].severity), (300, "stg_dup_signal_uid", "WARN"), "the finding container")
    ok(fs[0].location in t.duplicate_uids(), "the finding locates the shared uid")
    eq(staging._dup_findings(signals_table(cols)), [], "no rows -> no findings (the clean path)")


def test_load_io_list_no_match_returns_empty():
    """When no sheet matches the pattern, load_io_list returns ([], []) - the predicate stage() turns into the
    blocking stg_no_io_sheet FAIL (replacing the old raise SystemExit). The FAIL emission + no-write is verified
    end-to-end by the data-dependent parity smoke; this pins the load-bearing no-match branch."""
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        wb = Workbook(); wb.active.title = "NotAnIoSheet"; wb.active["A1"] = "x"; wb.save(p)
        params = {"iolist_params": {"sheets": "^__no_match__$", "header_row": 1}}
        rows, matched = staging.load_io_list(params, {}, p)
        eq(matched, [], "no sheet matched the pattern -> empty match list")
        eq(rows, [], "no rows read")


if __name__ == "__main__":
    import sys
    sys.exit(run("staging", [
        ("identity_flds_and_tag", test_identity_flds_and_tag),
        ("combined_fld_appends_ce_when_different", test_combined_fld_appends_ce_when_different),
        ("read_view_drops_skip_and_struck", test_read_view_drops_skip_and_struck),
        ("read_view_keeps_struck_when_not_excluding", test_read_view_keeps_struck_when_not_excluding),
        ("is_sorter_area_number_to_name", test_is_sorter_area_number_to_name),
        ("node_address_ranges_positional", test_node_address_ranges_positional),
        ("node_address_range_empty_when_no_addressed_rows", test_node_address_range_empty_when_no_addressed_rows),
        ("dup_signal_uid_findings", test_dup_signal_uid_findings),
        ("load_io_list_no_match_returns_empty", test_load_io_list_no_match_returns_empty),
    ]))
