"""Phase 400c/400e - the IF_*.xlsx projection + the lossless I/O List insertion (domain.interface_xlsx)."""
import os
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table as XlTable
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo, TableFormula

from _harness import run, eq, ok
from pipeline4.domain import interface_xlsx
from pipeline4.io import xlsx_edit


def test_coerce_int_and_safe_name():
    eq(interface_xlsx._coerce_int("10"), 10)
    eq(interface_xlsx._coerce_int("10.0"), 10, "a float-read int string coerces")
    eq(interface_xlsx._coerce_int(""), None)
    eq(interface_xlsx._coerce_int("None"), None, "a stringified None -> None (the WORD bit)")
    eq(interface_xlsx._safe_name('SORTER+DIAG-02'), "SORTER+DIAG-02")
    eq(interface_xlsx._safe_name('A/B:C'), "A_B_C", "filesystem-unsafe chars replaced")


def _synthetic_sheet():
    wb = Workbook()
    ws = wb.active
    headers = ["Category", "Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit",
               "Signal Name Side 1", "Expression Side 1"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    ws.cell(2, 1, "TEMPLATE"); ws.cell(2, 3, ">"); ws.cell(2, 4, 0)     # one template data row
    ws.add_table(XlTable(displayName="T", ref="A1:G2"))
    return wb, ws


def _col(ws, header):
    return {str(ws.cell(1, c).value).strip(): c for c in range(1, ws.max_column + 1)}[header]


def test_append_custom_rows_bool_block_and_padding():
    wb, ws = _synthetic_sheet()
    elems = [{"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 0,
              "signal_name": "PNC_a", "expression": '"DB"."a"'},
             {"category": "X", "direction": "Q", "data_type": "BOOL", "offset_byte": 10, "bit": 1,
              "signal_name": "PNC_b", "expression": '"DB"."b"'}]
    interface_xlsx._append_custom_rows(ws, elems)
    cat, off, bit, s1, e1 = (_col(ws, "Category"), _col(ws, "I/O Offset Byte"), _col(ws, "I/O Bit"),
                             _col(ws, "Signal Name Side 1"), _col(ws, "Expression Side 1"))
    # the 2 real signals on rows 3-4 (low bits of byte 10), then 14 blank padding rows -> a full 2-byte block
    eq((ws.cell(3, cat).value, ws.cell(3, off).value, ws.cell(3, bit).value), ("X", 10, 0))
    eq((ws.cell(3, s1).value, ws.cell(3, e1).value), ("PNC_a", '"DB"."a"'))
    eq((ws.cell(4, off).value, ws.cell(4, bit).value), (10, 1), "second signal, bit 1")
    eq(ws.cell(5, e1).value in (None, ""), True, "row 5 is blank padding (addressed, no signal)")
    eq((ws.cell(18, cat).value, ws.cell(18, off).value, ws.cell(18, bit).value), ("X", 11, 7),
       "the block runs a full 16 bits (byte 10 + 11), last padding row at 11.7")
    eq(ws.tables["T"].ref, "A1:G18", "the table ref extends to the last written row")


def test_append_custom_rows_word_row_and_separator():
    wb, ws = _synthetic_sheet()
    elems = [{"category": "SPD", "direction": "Q", "data_type": "WORD", "offset_byte": 12, "bit": None,
              "signal_name": "PNC_speed", "expression": '"DB"."speed"'}]
    interface_xlsx._append_custom_rows(ws, elems)
    dt, off, bit = _col(ws, "Data Type"), _col(ws, "I/O Offset Byte"), _col(ws, "I/O Bit")
    eq((ws.cell(3, dt).value, ws.cell(3, off).value, ws.cell(3, bit).value), ("WORD", 12, None),
       "a WORD is one row, even byte, no bit")
    ok(ws.tables["T"].ref in ("A1:G3", "A1:G4"), "one word row (+ a trailing separator)")


# =================================================================================================== #
# Phase 400e - the lossless IF_ sheet insertion + the Excel-independent address-cache seeding
# =================================================================================================== #
def test_io_address_mirror():
    f = interface_xlsx._io_address
    isy = "I<base+offset>/.<bit>"
    eq(f(isy, 10000, ">", "BOOL", 0, 0), "Q10000.0", "BOOL output -> Q with bit")
    eq(f(isy, 10000, "<", "BOOL", 5, 3), "I10005.3", "BOOL input -> I, base+offset.bit")
    eq(f(isy, 10000, ">", "WORD", 26, None), "Q10026", "WORD -> byte only (TEXTBEFORE '/'), no bit")
    eq(f(isy, 10000, "", "BOOL", 0, 0), "", "no direction -> no address")
    eq(f(isy, None, ">", "BOOL", 0, 0), "", "no base -> no address")
    eq(f(isy, 10000, ">", "BOOL", None, 0), "", "no offset -> no address")


def test_resolve_num_chain():
    wb = Workbook(); ws = wb.active
    ws["C2"], ws["D2"] = 1, 0
    ws["C3"], ws["D3"] = "=C2", "=D2+1"        # the interface table's two offset/bit chain shapes
    ws["C4"], ws["D4"] = "=C3", "=D3+1"
    ws["C5"], ws["C6"] = "12.0", "x"
    memo = {}
    eq(interface_xlsx._resolve_num(ws, "C4", memo), 1, "=C3 -> =C2 -> 1")
    eq(interface_xlsx._resolve_num(ws, "D4", memo), 2, "=D3+1 -> (=D2+1)+1 -> 2")
    eq(interface_xlsx._resolve_num(ws, "C5", memo), 12, "float-string literal -> int")
    ok(interface_xlsx._resolve_num(ws, "C6", memo) is None, "non-numeric -> None")


def _make_if_addr_sheet(path, base=10000):
    """A minimal inserted-IF_ sheet: the address-relevant headers + an Input Format / Base Address on
    row 2, a (placeholder) LET-formula address cell per data row, and an offset/bit chain (=C3 / =D3+1)."""
    wb = Workbook(); ws = wb.active; ws.title = "IF_T"
    cols = ["Data Type", "Direction </>", "I/O Offset Byte", "I/O Bit", "I/O Address Side 1",
            "Base Address", "Input Format", "Signal Name Side 1"]     # A..H ; offset=C bit=D addr=E
    for i, h in enumerate(cols, start=1):
        ws.cell(1, i, h)
    LET = "=$A$1"                                                     # any formula -> data_type 'f'

    def setrow(r, dt, d, off, bit, sig):
        ws.cell(r, 1, dt); ws.cell(r, 2, d); ws.cell(r, 3, off); ws.cell(r, 4, bit)
        ws.cell(r, 5, LET); ws.cell(r, 8, sig)
    setrow(2, "BOOL", ">", 0, 0, "SIG1")            # Q<base>.0
    setrow(3, "BOOL", ">", 1, 0, "SIG2")            # Q<base+1>.0
    setrow(4, "BOOL", ">", "=C3", "=D3+1", "SIG3")  # CHAIN -> off 1, bit 1 -> Q<base+1>.1
    setrow(5, "WORD", ">", 4, None, "SIG4")         # WORD -> Q<base+4>, no bit
    setrow(6, "BOOL", "<", 0, 0, "SIG5")            # I<base>.0
    ws.cell(7, 5, LET)                              # blank separator: LET present, no direction -> ''
    ws.cell(2, 6, base); ws.cell(2, 7, "I<base+offset>/.<bit>")
    wb.save(path)


def test_interface_address_caches():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(p)
        cache = interface_xlsx._interface_address_caches(p)
        eq(cache.get("E2"), ("str", "Q10000.0"), "BOOL output bit 0")
        eq(cache.get("E3"), ("str", "Q10001.0"), "BOOL output byte+1")
        eq(cache.get("E4"), ("str", "Q10001.1"), "offset/bit CHAIN resolved (=C3 / =D3+1)")
        eq(cache.get("E5"), ("str", "Q10004"), "WORD: byte only, no bit")
        eq(cache.get("E6"), ("str", "I10000.0"), "BOOL input -> I")
        ok("E7" not in cache, "blank separator row (no direction) is not seeded")


def test_insert_seeds_interface_address_cache():
    # the phase-510 bug: the inserted IF_ address LET has no Excel cache, so a data_only reader sees
    # None and skips. insert SEEDS a computed cache that data_only sees (no Excel), keeping the formula.
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_T.xlsx"); _make_if_addr_sheet(ifp)
        iol = os.path.join(d, "iol.xlsx")
        wb = Workbook(); wb.active.title = "NET SAFETY 50"; wb.active["A1"] = "src"; wb.save(iol)
        acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ok(any("seeded" in a and "IF_SORTER-01" in a for a in acts), "seeding is reported")
        ws = load_workbook(iol, data_only=True)["IF_SORTER-01"]
        eq(ws["E2"].value, "Q10000.0", "data_only reader sees the seeded address (no Excel needed)")
        eq(ws["E4"].value, "Q10001.1", "the chained-offset address is seeded too")
        live = load_workbook(iol)["IF_SORTER-01"]["E2"].value                # formula kept for Excel
        ok(isinstance(live, str) and live.startswith("="), "the LET formula survives for an engineer")


def _make_if_with_table(path):
    """A minimal IF_ sheet whose table carries a source dataDxfId + a calculatedColumnFormula (both must
    be stripped on copy, else Excel drops the table), plus a tiny array-free body."""
    wb = Workbook(); ws = wb.active; ws.title = "SORTER"
    hdr = ["Category", "Description", "Data Type", "Direction </>", "I/O Offset Byte",
           "I/O Bit", "I/O Address Side 1", "Signal Name Side 1"]
    for i, h in enumerate(hdr, start=1):
        ws.cell(1, i, h)
    ws.cell(2, 3, "BOOL"); ws.cell(2, 4, ">"); ws.cell(2, 5, 0); ws.cell(2, 6, 0)
    ws.cell(2, 7, "=DT[[#This Row],[I/O Offset Byte]]")
    cols = [TableColumn(id=i + 1, name=h) for i, h in enumerate(hdr)]
    cols[6].dataDxfId = 99                                                    # out-of-range dxf ref
    cols[6].calculatedColumnFormula = TableFormula(attr_text="DT[[#This Row],[I/O Offset Byte]]")
    t = Table(displayName="DT", ref="A1:H2"); t.tableColumns = cols
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    wb.save(path)


def test_insert_lossless_idempotent_and_table_clean():
    import warnings as _w
    _w.simplefilter("ignore")
    from openpyxl.worksheet.formula import ArrayFormula
    with tempfile.TemporaryDirectory() as d:
        ifp = os.path.join(d, "IF_SORTER-01.xlsx"); _make_if_with_table(ifp)
        iol = os.path.join(d, "iolist.xlsx")
        wb = Workbook(); ws = wb.active; ws.title = "NET SAFETY 50"
        ws["A1"] = "src"; ws["B1"] = "=A1"          # a real formula (its cache must survive)
        ws["C2"] = ArrayFormula("C2:C4", "=A1")     # a dynamic array on an untouched sheet -> must be frozen
        wb.save(iol)

        acts = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])
        ok(any("inserted" in a and "IF_SORTER-01" in a for a in acts), "sheet inserted")
        wb2 = load_workbook(iol)
        ok("IF_SORTER-01" in wb2.sheetnames and "NET SAFETY 50" in wb2.sheetnames, "IF_ added, original kept")
        ok("SORTER-01" not in wb2.sheetnames, "the un-prefixed name is NOT used")
        ok("DT_IF_SORTER_01" in list(wb2["IF_SORTER-01"].tables), "table renamed per-instance")
        eq(wb2["NET SAFETY 50"]["B1"].value, "=A1", "the existing formula survives (lossless)")
        wb2.close()
        with zipfile.ZipFile(iol) as z:
            tx = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/tables/")]
            nspart = xlsx_edit._sheet_name_to_part(open(iol, "rb").read())["NET SAFETY 50"]
            ok('t="array"' not in z.read(nspart).decode("utf-8"), "the flattened array was frozen (no overlap)")
        ok(tx and not any("DxfId" in x for x in tx), "source dataDxfId stripped (else Excel drops the table)")
        ok(not any("calculatedColumnFormula" in x for x in tx), "stale calc-column formula stripped")

        acts2 = interface_xlsx.insert_sheets_into_iolist(iol, [("IF_SORTER-01", ifp)])   # idempotent
        ok(any("already present" in a for a in acts2), "second run skips the existing sheet")
        eq(load_workbook(iol).sheetnames.count("IF_SORTER-01"), 1, "no duplicate sheet")


if __name__ == "__main__":
    import sys
    sys.exit(run("interface_xlsx", [
        ("coerce_int_and_safe_name", test_coerce_int_and_safe_name),
        ("append_custom_rows_bool_block_and_padding", test_append_custom_rows_bool_block_and_padding),
        ("append_custom_rows_word_row_and_separator", test_append_custom_rows_word_row_and_separator),
        ("io_address_mirror", test_io_address_mirror),
        ("resolve_num_chain", test_resolve_num_chain),
        ("interface_address_caches", test_interface_address_caches),
        ("insert_seeds_interface_address_cache", test_insert_seeds_interface_address_cache),
        ("insert_lossless_idempotent_and_table_clean", test_insert_lossless_idempotent_and_table_clean),
    ]))
