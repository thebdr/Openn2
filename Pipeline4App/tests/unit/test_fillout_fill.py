"""ph200 / 210 - the in-place AB/AC fill (domain/fillout/fill) over a synthetic raw I/O List."""
import os
import tempfile

from openpyxl import Workbook, load_workbook

from _harness import run, eq, ok
from pipeline4.core import config
from pipeline4.domain.fillout import fill


def _sandboxed(fn):
    """Run `fn` with `config.database_dir` pointed at a throwaway dir: the domain build/fill paths
    SAVE the SSOT, and un-redirected that overwrote the builtin Shared/Database (the parity
    environment) with this module's synthetic staging on every gate run."""
    def wrapped():
        original = config.database_dir
        with tempfile.TemporaryDirectory() as sandbox:
            config.database_dir = lambda: sandbox
            try:
                fn()
            finally:
                config.database_dir = original
    return wrapped


_SHEET = "NETSAFETY"


def _raw_iolist(path):
    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET
    for col, head in (("F", "ID"), ("G", "Bit"), ("K", "Desc1"), ("L", "Desc2"), ("R", "Type"), ("X", "Mnem")):
        ws[f"{col}1"] = head                                   # AA..AG headers left BLANK (a raw doc)
    ws["G2"] = "I0.0"; ws["K2"] = "EMERGENCY PUSH-BUTTON PRESSED"; ws["L2"] = "CH1"      # -> E1/2 (Mode-1)
    ws["F3"] = "5"; ws["R3"] = "PLC"; ws["K3"] = "MAIN PLC NODE"                          # node -> PLC
    ws["G4"] = "I0.1"; ws["K4"] = "SOMETHING NOT RECOGNISED AT ALL"                       # -> <input required>
    ws["G5"] = "I0.2"; ws["K5"] = "DOOR OPEN"; ws["L5"] = "CH2"; ws["AB5"] = "MANUAL"     # Mode-2 (computed DI2/2)
    wb.save(path)


def _params(path):
    return {"iolist_path": path, "iolist_params": {"sheets": _SHEET, "header_row": 1}}


def test_fill_ab_ac_modes_and_unresolved():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        _raw_iolist(p)
        res = fill.fill_script_type(_params(p))
        wb = load_workbook(p)
        ws = wb[_SHEET]
        eq(ws["AB2"].value, "E1/2", "Mode-1: AB written from the rules")
        eq(ws["AC2"].value, "E1/2", "AC mirrors the computed value")
        eq(ws["AB3"].value, "PLC", "node AB = Type col verbatim")
        eq(ws["AB4"].value, "<input required>", "unresolved AB -> sentinel")
        eq(ws["AB5"].value, "MANUAL", "Mode-2: a human AB value is preserved")
        eq(ws["AC5"].value, "DI2/2", "AC still shows what the rules computed (Mode-2)")
        eq(ws["AB1"].value, "Script Type", "a blank output header is written")
        ok("_UnresolvedIndex" in wb.sheetnames, "the _UnresolvedIndex sheet is added")
        eq((res["filled"], res["mismatch"], res["unresolved"]), (2, 1, 1), "counts: 2 filled, 1 kept, 1 unresolved")
        eq([f.type for f in res["findings"] if f.severity == "FAIL"], ["fill_unresolved"], "one blocking FAIL")
        eq([f.type for f in res["findings"] if f.type == "fill_type_mismatch"][0:1], ["fill_type_mismatch"], "Mode-2 audit WARN")
        ok(res["backup"], "the raw doc changed -> the backup is KEPT")


def test_fill_is_idempotent_noop_drops_backup():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "io.xlsx")
        _raw_iolist(p)
        fill.fill_script_type(_params(p))                      # first run fills it
        res2 = fill.fill_script_type(_params(p))               # second run changes nothing
        eq(res2["backup"], "", "no-op re-run deletes its backup")
        eq(res2["filled"], 0, "nothing re-filled (AB already present)")


if __name__ == "__main__":
    import sys
    sys.exit(run("fillout_fill", [
        ("fill_ab_ac_modes_and_unresolved", _sandboxed(test_fill_ab_ac_modes_and_unresolved)),
        ("fill_is_idempotent_noop_drops_backup", _sandboxed(test_fill_is_idempotent_noop_drops_backup)),
    ]))
