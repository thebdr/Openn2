"""Phase 400 (Interfaces) - clean-room port of PL3's interfaces.

Builds one `IF_<instance>.xlsx` per IOC signal from the MachineInterfaces template, mirrors flagged
signals into a custom byte-packed block, and (optionally) inserts each interface as an `IF_` sheet into
the I/O List. The mirrored signals + their byte layout land in the **`interfaces`** SSOT table; the
`IF_*.xlsx` are its projection.

This first cut (400a) is the per-signal **`interface_tagname`** (the Signal Name Side 1 base): the
type's template (the relocated `chain_reactions/interface_tagnames.csv`) resolved against the signal,
keeping `{interface_name}`/`{interface_id}` for the generator. It runs AFTER 520 (it reads the written
`name_in_db` for the door types' `{db_element}`), matching DESIGN section 9's `… 520 → 600 → 400` order.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from pipeline4.core import config
from pipeline4.core.database import Database
from pipeline4.core.table import Table
from pipeline4.domain import identity
from pipeline4.domain.signals import signals_table

TRIGGER_TYPE = "IOC"
GENERIC_SHEET = "<GENERIC>"
_DIAG_RE = re.compile(r"\+\s*DIAG", re.IGNORECASE)   # the +DIAG marker in an IOC Index
_DIR_TO_IO = {"<": "I", ">": "Q"}


def annotate_interface_tagnames(database: Database, tagnames: dict | None = None) -> None:
    """Write `interface_tagname` onto every signal from its resolved type's template (keyed by `type_id`).
    Needs the 520 write-back (`name_in_db`) + the staged `name_in_tagtable`."""
    tagnames = tagnames if tagnames is not None else config.load_interface_tagnames()
    for row in database["signals"]:
        type_id = ((row.get("type") or {}).get("type_id") or "").strip().upper()
        row["interface_tagname"] = identity.interface_tagname(row, tagnames.get(type_id, ""))


def build_tagnames(database: Database | None = None) -> Database:
    """400a entry: load the (post-520) Database, annotate `interface_tagname`, save. Returns the Database."""
    if database is None:
        colmap = config.load_column_map("IoList")
        database = Database([signals_table([m["canonical"] for m in colmap])]).load(config.database_dir())
    annotate_interface_tagnames(database)
    database.save(config.database_dir())
    return database


# --- 400b: the interface instances + the mirrored elements (the SSOT tables) --------------------- #
def interfaces_table() -> Table:
    """One row per IOC interface instance (its identity + plug values + chosen template sheet)."""
    return Table(
        "interfaces",
        columns=["uid", "instance", "machine_type", "interface_id", "is_diag", "base_address",
                 "base_node", "device", "ip", "template_sheet", "source"],
        json_columns=["is_diag"],
        key_columns=["instance"],
    )


def interface_elements_table() -> Table:
    """One row per mirrored element on an interface (the custom data block). `expression` = the binding
    (`plc_binding` / a rule member); `signal_name` = the interface tag name (`{interface_name}`/
    `{interface_id}` filled). `(interface, direction, category, expression)` is unique (the mirror dedup)."""
    return Table(
        "interface_elements",
        columns=["uid", "interface", "category", "description", "functional_unit", "location", "device",
                 "data_type", "direction", "offset_byte", "bit", "signal_name", "expression",
                 "diag_cabinet", "swp_cabinet", "diag_bit", "source", "source_signal"],
        json_columns=[],
        key_columns=["interface", "direction", "category", "expression"],
    )


@dataclass
class _Elem:
    """One element to mirror onto an interface. `mirror_name` -> Expression Side 1 + the dedup key;
    `signal_name` -> Signal Name Side 1 (the interface_tagname base); offset_byte/bit filled by allocate."""
    script_type: str
    mirror_name: str
    signal_name: str = ""
    description: str = ""
    fu: str = ""
    loc: str = ""
    dev: str = ""
    data_type: str = "BOOL"
    direction: str = "Q"
    source: str = "mirror"
    diag_cabinet: str = ""
    swp_cabinet: str = ""
    diag_bit: str = ""
    src_row: dict = field(default_factory=dict)
    offset_byte: object = None
    bit: object = None


def parse_instance(index_field) -> tuple:
    """'SORTER-01' -> ('SORTER', '01'): machine type = leading letters, index = trailing digits. A
    '+DIAG' marker is stripped first ('SORTER+DIAG-02' -> ('SORTER', '02'))."""
    s = _DIAG_RE.sub("", str(index_field or "")).strip()
    type_match = re.match(r"^[A-Za-z]+", s)
    digit_match = re.search(r"(\d+)\s*$", s)
    return (type_match.group(0) if type_match else "", digit_match.group(1) if digit_match else "")


def _detect_diag(index_field) -> bool:
    return bool(_DIAG_RE.search(str(index_field or "")))


def _to_int(value):
    if value is None:
        return None
    s = str(value).strip()
    if not s:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return int(float(s))
        except ValueError:
            return None


def _idx_key(value):
    """Canonical comparable interface index: the int when numeric ('01' -> 1), else the stripped string."""
    n = _to_int(value)
    if n is not None:
        return n
    s = str(value or "").strip()
    return s or None


def _index_in(index, mapping) -> bool:
    """True iff the `|`-list `mapping` contains `index`, normalized via _idx_key so '01' == '1'."""
    want = _idx_key(index)
    if want is None:
        return False
    return any(_idx_key(tok) == want for tok in str(mapping or "").split("|") if tok.strip())


def find_interfaces(rows) -> tuple:
    """One record per IOC row of the staged signals. Returns (records, warnings); an IOC row with no
    Index is warned + skipped."""
    records, warnings = [], []
    for r in rows or []:
        if str(r.get("script_type", "") or "").strip().upper() != TRIGGER_TYPE:
            continue
        instance = str(r.get("index", "") or "").strip()
        loc = f"{r.get('source_sheet', '')}!{r.get('source_row', '')}"
        if not instance:
            warnings.append(f"IOC row {loc} has no Index - skipped")
            continue
        machine_type, index = parse_instance(instance)
        records.append({
            "instance": instance, "machine_type": machine_type, "index": index,
            "is_diag": _detect_diag(instance),
            "base": str(r.get("bit", "") or "").strip(),
            "base_node": str(r.get("id_node", "") or "").strip(),
            "device": str(r.get("device", "") or "").strip(),
            "ip": str(r.get("profinet_ip", "") or "").strip(), "source": loc,
        })
    return records, warnings


def template_sheet_names(template_path) -> list:
    wb = load_workbook(template_path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def choose_sheet(sheet_names, machine_type) -> str:
    """Exact (case-insensitive) machine-type sheet, else the <GENERIC> fallback."""
    for name in sheet_names:
        if name.upper() == str(machine_type or "").upper():
            return name
    return GENERIC_SHEET


def _find_data_table(ws):
    """The (name, c1, r1, c2, r2, {header: col}) of the sheet's data table (the one whose header row
    carries 'Category'); None if absent. Robust to table renames."""
    for name in list(ws.tables):
        c1, r1, c2, r2 = range_boundaries(ws.tables[name].ref)
        hdr = {str(ws.cell(r1, c).value).strip(): c
               for c in range(c1, c2 + 1) if ws.cell(r1, c).value not in (None, "")}
        if "Category" in hdr:
            return name, c1, r1, c2, r2, hdr
    return None


def template_last_used_byte(template_path, sheet_name) -> dict:
    """Max I/O Offset Byte per direction among the template sheet's existing rows ('<'=I, '>'=Q), from
    cached values. Computed fresh so a hand-edited template is honored. {'I': int|-1, 'Q': int|-1}."""
    last = {"I": -1, "Q": -1}
    wb = load_workbook(template_path, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return last
        found = _find_data_table(wb[sheet_name])
        if not found:
            return last
        _name, _c1, r1, _c2, r2, hdr = found
        dcol, ocol = hdr.get("Direction </>"), hdr.get("I/O Offset Byte")
        if not dcol or not ocol:
            return last
        ws = wb[sheet_name]
        for r in range(r1 + 1, r2 + 1):
            d = _DIR_TO_IO.get(str(ws.cell(r, dcol).value or "").strip())
            o = ws.cell(r, ocol).value
            if d and isinstance(o, (int, float)):
                last[d] = max(last[d], int(o))
    finally:
        wb.close()
    return last


def _rule_tagname(rule, row, member, direction) -> str:
    """A rule element's interface tag name: the rule's interface_tagname template with {member} +
    {direction} bound (keeping {interface_name}/{interface_id}); falls back to the member."""
    ctx = dict(row)
    ctx["member"] = member
    ctx["direction"] = direction
    return identity.interp_keep(rule.get("interface_tagname", ""), ctx) or member


def collect_mirror_set(rows, *, index, is_diag, diag_rules, if_rules) -> tuple:
    """Ordered list of _Elem to mirror onto interface `index`, plus warnings:
      (a) signals whose interface_mapping lists this index (or every in_diag signal when is_diag) -> Q;
      (c) their diagnosis_logic_rules followers -> Q (DB members, TIA-qualified);
      (d) their interface_elements followers -> the rule's direction (the only I source).
    Deduped on (direction, script_type, mirror_name). NOTE: the +DIAG in_diag auto-mirror needs the
    per-type `in_diag` (relocated with ph600) - until then a +DIAG interface only gets the (a) mappings."""
    warnings, elems, seen = [], [], set()

    def _add(e):
        if not e.mirror_name:
            return False
        key = (e.direction, e.script_type, e.mirror_name)
        if key in seen:
            return False
        seen.add(key)
        elems.append(e)
        return True

    direct = []
    for r in rows or []:
        st = str(r.get("script_type", "") or "").strip()
        if st.upper() == TRIGGER_TYPE:                      # never mirror the IOC rows themselves
            continue
        picked = _index_in(index, r.get("interface_mapping"))
        if not picked and is_diag and (r.get("type") or {}).get("in_diag"):
            picked = True
        if not picked:
            continue
        name = str(r.get("plc_binding", "") or "")          # _mirror_name = the stored plc_binding
        if not name:
            warnings.append(f"{st} {identity.fld(r)} "
                            f"({r.get('source_sheet', '')}!{r.get('source_row', '')}): "
                            f"no db_element or tag - not mirrored")
            continue
        e = _Elem(script_type=st, mirror_name=name,
                  signal_name=str(r.get("interface_tagname", "") or ""),
                  description=identity.interp((r.get("type") or {}).get("tag_name", ""), r),
                  fu=str(r.get("functional_unit", "") or ""), loc=str(r.get("location", "") or ""),
                  dev=str(r.get("device", "") or ""), direction="Q", source="mirror",
                  diag_cabinet=str(r.get("diag_cabinet", "") or ""),
                  swp_cabinet=str(r.get("swp_cabinet", "") or ""),
                  diag_bit=str(r.get("diag_bit", "") or ""), src_row=r)
        if _add(e):
            direct.append(e)

    for e in direct:                                        # (c) diagnosis-logic-rule followers -> Q
        r = e.src_row
        for rule in diag_rules:
            if e.script_type in rule["required_types"]:
                member = identity.interp(rule["member"], r)
                if not member:
                    continue
                db = rule.get("db_name", "")
                nm = f'"{db}"."{member}"' if db else f'"{member}"'
                _add(_Elem(script_type=rule["name"], mirror_name=nm,
                           signal_name=_rule_tagname(rule, r, member, "Q"), description=member,
                           fu=e.fu, loc=e.loc, dev=e.dev, direction="Q", source="follow",
                           diag_cabinet=e.diag_cabinet, swp_cabinet=e.swp_cabinet,
                           diag_bit=e.diag_bit, src_row=r))

    for e in direct:                                        # (d) interface_elements followers (I or Q)
        r = e.src_row
        for rule in if_rules:
            if e.script_type in rule["required_types"]:
                member = identity.interp(rule["member"], r)
                if not member:
                    continue
                _add(_Elem(script_type=rule["script_type"], mirror_name=member,
                           signal_name=_rule_tagname(rule, r, member, rule["direction"]), description=member,
                           fu=e.fu, loc=e.loc, dev=e.dev, data_type=rule["data_type"],
                           direction=rule["direction"], source="if_rule",
                           diag_cabinet=e.diag_cabinet, swp_cabinet=e.swp_cabinet,
                           diag_bit=e.diag_bit, src_row=r))
    return elems, warnings


def allocate_bytes(elems, last, gap) -> None:
    """Assign (offset_byte, bit) per direction independently. The custom area starts 2-byte-aligned at
    last[dir] + 1 + gap. Each BOOL script_type group reserves a FULL 2-byte block (padded to a 16-bit
    multiple), its signals on the low bits in order; each WORD element gets its own even-aligned 2-byte
    slot (bit None). Never two script_types in one byte."""
    for direction in ("Q", "I"):
        de = [e for e in elems if e.direction == direction]
        if not de:
            continue
        byte = last.get(direction, -1) + 1 + gap
        if byte % 2:
            byte += 1
        order, groups = [], {}
        for e in de:
            if e.script_type not in groups:
                groups[e.script_type] = []
                order.append(e.script_type)
            groups[e.script_type].append(e)
        for st in order:
            g = groups[st]
            if g[0].data_type == "WORD":
                for e in g:
                    if byte % 2:
                        byte += 1
                    e.offset_byte, e.bit = byte, None
                    byte += 2
            else:
                if byte % 2:
                    byte += 1
                for i, e in enumerate(g):
                    e.offset_byte, e.bit = byte + i // 8, i % 8
                byte += ((len(g) + 15) // 16) * 2


def build_interfaces(database: Database | None = None, template_path: str | None = None) -> tuple:
    """400b: per IOC instance, collect its mirror set + lay out the bytes -> the `interfaces` +
    `interface_elements` tables. Runs `annotate_interface_tagnames` first (the Signal Name Side 1 base
    needs name_in_db/name_in_tagtable). Saves the Database. Returns (database, warnings)."""
    if database is None:
        colmap = config.load_column_map("IoList")
        database = Database([signals_table([m["canonical"] for m in colmap])]).load(config.database_dir())
    annotate_interface_tagnames(database)
    rows = list(database["signals"])
    template_path = template_path or config.INTERFACE_TEMPLATE
    records, warnings = find_interfaces(rows)
    diag_rules, if_rules = config.load_diagnosis_logic_rules(), config.load_interface_elements()
    sheet_names = template_sheet_names(template_path)

    itab, etab = interfaces_table(), interface_elements_table()
    for rec in records:
        sheet = choose_sheet(sheet_names, rec["machine_type"])
        elems, w = collect_mirror_set(rows, index=rec["index"], is_diag=rec["is_diag"],
                                      diag_rules=diag_rules, if_rules=if_rules)
        warnings.extend(w)
        allocate_bytes(elems, template_last_used_byte(template_path, sheet), config.INTERFACE_CUSTOM_GAP)
        itab.add(instance=rec["instance"], machine_type=rec["machine_type"], interface_id=rec["index"],
                 is_diag=rec["is_diag"], base_address=rec["base"], base_node=rec["base_node"],
                 device=rec["device"], ip=rec["ip"], template_sheet=sheet, source=rec["source"])
        fill = {"interface_name": rec["machine_type"], "interface_id": rec["index"]}
        for e in elems:
            etab.add(interface=rec["instance"], category=e.script_type, description=e.description,
                     functional_unit=e.fu, location=e.loc, device=e.dev, data_type=e.data_type,
                     direction=e.direction, offset_byte=e.offset_byte, bit=e.bit,
                     signal_name=identity.interp_keep(e.signal_name, fill), expression=e.mirror_name,
                     diag_cabinet=e.diag_cabinet, swp_cabinet=e.swp_cabinet, diag_bit=e.diag_bit,
                     source=e.source, source_signal=(e.src_row or {}).get("uid", ""))

    for table in (itab, etab):
        if table.name in database:
            database[table.name].rows = table.rows
        else:
            database.add_table(table)
    database.save(config.database_dir())
    return database, warnings
