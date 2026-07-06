"""Phase 510 - project the resolved I/O signals + the interface_elements to the PLCTags.xlsx BuilderData
surface (what OP4 imports). A PURE projection of the `signals` + `interface_elements` SSOT tables - no
I/O List read-back: the interface addresses are the `io_address_side1` column phase 400 already STORED
(equal to the value 400e seeds into the inserted IF_ sheets - verified). Clean-room port of PL3's
signals.generate_io_tags / build_io_tags / interface_tags / write_plc_tags.

Two tag sources, mixed into one workbook (sorted by Path = tag table):
  (a) resolved I/O signals  - Name=name_in_tagtable, Path=tagtable, Address=%-bit, Comment=io_comment,
      Data Type=Bool (PL3 hard-codes Bool for direct I/O points);
  (b) interface tags        - one per interface_elements row: Name=signal_name, Path=IF_<instance>,
      Address=%io_address_side1, Comment=description, Data Type=Bool/Word.

The DUPLICATE-TAG GATE: two tags landing on the same (tag table, name) [case-insensitive] break the TIA
import, so each 2nd+ occurrence is an `iotag_duplicate` FAIL linking BOTH producing I/O-List rows
(location = the duplicate's row, location2 = the first occurrence's) and the workbook is NOT written
(the raw-FAIL guard). The same name on TWO tables stays legal: one signal mirrors into several IF_
tables by design.

Output -> `config.io_tags_dir()`/PLCTags.xlsx (sheets "PLC Tags" + "TagTable Properties"; all values text).
"""
from __future__ import annotations

import os
from collections import Counter

from openpyxl import Workbook

from pipeline4.core import config, run
from pipeline4.core.database import Database
from pipeline4.core.finding import Finding, record_standalone
from pipeline4.domain import identity
from pipeline4.domain.signals import signals_table


def _f(type: str, severity: str, detail: str, location: str = "", source_uid: str = "",
       location2: str = "", doc: str = "", doc2: str = "") -> Finding:
    """A phase-510 Finding - the I/O-tag projection report container: the `iotag_no_address` WARN (a named
    interface element with no resolved I/O address) + the `iotag_duplicate` FAIL (its location/location2
    pair links the duplicate's I/O-List row and the first occurrence's)."""
    return Finding(phase=510, type=type, severity=severity, detail=detail,
                   location=location, source_uid=source_uid, location2=location2, doc=doc, doc2=doc2)


TAG_TABLE_FILE = "PLCTags.xlsx"
TAG_COLUMNS = ["Name", "Path", "Data Type", "Logical Address", "Comment",
               "Hmi Visible", "Hmi Accessible", "Hmi Writeable", "Typeobject ID", "Version ID"]
PROP_COLUMNS = ["Path", "BelongsToUnit", "Accessibility"]
IF_PREFIX = "IF_"

# the TIA data-type spelling: a known type -> its TIA casing; an unknown type capitalizes; blank -> Bool.
_TIA_DTYPE = {"bool": "Bool", "word": "Word", "dword": "DWord", "byte": "Byte",
              "int": "Int", "dint": "DInt", "real": "Real", "char": "Char"}


def _tia_dtype(raw) -> str:
    s = str(raw if raw is not None else "").strip()
    if not s:
        return "Bool"
    return _TIA_DTYPE.get(s.lower(), s.capitalize())


def _logical_address(addr) -> str:
    """%-prefix an I/Q address (e.g. `I1.0` -> `%I1.0`, `Q10000.0` -> `%Q10000.0`); '' for a blank.
    Idempotent: an already-%-prefixed value is returned unchanged."""
    s = str(addr if addr is not None else "").strip()
    if not s:
        return ""
    return s if s.startswith("%") else "%" + s


# --- the two tag sources (each tag = {name, path, data_type, address, comment} + the source link:
# `location` = the producing I/O-List row's Sheet!Cell, `source_uid` = the producing entity's uid) ---- #
def io_signal_tags(database) -> list:
    """One tag per taggable I/O signal (is_io_signal + a non-empty name_in_tagtable). Data Type is Bool
    (PL3 hard-codes Bool for direct I/O points); the comment resolves the type's io_comment template."""
    tags = []
    for row in database["signals"]:
        if not identity.is_io_signal(row):
            continue
        name = str(row.get("name_in_tagtable") or "").strip()
        if not name:
            continue
        path = str(row.get("tagtable") or "").strip() or str(row.get("script_type") or "").strip()
        tags.append({"name": name, "path": path, "data_type": "Bool",
                     "address": _logical_address(row.get("bit")), "comment": identity.tag_comment(row),
                     "location": str(row.get("source_cell") or "").strip(),
                     "source_uid": str(row.get("uid") or "")})
    return tags


def interface_tags(database) -> tuple:
    """One tag per `interface_elements` row (the SSOT mirror block) - NOT a read-back of the IF_ sheets.
    Path = the IF_<instance> sheet name (the tag-table convention OP4 also sees). Address = the STORED
    `io_address_side1`, %-prefixed. A named element with no resolved address is a `iotag_no_address` WARN +
    skipped (mirrors PL3's per-row warn); a blank-named element is silently skipped. Returns (tags, findings).
    A mirror element's `location` follows its `source_signal` back to the producing I/O-List row; a
    template-native element (no source signal) falls back to the logical `<interface>/<name>` locator."""
    tags, findings = [], []
    if "interface_elements" not in database:
        return tags, findings
    cells = ({str(r.get("uid") or ""): str(r.get("source_cell") or "").strip()
              for r in database["signals"]} if "signals" in database else {})
    for e in database["interface_elements"]:
        name = str(e.get("signal_name") or "").strip()
        if not name:
            continue
        address = _logical_address(e.get("io_address_side1"))
        if not address:
            findings.append(_f("iotag_no_address", "WARN", "no resolved I/O address - tag skipped",
                               f"{e.get('interface')}/{name}", str(e.get("source_signal", ""))))
            continue
        tags.append({"name": name, "path": f"{IF_PREFIX}{e.get('interface')}",
                     "data_type": _tia_dtype(e.get("data_type")), "address": address,
                     "comment": str(e.get("description") or ""),
                     "location": cells.get(str(e.get("source_signal") or "")) or f"{e.get('interface')}/{name}",
                     "source_uid": str(e.get("uid") or "")})
    return tags, findings


def duplicate_findings(tags) -> list:
    """One `iotag_duplicate` FAIL per 2nd+ tag landing on an already-used (Path, Name) - a tag table
    cannot hold two tags with the same name (case-insensitive, TIA's uniqueness rule: the import breaks),
    and the collision is a source-document defect staging cannot see (two I/O-List rows with the same FLD
    have distinct uids - the uid includes source_cell). The finding links BOTH producing rows: `location` =
    the duplicate's I/O-List row, `location2` = the FIRST occurrence's (the log renders `<dup> vs <first>`,
    both clickable - the first-occurrence link lives THERE, not in the detail); the detail carries the
    tag's total occurrence count (`Count: <n>`). The same name on TWO tables is deliberately NOT flagged -
    one signal mirrors into several IF_ tables by design."""
    counts = Counter((t["path"], str(t["name"]).casefold()) for t in tags)
    findings, first, doc = [], {}, None
    for t in tags:
        key = (t["path"], str(t["name"]).casefold())
        head = first.setdefault(key, t)
        if head is t:
            continue
        if doc is None:                               # resolve the I/O List basename once, on the first hit
            doc = os.path.basename(str(config.load_params().get("iolist_path") or ""))
        loc = t.get("location") or f"{t['path']}/{t['name']}"
        loc2 = head.get("location") or f"{head['path']}/{head['name']}"
        findings.append(_f("iotag_duplicate", "FAIL",
                           f"duplicate tag '{t['name']}' in table '{t['path']}' - Count: {counts[key]}",
                           loc, t.get("source_uid", ""), location2=loc2,
                           doc=doc if "!" in loc else "", doc2=doc if "!" in loc2 else ""))
    return findings


# --- the workbook writer ------------------------------------------------------------------------- #
def _append_text_row(ws, values) -> None:
    """Append a row; force data_type='s' on any cell starting with =/+/- so a tag/FLD/address isn't read
    as a formula (e.g. a Name like `... [ =S1+MS1.CC1-F09001 ]`)."""
    ws.append(values)
    r = ws.max_row
    for c, v in enumerate(values, 1):
        if isinstance(v, str) and v[:1] in ("=", "+", "-"):
            ws.cell(r, c).data_type = "s"


def _group_tables(tags) -> dict:
    """{path: [tag, ...]} preserving insertion order within each table."""
    out: dict = {}
    for t in tags:
        out.setdefault(t["path"], []).append(t)
    return out


def write_plc_tags(tags, out_dir) -> str:
    """Write PLCTags.xlsx: a "PLC Tags" sheet (one row per tag, sorted by Path) + a "TagTable Properties"
    sheet (one row per distinct Path). All values text; Hmi flags the literal 'True'. Returns the path."""
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, TAG_TABLE_FILE)
    tables = _group_tables(tags)
    wb = Workbook()
    ws = wb.active
    ws.title = "PLC Tags"
    ws.append(TAG_COLUMNS)
    for table in sorted(tables):
        for t in tables[table]:
            _append_text_row(ws, [t["name"], table, t["data_type"], t["address"], t["comment"],
                                  "True", "True", "True", "", ""])
    props = wb.create_sheet("TagTable Properties")
    props.append(PROP_COLUMNS)
    for table in sorted(tables):
        _append_text_row(props, [table, "", ""])
    wb.save(path)
    wb.close()
    return path


def project(database: Database | None = None, out_dir: str | None = None) -> dict:
    """Project the two SSOT sources to `out_dir`/PLCTags.xlsx (out_dir defaults to `config.io_tags_dir()`).
    Returns {'path', 'total', 'io_count', 'iface_count', 'tables', 'findings'} - `findings` = the
    `iotag_no_address` WARNs + the `iotag_duplicate` FAILs (the caller renders/gates them). On a duplicate
    tag the workbook is NOT written (`run.has_blocking` raw-FAIL guard - never hand OP4 a broken import
    surface; 'path' comes back '') and the FAILs alone are RECORDED to `validation_issues.csv`
    (`record_standalone`, the ph200 seam - the log's [FAIL]->Findings jump needs the row); the WARNs stay
    render-only as before."""
    if database is None:
        colmap = config.load_column_map("IoList")
        database = Database([signals_table([m["canonical"] for m in colmap])]).load(config.database_dir())
    out_dir = out_dir or config.io_tags_dir()
    io = io_signal_tags(database)
    iface, findings = interface_tags(database)
    tags = io + iface
    findings += duplicate_findings(tags)
    if run.has_blocking(findings):                    # duplicate tags -> record the FAILs, write NOTHING
        record_standalone([f for f in findings if f.type == "iotag_duplicate"])
        path = ""
    else:
        path = write_plc_tags(tags, out_dir)
    return {"path": path, "total": len(tags), "io_count": len(io), "iface_count": len(iface),
            "tables": sorted({t["path"] for t in tags}), "findings": findings}
