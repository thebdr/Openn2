"""110 - Validate I/O List standalone. Clean-room port of PL3's `domain/validation/iolist.py`.

Reads the I/O List workbook (raw rows, all sheets matching `iolist_params.sheets`) through the ONE shared
reader. Column access + the header glossary come from `column_map("IoList")`; only columns NOT flagged
`preliminary_check_exclude` (the customer-authored A-Z) are validated. Emits one `Finding` per finding; the
`type` slug is the code-traceable log-type index + the message key. PL4: takes the resolved `params` (no
`ctx`), reads config via `get_param`, EN messages (no i18n).
"""
from __future__ import annotations

import os
import re

from openpyxl.utils import column_index_from_string as _ci, get_column_letter

from pipeline4.core import config
from pipeline4.core.model import InfoBlock
from pipeline4.io import workbook as wbk
from pipeline4.io.workbook import FileLockedError
from pipeline4.domain.validation import address as adr, model as vm

_BAD_PNAME = set("_./=*")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "").replace("\r", " ").replace("\n", " ")).strip()


def _clean(s) -> str:
    """Whitespace-removed, upper-cased (permanent-part comparison)."""
    return "".join(str(s or "").split()).upper()


def _is_err(v) -> bool:
    """An Excel error literal (#REF!, #N/A, #VALUE!, ...)."""
    return str(v or "").strip().upper().startswith("#")


def _row_info(v, col, r) -> InfoBlock:
    fld = vm.fld_text(v.text(r, col["functional_unit"]), v.text(r, col["location"]), v.text(r, col["device"]))
    ti = "-".join(p for p in (v.text(r, col["script_type"]), v.text(r, col["index"])) if p)
    return InfoBlock(bit=v.text(r, col["bit"]), fld=fld, desc_l1=v.text(r, col["desc_l1"]),
                     desc_l1b=v.text(r, col["desc_l1b"]), drawing=v.text(r, col["drawing"]), type_index=ti)


def run_iolist(params: dict) -> list:
    """Validate the I/O List of `params['iolist_path']`. Returns a list of phase-110 Findings."""
    out = []
    io_path = params.get("iolist_path") or ""
    doc = os.path.basename(io_path)
    if not io_path or not os.path.exists(io_path):
        return [vm.entry("FAIL", 110, "iolist_missing", path=io_path)]

    colmap = config.load_column_map("IoList")
    col = {m["canonical"]: m["column"] for m in colmap}
    check_cols = [m for m in colmap if not m["preliminary_check_exclude"]]
    mapped_idx = {_ci(m["column"]) for m in colmap}
    expected_norm = {_norm(m["expected_header"]).lower() for m in colmap if m["expected_header"]}

    sheet_pat = config.get_param(params, "iolist_params.sheets")
    hr = int(config.get_param(params, "iolist_params.header_row", 1) or 1)
    strike_error = str(config.get_param(params, "validation_params.global.strike_handling", "exclude")).lower() == "error"
    perm = {_clean(p) for p in (config.get_param(params, "validation_params.iolist.permanent_parts", []) or [])}

    try:
        views = wbk.open_sheets(io_path, sheet_pat, hr, doc_label=doc)
        all_names = wbk.available_sheets(io_path)
    except FileLockedError:
        return [vm.entry("FAIL", 110, "iolist_locked")]
    except Exception as e:  # noqa: BLE001
        return [vm.entry("FAIL", 110, "iolist_open_error", error=str(e))]

    matched = [v.name for v in views]
    if not matched:
        return [vm.entry("FAIL", 110, "no_sheet", pattern=str(sheet_pat))]
    for n in all_names:
        out.append(vm.entry("INFO", 110, "proc_sheet" if n in matched else "skip_sheet", sheet=n))

    seen_ip, seen_fld, total_nodes = {}, {}, 0
    for v in views:
        # --- header check (mapped, non-excluded columns) --- #
        for m in check_cols:
            actual = v.header(m["column"])
            if (actual or m["required"]) and not v.header_matches(m["column"], m["expected_header"]):
                out.append(vm.entry("FAIL", 110, "col_wrong", actual=actual or "(empty)",
                                    n=_ci(m["column"]), expected=m["expected_header"],
                                    location=f"{v.name}!{m['column']}{hr}", doc=doc))
        # --- unexpected / duplicated headers (columns outside the map) --- #
        for c in range(1, v.max_col + 1):
            if c in mapped_idx:
                continue
            hdr = v.header(c)
            if not hdr:
                continue
            loc = f"{v.name}!{get_column_letter(c)}{hr}"
            if _norm(hdr).lower() in expected_norm:
                out.append(vm.entry("FAIL", 110, "col_duplicated", n=c, name=hdr, location=loc, doc=doc))
            else:
                out.append(vm.entry("WARN", 110, "col_unexpected", n=c, name=hdr, location=loc, doc=doc))

        # --- per-row checks --- #
        nodes = 0
        for r in v.data_rows():
            if all(v.text(r, m["column"]) == "" for m in colmap):
                continue
            before = len(out)            # to emit a per-row PASS when the row raises no finding
            info = _row_info(v, col, r)
            g, f = v.text(r, col["bit"]), v.text(r, col["id_node"])
            ip = v.cell(r, col["profinet_ip"])
            ip_t = "" if ip is None else str(ip).strip()
            pname = v.text(r, col["profinet_name"])
            thw = v.text(r, col["type_hw"]).strip().upper()
            fu, lo, de = v.text(r, col["functional_unit"]), v.text(r, col["location"]), v.text(r, col["device"])

            if strike_error and v.row_struck(r):
                out.append(vm.entry("FAIL", 110, "struck_row", location=v.location(r, col["bit"]), doc=doc, info=info))
            # Profinet IP
            if _is_err(ip_t):
                out.append(vm.entry("FAIL", 110, "ip_error", location=v.location(r, col["profinet_ip"]), doc=doc, info=info))
            elif ip_t and fu and "." in ip_t:
                nodes += 1
                if ip_t in seen_ip and not ip_t.endswith(".1"):
                    out.append(vm.entry("FAIL", 110, "ip_duplicated", ip=ip_t, first=seen_ip[ip_t],
                                        location=v.location(r, col["profinet_ip"]), doc=doc, info=info))
                else:
                    seen_ip.setdefault(ip_t, v.location(r, col["profinet_ip"]))
                fk = vm.key(fu, lo, de)
                if fk and fk in seen_fld:
                    out.append(vm.entry("FAIL", 110, "dup_fld", first=seen_fld[fk],
                                        location=v.location(r, col["profinet_name"]), doc=doc, info=info))
                elif fk:
                    seen_fld[fk] = v.location(r, col["profinet_name"])
            # G (Bit) excludes F (ID)
            if g and f:
                out.append(vm.entry("FAIL", 110, "fg_exclusion", location=v.location(r, col["bit"]), doc=doc, info=info))
            # address format
            if g and not adr.format_ok(g):
                out.append(vm.entry("FAIL", 110, "addr_format", location=v.location(r, col["bit"]), doc=doc, info=info))
            # permanent part (A / W)
            if thw in ("A", "W"):
                pp = v.text(r, col["desc_l1"])
                if pp == "":
                    out.append(vm.entry("FAIL", 110, "pp_empty", location=v.location(r, col["desc_l1"]), doc=doc, info=info))
                elif _clean(pp) not in perm:
                    out.append(vm.entry("FAIL", 110, "pp_unknown", value=pp,
                                        location=v.location(r, col["desc_l1"]), doc=doc, info=info))
            # T.S. ref (A / W / PA / PW)
            if thw in ("A", "W", "PA", "PW") and "TS_" not in v.text(r, col["ts_ref"]).upper():
                out.append(vm.entry("FAIL", 110, "ts_missing", location=v.location(r, col["ts_ref"]), doc=doc, info=info))
            # Profinet node identity (PA / PW)
            if thw in ("PA", "PW"):
                if not ip_t or "." not in ip_t:
                    out.append(vm.entry("FAIL", 110, "ip_missing", location=v.location(r, col["profinet_ip"]), doc=doc, info=info))
                if (not pname.startswith("n")) or any(ch in pname for ch in _BAD_PNAME):
                    out.append(vm.entry("FAIL", 110, "pname_wrong", location=v.location(r, col["profinet_name"]), doc=doc, info=info))

            # the row raised no finding -> a PASS line for the full report (dropped from errors-only)
            if len(out) == before:
                out.append(vm.entry("PASS", 110, "row_ok", location=v.location(r, col["functional_unit"]), doc=doc, info=info))

        if nodes > 126:
            out.append(vm.entry("FAIL", 110, "too_many_nodes", n=nodes, location=f"{v.name}!A1", doc=doc))
        elif nodes > 64:
            out.append(vm.entry("WARN", 110, "nodes_over_64", n=nodes, location=f"{v.name}!A1", doc=doc))
        total_nodes += nodes

    if views:
        views[0].close()
    out.append(vm.entry("INFO", 110, "iolist_summary", sheets=len(matched), nodes=total_nodes))
    return out
