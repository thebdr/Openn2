"""The Tk-free logic behind the Files tab (GUI M5): which files to show (the `files_tab` config section -
regex include/exclude per section, ${placeholder} roots), which viewer a file gets, the file-tree
structure, and the read helpers (CSV + xlsx). The Tk view (`gui/files_panel.py`) renders the tree these
functions build and calls the readers; keeping the logic here makes it unit-testable without a display.
VIEW-ONLY (decision #3: a table edit must go through the codec, so this previews + opens externally).

Visibility is CONFIG-DRIVEN (UI_REFRESH_PLAN E): `sections_from_config` compiles app_config.yaml's
`files_tab.sections` into `{title, roots, include, exclude}` dicts; `populate(root, include, exclude)`
matches each regex (case-insensitive `re.search`) against the file's path RELATIVE to its root. A shown
file with no dedicated viewer still gets the header + Open-externally pane - nothing is silently hidden
by extension anymore.
"""
from __future__ import annotations

import csv as _csv
import os
import re

from pipeline4.core import config

# extension -> viewer kind (the ROUTING only - visibility is the sections' regex filters).
_GRID_EXT = (".csv",)
_XLSX_EXT = (".xlsx", ".xlsm", ".xls")
_TEXT_EXT = (".yaml", ".yml", ".json", ".xml", ".scl", ".db", ".txt", ".md", ".log")

XLSX_MAX_ROWS, XLSX_MAX_COLS = 3000, 80   # cap a huge sheet so the preview stays responsive (PL3 parity)

_PLACEHOLDER_RE = re.compile(r"^\$\{([A-Za-z_]+)\}$")


def compile_filters(patterns) -> tuple:
    """`(compiled, warnings)`: each pattern -> `re.compile(p, re.IGNORECASE)`; a bad regex becomes a
    warning line and is skipped (the section still works with the remaining patterns)."""
    compiled, warnings = [], []
    for pattern in patterns or []:
        try:
            compiled.append(re.compile(str(pattern), re.IGNORECASE))
        except re.error as error:
            warnings.append(f"bad regex {str(pattern)!r}: {error}")
    return compiled, warnings


def matches(rel: str, include, exclude) -> bool:
    """Show the file at relative path `rel`? ANY include hit (an EMPTY include list = everything) and NO
    exclude hit - `re.search` semantics on the /-separated relative path."""
    if include and not any(rx.search(rel) for rx in include):
        return False
    return not any(rx.search(rel) for rx in exclude)


def placeholder_map(params: dict) -> dict:
    """The `${name}` -> path table a files_tab root resolves through, against the ACTIVE config/project."""
    return {
        "config_project": config.config_project_dir(),
        "user_input": config.user_input_dir(),
        "database_dir": config.database_dir(),
        "output_root": config.output_root(),
        "reports": config.validation_report_dir(),
        "iolist": str(params.get("iolist_path") or ""),
        "matrix": str(params.get("matrix_path") or ""),
        "project_root": config.active_project() or "",
    }


def resolve_root(token, mapping: dict) -> tuple:
    """A files_tab root entry -> `(path-or-None, warning-or-None)`. `${unknown}` warns; a KNOWN but empty
    placeholder (e.g. no I/O List configured) skips silently; a literal path passes through as-is."""
    text = str(token or "").strip()
    matched = _PLACEHOLDER_RE.match(text)
    if not matched:
        return (text or None), None
    name = matched.group(1)
    if name not in mapping:
        return None, f"unknown placeholder ${{{name}}} (known: {', '.join(sorted(mapping))})"
    return (mapping[name] or None), None


def sections_from_config(specs, params: dict) -> tuple:
    """The raw `files_tab.sections` list -> `([{title, roots, include, exclude}], warnings)`. Tolerant:
    a malformed entry degrades to a warning line, never an exception (the rest of the tab still works)."""
    mapping = placeholder_map(params or {})
    sections, warnings = [], []
    for spec in specs or []:
        if not isinstance(spec, dict):
            warnings.append(f"section entry is not a mapping: {spec!r}")
            continue
        title = str(spec.get("title") or "Untitled section")
        roots = []
        for token in spec.get("roots") or []:
            path, warning = resolve_root(token, mapping)
            if warning:
                warnings.append(f"{title}: {warning}")
            elif path:
                roots.append(path)
        include, bad_inc = compile_filters(spec.get("include") or [".*"])
        exclude, bad_exc = compile_filters(spec.get("exclude") or [])
        warnings.extend(f"{title}: {w}" for w in bad_inc + bad_exc)
        sections.append({"title": title, "roots": roots, "include": include, "exclude": exclude})
    return sections, warnings


def viewer_kind(name: str) -> str | None:
    """The viewer a file gets by extension: 'csv' | 'xlsx' | 'text', or None when unsupported."""
    low = name.lower()
    if low.endswith(_GRID_EXT):
        return "csv"
    if low.endswith(_XLSX_EXT):
        return "xlsx"
    if low.endswith(_TEXT_EXT):
        return "text"
    return None


def _node(name, path, is_dir, children=None) -> dict:
    return {"name": name, "path": os.path.abspath(path), "is_dir": is_dir, "children": children or []}


def populate(root: str, include=(), exclude=(), _base: str | None = None) -> list:
    """The tree nodes under `root` (a dir -> its entries, recursive; a file root -> a single node),
    filtered by the section's include/exclude regexes against each file's path RELATIVE to `root`
    (/-separated). Directories are listed first-sorted folders-then-files; dotfiles + `__pycache__` are
    structural skips; a folder whose subtree shows no file is PRUNED. Returns node dicts (Tk-free)."""
    if not root or not os.path.exists(root):
        return []
    base = _base or (root if os.path.isdir(root) else os.path.dirname(root))

    def rel_of(path: str) -> str:
        return os.path.relpath(path, base).replace(os.sep, "/")

    if os.path.isfile(root):
        name = os.path.basename(root)
        return [_node(name, root, False)] if matches(rel_of(root), include, exclude) else []
    try:
        entries = sorted(os.scandir(root), key=lambda e: (e.is_file(), e.name.lower()))
    except OSError:
        return []
    nodes = []
    for e in entries:
        if e.name.startswith(".") or e.name == "__pycache__":
            continue
        if e.is_dir():
            children = populate(e.path, include, exclude, _base=base)
            if children:                                # prune empty branches
                nodes.append(_node(e.name + "/", e.path, True, children))
        elif matches(rel_of(e.path), include, exclude):
            nodes.append(_node(e.name, e.path, False))
    return nodes


def sniff_delim(text: str) -> str:
    """The delimiter of a CSV's first non-empty line: ';' when it has more semicolons than commas, else
    ',' (the SSOT tables are comma; hand-authored config CSVs are sometimes semicolon)."""
    line = next((ln for ln in text.splitlines() if ln.strip()), "")
    return ";" if line.count(";") > line.count(",") else ","


def read_csv_rows(path: str) -> list:
    """A CSV file -> a list of string rows (delimiter sniffed). Read-only; cells stay raw text (JSON cells
    show as their JSON string, which is exactly what's on disk)."""
    with open(path, encoding="utf-8-sig", newline="") as handle:
        text = handle.read()
    return list(_csv.reader(text.splitlines(), delimiter=sniff_delim(text)))


def read_xlsx(path: str, sheet: str | None = None):
    """(sheetnames, rows) for an xlsx/xlsm: the sheet names + the chosen sheet's cells as strings (cached
    values, read-only), capped to XLSX_MAX_ROWS x XLSX_MAX_COLS. `sheet=None` reads the first sheet."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        names = list(wb.sheetnames)
        if not names:
            return [], []
        name = sheet if sheet in names else names[0]
        rows = []
        for i, row in enumerate(wb[name].iter_rows(values_only=True)):
            if i >= XLSX_MAX_ROWS:
                break
            rows.append(["" if c is None else str(c) for c in row[:XLSX_MAX_COLS]])
        return names, rows
    finally:
        wb.close()
