# Generates synthetic customer documents for the pipeline test:
#   TestInput\IoList_test.xlsx  - I/O List, header row 2, Type/DiagCabinet/DiagBit
#                                 in columns AB/AC/AD, one struck-through row,
#                                 one deliberate violation per validation rule
#   TestInput\CE_test.xlsx      - C&E matrix with matching + broken rows
# Usage: python New-TestInputs.py <targetFolder>
import sys, os
from openpyxl import Workbook
from openpyxl.styles import Font

target = sys.argv[1]
# "clean" omits the deliberate-violation rows + struck row, for the export test
clean = len(sys.argv) > 2 and sys.argv[2] == "clean"
os.makedirs(target, exist_ok=True)

# ---------------- I/O List ----------------
io_headers = [
    "Description Module", "Manufacturer", "Part No.", "Cod. Fives", "Slot", "ID", "Bit",
    "Normal condition", "Connector", "Pin No.", "Description language 1", "Desc L1 part2",
    "Description language 2", "Desc L2 part2", "Functional unit", "Location", "Device",
    "Reserved1", "Drawing name", "Sheet", "T.S. ref.", "Profinet IP", "Profinet name",
    "Mnemonic", "Alarm filter", "Tag filter", "Reserved2",
    "Type", "Diagnosis Cabinet", "Diagnosis Bit",   # AB, AC, AD
]
def io_row(module="", partno="", slot="", node="", bit="", desc="", fu="=SE01",
           loc="+ZP001.EL001", dev="", ip="", pn="", typ="", diagcab="", diagbit=""):
    return [module, "SIEMENS" if partno else "", partno, "", slot, node, bit, "", "", "",
            desc, "", "", "", fu, loc, dev, "", "8FSN_Q0005", "300", "", ip, pn, "", "", "",
            "", typ, diagcab, diagbit]

rows = [
    # hardware heads + card
    io_row(module="CPU 1517F-3 PN/DP", partno="6ES7517-3FP00-0AB0", node=1, dev="-K30001", ip="192.168.50.1", pn="plc1"),
    io_row(module="IM 155-6 PN ST",    partno="6ES7155-6AU01-0BN0", node=5, dev="-K40001", ip="192.168.50.5", pn="n0005"),
    io_row(module="F-DI 8x24VDC HF",   partno="6ES7136-6BA01-0CA0", slot=1, node=5, dev="-K40019"),
    # clean signals
    io_row(node=5, bit=0, desc="EMERGENCY PUSH BUTTON CH1", dev="-S31001", typ="E1/2"),
    io_row(node=5, bit=1, desc="EMERGENCY PUSH BUTTON CH2", dev="-S31001", typ="E2/2"),
    io_row(node=5, bit=2, desc="DOOR CLOSED CH1",           dev="-S40001", typ="DI1/2"),
    io_row(node=5, bit=3, desc="DOOR CLOSED CH2",           dev="-S40001", typ="DI2/2"),
    io_row(node=5, bit=4, desc="CONTACTOR OUTPUT",          dev="-K50001", typ="KQ"),
    io_row(node=5, bit=5, desc="ALARM PUMP",                dev="-B60001", typ="A",  diagcab="+CAB01", diagbit=12),
    io_row(node=5, bit=6, desc="AREA 2 FEEDBACK",           dev="-FA0002", typ="FA2"),
]
if not clean:
    rows += [
        # struck row (predisposition - must be excluded)
        io_row(node=5, bit=7, desc="FUTURE EMERGENCY",          dev="-S99901", typ="E1/2"),
        # one violation per rule
        io_row(node=5, bit=8, desc="UNKNOWN TYPE",              dev="-S77001", typ="XX"),                       # V101
        io_row(module="IM 155-6 PN ST", partno="6ES7155-6AU01-0BN0", node=9, dev="-K41001",
               ip="10.0.0.9", pn="n0009"),                                                                      # V105
        io_row(node=5, bit=9, desc="DIAG BIT RANGE",            dev="-B60002", typ="W", diagcab="+CAB01", diagbit=99),  # V108
        io_row(node=5, bit=0, desc="EMERGENCY PUSH BUTTON CH1", dev="-S31001", typ="E1/2"),                     # V102+V103 (duplicate)
        io_row(node=5, bit=10, desc="BAD TAG NAME",             dev="S88001", typ="W"),                         # V104
    ]

def append_literal(ws, values):
    # openpyxl treats strings starting with '=' as formulas; force literal text
    ws.append(values)
    for c, v in enumerate(values, 1):
        if isinstance(v, str) and v.startswith("="):
            ws.cell(row=ws.max_row, column=c).data_type = "s"

wb = Workbook(); ws = wb.active; ws.title = "IO List"
ws.append(["I/O LIST - synthetic test document"])
ws.append(io_headers)
struck_row_index = None
for i, r in enumerate(rows):
    append_literal(ws, r)
    if r[16] == "-S99901":
        struck_row_index = ws.max_row
if struck_row_index is not None:
    for c in range(1, len(io_headers) + 1):
        ws.cell(row=struck_row_index, column=c).font = Font(strike=True)
wb.save(os.path.join(target, "IoList_test.xlsx"))

# ---------------- C&E matrix ----------------
ce_headers = ["PLC-F", "VALIDATED", "POSITION", "NOTE", "MODULE", "BIT (ADDRESS)",
              "SLOT", "PIN No.", "DESCRIPTION", "FUNCTIONAL UNIT", "LOCATION", "DEVICE", "AREA"]
ce_rows = [
    ["OK", "OK",     "OK", "",            "=SE01+ZP001.EL001", "I41.0", "-K40019", "1.5", "EMERGENCY PUSH BUTTON", "=SE01", "+ZP001.EL001", "-S31001", "S1"],
    ["OK", "OK",     "OK", "",            "",                  "I41.2", "-K40019", "2.6", "DOOR CLOSED",           "=SE01", "+ZP001.EL001", "-S40001", "S1"],
]
if not clean:
    ce_rows += [
        ["OK", "OK",     "OK", "",        "",                  "I99.0", "-K40019", "3.1", "GHOST DEVICE",          "=SE01", "+ZP001.EL001", "-S66666", "S1"],   # V301
        ["OK", "BYPASS", "OK", "",        "",                  "I41.5", "-K40019", "4.0", "ALARM PUMP",            "=SE01", "+ZP001.EL001", "-B60001", "S1"],   # V203
        ["OK", "OK",     "OK", "",        "",                  "X1.2",  "-K40019", "5.0", "BAD ADDRESS",           "=SE01", "+ZP001.EL001", "-S40001", "S1"],   # V201
        ["OK", "OK",     "OK", "",        "",                  "I41.6", "-K40019", "6.0", "AREA 2 FEEDBACK",       "=SE01", "+ZP001.EL001", "-FA0002", "ZZ"],   # V303
    ]
wb2 = Workbook(); ws2 = wb2.active; ws2.title = "C&E"
ws2.append(["CAUSE & EFFECT - synthetic test document"])
ws2.append(ce_headers)
for r in ce_rows:
    append_literal(ws2, r)
wb2.save(os.path.join(target, "CE_test.xlsx"))

print("test inputs written to " + target)
print("expected Errors: V101 V102 V103 V105 V108 V201 V203 V301 (8 findings)")
print("expected Warnings: V104 V302(-K50001) V303 (3 findings)")
