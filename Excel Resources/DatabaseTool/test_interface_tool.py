#!/usr/bin/env python3
"""Self-contained test for interface_tool: builds a synthetic I/O List and a
multi-sheet interface template (one type sheet + <GENERIC>), runs the generator,
and checks instance parsing, sheet selection + GENERIC fallback, base/index
plugging, <index> token replacement, strikethrough exclusion and
preserve-on-rerun. Run: python test_interface_tool.py
"""
import os
import shutil
import tempfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import interface_tool as it

failures = 0


def check(name, ok):
    global failures
    print(("PASS  " if ok else "FAIL  ") + name)
    if not ok:
        failures += 1


def make_io_list(path, rows, struck_rows=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "NET SAFETY 50"
    headers = ["Device", "Profinet IP", "ID", "Bit", "Script Type", "Index"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for i, row in enumerate(rows, start=2):
        ws.cell(i, 1, row["device"])
        ws.cell(i, 2, row["ip"])
        ws.cell(i, 3, row.get("id", ""))
        ws.cell(i, 4, row.get("bit", ""))
        ws.cell(i, 5, row["script"])
        ws.cell(i, 6, row["index"])
        if (i - 2) in struck_rows:
            for c in range(1, 7):
                ws.cell(i, c).font = Font(strike=True)
    wb.save(path)


def make_template(path):
    """Two sheets (SORTER + <GENERIC>) with the real plug columns/tokens."""
    wb = Workbook()
    for name in ("SORTER", "<GENERIC>"):
        ws = wb.create_sheet(name)
        for c, h in enumerate(["Category", "Signal Name", "Side", "Index", "Base Node", "Base Address"], 1):
            ws.cell(1, c, h)
        if name == "SORTER":
            ws.append(["WATCHDOG", "PNC_Q_Sorter<index> HB", 1, "01", 99, 10000])  # Side 1 (our PLC)
            ws.append(["STATE", "PNC_I_Sorter<index> ST", 2, "01", 99, 0])          # Side 2
    del wb["Sheet"]
    wb.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="iftest_")
    try:
        io_path = os.path.join(tmp, "io.xlsx")
        out = os.path.join(tmp, "out")
        tpl = os.path.join(tmp, "tpl.xlsx")
        make_template(tpl)

        # parse_instance unit checks
        check("parse SORTER-01", it.parse_instance("SORTER-01") == ("SORTER", "01"))
        check("parse PALLETIZER-12", it.parse_instance("PALLETIZER-12") == ("PALLETIZER", "12"))

        make_io_list(io_path, [
            {"device": "-K66201", "ip": "192.168.50.6", "id": "6", "bit": "88888", "script": "IOC", "index": "SORTER-01"},
            {"device": "-S31001", "ip": "", "bit": "I0.0", "script": "E1/2", "index": "0001"},   # not IOC
            {"device": "-K66203", "ip": "192.168.50.8", "id": "8", "bit": "30000", "script": "IOC", "index": "PALLETIZER-03"},
            {"device": "-K69999", "ip": "192.168.50.9", "id": "9", "bit": "40000", "script": "IOC", "index": "GHOST-99"},  # struck
        ], struck_rows={3})

        n = it.generate(io_path, "NET SAFETY 50", 1, out, tpl)
        check("creates one file per non-struck IOC row", n == 2)
        files = sorted(os.listdir(out))
        check("instance file names", files == ["IF_PALLETIZER-03.xlsx", "IF_SORTER-01.xlsx"])
        check("struck IOC excluded", "IF_GHOST-99.xlsx" not in files)

        # SORTER-01 uses the SORTER sheet, base + index plugged, token replaced
        wb = load_workbook(os.path.join(out, "IF_SORTER-01.xlsx"))
        check("single sheet named for instance", wb.sheetnames == ["SORTER-01"])
        ws = wb.active
        check("base address plugged into Side-1 row (distinct value)", ws["F2"].value == 88888)
        check("Side-2 base address untouched", ws["F3"].value == 0)
        check("base node plugged from ID into Side-1", ws["E2"].value == 6)
        check("Index column set", ws["D2"].value == "01" and ws["D3"].value == "01")
        check("<index> token replaced in signal names",
              ws["B2"].value == "PNC_Q_Sorter01 HB" and "<index>" not in ws["B3"].value)
        wb.close()

        # PALLETIZER-03 falls back to <GENERIC> (no PALLETIZER sheet)
        wb = load_workbook(os.path.join(out, "IF_PALLETIZER-03.xlsx"))
        check("unknown type uses GENERIC structure (no SORTER signals)",
              wb.active["A2"].value is None)  # GENERIC had no data rows
        wb.close()

        # preserve-on-rerun
        marker = os.path.join(out, "IF_SORTER-01.xlsx")
        wb = load_workbook(marker)
        wb.active["A20"] = "HAND ADDED"
        wb.save(marker)
        wb.close()
        n2 = it.generate(io_path, "NET SAFETY 50", 1, out, tpl)
        check("rerun creates nothing (all preserved)", n2 == 0)
        wb = load_workbook(marker)
        check("hand-added data survives rerun", wb.active["A20"].value == "HAND ADDED")
        wb.close()

        print("ALL CHECKS PASS" if failures == 0 else f"{failures} CHECK(S) FAILED")
        raise SystemExit(0 if failures == 0 else 1)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
