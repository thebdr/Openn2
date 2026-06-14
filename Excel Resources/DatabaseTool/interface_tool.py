#!/usr/bin/env python3
"""Generate per-machine interface I/O tables from the I/O List.

A row whose Script Type = "IOC" marks a machine interface; its Index has the
form MACHINETYPE-nn (e.g. SORTER-01). The machine type (SORTER) selects a sheet
in the interface template workbook (one sheet per type, plus a "<GENERIC>"
fallback when the type has no dedicated sheet); the trailing digits (01) are the
instance index.

For each IOC instance the template workbook is copied, every sheet except the
chosen type sheet is removed, and two things from the I/O List are plugged in:
  - Base Address  (the I/Q start byte, Bit column of the IOC row) into the
                   Side-1 cell of the "Base Address" column
  - Index         (the digits after the machine type) into the "Index" column
                   and in place of every "<index>" token in the sheet
The template's own LET/SUBSTITUTE formulas and Excel tables are preserved
(openpyxl round-trips them); the engineer then fills/extends the table by hand.
Output: <out>/IF_<instance>.xlsx, one per instance, PRESERVED if it exists.

Strikethrough source rows mean "predisposition, not used" and are skipped.

Usage:
  python interface_tool.py <io_list.xlsx> [--sheet "NET SAFETY 50"]
         [--header-row 1] [--out <dir>] [--template <xlsx>]
"""
from __future__ import annotations
import argparse
import os
import re
import shutil
from openpyxl import load_workbook

TRIGGER_TYPE = "IOC"
GENERIC_SHEET = "<GENERIC>"
INDEX_TOKEN = "<index>"
DEFAULT_TEMPLATE = os.path.join(os.path.dirname(__file__), "Templates", "TEMPLATE_INTERFACES_v0.0.xlsx")


def parse_instance(index_field: str) -> tuple[str, str]:
    """'SORTER-01' -> ('SORTER', '01'). Machine type = leading letters,
    index = trailing digits."""
    s = str(index_field).strip()
    type_match = re.match(r"^[A-Za-z]+", s)
    digit_match = re.search(r"(\d+)\s*$", s)
    machine_type = type_match.group(0) if type_match else ""
    index = digit_match.group(1) if digit_match else ""
    return machine_type, index


def _header_index(ws, header_row: int) -> dict:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[header_row]
        if cell.value not in (None, "")
    }


def _row_struck(ws, row_index: int, last_col: int) -> bool:
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row_index, column=c)
        if cell.value not in (None, "") and cell.font and cell.font.strike:
            return True
    return False


def find_interfaces(io_path: str, sheet: str, header_row: int) -> list[dict]:
    """One record per IOC row: instance, machine_type, index, base, device, ip, source_row."""
    wb = load_workbook(io_path, data_only=True)  # not read_only: need font.strike
    ws = wb[sheet]
    cols = _header_index(ws, header_row)
    for required in ("Script Type", "Index"):
        if required not in cols:
            raise SystemExit(f"column '{required}' not found in sheet '{sheet}' header row {header_row}")
    script_c, index_c = cols["Script Type"], cols["Index"]
    bit_c = cols.get("Bit")
    device_c = cols.get("Device")
    ip_c = cols.get("Profinet IP")
    last_col = ws.max_column

    interfaces = []
    for r in range(header_row + 1, ws.max_row + 1):
        script = ws.cell(row=r, column=script_c).value
        if script is None or str(script).strip().upper() != TRIGGER_TYPE:
            continue
        if _row_struck(ws, r, last_col):
            continue
        instance = str(ws.cell(row=r, column=index_c).value or "").strip()
        if not instance:
            print(f"  WARNING: IOC row {r} has no Index - skipped")
            continue
        machine_type, index = parse_instance(instance)
        interfaces.append({
            "instance": instance,
            "machine_type": machine_type,
            "index": index,
            "base": str(ws.cell(row=r, column=bit_c).value or "").strip() if bit_c else "",
            "device": str(ws.cell(row=r, column=device_c).value or "").strip() if device_c else "",
            "ip": str(ws.cell(row=r, column=ip_c).value or "").strip() if ip_c else "",
            "source_row": r,
        })
    wb.close()
    return interfaces


def choose_sheet(sheet_names: list[str], machine_type: str) -> str:
    """Exact (case-insensitive) type match, else the generic fallback."""
    for name in sheet_names:
        if name.upper() == machine_type.upper():
            return name
    return GENERIC_SHEET


def _safe_name(text: str) -> str:
    bad = '\\/:*?"<>[]'
    return "".join("_" if ch in bad else ch for ch in text)[:31]


def _plug(ws, base: str, index: str) -> None:
    """Plugs the base address (Side-1 row) and index, and replaces <index> tokens."""
    headers = _header_index(ws, 1)
    side_c = headers.get("Side")
    index_c = headers.get("Index")
    base_c = headers.get("Base Address")

    base_num = None
    try:
        base_num = int(str(base).strip())
    except (ValueError, TypeError):
        pass

    # base address into the Side-1 row of the Base Address column
    if base_c and base_num is not None:
        target_row = None
        if side_c:
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(r, side_c).value).strip() in ("1", "1.0"):
                    target_row = r
                    break
        if target_row is None:  # no Side column: first non-empty Base Address cell
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, base_c).value not in (None, ""):
                    target_row = r
                    break
        if target_row:
            ws.cell(target_row, base_c).value = base_num

    # index into every populated Index-column cell
    if index_c and index:
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, index_c).value not in (None, ""):
                ws.cell(r, index_c).value = index

    # replace <index> token in every text cell
    if index:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and INDEX_TOKEN in cell.value:
                    cell.value = cell.value.replace(INDEX_TOKEN, index)


def generate(io_path: str, sheet: str, header_row: int, out_dir: str, template_path: str) -> int:
    if not os.path.exists(template_path):
        raise SystemExit(f"interface template not found: {template_path}")
    os.makedirs(out_dir, exist_ok=True)

    template_sheets = load_workbook(template_path).sheetnames
    interfaces = find_interfaces(io_path, sheet, header_row)
    if not interfaces:
        print("  no IOC interface rows found")
        return 0

    created = 0
    for itf in interfaces:
        target = os.path.join(out_dir, f"IF_{_safe_name(itf['instance'])}.xlsx")
        if os.path.exists(target):
            print(f"  preserved (exists): {os.path.basename(target)}")
            continue

        chosen = choose_sheet(template_sheets, itf["machine_type"])
        shutil.copyfile(template_path, target)
        wb = load_workbook(target)
        for name in list(wb.sheetnames):
            if name != chosen:
                del wb[name]
        ws = wb[chosen]
        ws.title = _safe_name(itf["instance"])
        _plug(ws, itf["base"], itf["index"])
        wb.save(target)
        wb.close()

        created += 1
        note = "" if chosen.upper() == itf["machine_type"].upper() else f" [no '{itf['machine_type']}' sheet -> {GENERIC_SHEET}]"
        print(f"  created: {os.path.basename(target)}  <- IOC row {itf['source_row']} "
              f"({itf['device']}, base {itf['base'] or '?'}){note}")
    return created


def main() -> None:
    ap = argparse.ArgumentParser(description="Generate interface I/O tables from IOC rows of the I/O List")
    ap.add_argument("io_list")
    ap.add_argument("--sheet", default="NET SAFETY 50")
    ap.add_argument("--header-row", type=int, default=1)
    ap.add_argument("--out", default="Output/Interfaces")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    args = ap.parse_args()

    print(f"interface generation from {args.io_list} [{args.sheet}]")
    print(f"template: {args.template}")
    n = generate(args.io_list, args.sheet, args.header_row, args.out, args.template)
    print(f"done: {n} interface table(s) created, others preserved")


if __name__ == "__main__":
    main()
