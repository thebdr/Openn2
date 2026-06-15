#!/usr/bin/env python3
"""Synthetic test for I/O tag + DB generation (template-based names/comments).
Run: python test_outputs.py"""
import os
import shutil
import tempfile
from openpyxl import load_workbook
from safetydb import outputs

failures = 0


def check(name, ok, detail=""):
    global failures
    print(("PASS  " if ok else "FAIL  ") + name + (f"  [{detail}]" if detail else ""))
    if not ok:
        failures += 1


def row(script, fu, loc, dev, bit, d1="", d2="", index="", drawing="", tag_name="",
        db_kind="", db_names=None, db_element="", in_diag=False, diag_desc="", category=None):
    return {
        "script_type": script, "functional_unit": fu, "location": loc, "device": dev,
        "bit": bit, "desc_l1": d1, "desc_l1b": d2, "index": index, "drawing": drawing,
        "_type": {
            "type_id": script, "category": category or ("Interface" if script == "IOC" else "Safety"),
            "tag_name": tag_name, "db_kind": db_kind, "db_names": db_names or [],
            "db_element": db_element, "in_diagnosis": in_diag, "diagnosis_logic": "",
            "diag_desc": diag_desc, "io_comment": "[ {script_type} {index} {drawing} ]",
        },
    }


FLD = " [ {functional_unit}{location}{device} ]"


def main():
    tmp = tempfile.mkdtemp(prefix="outtest_")
    try:
        rows = [
            row("E1/2", "=S1", "+MS1.CC1", "-S67001", "I20.0", index="0001", drawing="DWG1",
                tag_name="Emergency Push Button" + FLD, db_kind="safe_db", db_names=["01_Pushbutton"],
                db_element="Emergency Push Button" + FLD),
            row("A", "=S1", "+MS1.CC1", "-F09001", "I1.0", d1="CIRCUIT BREAKER", d2="400V",
                index="0007", drawing="DWG2", tag_name="{desc_l1} {desc_l1b}" + FLD, in_diag=True),
            row("KQ", "=S1", "+MS1.CC1", "-K1", "Q5.0", index="0003", drawing="DWG3",
                tag_name="Contactor Output" + FLD, db_kind="safe_db",
                db_names=["01_Pushbutton", "01_Pushbutton_RAW"], db_element="Contactor Feedback Error" + FLD),
            row("E2/2", "=S1", "+MS1.CC1", "-S67001", "I20.4", index="0001", tag_name=""),   # untagged ch2
            row("IOC", "=S1", "+MC1.CC1", "-K66201", "10000", index="SORTER-01"),            # interface (excluded)
            row("A", "=S1", "+MS1.CC1", "-X1", "10000", d1="BASE", d2="ADDR", in_diag=True,  # no I/Q addr (excluded)
                tag_name="{desc_l1} {desc_l1b}" + FLD),
        ]
        rows[1]["diag_desc"] = "CIRCUIT BREAKER 400V"   # staging would set this; here set it directly

        tables = outputs.build_io_tags(rows)
        check("one table per type (interface + untagged ch2 + non-IO excluded)",
              set(tables) == {"E1/2", "A", "KQ"}, str(sorted(tables)))
        check("counts (A: 1 taggable, not the base-addr row; E2/2 untagged absent)",
              len(tables["A"]) == 1 and len(tables["E1/2"]) == 1 and len(tables["KQ"]) == 1)

        e = tables["E1/2"][0]
        check("tag name from tag_name template",
              e["name"] == "Emergency Push Button [ =S1+MS1.CC1-S67001 ]", e["name"])
        check("comment from io_comment template", e["comment"] == "[ E1/2 0001 DWG1 ]", e["comment"])
        check("logical address %-prefixed", e["address"] == "%I20.0", e["address"])
        a = tables["A"][0]
        check("alarm tag name from template (desc + FLD)",
              a["name"] == "CIRCUIT BREAKER 400V [ =S1+MS1.CC1-F09001 ]", a["name"])

        outputs.write_io_tags(tables, tmp)
        xlsx = os.path.join(tmp, "IoTags", "PLCTags.xlsx")
        wb = load_workbook(xlsx)
        check("two sheets: PLC Tags + TagTable Properties",
              wb.sheetnames == ["PLC Tags", "TagTable Properties"], str(wb.sheetnames))
        check("PLC Tags header matches TIA export", [c.value for c in wb["PLC Tags"][1]] == outputs.TAG_COLUMNS)
        paths = {r[0].value for r in wb["TagTable Properties"].iter_rows(min_row=2)}
        check("TagTable Properties lists the distinct tables", paths == {"E1/2", "A", "KQ"}, str(paths))

        signal_types = {r["script_type"].upper(): r["_type"] for r in rows}
        dbs = outputs.build_dbs(rows, signal_types)
        check("DBs by db_names incl. the '|' multi (-> two)",
              set(dbs) == {"01_Pushbutton", "01_Pushbutton_RAW"}, str(sorted(dbs)))
        check("shared DB is fail-safe", dbs["01_Pushbutton"]["kind"] == "safe_db")
        mnames = lambda d: [m["name"] for m in dbs[d]["members"]]
        check("every DB opens with ALWAYS_FALSE + ALWAYS_TRUE",
              mnames("01_Pushbutton")[:2] == ["ALWAYS_FALSE", "ALWAYS_TRUE"])
        check("01_Pushbutton combines E1/2 + KQ db_element members",
              "Emergency Push Button [ =S1+MS1.CC1-S67001 ]" in mnames("01_Pushbutton")
              and "Contactor Feedback Error [ =S1+MS1.CC1-K1 ]" in mnames("01_Pushbutton"))
        check("KQ member in both its DBs",
              "Contactor Feedback Error [ =S1+MS1.CC1-K1 ]" in mnames("01_Pushbutton_RAW"))
        kq = next(m for m in dbs["01_Pushbutton"]["members"] if m["name"].startswith("Contactor"))
        check("DB member keeps the io_comment", kq["comment"] == "[ KQ 0003 DWG3 ]", kq["comment"])
        outputs.write_dbs(dbs, tmp)
        body = open(os.path.join(tmp, "DBs", "01_Pushbutton.db"), encoding="utf-8").read()
        check("DB source: DATA_BLOCK + constant + //-commented member",
              'DATA_BLOCK "01_Pushbutton"' in body and '"ALWAYS_FALSE" : Bool;' in body
              and '"Contactor Feedback Error [ =S1+MS1.CC1-K1 ]" : Bool; //[ KQ 0003 DWG3 ]' in body)

        # diagnosis List_IO (config-driven columns): only in_diagnosis types (A)
        diag = outputs.build_diagnosis_list_io(rows)
        check("diagnosis includes only in-diagnosis types", {d["DevType"] for d in diag} == {"A"})
        check("diagnosis row has the config columns incl. PLC_Binding + Diag Desc",
              diag[0]["Device"] == "-F09001" and "PLC_Binding" in diag[0]
              and diag[0]["Diag Desc"] == "CIRCUIT BREAKER 400V")
        n = outputs.write_diagnosis_list_io(diag, tmp)
        check("diagnosis CSV written (address-less rows included)",
              n == 2 and os.path.exists(os.path.join(tmp, "Diagnosis", "List_IO.csv")))

        print("ALL CHECKS PASS" if failures == 0 else f"{failures} CHECK(S) FAILED")
        raise SystemExit(0 if failures == 0 else 1)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
