#!/usr/bin/env python3
"""blockshells.py - the per-template shell spreadsheet + the keep/fill/override directives.

`SoftwareBlocks.xlsm` (generated in the **output folder**) holds one **shell sheet per
template**: the SoftwareBlocksBuilder format laid out in cells - a `$ template=...` row, a
`$ <mode>` directive row, the `%` header (`TemplateType,#Templates Capacity,#Templates
Index,#Elements Needed,!!key$$...`), and the Capacity table - with no `@` data. The
directive (cell B2) is one of:

  keep      - the Python builder generates the CSV; the sheet is ignored (a reference).
  fill      - the Python builder generates the CSV AND writes its @ rows into the sheet
              (so you can see / tweak the generated data).
  override  - the generation reads the @ rows FROM the sheet (you fill them by hand)
              instead of running the builder - for cases not worth automating.

Generating the shells is additive: existing sheets (your manual overrides) are preserved;
only missing templates get a fresh shell.
"""
from __future__ import annotations
import os
import re

import openpyxl

_HERE = os.path.dirname(os.path.abspath(__file__))
MODES = ("keep", "fill", "override")


def _output_dir(params: dict | None = None) -> str:
    from safetydb import config  # lazy: avoid an import cycle
    params = params if params is not None else config.load_params()
    d = params.get("output_dir", "Output")
    return d if os.path.isabs(d) else os.path.normpath(os.path.join(_HERE, d))


def shell_path(params: dict | None = None) -> str:
    """The shell workbook path, in the output folder (`Output/SoftwareBlocks.xlsm`)."""
    return os.path.join(_output_dir(params), "SoftwareBlocks.xlsm")


def _sheet_name(stem: str) -> str:
    """Short, valid sheet name from a template stem (<=31 chars, no special chars)."""
    s = re.sub(r"[\\/?*\[\]:]", "_", stem.rsplit("--", 1)[-1])
    return s[:31] or "template"


def _stem_of(ws) -> str | None:
    """The template stem recorded in the sheet's `$ template=...\\<stem>.xml` row."""
    v = str(ws["B1"].value or "")
    m = re.search(r"([^\\/]+)\.xml\s*$", v)
    return m.group(1) if m else None


def _find_sheet(wb, stem: str):
    for name in wb.sheetnames:
        if _stem_of(wb[name]) == stem:
            return wb[name]
    return None


def _is_mode_row(cells) -> bool:
    return len(cells) >= 2 and cells[0] == "$" and str(cells[1]).strip().lower() in MODES


def _shell_rows(stem: str, db: list, mode: str = "keep") -> list:
    """The shell grid for one template: the structure-only CSV grid with a $ mode row
    inserted after the $ template row, and no @ data."""
    import softwareblocks  # lazy: avoid an import cycle
    grid, _n = softwareblocks.build_template_csv(stem, lambda _db: [], db)
    return [grid[0], ["$", mode]] + grid[1:]


def generate_shells(db: list, path: str | None = None) -> dict:
    """Create/refresh the shell workbook, one sheet per BUILDERS template. Additive:
    existing sheets are left untouched. Returns {created, kept, path}."""
    import block_builders
    path = path or shell_path()
    os.makedirs(os.path.dirname(path), exist_ok=True)
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path, keep_vba=True)
    else:
        wb = openpyxl.Workbook()
    blank = wb["Sheet"] if (wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1) else None

    created = kept = 0
    for stem in block_builders.BUILDERS:
        if _find_sheet(wb, stem) is not None:
            kept += 1
            continue
        ws = wb.create_sheet(_sheet_name(stem))
        for row in _shell_rows(stem, db):
            ws.append(row)
        created += 1

    if blank is not None and created:
        wb.remove(blank)
    wb.save(path)
    return {"created": created, "kept": kept, "path": path}


def modes(path: str | None = None) -> dict:
    """{template stem -> mode} read from each sheet's directive (default 'keep')."""
    path = path or shell_path()
    if not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        stem = _stem_of(ws)
        if stem:
            m = str(ws["B2"].value or "keep").strip().lower()
            out[stem] = m if m in MODES else "keep"
    wb.close()
    return out


def override_grid(stem: str, path: str | None = None) -> list:
    """The CSV grid read from the template's sheet (the $ mode row dropped)."""
    path = path or shell_path()
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path)
    ws = _find_sheet(wb, stem)
    grid = []
    if ws is not None:
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if _is_mode_row(cells):
                continue
            grid.append(cells)
    wb.close()
    return grid


def write_fill(stem: str, grid: list, path: str | None = None) -> None:
    """Write a freshly-built grid (the @ rows) into the template's sheet, keeping its
    current mode directive."""
    path = path or shell_path()
    if not os.path.exists(path):
        return
    wb = openpyxl.load_workbook(path, keep_vba=True)
    ws = _find_sheet(wb, stem)
    if ws is None:
        wb.close()
        return
    mode = str(ws["B2"].value or "fill").strip().lower()
    ws.delete_rows(1, ws.max_row)
    for row in ([grid[0], ["$", mode]] + grid[1:]):
        ws.append(row)
    wb.save(path)


def main():
    import argparse
    import sys as _sys
    if _HERE not in _sys.path:
        _sys.path.insert(0, _HERE)
    from safetydb import config, staging
    ap = argparse.ArgumentParser(description="Generate the per-template shell workbook (Output/SoftwareBlocks.xlsm).")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    db, _w = staging.load_io_list(config.load_params(), config.load_signal_types())
    r = generate_shells(db, args.out)
    print(f"shells: +{r['created']} created, {r['kept']} kept  ->  {r['path']}")


if __name__ == "__main__":
    main()
