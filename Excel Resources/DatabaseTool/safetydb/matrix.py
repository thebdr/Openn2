"""matrix.py - annotate staged rows with the safety AREAs they belong to.

Derived from the Cause & Effect workbook (not the I/O List). `staging.load_io_list`
calls `annotate_areas`, so every staged row carries `matrix_areas` as part of the
pipeline; best-effort - '' when the C&E document is absent.

  * input signal  (I-address): its row in the CAUSE&EFFECT MATRIX carries an "X"
    under the EFFECT/area columns (column S onward; the area name is that column's
    header on the matrix header row, e.g. S2="AREA 1", T2="AREA 2").
  * output signal (Q-address): the AREA n sheets list outputs by their Q address
    (column C); the signal belongs to every AREA sheet whose DIGITAL OUTPUT matches.

`matrix_areas` is the '|'-joined list of area names (e.g. "AREA 1|AREA 2"); area
names are consistent between the matrix headers and the AREA sheet names.
"""
from __future__ import annotations
import os

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from . import config

EFFECT_START_COL = "S"  # the CAUSE&EFFECT MATRIX EFFECT block begins at column S


def _norm(value) -> str:
    """Address/marker key: whitespace removed, upper-cased."""
    return "".join(str(value or "").split()).upper()


def area_lookup(params: dict) -> tuple[dict, dict]:
    """Returns (inputs, outputs):
      inputs:  normalised matrix address -> [area names] (X under the S+ headers)
      outputs: normalised Q address       -> [AREA sheet names that list it]
    Empty dicts if the C&E document is missing."""
    ce = params["ce"]
    if not os.path.exists(ce["path"]):
        return {}, {}

    wb = load_workbook(ce["path"], data_only=True, read_only=True)
    inputs: dict[str, list] = {}
    outputs: dict[str, list] = {}

    # --- CAUSE&EFFECT MATRIX: input address -> areas marked X ---
    sheet = ce["matrix_sheet"]
    if sheet in wb.sheetnames:
        ws = wb[sheet]
        hr, dr = ce["matrix_header_row"], ce["matrix_data_row"]
        addr_c = column_index_from_string(
            {m["canonical"]: m["column"] for m in config.load_column_map("CE")}["address"])
        # area columns = S onward with a non-empty header on the header row
        area_cols = []  # (col_index, area_name)
        for c in range(column_index_from_string(EFFECT_START_COL), ws.max_column + 1):
            name = ws.cell(hr, c).value
            if name not in (None, ""):
                area_cols.append((c, str(name).strip()))
        for r in range(dr, ws.max_row + 1):
            addr = _norm(ws.cell(r, addr_c).value)
            if not addr:
                continue
            areas = [nm for c, nm in area_cols
                     if str(ws.cell(r, c).value or "").strip().upper() == "X"]
            if areas:
                inputs[addr] = areas

    # --- AREA n sheets: output Q address -> sheet names ---
    out_addr_c = column_index_from_string(
        {m["canonical"]: m["column"] for m in config.load_column_map("AREA")}["address"])
    adr = ce["area_data_row"]
    for s in wb.sheetnames:
        if s.strip().upper().startswith("AREA") and any(ch.isdigit() for ch in s):
            a = wb[s]
            for r in range(adr, a.max_row + 1):
                q = _norm(a.cell(r, out_addr_c).value)
                if q:
                    outputs.setdefault(q, []).append(s.strip())

    wb.close()
    return inputs, outputs


def annotate_areas(params: dict, rows: list) -> list:
    """Set row['matrix_areas'] ('|'-joined area names) on every staged row, by its
    I/Q address. Paired channels of one device (same pair_key + FUNCTIONAL UNIT +
    LOCATION + DEVICE) share the union of their areas, so channel-2 inherits the
    area its channel-1 sibling is marked in (the matrix lists the device once).
    Rows without an I/Q bit get ''. Returns the same list."""
    inputs, outputs = area_lookup(params)
    found = {}  # id(row) -> [areas]
    for row in rows:
        bit = _norm(row.get("bit"))
        head = bit[:1]
        if head == "I":
            found[id(row)] = list(inputs.get(bit, []))
        elif head == "Q":
            found[id(row)] = list(outputs.get(bit, []))
        else:
            found[id(row)] = []

    # pair inheritance: union areas across each paired device's channels
    groups: dict = {}
    for row in rows:
        pk = ((row.get("_type") or {}).get("pair_key") or "").strip().upper()
        if not pk:
            continue
        key = (pk, _norm(row.get("functional_unit")), _norm(row.get("location")), _norm(row.get("device")))
        groups.setdefault(key, []).append(row)
    for members in groups.values():
        union = []
        for m in members:
            for a in found[id(m)]:
                if a not in union:
                    union.append(a)
        if union:
            for m in members:
                found[id(m)] = union

    for row in rows:
        row["matrix_areas"] = "|".join(found[id(row)])
    return rows
